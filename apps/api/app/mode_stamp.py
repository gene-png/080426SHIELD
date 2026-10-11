"""Whether an assessment's AI suggestions came from a live model or test data (#646).

Fixture output (D-017) is canned test data. A deliverable or a dashboard built
on it read exactly like one built on a live model. This module answers, for ONE
subject -- the assessment (or Tech Debt capability list) a deliverable or a
dashboard is built from -- which of five states holds, and says it in one
sentence every surface prints.

DERIVED, never stored: from the subject's COMPLETED `ai_runs` (#645), whose
`mode` is what the provider actually served (`_mode_for(provider_serves)`), so
a stored key that forces live (D-037) is recorded as live. Only COMPLETED runs:
a failed or reaped run applied nothing (the completion compare-and-swap). The
unit is the RUN, not the call: a batched run is many calls in one mode.

ONE function per subject kind, CALLED by every surface -- the exporters, the
client dashboards and the admin workspace -- never copied.

The population is the subject, not the service. A run on a discarded draft, or
on another version, says nothing about this one.

`unknown` is its own state, never read as `none`: AI calls made before #756
carry no run, so where any such call could belong to this subject, nothing on
record says which mode drafted it.
"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Literal

from sqlalchemy import exists, select
from sqlalchemy.orm import Session

from app.logging import get_logger
from app.models.ai_run import AiRun, AiRunStatus
from app.models.llm_call import LLMCall, LLMCallMode, LLMCallStatus
from app.models.service import Service, ServiceKind

_log = get_logger(__name__)

AiModeState = Literal["live", "fixture", "mixed", "none", "unknown"]
Subject = Literal["assessment", "capability list", "register"]


@dataclass(frozen=True)
class AiModeStamp:
    """How many completed AI runs on one subject were live, and how many offline.

    `recorded=False` is "nothing on record says": either nobody looked (a
    context built without a lookup) or AI calls exist that no run accounts for.
    """

    live_runs: int
    fixture_runs: int
    recorded: bool = True
    subject: Subject = "assessment"

    @property
    def state(self) -> AiModeState:
        if not self.recorded:
            return "unknown"
        if self.fixture_runs and self.live_runs:
            return "mixed"
        if self.fixture_runs:
            return "fixture"
        if self.live_runs:
            return "live"
        return "none"

    @property
    def is_warning(self) -> bool:
        """True where a reader must not take AI-drafted values as analysis."""
        return self.state in ("fixture", "mixed", "unknown")

    def sentence(self) -> str:
        """The line every surface prints. Its wording depends only on `state`."""
        what = self.subject
        state = self.state
        if state == "fixture":
            return (
                f"OFFLINE TEST DATA: every AI suggestion in this {what} came from "
                "built-in test data, not a live AI model. Treat AI-drafted values as "
                "placeholders, not analysis."
            )
        if state == "mixed":
            total = self.live_runs + self.fixture_runs
            return (
                f"OFFLINE TEST DATA: {self.fixture_runs} of the {total} AI runs on this "
                f"{what} used built-in test data, not a live AI model. Values they "
                "drafted are placeholders, not analysis."
            )
        if state == "live":
            return f"AI suggestions in this {what} came from a live AI model."
        if state == "none":
            return f"No AI suggestions were used in this {what}."
        return (
            f"It is not recorded whether the AI suggestions in this {what} came from a "
            "live AI model or from offline test data."
        )

    def as_api(self) -> dict[str, Any]:
        """The shape every API surface carries (`AiSource`)."""
        return {
            "state": self.state,
            "sentence": self.sentence(),
            "live_runs": self.live_runs,
            "fixture_runs": self.fixture_runs,
        }


#: For a context built without a lookup. Renders "not recorded", never live.
UNKNOWN_AI_MODE = AiModeStamp(live_runs=0, fixture_runs=0, recorded=False)

#: The Risk Register's, always, until Risk runs through the run framework (#504).
UNKNOWN_AI_MODE_REGISTER = AiModeStamp(
    live_runs=0, fixture_runs=0, recorded=False, subject="register"
)


def _run_counts(runs: list[AiRun]) -> tuple[int, int]:
    live = sum(1 for r in runs if r.mode == LLMCallMode.LIVE)
    return live, len(runs) - live


def _unattributed_calls(
    db: Session, *, service_id: uuid.UUID, call_purpose: str, since: datetime | None
) -> bool:
    """Completed AI calls of this purpose for this service that no run accounts
    for (made before #756), at or after `since` when it is given."""
    conditions = [
        LLMCall.service_id == service_id,
        LLMCall.purpose == call_purpose,
        LLMCall.status == LLMCallStatus.COMPLETED,
        LLMCall.ai_run_id.is_(None),
    ]
    if since is not None:
        conditions.append(LLMCall.created_at >= since)
    return bool(db.execute(select(exists().where(*conditions))).scalar())


def ai_mode_for_assessment(
    db: Session,
    *,
    service_id: uuid.UUID,
    run_purpose: str,
    call_purpose: str,
    assessment_id: uuid.UUID,
    assessment_created_at: datetime,
) -> AiModeStamp:
    """The stamp for one assessment: its own completed runs.

    `unknown` when the service holds pre-#756 AI calls of this purpose made at
    or after the assessment was created: they may have drafted it, and nothing
    says in which mode.
    """
    runs = list(
        db.execute(
            select(AiRun).where(
                AiRun.service_id == service_id,
                AiRun.purpose == run_purpose,
                AiRun.subject_id == assessment_id,
                AiRun.status == AiRunStatus.COMPLETED,
            )
        ).scalars()
    )
    live, fixture = _run_counts(runs)
    recorded = not _unattributed_calls(
        db, service_id=service_id, call_purpose=call_purpose, since=assessment_created_at
    )
    stamp = AiModeStamp(live_runs=live, fixture_runs=fixture, recorded=recorded)
    _log.info(
        "ai_mode_stamp.assessment",
        service_id=str(service_id),
        assessment_id=str(assessment_id),
        state=stamp.state,
        live_runs=live,
        fixture_runs=fixture,
    )
    return stamp


def ai_mode_for_capability_list(
    db: Session,
    *,
    service_id: uuid.UUID,
    run_purpose: str,
    call_purpose: str,
    capability_list_id: uuid.UUID,
) -> AiModeStamp:
    """The stamp for one Tech Debt capability list.

    A list is written by exactly one extraction run, whose stored result names
    it (`capability_list_id`). Where no completed run names it, the list was
    extracted before #756 if the service holds any unattributed extraction
    call -- `unknown` -- and by no AI at all otherwise.
    """
    candidates = db.execute(
        select(AiRun).where(
            AiRun.service_id == service_id,
            AiRun.purpose == run_purpose,
            AiRun.status == AiRunStatus.COMPLETED,
        )
    ).scalars()
    runs = [
        r
        for r in candidates
        if isinstance(r.result, dict)
        and r.result.get("capability_list_id") == str(capability_list_id)
    ]
    live, fixture = _run_counts(runs)
    recorded = bool(runs) or not _unattributed_calls(
        db, service_id=service_id, call_purpose=call_purpose, since=None
    )
    stamp = AiModeStamp(
        live_runs=live, fixture_runs=fixture, recorded=recorded, subject="capability list"
    )
    _log.info(
        "ai_mode_stamp.capability_list",
        service_id=str(service_id),
        capability_list_id=str(capability_list_id),
        state=stamp.state,
        live_runs=live,
        fixture_runs=fixture,
    )
    return stamp


#: (run purpose, `llm_calls` purpose) per service kind. Tech Debt's run is
#: `tech_debt_extract`; its calls keep the historical `extract.capabilities`.
_PURPOSES: dict[ServiceKind, tuple[str, str]] = {
    ServiceKind.ATTACK_COVERAGE: ("mitre_map", "mitre_map"),
    ServiceKind.NIST_CSF: ("csf_score", "csf_score"),
    ServiceKind.ZERO_TRUST_CISA: ("zt_score", "zt_score"),
    ServiceKind.ZERO_TRUST_DOD: ("zt_score", "zt_score"),
    ServiceKind.TECH_DEBT: ("tech_debt_extract", "extract.capabilities"),
}


def ai_mode_for(db: Session, service: Service, subject: Any) -> AiModeStamp:
    """THE entry point every surface calls: the stamp for `subject` -- an
    assessment of `service`, or for Tech Debt a capability list. Raises for a
    service kind with no AI run (Risk is not a service, #504)."""
    if service.kind not in _PURPOSES:
        raise ValueError(f"no AI run purpose for service kind {service.kind!r}")
    run_purpose, call_purpose = _PURPOSES[service.kind]
    if service.kind == ServiceKind.TECH_DEBT:
        return ai_mode_for_capability_list(
            db,
            service_id=service.id,
            run_purpose=run_purpose,
            call_purpose=call_purpose,
            capability_list_id=subject.id,
        )
    return ai_mode_for_assessment(
        db,
        service_id=service.id,
        run_purpose=run_purpose,
        call_purpose=call_purpose,
        assessment_id=subject.id,
        assessment_created_at=subject.created_at,
    )


# ---------------------------------------------------------------------------
# Rendering. ONE definition per format, so every deliverable prints the same
# words in the same place: directly under the title, before any figure.
# ---------------------------------------------------------------------------

XLSX_SHEET_TITLE = "AI source"


def pdf_paragraph(stamp: AiModeStamp, style: Any) -> Any:
    """The stamp as a reportlab Paragraph, bold when it is a warning. Escaped:
    a Paragraph is markup."""
    from html import escape

    from reportlab.platypus import Paragraph

    text = escape(stamp.sentence())
    return Paragraph(f"<b>{text}</b>" if stamp.is_warning else text, style)


def add_docx_paragraph(doc: Any, stamp: AiModeStamp) -> None:
    """The stamp as a DOCX paragraph, bold when it is a warning. Through
    `app.docx_export`, which strips what XML 1.0 refuses (#993)."""
    from app.docx_export import add_paragraph

    add_paragraph(doc, stamp.sentence(), bold=stamp.is_warning)


def xlsx_rows(stamp: AiModeStamp) -> tuple[list[Any], list[list[Any]]]:
    """The "AI source" sheet's heading row and data rows: ONE definition, for
    `add_xlsx_sheet` and for the CSF Playbook, whose every sheet carries the
    approval banner above its heading (#294)."""
    return (
        ["AI source", stamp.state],
        [
            [stamp.sentence()],
            ["Live AI runs", stamp.live_runs if stamp.recorded else "not recorded"],
            ["Offline test-data runs", stamp.fixture_runs if stamp.recorded else "not recorded"],
        ],
    )


def add_xlsx_sheet(wb: Any, stamp: AiModeStamp) -> None:
    """The stamp as its own LAST sheet.

    Last, not first: a workbook's first sheet is what `wb.active` opens and
    what existing readers index by row, so a sheet inserted before it would
    move their data. The PDF and DOCX carry the stamp under the title.
    """
    from openpyxl.styles import Font

    from app.xlsx_export import safe_text_row  # #993: the #972 strip, twin of the DOCX

    ws = wb.create_sheet(XLSX_SHEET_TITLE)
    heading, rows = xlsx_rows(stamp)
    safe_text_row(ws, heading)
    for row in rows:
        safe_text_row(ws, row)
    ws.cell(row=1, column=1).font = Font(bold=True)
    ws.cell(row=2, column=1).font = Font(bold=stamp.is_warning)
    ws.column_dimensions["A"].width = 24
