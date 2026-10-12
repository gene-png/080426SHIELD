"""Shared document text extraction (N2, for #1071; design note rev 4, 15.2).

ONE reader set, ONE set of limits and ONE "not read" vocabulary, for EV-1's
evidence-link job (P4) and EV-2's client summary. Both import this module;
neither reimplements a reader ("same as" enforced by calling).

Readers: PDF (pypdf: page text, annotation text and filled AcroForm values),
DOCX (its XML,
walked directly: body, headers, footers, footnotes, endnotes, comments, content
controls, text boxes, tracked insertions, nested tables), XLSX (openpyxl, XML
parsed through defusedxml) and plain text. Images, RTF, legacy Office files,
other archives and unrecognised binaries are `not_read: unsupported_type` until
a reader exists for them.

**"I could not read it" and "it was empty" are different results.** A file
that cannot be read is `status="not_read"` with a typed reason and NO text; a
file that was read and holds nothing is `status="read"` with zero characters.
`DocText` refuses any shape that blurs the two (CLAUDE.md, the silent-success
rule: the "could not look" branch must never share the "nothing there" branch).
What a reader skips INSIDE a read file is counted, never dropped:
`pages_without_text` (PDF pages with no text layer) and `parts_not_read`
(document parts no reader opened, such as an XFA form or a DOCX altChunk).

The type is decided from the file's MAGIC BYTES. The declared type (the
upload's mime and the title's extension) is a claim checked against them: a
declaration that disagrees with the bytes is `unreadable`, never trusted.

Extraction STOPS EARLY at `max_chars`, at `max_pages` and at `max_seconds`
(checked per page, row and paragraph); `truncated` says so, and a time limit
reached before any text is `too_large`.

Extraction stops at a PIECE boundary (a page, row, paragraph or text line),
so `text` may run past `max_chars` by the last piece read. A piece that,
with its separating newline, is longer than `max_chars` is instead cut at the
last whitespace inside what is left of the limit. So `len(text)` is at most
`2 * max_chars`, separators included, and `chunk_text` refuses anything
longer (`TextOverCeiling`).

A cut can leave the START of a name or a number at the end of the text,
which the redactor (whole literals and whole shapes only) cannot see. When
the source was cut, or its truncation is unknown, `chunk_text` drops such a
trailing fragment BEFORE redacting, and counts it on the plan.

**Redaction runs over the WHOLE text, before chunking, inside `chunk_text`.**
It takes the raw text and the same inputs as `app.ai.redact.redact_for_ai`,
calls it once on the whole text, and only then cuts chunks; it refuses mode
"off". Redacting chunk by chunk is NOT safe: the redactor matches whole
literals and whole shapes, so a boundary inside a name, a phone number or an
address leaves fragments it cannot see (review round 3 of #1089).

**A signature cut makes the plan INCOMPLETE.** The redactor's signature rule
deletes from a sign-off to the END of the text. `ChunkPlan.signature_cut` and
`complete` say so, with the characters and lines removed, so a consumer never
reads the missing tail (a policy after a cover letter, an earlier message in
an email thread) as absent evidence (review round 4 of #1089).

Wired to nothing yet: no route and no AI call reads it.
"""

from __future__ import annotations

import functools
import io
import re
import time
import zipfile
from collections.abc import Callable, Iterable, Iterator, Mapping
from dataclasses import dataclass
from itertools import zip_longest
from typing import TYPE_CHECKING, Literal, Protocol, get_args

import openpyxl.xml
from defusedxml.ElementTree import fromstring as safe_fromstring
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError
from pypdf.generic import NameObject

from app.ai.redact import PLACEHOLDER_SIGNATURE, _redact_signature_blocks, redact_for_ai
from app.config import RedactionMode, Settings, get_settings
from app.logging import get_logger
from app.models.artifact import Artifact

if TYPE_CHECKING:
    # Annotation only: every parse goes through defusedxml (`safe_fromstring`).
    from xml.etree.ElementTree import Element  # nosec B405

_log = get_logger(__name__)

NotReadReason = Literal[
    "unsupported_type",
    "no_text_layer",
    "encrypted",
    "unreadable",
    "too_large",
    "storage_missing",
]

#: The whole "not read" vocabulary, for the one copy map both consumers render.
NOT_READ_REASONS: tuple[str, ...] = get_args(NotReadReason)


@dataclass(frozen=True)
class DocText:
    """What one file yielded.

    `chars_total` is `len(text)`, the characters EXTRACTED. When `truncated`
    is False, `text` is the whole document's text. When it is True, it is
    not: extraction stopped after the piece (page, row, paragraph or line)
    that passed `max_chars`, or at the page or time limit, so `chars_total`
    is a lower bound on the document's length.

    Where the stop falls: normally after a whole piece, so `text` may run past
    `max_chars` by up to one piece. A piece that would itself take more than
    `max_chars` (one huge line, paragraph or row) IS cut inside, at the last
    whitespace within the limit, or mid-token where the span has none. Either
    way `len(text) <= 2 * max_chars`. A stop can split a name or a number,
    across two pieces or at an in-piece cut; `chunk_text` drops such a
    trailing fragment before redacting (see the module docstring).

    `pages_without_text` counts PDF pages read that gave no text (a partly
    scanned file). `parts_not_read` counts content in a read file that no
    reader opened or that is deliberately not quoted: in a PDF, an XFA form,
    an unreadable AcroForm, an embedded file or a media annotation; in a DOCX,
    an altChunk, an embedded object, a picture, SmartArt, a chart, any other
    drawing or VML graphic that gives no text, an OMML equation, a symbol
    (`w:sym`) with no known character, a hidden run, a ruby (pronunciation)
    annotation, or an unparseable header, footer, footnotes, endnotes or
    comments part; in an XLSX, each cell comment, each worksheet drawing and
    each chartsheet. Both are additions to section 15.2's shape, and both are
    0 when nothing was skipped.

    A file whose only content is of those kinds is never an empty `read`: it
    is `not_read: no_text_layer`, meaning "no readable text; the content is in
    parts not read" (pictures, charts, equations), not only "a scan".
    """

    status: Literal["read", "not_read"]
    reason: NotReadReason | None
    text: str
    chars_total: int
    truncated: bool
    pages_without_text: int = 0
    parts_not_read: int = 0

    def __post_init__(self) -> None:
        if self.status == "read":
            if self.reason is not None:
                raise ValueError(f"a read result carries no reason, got {self.reason!r}")
            return
        if self.status != "not_read":
            raise ValueError(f"unknown status {self.status!r}")
        if self.reason not in NOT_READ_REASONS:
            raise ValueError(
                f"not_read needs a reason from {NOT_READ_REASONS}, got {self.reason!r}"
            )
        if (
            self.text
            or self.chars_total
            or self.truncated
            or self.pages_without_text
            or self.parts_not_read
        ):
            raise ValueError("a not_read result carries no text and no counts")


class TextOverCeiling(ValueError):
    """`chunk_text` was handed more than `2 * max_chars` characters
    (separators included), which `document_text` never returns. The redactor
    runs over the whole text in one call, so an unbounded text is an
    unbounded redaction."""


class RedactionOffRefused(ValueError):
    """`chunk_text` was asked to chunk with redaction mode "off". Text cut
    for an AI payload is never unredacted, whatever the configured mode."""


@dataclass(frozen=True)
class ChunkPlan:
    """What `chunk_text` produced from one document's text.

    `chunks` are cut from the REDACTED text; they do not overlap and join back
    to it, and `offsets[i]` is where `chunks[i]` starts in it. `sent` is
    `len(chunks)`; `skipped` is how many more the text needed past the cap.

    `redaction_counts` is `redact_for_ai`'s count for the WHOLE text, skipped
    chunks included. A caller wiring this to a live AI call records it: the
    egress client's own redaction pass will find almost nothing left in each
    chunk, so its per-call count is not the document's.

    `signature_cut` is True when the redactor's signature rule deleted the
    text from a sign-off to the end; `signature_chars_removed` and
    `signature_lines_removed` measure what went.

    `complete` is THREE-VALUED. It is True only when there was no signature
    cut, no chunk was skipped, AND the source is known not to be truncated
    (`source_truncated is False`). It is False when any of those failed. It is
    None when nothing failed but the caller did not know whether the source
    was cut (`source_truncated is None`). Only True means the chunks are the
    whole document; under False or None, nothing missing from them may be
    read as absent from it.
    """

    chunks: tuple[str, ...]
    sent: int
    skipped: int
    offsets: tuple[int, ...]
    redaction_counts: Mapping[str, int]
    signature_cut: bool
    signature_chars_removed: int
    signature_lines_removed: int
    source_truncated: bool | None
    #: Trailing tokens dropped from a cut source before redaction: the start
    #: of a name hint or the legal name, or a token with a digit (a partial
    #: phone number or ID), which the redactor could not have matched.
    tail_tokens_dropped: int
    tail_chars_dropped: int

    @property
    def complete(self) -> bool | None:
        if self.signature_cut or self.skipped or self.source_truncated:
            return False
        if self.source_truncated is None:
            return None
        return True


@dataclass(frozen=True)
class ExtractLimits:
    """The one set of limits. Built from settings by `extract_limits`."""

    max_bytes: int
    max_unzipped_bytes: int
    max_chars: int
    chunk_chars: int
    max_chunks: int
    max_pages: int
    max_seconds: float


def extract_limits(settings: Settings | None = None) -> ExtractLimits:
    """The limits as configured (`shield_document_*` in app/config.py)."""
    s = settings if settings is not None else get_settings()
    return ExtractLimits(
        max_bytes=s.shield_document_max_bytes,
        max_unzipped_bytes=s.shield_document_max_unzipped_bytes,
        max_chars=s.shield_document_max_chars,
        chunk_chars=s.shield_document_chunk_chars,
        max_chunks=s.shield_document_max_chunks,
        max_pages=s.shield_document_max_pages,
        max_seconds=s.shield_document_max_seconds,
    )


class _Readable(Protocol):
    def get(self, key: str) -> bytes: ...


Kind = Literal["pdf", "docx", "xlsx", "txt", "image", "ole", "zip", "rtf"]

#: Kinds with no text reader yet.
_UNSUPPORTED: frozenset[str] = frozenset({"image", "ole", "zip", "rtf"})

_MIME_KIND: dict[str, Kind] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "txt",
    "text/csv": "txt",
    "text/markdown": "txt",
    "application/msword": "ole",
    "application/vnd.ms-excel": "ole",
    "application/zip": "zip",
    "application/rtf": "rtf",
    "text/rtf": "rtf",
}

_EXT_KIND: dict[str, Kind] = {
    "pdf": "pdf",
    "docx": "docx",
    "xlsx": "xlsx",
    "txt": "txt",
    "csv": "txt",
    "md": "txt",
    "png": "image",
    "jpg": "image",
    "jpeg": "image",
    "gif": "image",
    "tif": "image",
    "tiff": "image",
    "bmp": "image",
    "webp": "image",
    "doc": "ole",
    "xls": "ole",
    "zip": "zip",
    "rtf": "rtf",
}

_IMAGE_MAGIC = (
    b"\x89PNG\r\n\x1a\n",
    b"\xff\xd8\xff",
    b"GIF87a",
    b"GIF89a",
    b"II*\x00",
    b"MM\x00*",
)
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_UTF32_BOMS = (b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")
_UTF16_BOMS = (b"\xff\xfe", b"\xfe\xff")


class _NotRead(Exception):
    """Internal: a reader decided the file is not read, and why."""

    def __init__(self, reason: NotReadReason, detail: str) -> None:
        super().__init__(detail)
        self.reason = reason
        self.detail = detail


@dataclass
class _Tally:
    """What a reader skipped inside a file it did read, and its deadline."""

    deadline: float
    pages_without_text: int = 0
    parts_not_read: int = 0
    stopped_early: bool = False
    timed_out: bool = False

    def out_of_time(self) -> bool:
        """True once the time limit has passed; it then marks the read as
        stopped early. Readers call it per page, row and paragraph, so input
        that yields no text (empty rows, empty paragraphs) is still bounded."""
        if _clock() > self.deadline:
            self.stopped_early = True
            self.timed_out = True
            return True
        return False


#: The clock the time limit reads; a module attribute so a test can drive it.
_clock = time.monotonic


def document_text(storage: _Readable, artifact: Artifact, *, limits: ExtractLimits) -> DocText:
    """Extract one artifact's text, or say exactly why it was not read.

    The text is UNREDACTED client content. It must reach an AI provider only
    through `chunk_text` (which redacts it whole) or the egress client in
    app/ai/llm.py, never directly.

    A storage MISS (`FileNotFoundError`, which every backend raises for one) is
    `storage_missing`. Any other storage failure propagates: an unreachable
    store is not a fact about the file and must not be recorded as one.
    """
    try:
        data = storage.get(artifact.file_storage_key)
    except FileNotFoundError:
        return _not_read(artifact, "storage_missing", "no object at the storage key")
    tally = _Tally(deadline=_clock() + limits.max_seconds)
    try:
        text, past_max_chars = _collect(
            _extract(data, artifact, limits, tally), limits.max_chars, tally
        )
    except _NotRead as nr:
        return _not_read(artifact, nr.reason, nr.detail)
    if not text.strip():
        if tally.timed_out:
            return _not_read(
                artifact, "too_large", f"no text before the {limits.max_seconds}s time limit"
            )
        if tally.parts_not_read:
            # The file's only content is in parts no reader opened (pictures,
            # charts, an embedded object): not an empty document.
            return _not_read(
                artifact,
                "no_text_layer",
                f"no text; {tally.parts_not_read} part(s) not read",
            )
    truncated = past_max_chars or tally.stopped_early
    result = DocText(
        status="read",
        reason=None,
        text=text,
        chars_total=len(text),
        truncated=truncated,
        pages_without_text=tally.pages_without_text,
        parts_not_read=tally.parts_not_read,
    )
    _log.info(
        "documents.text.read",
        artifact_id=str(artifact.id),
        chars_total=result.chars_total,
        truncated=truncated,
        pages_without_text=result.pages_without_text,
        parts_not_read=result.parts_not_read,
    )
    return result


def _collect(pieces: Iterator[str], max_chars: int, tally: _Tally) -> tuple[str, bool]:
    """Join a reader's pieces, sanitised, and STOP once past `max_chars` or
    past the time limit: no further page, row or part is extracted. Returns
    the text read and whether `max_chars` stopped it (the time limit marks
    `tally.stopped_early`).

    The time limit is checked BETWEEN pieces (and between PDF pages), so one
    piece that hangs inside a parser is not interrupted: bounding that needs
    the parse in a killable child process, as #552 does for redaction."""
    parts: list[str] = []
    count = 0
    try:
        for piece in pieces:
            clean = _sanitize(piece)
            sep = 1 if parts else 0
            if sep + len(clean) > max_chars:
                # A piece that, with its separator, is longer than the whole
                # limit (a one-line text file, a huge paragraph or row): cut
                # it to what is left of the limit, at the last whitespace. The
                # text before it is at most max_chars and any piece kept whole
                # adds at most max_chars more, separator included, so
                # len(text) <= 2 * max_chars.
                kept = _cut_at_whitespace(clean, max_chars - count - sep)
                if kept:
                    parts.append(kept)
                return "\n".join(parts), True
            count += len(clean) + sep
            parts.append(clean)
            if count > max_chars:
                return "\n".join(parts), True
            if tally.out_of_time():
                break
    finally:
        close = getattr(pieces, "close", None)
        if close is not None:
            close()
    return "\n".join(parts), False


def _cut_at_whitespace(text: str, budget: int) -> str:
    """At most `budget` characters of `text`, ending before the last
    whitespace in that span; a span with no whitespace is cut at `budget`, and
    that one token is then split (a name or number in it can be split too)."""
    if budget <= 0:
        return ""
    window = text[:budget]
    last = max(window.rfind(" "), window.rfind("\n"), window.rfind("\t"))
    return window[:last] if last > 0 else window


def _not_read(artifact: Artifact, reason: NotReadReason, detail: str) -> DocText:
    _log.warning(
        "documents.text.not_read",
        artifact_id=str(artifact.id),
        reason=reason,
        detail=detail,
    )
    return DocText(status="not_read", reason=reason, text="", chars_total=0, truncated=False)


def _extract(
    data: bytes, artifact: Artifact, limits: ExtractLimits, tally: _Tally
) -> Iterator[str]:
    if len(data) > limits.max_bytes:
        raise _NotRead("too_large", f"{len(data)} bytes over the {limits.max_bytes} byte limit")
    detected = _detect(data, limits)
    declared = _declared(artifact)
    mismatched = sorted(d for d in declared if d != detected)
    if mismatched:
        raise _NotRead("unreadable", f"declared {mismatched} but the bytes are {detected}")
    if detected in _UNSUPPORTED:
        raise _NotRead("unsupported_type", f"{detected} has no text reader")
    if detected == "pdf":
        return _read_pdf(data, limits, tally)
    if detected == "docx":
        return _read_docx(data, tally)
    if detected == "xlsx":
        if not openpyxl.xml.DEFUSEDXML:
            # Not a fact about the file: the api is missing a declared runtime
            # dependency, and must not parse a client's XML without it.
            raise RuntimeError("defusedxml is not importable; refusing to parse XLSX")
        return _read_xlsx(data, tally)
    return _read_txt(data, declared_text="txt" in declared)


def _declared(artifact: Artifact) -> set[Kind]:
    """The kinds the upload CLAIMS, from its mime and its title's extension.
    A generic mime (octet-stream) or no known extension claims nothing."""
    kinds: set[Kind] = set()
    mime = (artifact.mime_type or "").split(";")[0].strip().lower()
    if mime in _MIME_KIND:
        kinds.add(_MIME_KIND[mime])
    elif mime.startswith("image/"):
        kinds.add("image")
    title = artifact.title or ""
    if "." in title:
        ext = title.rsplit(".", 1)[1].lower()
        if ext in _EXT_KIND:
            kinds.add(_EXT_KIND[ext])
    return kinds


def _detect(data: bytes, limits: ExtractLimits) -> Kind:
    """The kind the BYTES are. Anything without a known signature is a text
    candidate; `_read_txt` decides whether it really is text."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(_OLE_MAGIC):
        return "ole"
    if data.startswith(b"{\\rtf"):
        return "rtf"
    if (
        data.startswith(_IMAGE_MAGIC)
        or (data[:4] == b"RIFF" and data[8:12] == b"WEBP")
        or _is_bmp(data)
    ):
        return "image"
    if data.startswith(b"PK\x03\x04"):
        return _detect_zip(data, limits)
    return "txt"


def _is_bmp(data: bytes) -> bool:
    """A BMP header, not just "BM": the stored file size matches, the four
    reserved bytes are zero, and the pixel offset lies inside the file.
    Text that happens to start "BM" ("BMS ID,Location") fails all three."""
    if len(data) < 26 or not data.startswith(b"BM"):
        return False
    size = int.from_bytes(data[2:6], "little")
    offset = int.from_bytes(data[10:14], "little")
    return size == len(data) and data[6:10] == b"\x00\x00\x00\x00" and 26 <= offset <= len(data)


def _detect_zip(data: bytes, limits: ExtractLimits) -> Kind:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, EOFError, ValueError) as exc:
        raise _NotRead("unreadable", f"zip container: {type(exc).__name__}: {exc}") from exc
    inflated = sum(i.file_size for i in infos)
    if inflated > limits.max_unzipped_bytes:
        raise _NotRead(
            "too_large",
            f"parts inflate to {inflated} bytes, over the {limits.max_unzipped_bytes} byte limit",
        )
    names = {i.filename for i in infos}
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    return "zip"


def _typed_reader(label: str) -> Callable:
    """Wrap a reader generator so any library failure inside it becomes
    `not_read: unreadable` with its type and message. `_NotRead` passes
    through; GeneratorExit (an early stop) is not an Exception and is not
    caught."""

    def deco(fn: Callable[..., Iterator[str]]) -> Callable[..., Iterator[str]]:
        @functools.wraps(fn)
        def wrapper(*args, **kwargs) -> Iterator[str]:
            try:
                yield from fn(*args, **kwargs)
            except _NotRead:
                raise
            except FileNotDecryptedError as exc:
                # pypdf tries the empty user password itself, which opens
                # owner-password-only files (any viewer opens those too). This
                # is raised only for a file that needs a password we lack.
                raise _NotRead("encrypted", f"{label}: {exc}") from exc
            except Exception as exc:  # noqa: BLE001 - typed and logged, never swallowed
                # Parsers raise a wide family on a damaged file (pypdf's stream
                # errors, zip and XML errors, KeyError/ValueError from deep in a
                # parse, RecursionError on absurd nesting). Each means the same
                # thing to a caller: this file was not read.
                raise _NotRead("unreadable", f"{label}: {type(exc).__name__}: {exc}") from exc

        return wrapper

    return deco


@_typed_reader("PDF")
def _read_pdf(data: bytes, limits: ExtractLimits, tally: _Tally) -> Iterator[str]:
    reader = PdfReader(io.BytesIO(data))
    pages = reader.pages
    total = len(pages)
    read_any = False
    timed_out = False
    for i in range(min(total, limits.max_pages)):
        if tally.out_of_time():
            timed_out = True
            break
        text = pages[i].extract_text() or ""
        if text.strip():
            read_any = True
            yield text
        else:
            tally.pages_without_text += 1
        for note in _pdf_annotations(pages[i], tally):
            read_any = True
            yield note
    capped = total > limits.max_pages
    if capped:
        tally.stopped_early = True
    if not timed_out and not tally.out_of_time():
        for value in _pdf_form_values(reader, tally):
            read_any = True
            yield value
        _count_pdf_attachments(reader, tally)
    else:
        timed_out = True
    if not read_any:
        if timed_out:
            raise _NotRead(
                "too_large",
                f"no text before the {limits.max_seconds}s time limit, {total} pages",
            )
        if capped:
            # The pages never looked at may hold text: not evidence of none.
            raise _NotRead(
                "too_large",
                f"no text in the first {limits.max_pages} of {total} pages; the rest is over "
                "the page limit",
            )
        raise _NotRead("no_text_layer", f"none of {total} pages has a text layer")


#: Annotations that carry no reviewer text of their own: a form widget (its
#: value is read as a form field), a link, and a popup (it repeats its
#: parent's text).
_PDF_ANNOT_SKIP = frozenset({"/Widget", "/Link", "/Popup"})
#: Annotations whose content is a file or media this reader does not open.
_PDF_ANNOT_NOT_READ = frozenset(
    {"/FileAttachment", "/RichMedia", "/Movie", "/Sound", "/Screen", "/3D"}
)


def _pdf_annotations(page: object, tally: _Tally) -> list[str]:
    """Annotation text on one page (comments, notes, free text, markup), read
    the way DOCX comments are, as "[Annotation] text" lines. Media and file
    attachments are counted as parts not read."""
    lines: list[str] = []
    try:
        annots = page.get("/Annots")  # type: ignore[attr-defined]
        if annots is None:
            return lines
        for ref in annots.get_object():
            annot = ref.get_object()
            subtype = str(annot.get("/Subtype", ""))
            if subtype in _PDF_ANNOT_SKIP:
                continue
            if subtype in _PDF_ANNOT_NOT_READ:
                tally.parts_not_read += 1
                continue
            contents = annot.get("/Contents")
            if contents is not None and str(contents).strip():
                lines.append(f"[Annotation] {contents}")
    except Exception as exc:  # noqa: BLE001 - counted and logged, never swallowed
        tally.parts_not_read += 1
        _log.warning("documents.text.pdf_annotations_not_read", error=type(exc).__name__)
    return lines


def _count_pdf_attachments(reader: PdfReader, tally: _Tally) -> None:
    """Files embedded in the PDF are not opened; each is counted."""
    try:
        tally.parts_not_read += len(reader.attachments)
    except Exception as exc:  # noqa: BLE001 - counted and logged, never swallowed
        tally.parts_not_read += 1
        _log.warning("documents.text.pdf_attachments_not_read", error=type(exc).__name__)


def _pdf_form_values(reader: PdfReader, tally: _Tally) -> Iterator[str]:
    """Filled AcroForm values, as "name: value" lines. Page text extraction
    does not see them: they live in the form's field dictionaries. An XFA form
    (data in XML streams) or a form pypdf cannot walk is counted as a part not
    read rather than dropped."""
    try:
        acro = reader.trailer["/Root"].get("/AcroForm")
        if acro is None:
            return
        acro = acro.get_object()
        if "/XFA" in acro:
            tally.parts_not_read += 1
        fields = reader.get_fields() or {}
    except Exception as exc:  # noqa: BLE001 - counted and logged, never swallowed
        tally.parts_not_read += 1
        _log.warning("documents.text.pdf_form_not_read", error=f"{type(exc).__name__}: {exc}")
        return
    lines: list[str] = []
    for name, field in fields.items():
        value = field.get("/V")
        if value is None:
            continue
        values = value if isinstance(value, list) else [value]
        shown = ", ".join(str(v)[1:] if isinstance(v, NameObject) else str(v) for v in values)
        if shown.strip():
            lines.append(f"{name}: {shown}")
    if lines:
        yield "[Form fields]"
        yield from lines


# -- DOCX: the package XML, walked directly ------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
#: Never text: tracked moves away (their runs use w:t; a tracked deletion uses
#: w:delText, which is never read), properties, and the legacy copy Word writes
#: beside a drawing (the mc:Choice copy is the one read).
_DOCX_SKIP = frozenset(
    {
        _W + "moveFrom",
        _W + "sectPr",
        _W + "pPr",
        _W + "rPr",
        _W + "tblPr",
        _W + "trPr",
        _W + "tcPr",
        _MC_FALLBACK,
    }
)
#: Content this reader does not open, each counted as a part not read: an
#: imported file, an embedded object, a picture, SmartArt and a chart.
_DOCX_NOT_READ = frozenset(
    {
        _W + "altChunk",
        _W + "object",
        "{http://schemas.openxmlformats.org/drawingml/2006/picture}pic",
        "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds",
        "{http://schemas.openxmlformats.org/drawingml/2006/chart}chart",
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara",
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath",
    }
)
#: Containers of graphics. One that gives no text and nothing counted inside
#: it (a VML picture, a chartex chart, ink, a shape) is counted itself.
_DOCX_GRAPHIC = frozenset({_W + "drawing", _W + "pict"})
#: Private Use Area: a symbol font's code point, with no known character.
_PUA = range(0xE000, 0xF900)
_DOCX_NOTES = ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml")


def _part_number(name: str) -> int:
    digits = re.findall(r"\d+", name.rsplit("/", 1)[1])
    return int(digits[0]) if digits else 0


@_typed_reader("DOCX")
def _read_docx(data: bytes, tally: _Tally) -> Iterator[str]:
    """Headers, then the body, then footers, footnotes, endnotes and comments.
    The main document part failing is `unreadable`; any other part failing is
    counted in `parts_not_read` and the rest is still read."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        body = safe_fromstring(z.read("word/document.xml"))
        headers = sorted(
            (n for n in names if re.fullmatch(r"word/header\d*\.xml", n)), key=_part_number
        )
        footers = sorted(
            (n for n in names if re.fullmatch(r"word/footer\d*\.xml", n)), key=_part_number
        )
        notes = [n for n in _DOCX_NOTES if n in names]
        for name in headers:
            yield from _docx_part(z, name, tally)
        yield from _docx_blocks(body, tally)
        for name in footers + notes:
            yield from _docx_part(z, name, tally)


def _docx_part(z: zipfile.ZipFile, name: str, tally: _Tally) -> Iterator[str]:
    try:
        root = safe_fromstring(z.read(name))
    except Exception as exc:  # noqa: BLE001 - counted and logged, never swallowed
        tally.parts_not_read += 1
        _log.warning("documents.text.docx_part_not_read", part=name, error=type(exc).__name__)
        return
    yield from _docx_blocks(root, tally)


def _docx_blocks(el: Element, tally: _Tally) -> Iterator[str]:
    """Paragraphs and tables in document order, through any wrapper (content
    controls, custom XML, footnote and comment elements)."""
    for child in el:
        if tally.out_of_time():
            return
        tag = child.tag
        if tag == _W + "sdt" and _docx_is_placeholder(child):
            continue
        if tag == _W + "p":
            yield from _docx_paragraph(child, tally)
        elif tag == _W + "tbl":
            yield from _docx_table(child, tally)
        elif tag in _DOCX_NOT_READ:
            tally.parts_not_read += 1
        elif tag not in _DOCX_SKIP:
            yield from _docx_blocks(child, tally)


def _docx_paragraph(p: Element, tally: _Tally) -> Iterator[str]:
    """One paragraph's runs as one line (through run-level content controls,
    hyperlinks, fields and tracked insertions), then any paragraphs nested in
    it (text boxes) on their own lines."""
    pieces: list[str] = []
    nested: list[str] = []

    def walk(el: Element) -> None:
        for child in el:
            tag = child.tag
            if tag == _W + "t":
                pieces.append(child.text or "")
            elif tag in (_W + "tab", _W + "ptab"):
                pieces.append("\t")
            elif tag in (_W + "br", _W + "cr"):
                pieces.append("\n")
            elif tag == _W + "noBreakHyphen":
                # "Jean-Pierre" must reach the redactor as the hint is written.
                pieces.append("-")
            elif tag == _W + "sym":
                symbol = _docx_symbol(child)
                if symbol:
                    pieces.append(symbol)
                else:
                    tally.parts_not_read += 1
            elif tag in _DOCX_GRAPHIC:
                before = (len(pieces), len(nested), tally.parts_not_read)
                walk(child)
                if (len(pieces), len(nested), tally.parts_not_read) == before:
                    tally.parts_not_read += 1
            elif tag == _W + "r" and _docx_is_hidden(child):
                # Hidden text is not what a reader of the document sees, so it
                # is not quoted as evidence; it is counted, not dropped.
                tally.parts_not_read += 1
            elif tag == _W + "rt":
                # Ruby text is a pronunciation guide over its base text (which
                # is read); excluded from the quote and counted.
                tally.parts_not_read += 1
            elif tag == _W + "sdt" and _docx_is_placeholder(child):
                pass
            elif tag == _W + "p":
                nested.extend(_docx_paragraph(child, tally))
            elif tag == _W + "tbl":
                nested.extend(_docx_table(child, tally))
            elif tag in _DOCX_NOT_READ:
                tally.parts_not_read += 1
            elif tag not in _DOCX_SKIP:
                walk(child)

    walk(p)
    line = "".join(pieces)
    if line.strip():
        yield line
    yield from nested


def _docx_table(tbl: Element, tally: _Tally) -> Iterator[str]:
    """One line per row, cells tab-joined; a cell's own paragraphs and nested
    tables are joined into it. Rows and cells inside content controls count."""
    for row in _docx_descendants(tbl, _W + "tr"):
        if tally.out_of_time():
            return
        cells = ["\n".join(_docx_blocks(cell, tally)) for cell in _docx_descendants(row, _W + "tc")]
        line = "\t".join(cells)
        if line.strip():
            yield line


def _docx_is_placeholder(sdt: Element) -> bool:
    """A content control still showing its placeholder ("Click or tap here to
    enter text.") holds no document content."""
    props = sdt.find(_W + "sdtPr")
    return props is not None and props.find(_W + "showingPlcHdr") is not None


def _docx_is_hidden(run: Element) -> bool:
    """A run formatted hidden (w:vanish), unless switched back off."""
    props = run.find(_W + "rPr")
    vanish = props.find(_W + "vanish") if props is not None else None
    if vanish is None:
        return False
    return vanish.get(_W + "val", "true").lower() not in ("0", "false", "off")


def _docx_symbol(sym: Element) -> str:
    """A w:sym's character where it is a real one, else "" (the caller counts
    it). A symbol-font code point in the Private Use Area (Wingdings' checked
    and empty boxes, F0FE and F0A8) has no known character; neither has a
    code that is not hex or not a code point at all."""
    try:
        code = int(sym.get(_W + "char", ""), 16)
        if code in _PUA:
            return ""
        return chr(code)
    except (ValueError, OverflowError):
        return ""


def _docx_descendants(el: Element, tag: str) -> Iterator[Element]:
    """Children with `tag`, looking through wrappers (a content control around
    a row or cell) but not into a match, so a nested table's rows stay in it."""
    for child in el:
        if child.tag == tag:
            yield child
        elif child.tag not in _DOCX_SKIP:
            yield from _docx_descendants(child, tag)


# -- XLSX ----------------------------------------------------------------------


@_typed_reader("XLSX")
def _read_xlsx(data: bytes, tally: _Tally) -> Iterator[str]:
    """Every worksheet's rows. A sheet's "[Sheet: name]" line is written only
    once it has a row, so a workbook of empty sheets is an empty read."""
    tally.parts_not_read += _xlsx_parts_not_read(data)
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        for ws_v, ws_f in zip(values.worksheets, formulas.worksheets, strict=True):
            # Read-only mode trusts the sheet's stored <dimension ref>, which
            # writers get wrong; reset it so every row is read.
            ws_v.reset_dimensions()
            ws_f.reset_dimensions()
            header: str | None = f"[Sheet: {ws_v.title}]"
            rows_v = ws_v.iter_rows(values_only=True)
            rows_f = ws_f.iter_rows(values_only=True)
            for row_v, row_f in zip_longest(rows_v, rows_f, fillvalue=()):
                if tally.out_of_time():
                    return
                line = _xlsx_row(row_v, row_f)
                if line:
                    if header is not None:
                        yield header
                        header = None
                    yield line
    finally:
        values.close()
        formulas.close()


_XLSX_COMMENTS = re.compile(r"xl/comments(\d*|/[^/]+)\.xml")
_XLSX_DRAWING = re.compile(r"xl/drawings/drawing\d+\.xml")
_XLSX_CHARTSHEET = re.compile(r"xl/chartsheets/sheet\d+\.xml")
_XLSX_CHARTSHEET_RELS = re.compile(r"xl/chartsheets/_rels/sheet\d+\.xml\.rels")


def _xlsx_parts_not_read(data: bytes) -> int:
    """What read-only openpyxl does not read, counted: each cell comment (the
    legacy comment parts; threaded comments are mirrored there by Excel and
    not counted twice), each worksheet drawing (pictures, charts, shapes) and
    each chartsheet. A chartsheet's own drawing is the chartsheet, counted once.
    A comments part that will not parse counts as one."""
    count = 0
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        chartsheet_drawings: set[str] = set()
        for name in names:
            if _XLSX_CHARTSHEET_RELS.fullmatch(name):
                for target in re.findall(rb'Target="([^"]+)"', z.read(name)):
                    chartsheet_drawings.add("xl/drawings/" + target.decode().rsplit("/", 1)[-1])
        for name in names:
            if _XLSX_COMMENTS.fullmatch(name):
                try:
                    root = safe_fromstring(z.read(name))
                    count += sum(1 for el in root.iter() if el.tag.endswith("}comment"))
                except Exception as exc:  # noqa: BLE001 - counted and logged
                    count += 1
                    _log.warning("documents.text.xlsx_comments_not_read", error=type(exc).__name__)
            elif _XLSX_CHARTSHEET.fullmatch(name) or (
                _XLSX_DRAWING.fullmatch(name) and name not in chartsheet_drawings
            ):
                count += 1
    return count


def _xlsx_row(row_v: tuple, row_f: tuple) -> str:
    """One row, tab-joined. A formula with no cached value keeps its formula
    text: a blank cell would read as "nothing there", which is not what the
    file holds."""
    cells: list[str] = []
    for value, formula in zip_longest(row_v, row_f, fillvalue=None):
        if value is None and isinstance(formula, str) and formula.startswith("="):
            value = formula
        cells.append("" if value is None else str(value))
    while cells and cells[-1] == "":
        cells.pop()
    return "\t".join(cells)


# -- TXT -----------------------------------------------------------------------


def _read_txt(data: bytes, *, declared_text: bool) -> Iterator[str]:
    """UTF-8 (with or without a BOM) and BOM-marked UTF-16 and UTF-32 are read;
    the UTF-32 BOMs are checked first, since UTF-32LE's begins with UTF-16LE's.
    Anything else is not decoded by guessing, which would put mojibake into a
    quote that then reads as the client's own words. A file DECLARED as text
    that is not decodable text is `unreadable`; bytes that declared nothing and
    match no format are `unsupported_type`, because no reader exists for them."""
    failure: NotReadReason = "unreadable" if declared_text else "unsupported_type"
    if data.startswith(_UTF32_BOMS):
        encoding = "utf-32"
    elif data.startswith(_UTF16_BOMS):
        encoding = "utf-16"
    elif b"\x00" in data:
        raise _NotRead(failure, "text: NUL bytes and no UTF-16/32 BOM, so binary")
    else:
        encoding = "utf-8-sig"
    try:
        text = data.decode(encoding)
    except UnicodeDecodeError as exc:
        raise _NotRead(failure, f"text: not {encoding} ({exc})") from exc
    # One piece per line, so the max_chars stop lands between lines, never
    # inside one. Joined back with "\n", the lines rebuild the text exactly.
    yield from _LINE_BREAKS.sub("\n", text).split("\n")


#: Line separators other than "\n", including "\r\n" as one: CR, VT, FF, the
#: file/group/record separators, NEL, and U+2028/U+2029.
_LINE_BREAKS = re.compile("\r\n|[\r\v\f\x1c\x1d\x1e\x85\u2028\u2029]")
#: Every other control character: C0 except tab, newline and the separators
#: above (so \x1f, the unit separator, IS removed), DEL, and C1 except NEL.
_CONTROLS = re.compile("[\x00-\x08\x0e-\x1b\x1f\x7f-\x84\x86-\x9f]")


def _sanitize(text: str) -> str:
    """One newline convention and no control characters, for every reader.
    A removed control becomes a space so the words either side stay apart."""
    return _CONTROLS.sub(" ", _LINE_BREAKS.sub("\n", text.lstrip("\ufeff")))


def chunk_text(
    text: str,
    *,
    chunk_chars: int,
    max_chunks: int,
    mode: RedactionMode,
    client_org_name: str | None,
    name_hints: Iterable[str],
    source_truncated: bool | None,
    max_chars: int,
) -> ChunkPlan:
    """Redact the WHOLE text with `redact_for_ai`, then split the result into
    chunks of at most `chunk_chars` for per-payload AI calls.

    `source_truncated` is the source's own truncation (`DocText.truncated`),
    or None when the caller does not know it; None never yields a complete
    plan. `max_chars` is the extraction limit the text came from; text longer
    than `2 * max_chars` is refused with `TextOverCeiling`.

    When the source was cut (`source_truncated` True or None), a trailing
    fragment the redactor cannot match is dropped BEFORE redacting: tokens
    forming a proper prefix of a name hint's or the legal name's words
    (case-insensitive, split on whitespace as the redactor splits them, the
    last token possibly itself cut), and up to four trailing tokens holding a
    digit (a partial phone number or ID). The plan counts what was dropped.

    Takes the same inputs `redact_for_ai` takes, and calls it, so there is no
    path to a chunk that has not been redacted (enforced by calling). Mode
    "off" is refused with `RedactionOffRefused`. Redacting per chunk is not
    safe (see the module docstring).

    A chunk ends at the last paragraph break in the second half of its
    window, else the last whitespace there, else at `chunk_chars` exactly.
    At most `max_chunks` are returned; `skipped` counts the rest.

    This calls the redactor in the calling process. A caller wiring it into
    a live AI path must run it inside #552's redaction isolation (the child
    process with a per-payload time budget), as every other egress does.
    """
    if not isinstance(text, str):
        raise TypeError(f"chunk_text needs str, got {type(text).__name__}")
    if mode == "off":
        raise RedactionOffRefused(
            'chunk_text refuses redaction mode "off": AI payload text is always redacted'
        )
    if chunk_chars < 1:
        raise ValueError(f"chunk_chars must be at least 1, got {chunk_chars}")
    if max_chunks < 1:
        raise ValueError(f"max_chunks must be at least 1, got {max_chunks}")
    if len(text) > 2 * max_chars:
        raise TextOverCeiling(
            f"{len(text)} characters is over the ceiling of {2 * max_chars} (2 * max_chars)"
        )
    hints = list(name_hints)
    tokens_dropped = 0
    chars_dropped = 0
    if source_truncated is not False:
        kept, tokens_dropped = _drop_cut_tail(text, [*hints, client_org_name or ""])
        chars_dropped = len(text) - len(kept)
        text = kept
    redacted, counts = redact_for_ai(
        text, mode=mode, client_org_name=client_org_name, name_hints=hints
    )
    cut_chars, cut_lines = _signature_cut(text, redacted, counts)
    chunks: list[str] = []
    offsets: list[int] = []
    start = 0
    while start < len(redacted):
        end = min(start + chunk_chars, len(redacted))
        if end < len(redacted):
            end = _chunk_end(redacted, start, end, chunk_chars)
        if end <= start:
            # Every pass must consume text. Without this a bad limit loops
            # forever instead of failing.
            raise RuntimeError(
                f"chunk_text made no progress at {start} (chunk_chars={chunk_chars})"
            )
        chunks.append(redacted[start:end])
        offsets.append(start)
        start = end
    plan = ChunkPlan(
        chunks=tuple(chunks[:max_chunks]),
        sent=min(len(chunks), max_chunks),
        skipped=max(len(chunks) - max_chunks, 0),
        offsets=tuple(offsets[:max_chunks]),
        redaction_counts=dict(counts),
        signature_cut=cut_chars > 0,
        signature_chars_removed=cut_chars,
        signature_lines_removed=cut_lines,
        source_truncated=source_truncated,
        tail_tokens_dropped=tokens_dropped,
        tail_chars_dropped=chars_dropped,
    )
    _log.info(
        "documents.text.chunked",
        sent=plan.sent,
        skipped=plan.skipped,
        redaction_counts=plan.redaction_counts,
        signature_cut=plan.signature_cut,
        signature_chars_removed=cut_chars,
        signature_lines_removed=cut_lines,
        source_truncated=source_truncated,
        complete=plan.complete,
        tail_tokens_dropped=tokens_dropped,
        tail_chars_dropped=chars_dropped,
    )
    return plan


#: The most trailing digit-bearing tokens dropped from a cut: enough for a
#: phone written in groups ("+1 410 555 01").
_MAX_DIGIT_TOKENS_DROPPED = 4
_TOKEN = re.compile(r"\S+")


def _drop_cut_tail(text: str, literals: list[str]) -> tuple[str, int]:
    """`text` without a trailing fragment the redactor could not match, and
    how many tokens went. Repeats until the tail is clean: a name prefix, then
    a digit-bearing token, then a name prefix again, and so on."""
    words = [[t.lower() for t in lit.split()] for lit in literals if lit.split()]
    dropped = 0
    digit_tokens = 0
    while True:
        spans = [m.span() for m in _TOKEN.finditer(text)]
        if not spans:
            return text, dropped
        tail = [text[a:b].lower() for a, b in spans]
        k = max((_name_prefix_len(tail, w) for w in words), default=0)
        if (
            k == 0
            and digit_tokens < _MAX_DIGIT_TOKENS_DROPPED
            and any(ch.isdigit() for ch in tail[-1])
        ):
            k = 1
            digit_tokens += 1
        if k == 0:
            return text, dropped
        text = text[: spans[-k][0]].rstrip()
        dropped += k


def _name_prefix_len(tail: list[str], words: list[str]) -> int:
    """The number of trailing tokens of `tail` that form a PROPER prefix of
    `words` (a name's words): every token but the last equal to the name's,
    and the last a start of the name's next word. 0 if none; the whole name,
    complete, is not a prefix (the redactor matches it)."""
    for k in range(min(len(words), len(tail)), 0, -1):
        last = tail[-1]
        if tail[len(tail) - k : -1] != words[: k - 1]:
            continue
        if not words[k - 1].startswith(last):
            continue
        if k == len(words) and last == words[-1]:
            continue
        return k
    return 0


def _signature_cut(text: str, redacted: str, counts: Mapping[str, int]) -> tuple[int, int]:
    """Characters and lines the redactor's signature rule removed from `text`.

    Measured by CALLING the redactor's own signature step on the same input
    (it is the first pass `redact_for_ai` runs, so it sees exactly this text)
    rather than by re-deriving where it cut. app/ai/redact.py is not changed
    here (#552's hold). If the step and `redact_for_ai`'s count disagree, the
    redactor has changed under this module, and that is raised, not guessed.
    """
    _, fired, removed = _redact_signature_blocks(text)
    if bool(fired) != bool(counts.get("signature_block", 0)):
        raise RuntimeError(
            "signature step and redact_for_ai disagree: "
            f"step fired={fired}, counts={dict(counts)}"
        )
    if fired and not redacted.endswith(PLACEHOLDER_SIGNATURE + "\n"):
        raise RuntimeError(
            "signature step fired but the redacted text does not end with its placeholder"
        )
    if not removed:
        return 0, 0
    tail = text[len(text) - removed :]
    return removed, len(tail.splitlines())


def _chunk_end(text: str, start: int, end: int, chunk_chars: int) -> int:
    """Where a full window should end: after the last paragraph break in its
    second half, else after the last whitespace there, else at `end`."""
    window = text[start:end]
    floor = chunk_chars // 2
    newline = window.rfind("\n")
    if newline >= floor:
        return start + newline + 1
    space = max(window.rfind(" "), window.rfind("\t"))
    if space >= floor:
        return start + space + 1
    return end
