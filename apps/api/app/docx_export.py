"""Shared python-docx helpers for service Word deliverables (Work Order C4).

Pure helpers over python-docx so each service's `render_docx` mirrors its PDF
(title, client, summary paragraphs, a findings table). python-docx is imported
lazily inside `new_document` so importing the exporter module stays cheap.

Every string reaches python-docx through `_clean` (#993). python-docx refuses a
string XML 1.0 cannot carry, so a client legal name such as "R&D\\x07 Labs" --
which the admin route stores -- failed every Word deliverable for that client.
`_clean` drops the characters the XLSX path drops, by CALLING
`app.xlsx_export.strip_illegal_characters` rather than restating the set, and
logs ONE warning per stripped string, `docx_export.control_characters_removed`,
naming the document part, the field and its position, and how many characters
went. NEVER the value.

A caller writing text into a document goes through a helper here, never
python-docx directly, or the guard does not reach it.
"""

from __future__ import annotations

import contextlib
import io
from collections.abc import Callable, Iterable, Sequence
from typing import Any

from app.logging import get_logger
from app.xlsx_export import strip_illegal_characters

_log = get_logger(__name__)


def _clean(
    text: str,
    *,
    part: str,
    field: str,
    locate: Callable[[], dict[str, int]] | None = None,
    **where: int,
) -> str:
    """`text` without the characters `app.xlsx_export.strip_illegal_characters`
    removes (that function is the exact set); a strip is logged, the value
    never is. `where` (or `locate`, called only when something was
    stripped) places the string -- paragraph, table, row, column, 1-based -- so
    the warning finds it without quoting it."""
    clean, removed = strip_illegal_characters(text)
    if removed:
        _log.warning(
            "docx_export.control_characters_removed",
            part=part,
            field=field,
            removed=removed,
            **where,
            **(locate() if locate else {}),
        )
    return clean


#: python-docx refuses a core property longer than this, and its error quotes
#: the value.
CORE_PROPERTY_LIMIT = 255


def _core_property(text: str, *, field: str) -> str:
    """A core property's value: cleaned, then cut to `CORE_PROPERTY_LIMIT`.

    A plain truncation of the FILE'S METADATA only (#993): the title is
    "<service title> — <legal name>", each of which may be 255 characters on
    its own. The visible heading is written by `add_title`, which has no limit
    and keeps the whole text. A cut logs one warning with the field and the
    original length, never the value."""
    clean = _clean(text, part="core_properties", field=field)
    if len(clean) > CORE_PROPERTY_LIMIT:
        _log.warning(
            "docx_export.core_property_truncated",
            field=field,
            length=len(clean),
            limit=CORE_PROPERTY_LIMIT,
        )
        return clean[:CORE_PROPERTY_LIMIT]
    return clean


def new_document(title: str, *, author: str = "SHIELD by Kentro") -> Any:
    from docx import Document

    doc = Document()
    props = doc.core_properties
    props.title = _core_property(title, field="title")
    props.author = _core_property(author, field="author")
    return doc


def add_title(doc: Any, title: str, subtitle: str | None = None) -> None:
    doc.add_heading(_clean(title, part="body", field="title"), level=0)
    if subtitle:
        doc.add_paragraph(_clean(subtitle, part="body", field="subtitle"))


def add_heading(doc: Any, text: str, level: int = 1) -> None:
    doc.add_heading(_clean(text, part="body", field="heading"), level=level)


def set_footer(doc: Any, text: str, *, rgb: tuple[int, int, int] | None = None) -> None:
    """Put `text` in the footer of every section, so Word repeats it per page.

    A body paragraph appears once wherever it was added; a section footer is
    the only thing in the format that reaches every page. This is what #294
    needs for the CSF Playbook's approval status, and it is written here rather
    than in `playbook_export` because it is a property of the FORMAT -- every
    other service's DOCX deliverable has the same page underneath it.

    `is_linked_to_previous = False` is set on EVERY section, not just the
    first. python-docx links a new section's footer to its predecessor by
    default, so a document that later grows a second section would silently
    keep inheriting -- which happens to be right today and is right by accident.
    Setting it per section makes the guarantee hold for a document shape nobody
    has written yet.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.text = ""
        run = para.add_run(_clean(text, part="footer", field="text"))
        run.font.size = Pt(7)
        run.bold = True
        if rgb is not None:
            run.font.color.rgb = RGBColor(*rgb)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_break(doc: Any) -> None:
    doc.add_page_break()


def shade_cell(cell: Any, hex_color: str) -> None:
    """Set a table cell's background fill (python-docx has no direct API)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc_pr.append(shd)


def add_paragraph(doc: Any, text: str, *, bold: bool | None = None) -> None:
    """One body paragraph. `bold=None` leaves the run's weight unset, which is
    what `doc.add_paragraph(text)` writes; True or False sets it."""
    clean = _clean(
        text, part="body", field="paragraph", locate=lambda: {"paragraph": len(doc.paragraphs) + 1}
    )
    if bold is None:
        doc.add_paragraph(clean)
        return
    run = doc.add_paragraph().add_run(clean)
    run.bold = bold


def add_paragraphs(doc: Any, lines: Iterable[str]) -> None:
    for line in lines:
        add_paragraph(doc, line)


def add_table(doc: Any, headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> Any:
    """Add a styled table. Returns the table so callers can shade cells."""
    table = doc.add_table(rows=1, cols=len(headers))
    # Fall back to the default style if the template lacks the named one.
    with contextlib.suppress(KeyError):
        table.style = "Light Grid Accent 1"
    number = len(doc.tables)  # this table's position in the document, 1-based
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = _clean(str(h), part="body", field="table_header", table=number, column=i + 1)
    for r, row in enumerate(rows, start=2):  # row 1 is the header
        cells = table.add_row().cells
        for i, val in enumerate(row):
            text = "" if val is None else str(val)
            cells[i].text = _clean(
                text, part="body", field="table_cell", table=number, row=r, column=i + 1
            )
    return table


def to_bytes(doc: Any) -> bytes:
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# MIME type + extension for Word deliverables.
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
