"""Shared helpers for the #646 deliverable-stamp tests.

`seed_completed_call` writes the row `LLMClient` writes for a finished call,
straight into the test's database (the fixtures point `DATABASE_URL` at it), so
a test can put fixture output behind a service without running its Run-AI.

The text readers return what a person opening the file would read.
"""

from __future__ import annotations

import io
import os
import uuid

from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.models.llm_call import LLMCall, LLMCallMode, LLMCallStatus

FIXTURE_LINE = "OFFLINE TEST DATA: every AI suggestion behind this document came"
NO_AI_LINE = "No AI call is recorded for this document"


def seed_completed_call(
    *,
    requested_by: str,
    mode: LLMCallMode = LLMCallMode.FIXTURE,
    service_id: str | None = None,
    client_id: str | None = None,
    purpose: str = "mitre_map",
) -> None:
    engine = create_engine(os.environ["DATABASE_URL"])
    try:
        with Session(engine) as db:
            db.add(
                LLMCall(
                    purpose=purpose,
                    prompt_version="v1",
                    provider="fixture" if mode == LLMCallMode.FIXTURE else "anthropic",
                    model="m",
                    mode=mode,
                    status=LLMCallStatus.COMPLETED,
                    requested_by=uuid.UUID(requested_by),
                    service_id=uuid.UUID(service_id) if service_id else None,
                    client_id=uuid.UUID(client_id) if client_id else None,
                )
            )
            db.commit()
    finally:
        engine.dispose()


def pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    # PDF text extraction breaks lines where the layout wrapped them; collapse
    # whitespace so a sentence can be matched whole.
    return " ".join("".join(page.extract_text() for page in reader.pages).split())


def docx_text(raw: bytes) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)


def xlsx_ai_sheet(raw: bytes) -> list[list[object]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames[-1] == "AI source", wb.sheetnames
    return [list(r) for r in wb[wb.sheetnames[-1]].iter_rows(values_only=True)]
