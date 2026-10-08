"""A stored Partial reason is shown only when its sentence is true of the computed
Detect / Prevent / Respond line (for #842; the advisor's ruling on #736, comment
6053562002, narrowing Q7).

Before this, a computed leaf kept its stored reason whatever the computed line
said, so "Detected, not blocked" could sit beside "Detect: awaiting review". The
computed and stored statuses were both Partial, so `review_queue` never listed
the row and nothing caught it.

The rule, as approved (the claims table in the #842 plan):

* `prevention_limited`: Detect in place and Prevent not in place;
* `recovery_absent`: Detect in place and Respond not in place;
* `detection_weak`, `periodic_not_continuous`: Detect in place;
* `missing_control_category`: at least one judged function not in place;
* `reach_limited`, `evasive_variant_uncovered`: always.

`awaiting_review` satisfies neither an "in place" nor a "not in place" claim.
Otherwise the row reads "Set by what is in place".

Every expected string and the spec table below are COPIED from the approved
text, never imported from the module under test.
"""

from __future__ import annotations

import io
import itertools
import uuid

import pytest
from sqlalchemy import update

from app.attack.analytics import compute
from app.attack.catalog import NOT_PREVENTABLE
from app.attack.computed import Capabilities, InPlace, effective_coverage
from app.attack.exporters import build_context, render_xlsx
from app.attack.pending import pending_codes
from app.attack.rules import parents_computed
from app.models.attack_assessment import (
    AttackAssessment,
    AttackAssessmentStatus,
    AttackCoverage,
)
from tests._attack_rows import standalone_rows
from tests.unit.test_attack_catalog_version_guard import (  # noqa: F401  (fixture)
    _auth,
    _register,
    _service_and_assessment,
    env,
)

pytestmark = pytest.mark.unit

#: The approved client wording, (label, sentence), copied byte for byte.
WORDING = {
    "missing_control_category": (
        "A control type is missing",
        "Something already defends against this technique, but a whole category of "
        "control, such as prevention or response, is not in place.",
    ),
    "reach_limited": (
        "Not covered everywhere",
        "Defended on most of your environment, but not on some systems, such as another "
        "operating system, a cloud or SaaS service, or unmanaged or off-network devices.",
    ),
    "detection_weak": (
        "Detection is unreliable",
        "There is a signal for this activity, but it is noisy or approximate, or depends "
        "on custom detection rules that still need to be written and tuned.",
    ),
    # #842, Option B (the advisor's ruling on #736, comment 6053562002).
    "prevention_limited": (
        "Detected, no blocking control",
        "This activity can be detected, but no control in place blocks it.",
    ),
    "evasive_variant_uncovered": (
        "Advanced variants not covered",
        "Common forms of this technique are covered; advanced forms, such as "
        "kernel-level, firmware, encrypted or novel variants, are not.",
    ),
    "recovery_absent": (
        "No recovery evidenced",
        "The attack would be detected, but no way to recover from it, such as backups "
        "or restore procedures, is evidenced.",
    ),
    "periodic_not_continuous": (
        "Checked periodically, not continuously",
        "This is found only when a scheduled scan runs, not as it happens.",
    ),
}
SET_BY_IN_PLACE = (
    "Set by what is in place",
    "This technique's status is computed from which of Detect, Prevent and Respond "
    "are in place.",
)

IN, NOT, AWAIT, CANNOT = "in_place", "not_in_place", "awaiting_review", "cannot_be_prevented"


def _spec_holds(code: str, detect: str, prevent: str, respond: str) -> bool:
    """The claims table, written from the approved plan, not from the code."""
    judged = [detect, respond] if prevent == CANNOT else [detect, prevent, respond]
    if code == "prevention_limited":
        return detect == IN and prevent == NOT
    if code == "recovery_absent":
        return detect == IN and respond == NOT
    if code in ("detection_weak", "periodic_not_continuous"):
        return detect == IN
    if code == "missing_control_category":
        return NOT in judged
    if code in ("reach_limited", "evasive_variant_uncovered"):
        return True
    raise AssertionError(f"no spec row for {code}")


MATRIX = [
    (code, d, p, r)
    for code in WORDING
    for d, p, r in itertools.product((IN, NOT, AWAIT), (IN, NOT, AWAIT, CANNOT), (IN, NOT, AWAIT))
]


def test_the_matrix_covers_every_combination_and_code() -> None:
    # 3 Detect x 4 Prevent x 3 Respond states, times the Partial codes.
    assert len({(d, p, r) for _, d, p, r in MATRIX}) == 3 * 4 * 3
    assert len(MATRIX) == 3 * 4 * 3 * len(WORDING)


@pytest.mark.parametrize(("code", "detect", "prevent", "respond"), MATRIX)
def test_a_stored_reason_shows_only_when_its_sentence_is_true(
    code: str, detect: str, prevent: str, respond: str
) -> None:
    from app.attack.partial_reasons import partial_reason

    caps = Capabilities(detect=InPlace(detect), prevent=InPlace(prevent), respond=InPlace(respond))
    got = partial_reason("partial", code, computed_parent=False, capabilities=caps)
    expected = WORDING[code] if _spec_holds(code, detect, prevent, respond) else SET_BY_IN_PLACE
    assert (got.label, got.sentence) == expected


@pytest.mark.parametrize("code", list(WORDING))
def test_a_row_whose_status_was_not_computed_keeps_its_stored_reason(code: str) -> None:
    from app.attack.partial_reasons import partial_reason

    got = partial_reason("partial", code, computed_parent=False)
    assert (got.label, got.sentence) == WORDING[code]


# --- one world, through the surfaces -------------------------------------------------

#: Facts about ATT&CK 19.2 (test_attack_not_preventable.py): T1057 has no
#: mitigation. The other rows need techniques WITH one (MITRE's list, pinned
#: against STIX in that test), or Prevent reads "cannot be prevented" and an
#: awaiting Prevent tool no longer matters.
NOT_PREVENTABLE_CODE = "T1057"

#: (stored reason, detection, prevention, response, tools awaiting review, expected)
ROWS = {
    # The issue's case: a detection citation awaits review.
    "awaiting_detect": (
        "prevention_limited",
        ["Tool D"],
        [],
        ["Tool R"],
        ["Tool D"],
        SET_BY_IN_PLACE,
    ),
    "not_blocked": (
        "prevention_limited",
        ["Tool D"],
        [],
        ["Tool R"],
        [],
        WORDING["prevention_limited"],
    ),
    "no_detection": (
        "detection_weak",
        [],
        ["Tool P"],
        ["Tool R"],
        [],
        SET_BY_IN_PLACE,
    ),
    "respond_awaiting": (
        "recovery_absent",
        ["Tool D"],
        ["Tool P"],
        ["Tool R"],
        ["Tool R"],
        SET_BY_IN_PLACE,
    ),
    "reach_kept": (
        "reach_limited",
        ["Tool D"],
        ["Tool P"],
        ["Tool R"],
        ["Tool D"],
        WORDING["reach_limited"],
    ),
    "only_awaiting_missing": (
        "missing_control_category",
        ["Tool D"],
        ["Tool P"],
        ["Tool R"],
        ["Tool P"],
        SET_BY_IN_PLACE,
    ),
}
#: The not-preventable row: Prevent reads "cannot be prevented", never "not in
#: place", so `prevention_limited` is not true of it (Gene's decision 3 on #554).
UNPREVENTABLE_ROW = (
    "prevention_limited",
    ["Tool D"],
    [],
    [],
    [],
    SET_BY_IN_PLACE,
)


def _citations(awaiting: list[str]) -> list[dict]:
    return [{"tool": t, "cited": t, "reason": "inferred", "cleared_at": None} for t in awaiting]


def _standalone_codes(n: int) -> list[str]:
    from tests._attack_rows import _standalone_codes as codes

    return sorted(c for c in codes() if c not in NOT_PREVENTABLE)[:n]


def _memory_world(status_rules: int | None):
    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=1,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=2,
        status_rules=status_rules,
    )
    codes = dict(zip(ROWS, _standalone_codes(len(ROWS)), strict=True))
    world = {**{codes[k]: v for k, v in ROWS.items()}, NOT_PREVENTABLE_CODE: UNPREVENTABLE_ROW}
    rows = [
        AttackCoverage(
            id=uuid.uuid4(),
            assessment_id=a.id,
            technique_code=code,
            status="partial",
            reason_code=reason,
            detection_tools=d,
            prevention_tools=p,
            response_tools=r,
            unconfirmed_citations=_citations(awaiting),
        )
        for code, (reason, d, p, r, awaiting, _) in world.items()
    ]
    eff = effective_coverage(a, rows)
    rollup = compute(
        {c.technique_code: c.status for c in eff},
        pending_codes(eff, parents_computed=parents_computed(a)),
    )
    ctx = build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=eff,
        rollup=rollup,
    )
    return ctx, world


def _why_partial(raw: bytes) -> dict[str, str]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))["Coverage"]
    sheet = [[c.value for c in r] for r in ws.iter_rows()]
    header, rows = sheet[0], sheet[1:]
    col = header.index("Why partial")
    return {r[0]: r[col] for r in rows}


def test_the_xlsx_shows_a_stored_reason_only_when_it_is_true() -> None:
    ctx, world = _memory_world(2)
    # The world is one the system produces: every row computes to Partial.
    assert {c.technique_code: c.status for c in ctx.coverage} == dict.fromkeys(world, "partial")
    cells = _why_partial(render_xlsx(ctx))
    for code, (*_, expected) in world.items():
        assert cells[code] == f"{expected[0]}: {expected[1]}", code


def test_an_assessment_approved_before_r3_keeps_every_stored_reason() -> None:
    """status_rules 1: nothing is computed, so nothing is checked."""
    ctx, world = _memory_world(1)
    cells = _why_partial(render_xlsx(ctx))
    for code, (reason, *_) in world.items():
        label, sentence = WORDING[reason]
        assert cells[code] == f"{label}: {sentence}", code


def test_the_client_dashboard_shows_a_stored_reason_only_when_it_is_true(
    env,  # noqa: F811
) -> None:
    """Through the route a client calls, on a released assessment under R3."""
    c, Sess = env
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    svc, a = _service_and_assessment(c, bearer)
    by_code = {row["technique_code"]: row for row in a["coverage"]}
    rows = standalone_rows(
        [r for r in a["coverage"] if r["technique_code"] not in NOT_PREVENTABLE], len(ROWS)
    )
    world = {row["technique_code"]: spec for row, spec in zip(rows, ROWS.values(), strict=True)}
    world[NOT_PREVENTABLE_CODE] = UNPREVENTABLE_ROW
    for code, (reason, d, p, r, _awaiting, _) in world.items():
        out = c.patch(
            f"/attack/coverage/{by_code[code]['id']}",
            headers=_auth(bearer),
            json={
                "status": "partial",
                "reason_code": reason,
                "detection_tools": d,
                "prevention_tools": p,
                "response_tools": r,
            },
        )
        assert out.status_code == 200, out.text
    # The world, not the outcome: some citations still await a human, as a run
    # that inferred them leaves them. The PATCH above confirmed everything.
    with Sess() as s:
        for code, (*_, awaiting, _) in world.items():
            if awaiting:
                s.execute(
                    update(AttackCoverage)
                    .where(AttackCoverage.id == uuid.UUID(by_code[code]["id"]))
                    .values(unconfirmed_citations=_citations(awaiting))
                )
        s.commit()
    r = c.post(f"/attack/assessments/{a['id']}/approve", headers=_auth(bearer))
    assert r.status_code == 200, r.text
    fin = c.post(f"/attack/services/{svc}/deliverables/finalize", headers=_auth(bearer))
    assert fin.status_code in (200, 201), fin.text
    rel = c.post(f"/attack/deliverables/{fin.json()['id']}/release", headers=_auth(bearer))
    assert rel.status_code == 200, rel.text

    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    dash = c.get(
        f"/clients/{client_id}/attack/{svc}/dashboard",
        headers=_auth(client["tokens"]["access_token"]),
    )
    assert dash.status_code == 200, dash.text
    body = dash.json()
    techs = {t["code"]: t for t in body["techniques"]}
    for code, (*_, expected) in world.items():
        assert techs[code]["status"] == "partial", (code, techs[code])
        assert techs[code]["partial_reason"] == {
            "label": expected[0],
            "sentence": expected[1],
        }, code
    # The by-reason table moves rows between reasons; it still adds up.
    counts = {(x["label"], x["sentence"]): x["count"] for x in body["partial_reasons"]}
    assert counts[SET_BY_IN_PLACE] == 5, counts
    assert counts[WORDING["prevention_limited"]] == 1, counts
    assert counts[WORDING["reach_limited"]] == 1, counts
    assert sum(counts.values()) == body["rollup"]["partial"]
