"""#1094: the "≥" is ON the XLSX savings value, not only beside it.

When estimated savings are a lower bound, the PDF, DOCX and plan card print
"≥" on the number. The XLSX cell was a plain number formatted "$#,##0" with the
"≥" only in the next column, so a reader who copied the number got an exact
looking figure. Ruled on #736 (6229200940): keep the cell NUMERIC (sums still
work) and give it the number format `"≥ "$#,##0` only when the savings are a
floor; an exact figure keeps `$#,##0`. The cell is read back from a workbook the
real renderer wrote, never from the format string the code holds.
"""

from __future__ import annotations

import io
import zipfile
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

from app.models import CapabilityList
from app.tech_debt.exporters import DeliverableContext, render_xlsx
from tests.unit.test_tech_debt_bundle_part_savings import (  # noqa: F401  (fixtures)
    SAVINGS,
    _xlsx_rows,
    released,
)
from tests.unit.test_tech_debt_dashboard import app_client  # noqa: F401  (fixture)

pytestmark = pytest.mark.unit

#: The format the advisor ruled for a floor, and the one an exact figure keeps.
FLOOR_FORMAT = '"≥ "$#,##0'
EXACT_FORMAT = "$#,##0"


def _savings_cell(raw: bytes):
    ws = load_workbook(io.BytesIO(raw)).worksheets[0]
    for row in ws.iter_rows():
        if row[0].value == "Estimated annual savings":
            return row[4]
    raise AssertionError("no 'Estimated annual savings' row in the workbook")


def _render(*, savings: float, known: bool) -> bytes:
    ctx = DeliverableContext(
        client_legal_name="Example Co",
        service_title="Technical Debt Review",
        cap_list=CapabilityList(version=1),
        items=[],
        total_cost=0.0,
        estimated_savings=savings,
        savings_cost_known=known,
    )
    return render_xlsx(ctx)


def test_a_floor_savings_cell_is_numeric_and_carries_the_marker_in_its_format() -> None:
    cell = _savings_cell(_render(savings=12345.0, known=False))
    assert cell.value == 12345.0, cell.value
    # openpyxl reads 12345.0 back as an int: what matters is that it is a number.
    assert isinstance(cell.value, (int, float)), "the cell must stay a number so sums work"
    assert cell.number_format == FLOOR_FORMAT, cell.number_format


def test_an_exact_savings_cell_keeps_the_plain_format() -> None:
    cell = _savings_cell(_render(savings=12345.0, known=True))
    assert cell.value == 12345.0, cell.value
    assert cell.number_format == EXACT_FORMAT, cell.number_format
    assert "≥" not in cell.number_format


def test_a_floor_of_zero_still_carries_the_marker() -> None:
    # The redundancy card marks a floor even at $0 (#1044); so does the file.
    cell = _savings_cell(_render(savings=0.0, known=False))
    assert cell.value == 0.0
    assert cell.number_format == FLOOR_FORMAT, cell.number_format


def test_a_released_floor_deliverable_marks_the_cell_the_client_downloads(
    released,  # noqa: F811
) -> None:
    raw = released["files"]["xlsx"]
    cell = _savings_cell(raw)
    assert cell.value == SAVINGS, cell.value
    assert cell.number_format == FLOOR_FORMAT, cell.number_format
    # The note beside it is unchanged: the marker adds no new words.
    assert (
        _xlsx_rows(raw)["Estimated annual savings"][5] == "≥ (one or more cut rows missing a cost)"
    )


# --- what the workbook FILE holds, read from its own XML ----------------------
#
# The tests above read the format back through openpyxl, the library that wrote
# it. These unzip the .xlsx and read the parts Excel and LibreOffice read: the
# sheet's cell (its type and value) and the styles part (the number format the
# cell's style points to). They pin what is actually written, with no spreadsheet
# program needed. They do NOT prove how Excel or LibreOffice draw it.

_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _cell_in_the_file(raw: bytes) -> tuple[ElementTree.Element, str, bytes]:
    """The savings cell's XML element, its number format code, and the styles part."""
    ref = _savings_cell(raw).coordinate
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        styles_bytes = z.read("xl/styles.xml")
        # The capability sheet is the workbook's first sheet.
        sheet_bytes = z.read("xl/worksheets/sheet1.xml")
    sheet = ElementTree.fromstring(sheet_bytes)  # noqa: S314 (our own output)
    styles = ElementTree.fromstring(styles_bytes)  # noqa: S314 (our own output)
    cell = sheet.find(f".//m:c[@r='{ref}']", _NS)
    assert cell is not None, f"no cell {ref} in sheet1.xml"
    codes = {
        int(n.get("numFmtId")): n.get("formatCode")
        for n in styles.findall("m:numFmts/m:numFmt", _NS)
    }
    xf = styles.findall("m:cellXfs/m:xf", _NS)[int(cell.get("s", "0"))]
    fmt_id = int(xf.get("numFmtId"))
    assert fmt_id in codes, f"the cell's format id {fmt_id} is a built-in, not a custom code"
    return cell, codes[fmt_id], styles_bytes


def test_a_floor_cell_in_the_file_is_numeric_and_its_format_code_is_the_marker() -> None:
    cell, code, styles_bytes = _cell_in_the_file(_render(savings=12345.0, known=False))
    # (a) a NUMBER: no shared-string or inline-string type, and a numeric <v>.
    assert cell.get("t") in (None, "n"), cell.attrib
    assert float(cell.find("m:v", _NS).text) == 12345.0
    # (b) the format code is the ruled literal, the quotes around "≥ " intact...
    assert code == FLOOR_FORMAT, code
    # ...and the file holds it as UTF-8 text with the quotes escaped as XML, not a
    # character reference or a lost glyph: this is the attribute Excel reads.
    assert 'formatCode="&quot;≥ &quot;$#,##0"'.encode() in styles_bytes
    assert code.startswith('"≥ "'), code


def test_an_exact_cell_in_the_file_keeps_the_plain_code_and_no_marker() -> None:
    cell, code, _ = _cell_in_the_file(_render(savings=12345.0, known=True))
    assert cell.get("t") in (None, "n"), cell.attrib
    assert float(cell.find("m:v", _NS).text) == 12345.0
    # (c) the plain code, and no marker anywhere in what the cell is formatted by.
    assert code == EXACT_FORMAT, code
    assert "≥" not in code
