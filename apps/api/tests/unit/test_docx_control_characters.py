"""Every service's Word deliverable is produced when client text carries a
control character, with the character removed (#993).

python-docx refuses any string that XML 1.0 cannot carry ("All strings must be
XML compatible ..."), so one stored legal name such as "R&D\\x07 Labs" -- which
`POST /admin/clients` accepts -- failed every DOCX export for that client: the
finalize of ATT&CK, CSF, Zero Trust and Tech Debt, the Risk Register export,
and the CSF Playbook export. The XLSX half already strips these characters
(`app/xlsx_export.py`, #972); the DOCX half now strips the same set, through
`app/docx_export.py`, and logs one warning per stripped string that never
carries the value.

Every check goes through the surface the client reaches: the legal name is set
through the admin route, the title through the route that opens the service,
model output through the fixture provider the route calls, and the documents
through each service's finalize or export route. The files are downloaded and
read back with python-docx. Expected strings are literals, and the positive
state is asserted first.

The CSF Playbook's "Client: ..." row in its XLSX is pinned here too: before
this fix the route raised at the DOCX half, so nothing could reach that cell
through the route (#736 6084741554).
"""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook

from app.ai.llm import LLMResponse
from tests.unit import test_export_escape as _escape
from tests.unit.test_csf_playbook_override_every_artifact import _export, _one_gap
from tests.unit.test_export_escape import _FINALIZE, _tenant

#: The escape suite's app: a fresh database, the admin route, fixture AI.
app_client = _escape.app_client

pytestmark = pytest.mark.unit

DIRTY_NAME = "R&D\x07 Labs"
CLEAN_NAME = "R&D Labs"
DIRTY_TITLE = "Ops\x07 Review"
CLEAN_TITLE = "Ops Review"
EVENT = "docx_export.control_characters_removed"


def _docx_text(raw: bytes) -> list[str]:
    """Every string a reader of the file meets: the document title property,
    body paragraphs, table cells, and section headers and footers."""
    from docx import Document

    doc = Document(io.BytesIO(raw))
    out = [doc.core_properties.title or ""]
    out += [p.text for p in doc.paragraphs]
    out += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    for s in doc.sections:
        out += [p.text for p in s.footer.paragraphs] + [p.text for p in s.header.paragraphs]
    return out


def _download(c, h, artifact_id: str) -> bytes:
    d = c.get(f"/artifacts/{artifact_id}/download", headers=h)
    assert d.status_code == 200, d.text
    return d.content


def _finalized_docx(app_client, service: str, *, name: str, title: str) -> bytes:
    c, provider = app_client
    bearer, _cid, h = _tenant(c, name)
    fin = _FINALIZE[service](c, provider, bearer, h, title)
    assert fin.status_code == 201, fin.text
    return _download(c, h, fin.json()["docx_artifact_id"])


def _events(out: str) -> list[dict]:
    lines = [json.loads(line) for line in out.splitlines() if line.startswith("{")]
    return [e for e in lines if e.get("event") == EVENT]


# --- the four finalize routes ------------------------------------------------------


@pytest.mark.parametrize("service", ("attack", "csf", "zt", "tech_debt"))
def test_a_legal_name_with_a_control_character_finalizes_the_docx(app_client, service) -> None:
    text = _docx_text(_finalized_docx(app_client, service, name=DIRTY_NAME, title="Plain title"))
    assert any(CLEAN_NAME in s for s in text), text  # the positive state first
    assert not any("\x07" in s for s in text)


@pytest.mark.parametrize("service", ("attack", "csf", "zt", "tech_debt"))
def test_the_document_title_property_carries_the_stripped_name(app_client, service) -> None:
    text = _docx_text(_finalized_docx(app_client, service, name=DIRTY_NAME, title="Plain title"))
    assert text[0] == f"Plain title — {CLEAN_NAME}"


@pytest.mark.parametrize("service", ("attack", "csf", "zt", "tech_debt"))
def test_a_service_title_with_a_control_character_finalizes_the_docx(app_client, service) -> None:
    text = _docx_text(_finalized_docx(app_client, service, name="Plain Client", title=DIRTY_TITLE))
    assert text[0] == f"{CLEAN_TITLE} — Plain Client"
    assert CLEAN_TITLE in text[1:], text  # the heading, as a paragraph of its own


def test_tech_debt_item_text_with_a_control_character_reaches_the_docx_table(
    app_client, monkeypatch
) -> None:
    """The extraction model's item name, vendor and category print in the
    Capability list table."""
    item = dict(_escape._TD_ITEM, name="Splunk\x07 SIEM", vendor="Cisco\x0b", category="Sec\x1fOps")
    monkeypatch.setattr(_escape, "_TD_ITEM", item)
    text = _docx_text(
        _finalized_docx(app_client, "tech_debt", name="Plain Client", title="Plain title")
    )
    for value in ("Splunk SIEM", "Cisco", "SecOps"):
        assert value in text, f"{value!r} is in no cell"


# --- the Risk Register export --------------------------------------------------------


def _risk_docx(app_client, *, name: str, title: str) -> bytes:
    from tests.unit.test_risk_register import _entries_payload, _entry, _seed_attack_and_zt

    c, provider = app_client
    bearer, cid, h = _tenant(c, name)
    _seed_attack_and_zt(c, bearer, cid)
    provider.register_static("risk_synthesize", LLMResponse(_entries_payload(_entry(title))))
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    ex = c.post(f"/risk/clients/{cid}/register/export", headers=h)
    assert ex.status_code == 200, ex.text
    return _download(c, h, ex.json()["docx_artifact_id"])


def test_the_risk_register_docx_with_a_control_character_name_exports(app_client) -> None:
    text = _docx_text(_risk_docx(app_client, name=DIRTY_NAME, title="Weak MFA"))
    assert text[0] == f"Risk Register — {CLEAN_NAME}"
    assert CLEAN_NAME in text[1:], text
    assert "Weak MFA" in text


def test_a_risk_entry_title_with_a_control_character_reaches_the_docx(app_client) -> None:
    # `_entry` writes JSON by hand, so the byte goes in as a JSON escape: the
    # model's response carries it exactly as a provider's would.
    text = _docx_text(_risk_docx(app_client, name="Plain Client", title="Weak\\u0007 MFA"))
    assert "Weak MFA" in text, text


# --- the CSF Playbook export ----------------------------------------------------------


def _playbook(app_client, name: str) -> dict[str, bytes]:
    c, _provider = app_client
    _bearer, _cid, h = _tenant(c, name)
    svc, _code = _one_gap(c, h)
    return _export(c, h, svc)


def test_the_playbook_exports_every_docx_for_a_control_character_name(app_client) -> None:
    files = _playbook(app_client, DIRTY_NAME)
    exec_text = _docx_text(files["exec_docx"])
    full_text = _docx_text(files["full_docx"])
    assert exec_text[0] == f"CSF 2.0 Executive Briefing — {CLEAN_NAME}"
    assert full_text[0] == f"CSF 2.0 Full Playbook — {CLEAN_NAME}"
    for text in (exec_text, full_text):
        assert f"Prepared for: {CLEAN_NAME}" in text


def test_the_playbook_workbook_client_row_is_stripped(app_client) -> None:
    """The one site PR #995 could not reach through the route: the About
    sheet's "Client: ..." row."""
    ws = load_workbook(io.BytesIO(_playbook(app_client, DIRTY_NAME)["xlsx"]))["About"]
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "Client: R&D Labs" in values, values


# --- the log ---------------------------------------------------------------------------


def test_each_stripped_string_logs_one_warning_without_the_value(app_client, capsys) -> None:
    """`app_client` first: the app's lifespan configures logging, and structlog
    renders to stdout, so `capsys` (not `caplog`) sees the line -- the lesson
    `test_export_escape.py` records."""
    capsys.readouterr()
    _finalized_docx(app_client, "csf", name="QZXJ\x07TOKEN", title="Plain title")
    out = capsys.readouterr().out
    events = _events(out)
    assert events, "no warning for a stripped string"  # the positive state first
    for e in events:
        assert e["level"] == "warning"
        assert e["removed"] == 1
        assert e["part"] and e["field"]
    # One per site: the title property, the subtitle under the heading.
    sites = sorted((e["part"], e["field"]) for e in events)
    assert sites == [("body", "subtitle"), ("core_properties", "title")], sites
    for fragment in ("QZXJ", "TOKEN", "\\u0007"):
        assert fragment not in out, f"the log carries part of the value: {fragment!r}"


def test_a_clean_document_logs_nothing(app_client, capsys) -> None:
    """A dirty export's warning is seen FIRST, in this test, so the capture is
    proven live before a clean one is shown to add nothing."""
    c, provider = app_client
    capsys.readouterr()
    bearer, _cid, h = _tenant(c, "Dirty\x07")
    fin = _FINALIZE["zt"](c, provider, bearer, h, "Plain title")
    assert fin.status_code == 201, fin.text
    assert _events(capsys.readouterr().out)
    created = c.post(
        "/admin/clients",
        headers={"Authorization": f"Bearer {bearer}"},
        json={"legal_name": "Clean Client"},
    )
    assert created.status_code in (200, 201), created.text
    clean = {"Authorization": f"Bearer {bearer}", "X-Client-Id": created.json()["id"]}
    fin = _FINALIZE["zt"](c, provider, bearer, clean, "Plain title")
    assert fin.status_code == 201, fin.text
    assert _events(capsys.readouterr().out) == []


# --- the sites only code-written text reaches today ---------------------------------
#
# No client or model text reaches these today; each is pinned by putting a byte
# into the code-written text the route renders, so deleting the strip at that
# site turns a named test red.


@pytest.mark.parametrize("hook", ["target_cap_sentences", "stage_above_max_sentences"])
def test_zt_catalog_sentences_are_stripped_in_the_docx(app_client, monkeypatch, hook) -> None:
    import app.zt.exporters as zt_exporters

    monkeypatch.setattr(zt_exporters, hook, lambda *a, **kw: ["DoD\x07 sentence."])
    text = _docx_text(_finalized_docx(app_client, "zt", name="Plain Client", title="Plain title"))
    assert "DoD sentence." in text, text


def test_the_playbook_footer_and_function_heading_are_stripped(app_client, monkeypatch) -> None:
    import app.csf.playbook_export as playbook

    monkeypatch.setattr(playbook, "_approval_notice", lambda approved: "Notice\x07 text")
    monkeypatch.setitem(playbook.FUNCTION_NAMES, "GV", "Gov\x07ern")
    full = _docx_text(_playbook(app_client, "Plain Client")["full_docx"])
    assert "Govern (GV)" in full, full  # the function-detail heading
    assert "Notice text" in full, full  # the footer, and the cover line


# --- U+FFFE and U+FFFF: refused by XML 1.0, missed by openpyxl's regex ----------------


@pytest.mark.parametrize("char", ["\ufffe", "\uffff"], ids=["U+FFFE", "U+FFFF"])
def test_a_noncharacter_legal_name_finalizes_the_docx_and_the_xlsx(
    app_client, capsys, char
) -> None:
    """Pydantic accepts these from JSON; python-docx and openpyxl both raised on
    them, so finalize failed in both formats (review of 185393fc). Each strip
    is logged with its count: a strip counted as zero would log nothing.

    `app_client` first: its lifespan configures logging, so `capsys` sees it."""
    c, provider = app_client
    bearer, _cid, h = _tenant(c, f"QZXJ{char}")
    capsys.readouterr()
    fin = _FINALIZE["csf"](c, provider, bearer, h, "Plain title")
    assert fin.status_code == 201, fin.text
    out = capsys.readouterr().out
    text = _docx_text(_download(c, h, fin.json()["docx_artifact_id"]))
    assert text[0] == "Plain title — QZXJ"  # the positive state first
    assert not any(char in s for s in text)
    wb = load_workbook(io.BytesIO(_download(c, h, fin.json()["xlsx_artifact_id"])))
    values = [cell.value for ws in wb.worksheets for row in ws.iter_rows() for cell in row]
    assert "QZXJ" in values, values
    assert not any(isinstance(v, str) and char in v for v in values)
    logged = [json.loads(line) for line in out.splitlines() if line.startswith("{")]
    for event in (EVENT, "xlsx_export.control_characters_removed"):
        found = [e for e in logged if e.get("event") == event]
        assert found, f"no {event} for the stripped legal name"
        assert all(e["removed"] == 1 for e in found), found
    assert "QZXJ" not in out, "the log carries the legal name"


# --- the core-properties length limit -------------------------------------------------


def test_a_long_title_and_name_finalize_with_the_core_title_capped(app_client, capsys) -> None:
    """python-docx refuses a core property over 255 characters, and its error
    quotes the value. Each part may be 255 on its own, so the joined title is
    capped; the visible heading keeps the whole title. The cut is logged once,
    with the length the code logs (after cleaning: 200 + 3 + 60), never the
    value."""
    title = "T" * 199 + "Z"  # 200 characters
    name = "N" * 59 + "Q"  # 60 characters
    capsys.readouterr()
    docx = _finalized_docx(app_client, "csf", name=name, title=title)
    out = capsys.readouterr().out
    text = _docx_text(docx)
    assert len(text[0]) == 255, len(text[0])
    assert text[0] == "T" * 199 + "Z — " + "N" * 52  # 200 + 3 + 52
    assert title in text[1:], "the visible heading lost part of the title"
    logged = [json.loads(line) for line in out.splitlines() if line.startswith("{")]
    cut = [e for e in logged if e.get("event") == "docx_export.core_property_truncated"]
    assert len(cut) == 1, cut
    assert (cut[0]["field"], cut[0]["length"], cut[0]["limit"]) == ("title", 263, 255)
    for run in ("TTTTTTTT", "NNNNNNNN"):
        assert run not in out, f"the log carries the value: {run!r}"


# --- the AI-mode stamp: its DOCX paragraph and its XLSX sheet ---------------------------


def test_the_ai_mode_stamp_is_stripped_in_the_docx_and_the_xlsx(app_client, monkeypatch) -> None:
    """Only code writes the stamp's sentence today, so a byte is put into it.
    Before #993 the XLSX sheet's bare `ws.append` raised first, so the DOCX
    paragraph could not be reached through the route."""
    from app.mode_stamp import AiModeStamp

    monkeypatch.setattr(AiModeStamp, "sentence", lambda self: "Stamp\x07 sentence.")
    c, provider = app_client
    bearer, _cid, h = _tenant(c, "Plain Client")
    fin = _FINALIZE["csf"](c, provider, bearer, h, "Plain title")
    assert fin.status_code == 201, fin.text
    assert "Stamp sentence." in _docx_text(_download(c, h, fin.json()["docx_artifact_id"]))
    ws = load_workbook(io.BytesIO(_download(c, h, fin.json()["xlsx_artifact_id"])))["AI source"]
    assert "Stamp sentence." in [cell.value for row in ws.iter_rows() for cell in row]
