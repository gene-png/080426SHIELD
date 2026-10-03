"""#804: the deliverable names the `consolidate` disposition "Cut, covered by
another tool", in all three formats.

The stored value stays `consolidate` (no migration); only what a reader sees
changed. `exporters._disposition_label` is the deliverable's one label map,
and its web twin is `apps/web/src/lib/tech_debt/dispositionLabels.ts`.
"""

from __future__ import annotations

import io

import pytest

from app.tech_debt.exporters import render_docx, render_pdf, render_xlsx
from tests.unit.test_exporters import context_with_items  # noqa: F401  (fixture)

pytestmark = pytest.mark.unit

NEW = "Cut, covered by another tool"


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    pages = PdfReader(io.BytesIO(raw)).pages
    return " ".join(" ".join((p.extract_text() or "") for p in pages).split())


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return " ".join(" ".join([p.text for p in doc.paragraphs] + cells).split())


def _xlsx_text(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    return " ".join(
        str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row
    )


@pytest.mark.parametrize(
    "render, read",
    [(render_pdf, _pdf_text), (render_docx, _docx_text), (render_xlsx, _xlsx_text)],
    ids=["pdf", "docx", "xlsx"],
)
def test_the_deliverable_names_the_covered_cut(
    context_with_items, render, read  # noqa: F811
) -> None:
    text = read(render(context_with_items))
    assert NEW in text, text[:500]
    assert "Consolidate" not in text
