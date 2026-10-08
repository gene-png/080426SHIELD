"""A bundle's components are parts of ONE license, not tools of their own (for #835).

`add_capability_components` (routes/tech_debt.py) writes each named part with
`parent_item_id` set and NO cost, because the parent row keeps the whole
license value. Before this, every reader below treated a part as one more
separately licensed, separately costed tool:

  - the client dashboard read the part's missing cost as a missing cost, so a
    fully costed inventory with one split bundle showed spend as a floor;
  - the deliverable called the same figure "Annual cost (may not be complete)"
    and counted the parts as "Capabilities reviewed";
  - two parts of one bundle in one category read as a redundancy of two;
  - the admin overlap view asked for a cost on every part ("Missing cost"),
    which invites the double count, and put the parts in a vendor bucket with
    their own parent.

The advisor's rulings (#736): Applications and "Capabilities reviewed" count
SOURCE items; a part DOES count in its own category, so Defender for Endpoint
beside CrowdStrike stays visible as a redundancy (the reason splitting exists,
UX finding 5), but the parts of one bundle and the bundle itself count once
per category; savings are unchanged.

Every assertion goes through the route or the downloaded file a reader gets,
never through `tech_debt/components.py` directly.
"""

from __future__ import annotations

import io
import json
import os

import pytest
from sqlalchemy import create_engine, text

from app.ai.llm import LLMResponse
from tests._ai_runs import tech_debt_extract
from tests.unit.test_tech_debt_dashboard import (  # noqa: F401  (fixture)
    _register,
    app_client,
)

pytestmark = pytest.mark.unit

BUNDLE = "Microsoft 365 E5"
#: Four uploaded rows, every one costed. 294,120 + 120,000 + 60,000 + 40,000.
SOURCE = [
    (BUNDLE, "Microsoft", "Productivity and Security Suite", 294120),
    ("CrowdStrike Falcon", "CrowdStrike", "EDR", 120000),
    ("Okta", "Okta", "IAM", 60000),
    ("Slack", "Salesforce", "Collaboration", 40000),
]
SPEND = 514120.0

#: The UX finding 5 split: each part lands beside a separately licensed tool.
SPLIT_BESIDE_TOOLS = [
    {"name": "Microsoft Defender for Endpoint", "category": "EDR"},
    {"name": "Microsoft Entra ID P2", "category": "IAM"},
]
#: Two parts in one category, and a third in the bundle's own category.
SPLIT_ONE_CATEGORY = [
    {"name": "Defender for Office 365", "category": "Email Security"},
    {"name": "Exchange Online Protection", "category": "Email Security"},
    {"name": "Microsoft Teams", "category": "Productivity and Security Suite"},
]


def _extract_and_split(c, provider, components: list[dict], *, source: list[tuple] = SOURCE):
    """Extract `source` and split the bundle into `components`. Returns the
    admin headers, the service id, the list, the items and the bundle's id."""
    provider.register(
        "extract.capabilities",
        lambda _p: LLMResponse(
            content=json.dumps(
                {
                    "items": [
                        {
                            "name": name,
                            "vendor": vendor,
                            "category": category,
                            "function": "f",
                            "annual_cost_usd": cost,
                            "license_count": 1,
                            "notes": None,
                            "confidence_pct": 90,
                            "source_row_index": i,
                        }
                        for i, (name, vendor, category, cost) in enumerate(source)
                    ]
                }
            )
        ),
    )
    admin = _register(c, "admin@example.com")
    h = {"Authorization": f"Bearer {admin['tokens']['access_token']}"}
    svc = c.post("/tech-debt/services", headers=h, json={"title": "TD"}).json()["id"]
    # One upload row per source item, so the reconciliation balances.
    csv = ("Tool,Cost\n" + "".join(f"{row[0]},1\n" for row in source)).encode()
    art = c.post(
        "/artifacts", headers=h, files={"file": ("inv.csv", io.BytesIO(csv), "text/csv")}
    ).json()["id"]
    lst = tech_debt_extract(c, svc, h, art)
    bundle_id = next(i["id"] for i in lst["items"] if i["name"] == BUNDLE)
    split = c.post(
        f"/tech-debt/capability-items/{bundle_id}/components",
        headers=h,
        json={"components": components},
    )
    assert split.status_code == 201, split.text
    items = split.json()["items"]
    assert sum(1 for i in items if i.get("parent_item_id") == bundle_id) == len(components)
    return h, svc, lst, items, bundle_id


def _release(
    c,
    provider,
    components: list[dict],
    *,
    source: list[tuple] = SOURCE,
    disposition: dict[str, str] | None = None,
    component_cost: float | None = None,
) -> dict:
    """Extract `source`, split the bundle, decide every row, approve, finalize
    and release. Returns what each reader sees.

    `component_cost` writes a LEGACY cost onto every part. Since #927 the PATCH
    route refuses that, so it is written straight to the table: the state a
    part could reach before this PR, which every reader must still handle."""
    h, svc, lst, items, bundle_id = _extract_and_split(c, provider, components, source=source)
    client = _register(c, "client@example.com")
    if component_cost is not None:
        _store_legacy_part_cost(bundle_id, component_cost, expected=len(components))
    for it in items:
        body: dict = {"disposition": (disposition or {}).get(it["name"], "keep")}
        r = c.patch(f"/tech-debt/capability-items/{it['id']}", headers=h, json=body)
        assert r.status_code == 200, r.text
    overlap = c.get(f"/tech-debt/services/{svc}/overlap-analysis", headers=h)
    assert overlap.status_code == 200, overlap.text
    assert c.post(f"/tech-debt/capability-lists/{lst['id']}/approve", headers=h).status_code == 200
    fin = c.post(f"/tech-debt/services/{svc}/deliverables/finalize", headers=h)
    assert fin.status_code == 201, fin.text
    deliv = fin.json()
    rel = c.post(f"/tech-debt/deliverables/{deliv['id']}/release", headers=h)
    assert rel.status_code == 200, rel.text
    files = {}
    for kind in ("xlsx", "pdf", "docx"):
        got = c.get(f"/artifacts/{deliv[f'{kind}_artifact_id']}/download", headers=h)
        assert got.status_code == 200, got.text
        files[kind] = got.content
    cid = client["user"]["client_id"]
    ch = {"Authorization": f"Bearer {client['tokens']['access_token']}", "X-Client-Id": cid}
    dash = c.get(f"/clients/{cid}/tech-debt/{svc}/dashboard", headers=ch)
    assert dash.status_code == 200, dash.text
    # The client's Results and Home pages render each deliverable's `summary`.
    results = c.get(f"/clients/{cid}/deliverables", headers=ch)
    assert results.status_code == 200, results.text
    mine = [d for d in results.json()["items"] if d["id"] == deliv["id"]]
    assert len(mine) == 1, results.text
    return {
        "dashboard": dash.json(),
        "overlap": overlap.json(),
        "files": files,
        "summary": mine[0]["summary"],
    }


@pytest.fixture()
def env(app_client):  # noqa: F811
    return app_client


def _store_legacy_part_cost(bundle_id: str, cost: float, *, expected: int) -> None:
    """A cost on a part as a pre-#927 PATCH could leave it (see `_release`)."""
    engine = create_engine(os.environ["DATABASE_URL"], future=True)
    with engine.begin() as conn:
        n = conn.execute(
            text("UPDATE capability_items SET annual_cost_usd = :c WHERE parent_item_id = :p"),
            {"c": cost, "p": bundle_id.replace("-", "")},
        ).rowcount
    engine.dispose()
    assert n == expected, f"wrote a legacy cost onto {n} parts, expected {expected}"


def _xlsx_cells(raw: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    return [str(x.value) for ws in wb for row in ws.iter_rows() for x in row if x.value is not None]


def _xlsx_total(raw: bytes) -> float:
    """The figure beside the workbook's cost label (column A names it)."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw)).active
    rows = [
        r
        for r in ws.iter_rows()
        if isinstance(r[0].value, str) and "annual cost" in r[0].value.lower()
    ]
    assert len(rows) == 1, [r[0].value for r in rows]
    return float(rows[0][4].value)


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return " ".join(
        " ".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(raw)).pages).split()
    )


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    return " ".join(" ".join(p.text for p in doc.paragraphs).split())


def _redundancies(dash: dict) -> dict[str, int]:
    return {r["category"]: r["count"] for r in dash["redundancies"]}


# --- spend -----------------------------------------------------------------


def test_a_split_bundle_leaves_a_fully_costed_spend_complete(env) -> None:
    c, provider = env
    dash = _release(c, provider, SPLIT_BESIDE_TOOLS)["dashboard"]
    assert dash["annual_spend_usd"] == SPEND
    assert (
        dash["spend_completeness"] == "complete"
    ), "a bundle's parts carry no cost by design; their missing cost is not a missing cost"


def test_a_costed_part_never_adds_to_spend(env) -> None:
    """The parent holds the license value. A LEGACY cost on a part (a PATCH
    could store one before #927) must not count the bundle twice."""
    c, provider = env
    got = _release(c, provider, SPLIT_BESIDE_TOOLS, component_cost=50000)
    dash = got["dashboard"]
    assert dash["annual_spend_usd"] == SPEND
    edr = next(s for s in dash["spend_by_category"] if s["category"] == "EDR")
    assert edr["total_usd"] == 120000.0
    assert "Total annual cost" in _xlsx_cells(got["files"]["xlsx"])
    # The TOTAL, not only its label: a part's typed cost must not reach it.
    assert _xlsx_total(got["files"]["xlsx"]) == SPEND
    assert "Total annual cost: $514,120 ·" in _pdf_text(got["files"]["pdf"])
    # The admin overlap view's total and its top-cost list read the same rule.
    assert got["overlap"]["total_cost"] == SPEND
    top = [t["name"] for t in got["overlap"]["top_cost_items"]]
    assert top[0] == BUNDLE
    assert not {p["name"] for p in SPLIT_BESIDE_TOOLS} & set(top)


def test_a_real_uncosted_source_row_still_reads_partial(env) -> None:
    """The guard skips PARTS, not every uncosted row."""
    c, provider = env
    source = [*SOURCE, ("Zoom", "Zoom", "Collaboration", None)]
    got = _release(c, provider, SPLIT_BESIDE_TOOLS, source=source)
    assert got["dashboard"]["spend_completeness"] == "partial"
    cells = _xlsx_cells(got["files"]["xlsx"])
    assert "Annual cost (may not be complete)" in cells
    assert "Total annual cost" not in cells


# --- counts ----------------------------------------------------------------


def test_applications_count_source_items_not_parts(env) -> None:
    c, provider = env
    dash = _release(c, provider, SPLIT_BESIDE_TOOLS)["dashboard"]
    assert dash["total_applications"] == 4
    assert dash["included_count"] == 4
    assert dash["bundle_part_count"] == 2
    # The inventory still lists every row, parts included.
    assert len(dash["items"]) == 6


def test_the_deliverable_calls_a_fully_costed_total_a_total(env) -> None:
    c, provider = env
    cells = _xlsx_cells(_release(c, provider, SPLIT_BESIDE_TOOLS)["files"]["xlsx"])
    assert "Total annual cost" in cells
    assert "Annual cost (may not be complete)" not in cells


def test_the_pdf_counts_source_items_as_capabilities_reviewed(env) -> None:
    c, provider = env
    text = _pdf_text(_release(c, provider, SPLIT_BESIDE_TOOLS)["files"]["pdf"])
    assert "Capabilities reviewed: 4 ·" in text
    assert "Capabilities reviewed: 6" not in text


def test_the_docx_counts_source_items_as_capabilities_reviewed(env) -> None:
    c, provider = env
    text = _docx_text(_release(c, provider, SPLIT_BESIDE_TOOLS)["files"]["docx"])
    assert "Capabilities reviewed: 4" in text
    assert "Capabilities reviewed: 6" not in text


def test_the_client_results_summary_counts_source_items(env) -> None:
    """`Deliverable.summary` is what the client's Results and Home pages show
    beside the files; it must agree with the PDF's "Capabilities reviewed"."""
    c, provider = env
    summary = _release(c, provider, SPLIT_BESIDE_TOOLS)["summary"]
    assert summary.startswith("4 capabilities reviewed; "), summary
    assert "6 capabilities" not in summary


# --- categories (ruling Q1 = B) ----------------------------------------------


def test_a_part_beside_a_separately_licensed_tool_is_still_a_redundancy(env) -> None:
    """UX finding 5: splitting exists so Defender shows up beside CrowdStrike."""
    c, provider = env
    dash = _release(c, provider, SPLIT_BESIDE_TOOLS)["dashboard"]
    assert _redundancies(dash) == {"EDR": 2, "IAM": 2}
    assert dash["redundant_category_count"] == 2
    edr = next(r for r in dash["redundancies"] if r["category"] == "EDR")
    assert {i["name"] for i in edr["items"]} == {
        "CrowdStrike Falcon",
        "Microsoft Defender for Endpoint",
    }


def test_two_parts_beside_one_tool_are_two_licenses(env) -> None:
    """Three EDR rows (CrowdStrike and two parts of one bundle) are a
    redundancy of two licenses, which the client card prints as "· 2 licenses"."""
    c, provider = env
    dash = _release(
        c,
        provider,
        [
            {"name": "Microsoft Defender for Endpoint", "category": "EDR"},
            {"name": "Microsoft Defender for Cloud Apps", "category": "EDR"},
        ],
    )["dashboard"]
    edr = next(r for r in dash["redundancies"] if r["category"] == "EDR")
    assert len(edr["items"]) == 3
    assert edr["count"] == 2


def test_parts_of_one_bundle_in_one_category_count_once(env) -> None:
    c, provider = env
    dash = _release(c, provider, SPLIT_ONE_CATEGORY)["dashboard"]
    # Positive control first: the categories exist and carry one license each.
    counts = {s["category"]: s["count"] for s in dash["spend_by_category"]}
    assert counts["Email Security"] == 1
    assert counts["Productivity and Security Suite"] == 1
    assert _redundancies(dash) == {}
    assert dash["sprawl_by_category"] == []
    assert dash["redundant_category_count"] == 0


# --- savings (unchanged by ruling) ------------------------------------------


def test_a_cut_part_with_no_cost_still_makes_savings_a_floor(env) -> None:
    c, provider = env
    dash = _release(
        c,
        provider,
        SPLIT_BESIDE_TOOLS,
        disposition={"Microsoft Defender for Endpoint": "cut"},
    )["dashboard"]
    assert dash["savings_cost_known"] is False
    # Spend is still complete: only the SAVINGS figure is a floor.
    assert dash["spend_completeness"] == "complete"


# --- admin overlap view -------------------------------------------------------


def test_the_overlap_view_does_not_ask_for_a_cost_on_a_part(env) -> None:
    c, provider = env
    overlap = _release(c, provider, SPLIT_BESIDE_TOOLS)["overlap"]
    assert overlap["no_cost_count"] == 0
    assert overlap["total_cost_label"] == "Total annual cost"
    edr = next(b for b in overlap["by_category"] if b["key"] == "EDR")
    assert edr["item_count"] == 2
    assert edr["cost_known"] is True


def test_the_overlap_view_counts_a_bundle_once_per_vendor_and_category(env) -> None:
    """Parts inherit the bundle's vendor, so without the license key Microsoft
    read as three subscriptions from one vendor."""
    c, provider = env
    overlap = _release(c, provider, SPLIT_ONE_CATEGORY)["overlap"]
    assert overlap["by_vendor"] == []
    assert overlap["by_category"] == []


# --- a cost on a part is refused (for #927) ---------------------------------

#: Approved verbatim by the advisor on #736 (comment 6056012075).
REFUSED = (
    "The annual cost of a bundle part is held by its bundle, Microsoft 365 E5. "
    "Enter the cost on the Microsoft 365 E5 row instead."
)


def _latest_items(c, h, svc: str) -> dict[str, dict]:
    r = c.get(f"/tech-debt/services/{svc}/capability-lists/latest", headers=h)
    assert r.status_code == 200, r.text
    return {i["name"]: i for i in r.json()["items"]}


@pytest.mark.parametrize("cost", [50000, 0])
def test_a_cost_on_a_part_is_refused_and_names_the_bundle(env, cost) -> None:
    c, provider = env
    h, svc, _lst, items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    part = next(i for i in items if i.get("parent_item_id") == bundle_id)
    r = c.patch(
        f"/tech-debt/capability-items/{part['id']}", headers=h, json={"annual_cost_usd": cost}
    )
    assert r.status_code == 422, r.text
    assert r.json()["error"]["reason"] == "component_cost_held_by_bundle"
    assert r.json()["error"]["message"] == REFUSED
    # Refused before anything is committed.
    assert _latest_items(c, h, svc)[part["name"]]["annual_cost_usd"] is None


def test_a_null_cost_on_a_part_is_allowed(env) -> None:
    """Null is the state a part is born in, and the way a legacy cost is cleared."""
    c, provider = env
    h, svc, _lst, items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    part = next(i for i in items if i.get("parent_item_id") == bundle_id)
    _store_legacy_part_cost(bundle_id, 50000, expected=len(SPLIT_BESIDE_TOOLS))
    r = c.patch(
        f"/tech-debt/capability-items/{part['id']}",
        headers=h,
        json={"annual_cost_usd": None, "notes": "cleared"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["annual_cost_usd"] is None
    assert _latest_items(c, h, svc)[part["name"]]["annual_cost_usd"] is None


def test_the_bundle_row_still_takes_a_cost(env) -> None:
    """The remedy the refusal names exists: the parent row's cost is editable."""
    c, provider = env
    h, svc, _lst, _items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    r = c.patch(
        f"/tech-debt/capability-items/{bundle_id}", headers=h, json={"annual_cost_usd": 300000}
    )
    assert r.status_code == 200, r.text
    assert float(_latest_items(c, h, svc)[BUNDLE]["annual_cost_usd"]) == 300000.0


def test_a_part_edit_without_a_cost_is_unaffected(env) -> None:
    c, provider = env
    h, _svc, _lst, items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    part = next(i for i in items if i.get("parent_item_id") == bundle_id)
    r = c.patch(f"/tech-debt/capability-items/{part['id']}", headers=h, json={"notes": "n"})
    assert r.status_code == 200, r.text


def test_the_refusal_without_a_loadable_bundle_names_no_bundle() -> None:
    """No writer can produce a part whose bundle row is gone (a row is never
    deleted on its own), so this branch is pinned on the message builder: the
    route must not invent a name."""
    from app.routes.tech_debt import _component_cost_refusal

    detail = _component_cost_refusal(None).detail
    assert detail == {
        "reason": "component_cost_held_by_bundle",
        "message": (
            "The annual cost of a bundle part is held by its bundle. "
            "Enter the cost on the bundle's row instead."
        ),
    }
    assert _component_cost_refusal("Microsoft 365 E5").detail["message"] == REFUSED


def test_a_blank_bundle_name_gets_the_unnamed_refusal(env) -> None:
    """The item PATCH accepts a whitespace-only name, so a bundle can be
    renamed "   ". The refusal must not interpolate a blank: it uses the
    approved fallback instead."""
    c, provider = env
    h, _svc, _lst, items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    renamed = c.patch(f"/tech-debt/capability-items/{bundle_id}", headers=h, json={"name": "   "})
    assert renamed.status_code == 200, renamed.text
    part = next(i for i in items if i.get("parent_item_id") == bundle_id)
    r = c.patch(
        f"/tech-debt/capability-items/{part['id']}", headers=h, json={"annual_cost_usd": 100}
    )
    assert r.status_code == 422, r.text
    assert r.json()["error"]["message"] == (
        "The annual cost of a bundle part is held by its bundle. "
        "Enter the cost on the bundle's row instead."
    )


def test_a_refused_part_patch_writes_none_of_its_fields(env) -> None:
    """A name-plus-cost PATCH on a part is refused whole: the name is not
    committed either. Moving the refusal below the commit turns this red;
    moving it only below the in-memory assignments does not, because the
    request then ends without a commit and nothing reaches the table."""
    c, provider = env
    h, svc, _lst, items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    part = next(i for i in items if i.get("parent_item_id") == bundle_id)
    r = c.patch(
        f"/tech-debt/capability-items/{part['id']}",
        headers=h,
        json={"name": "Renamed part", "annual_cost_usd": 100},
    )
    assert r.status_code == 422, r.text
    after = _latest_items(c, h, svc)
    assert part["name"] in after
    assert "Renamed part" not in after
    assert after[part["name"]]["annual_cost_usd"] is None
