"""#1343, the Risk Register third: a ZT finding whose target is a
per-capability target with no recorded origin says so, with the approved
sentence (#736 6220094031 item 5) in the approved form (6225462107):
`{code}: <sentence>`, one line per code.

Mirrors #915's capped-target note (`test_risk_zt_capped_target_disclosed.py`):
`generate` records the codes in provenance, one reader builds the lines, and
the consultant's register, the client's Risk dashboard and the three files all
state them.

The world: an ATT&CK gap (the register's gate) and a released CISA Zero Trust
assessment whose one answered capability sits at Stage 1 under its own target
of 3, against an engagement target of 2. The admin PATCH records that target as
the consultant's; the "not recorded" case clears that origin, which is what
every target stored before migration 0076 looks like. Expected text is written
here from the ruling, never imported.
"""

from __future__ import annotations

import io
import os
import uuid

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from tests._ai_mode import docx_text, pdf_text
from tests._attack_rows import first_standalone
from tests._risk_inputs import release
from tests.unit.test_risk_dashboard import (  # noqa: F401  (fixture)
    _register,
    app_client,
)
from tests.unit.test_zt_dashboard import _attach_intake_target

pytestmark = pytest.mark.unit

SENTENCE = "Who set this target is not recorded. An earlier AI run may have set it."


def _world(c, *, recorded: bool, maturity: int = 1) -> tuple[str, dict, dict, str]:
    from app.models.zt_assessment import ZtAnswer

    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    cid = client["user"]["client_id"]
    c.headers["X-Client-Id"] = cid
    h = {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}

    asvc = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": "A"})
    a = c.post(f"/attack/services/{asvc.json()['id']}/assessments", headers=h).json()
    cov = first_standalone(a["coverage"])
    r = c.patch(f"/attack/coverage/{cov['id']}", headers=h, json={"status": "gap"})
    assert r.status_code == 200, r.text
    assert c.post(f"/attack/assessments/{a['id']}/approve", headers=h).status_code == 200

    zsvc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": "ZT"})
    _attach_intake_target(zsvc.json()["id"], zt_stage=2)
    za = c.post(f"/zt/services/{zsvc.json()['id']}/assessments", headers=h).json()
    ans = za["answers"][0]
    r = c.patch(
        f"/zt/answers/{ans['id']}", headers=h, json={"maturity_stage": maturity, "target_stage": 3}
    )
    assert r.status_code == 200, r.text
    if not recorded:
        engine = create_engine(os.environ["DATABASE_URL"], future=True)
        with sessionmaker(bind=engine, future=True)() as s:
            row = s.get(ZtAnswer, uuid.UUID(ans["id"]))
            assert row.target_source == "consultant", "setup: the PATCH records it"
            row.target_source = None  # stored before 0076
            s.commit()
    assert c.post(f"/zt/assessments/{za['id']}/approve", headers=h).status_code == 200
    release(c, bearer, cid, "attack", asvc.json()["id"])
    release(c, bearer, cid, "zt", zsvc.json()["id"])
    return cid, h, client, ans["capability_code"]


def _flat(text: str) -> str:
    return " ".join(text.split())


def _files(c, h: dict, body: dict) -> tuple[str, str, list[str]]:
    from openpyxl import load_workbook

    def _bytes(key: str) -> bytes:
        dl = c.get(f"/artifacts/{body[key]}/download", headers=h)
        assert dl.status_code == 200, dl.text
        return dl.content

    wb = load_workbook(io.BytesIO(_bytes("xlsx_artifact_id")))
    summary = [str(v) for row in wb["Summary"].iter_rows(values_only=True) for v in row if v]
    return (
        _flat(pdf_text(_bytes("pdf_artifact_id"))),
        _flat(docx_text(_bytes("docx_artifact_id"))),
        summary,
    )


def test_the_register_records_the_code_and_states_it(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, _client, code = _world(c, recorded=False)
    line = f"{code}: {SENTENCE}"
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    assert g.json()["zt_unrecorded_target_notes"] == [line]
    # Read back: from the persisted provenance, not from the generate run.
    latest = c.get(f"/risk/clients/{cid}/register/latest", headers=h)
    assert latest.status_code == 200, latest.text
    body = latest.json()
    assert body["zt_unrecorded_target_notes"] == [line]
    assert list(body["unrecorded_target_codes"].values()) == [[code]]


def test_every_published_file_and_the_client_dashboard_state_it(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, client, code = _world(c, recorded=False)
    line = f"{code}: {SENTENCE}"
    assert c.post(f"/risk/clients/{cid}/register/generate", headers=h).status_code == 201
    pub = c.post(f"/risk/clients/{cid}/register/publish", headers=h)
    assert pub.status_code in (200, 201), pub.text
    pdf, docx, summary = _files(c, h, pub.json())
    assert line in pdf
    assert line in docx
    assert line in summary
    dash = c.get(
        f"/clients/{cid}/risk/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert dash.status_code == 200, dash.text
    assert dash.json()["zt_unrecorded_target_notes"] == [line]


def test_a_consultant_set_target_states_nothing(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, client, _code = _world(c, recorded=True)
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    # The positive state first: the codes were recorded, and there are none.
    assert list(g.json()["unrecorded_target_codes"].values()) == [[]]
    assert g.json()["zt_unrecorded_target_notes"] == []
    pub = c.post(f"/risk/clients/{cid}/register/publish", headers=h)
    assert pub.status_code in (200, 201), pub.text
    pdf, docx, _summary = _files(c, h, pub.json())
    assert "Total entries:" in pdf  # a real summary, so the absence means something
    assert SENTENCE not in pdf
    assert SENTENCE not in docx
    dash = c.get(
        f"/clients/{cid}/risk/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert dash.status_code == 200, dash.text
    assert dash.json()["zt_unrecorded_target_notes"] == []


def test_a_target_that_suppresses_a_finding_is_stated_too(app_client) -> None:  # noqa: F811
    # The same unrecorded target of 3, but the capability is AT it (stage 3):
    # the target decided that this is NOT a risk. Who set it is still unknown,
    # and the sentence is about that, so the register states it with no
    # finding to hang it on (PR A's deliverable lists it with or without a gap).
    c = app_client
    cid, h, _client, code = _world(c, recorded=False, maturity=3)
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    line = f"{code}: {SENTENCE}"
    assert g.json()["zt_unrecorded_target_notes"] == [line]
    pub = c.post(f"/risk/clients/{cid}/register/publish", headers=h)
    assert pub.status_code in (200, 201), pub.text
    pdf, docx, summary = _files(c, h, pub.json())
    # The register is real, and no ZT finding for this capability is in it.
    assert "Total entries:" in pdf
    assert f"ZT {code}: stage" not in pdf
    # ... yet the files state who set the target.
    assert line in pdf
    assert line in docx
    assert line in summary
