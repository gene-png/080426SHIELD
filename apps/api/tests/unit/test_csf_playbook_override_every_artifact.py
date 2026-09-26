"""One playbook export gives a gap ONE priority, in all five artifacts (#696).

The consultant's stored `priority_override` is the effective priority: the
spec (Sprint 5 T5, step 10) says a stored override wins, the gap-actions API
reports it as `effective_priority`, and the XLSX Action Plan printed it. The
exec and full PDF/DOCX were never handed the overrides, so the same export
counted, listed and ranked that gap at its computed priority instead.

Driven through `/playbook/export` and the artifact download, because the defect
was which rows the ROUTE hands each renderer; a renderer called directly with
rows the test built cannot see that.
"""

from __future__ import annotations

import io
import os
import re
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.csf.catalog import SUBCATEGORIES

pytestmark = pytest.mark.unit

KINDS = {"xlsx", "exec_pdf", "exec_docx", "full_pdf", "full_docx"}


@pytest.fixture()
def app_client(tmp_path) -> Iterator[tuple[TestClient, dict]]:
    url = f"sqlite:///{tmp_path / 'shield-csf696.db'}"
    os.environ["DATABASE_URL"] = url
    api_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(api_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(api_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")
    engine = create_engine(url, future=True)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    from app.db.session import get_db
    from app.main import create_app
    from app.models.client import Client as _Client
    from app.models.client_domain import ClientDomain as _ClientDomain
    from app.routes.artifacts import _storage_dep
    from app.storage.local import LocalFilesystemStorage

    def override_get_db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_storage_dep] = lambda: LocalFilesystemStorage(tmp_path / "storage")
    seed = TestSession()
    tenant = _Client(legal_name="Test Tenant")
    seed.add(tenant)
    seed.flush()
    seed.add(_ClientDomain(client_id=tenant.id, domain="example.com"))
    seed.commit()
    with TestClient(app, headers={"X-Client-Id": str(tenant.id)}) as c:
        r = c.post(
            "/auth/register",
            json={
                "email": "admin@example.com",
                "password": "correct horse battery staple!",
                "display_name": "A",
            },
        )
        assert r.status_code == 201, r.text
        yield c, {"Authorization": f"Bearer {r.json()['tokens']['access_token']}"}


def _one_gap(c: TestClient, h: dict, code: str | None = None) -> tuple[str, str]:
    """A CSF service whose only gap is `code`, SUBCATEGORIES[0] by default
    (target L4 on every tier, nothing scored). Returns (service_id, gap code)."""
    svc = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()["id"]
    c.post(f"/csf/services/{svc}/assessments", headers=h)
    c.post(
        f"/csf/services/{svc}/profiles/seed", headers=h, json={"tiers": ["high", "moderate", "low"]}
    )
    code = code or SUBCATEGORIES[0].code
    for tier in ("high", "moderate", "low"):
        rows = c.get(f"/csf/services/{svc}/profile/{tier}", headers=h).json()["rows"]
        sid = next(r["id"] for r in rows if r["subcategory_code"] == code)
        res = c.patch(f"/csf/dimension-scores/{sid}", headers=h, json={"target_level": 4})
        assert res.status_code == 200, res.text
    gaps = [
        s["subcategory_code"]
        for s in c.get(f"/csf/services/{svc}/enterprise-profile", headers=h).json()["subcategories"]
        if s["gap"]
    ]
    assert gaps == [code], f"the setup must produce exactly one gap, got {gaps}"
    return svc, code


def _default_priority(c: TestClient, h: dict, svc: str, code: str) -> str:
    ent = c.get(f"/csf/services/{svc}/enterprise-profile", headers=h).json()
    (row,) = [s for s in ent["subcategories"] if s["subcategory_code"] == code]
    assert row["priority"] in {"P1", "P2", "P3"}
    return row["priority"]


def _override(c: TestClient, h: dict, svc: str, code: str, value: str) -> None:
    res = c.put(
        f"/csf/services/{svc}/gap-actions/{code}", headers=h, json={"priority_override": value}
    )
    assert res.status_code == 200, res.text


def _export(c: TestClient, h: dict, svc: str) -> dict[str, bytes]:
    ex = c.post(f"/csf/services/{svc}/playbook/export", headers=h)
    assert ex.status_code == 200, ex.text
    arts = {a["kind"]: a["artifact_id"] for a in ex.json()["artifacts"]}
    assert set(arts) == KINDS
    out = {}
    for kind, art_id in arts.items():
        dl = c.get(f"/artifacts/{art_id}/download", headers=h)
        assert dl.status_code == 200, f"{kind}: {dl.status_code}"
        out[kind] = dl.content
    return out


def _flat_text(kind: str, raw: bytes) -> str:
    """Every word of a PDF or DOCX, whitespace collapsed (PDF text wraps)."""
    from docx import Document
    from pypdf import PdfReader

    if kind.endswith("pdf"):
        text = " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
    else:
        doc = Document(io.BytesIO(raw))
        cells = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        text = " ".join([p.text for p in doc.paragraphs] + cells)
    return re.sub(r"\s+", " ", text)


def _xlsx_priority(raw: bytes, sheet: str, code: str) -> str | None:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))[sheet]
    rows = list(ws.iter_rows(values_only=True))
    (header_at,) = [i for i, r in enumerate(rows) if r and r[0] == "Subcategory"]
    col = rows[header_at].index("Priority")
    (row,) = [r for r in rows[header_at + 1 :] if r[0] == code]
    return row[col] or None


def _split(p1: int, p2: int, p3: int) -> str:
    # The overview's priority breakdown, as every PDF and DOCX renderer prints it.
    return f"{p1} Priority 1 (critical), {p2} Priority 2, and {p3} Priority 3."


def _overridden(app_client) -> tuple[str, str, str, dict[str, bytes]]:
    """(gap code, computed priority, override, exported files)."""
    c, h = app_client
    svc, code = _one_gap(c, h)
    default = _default_priority(c, h, svc, code)
    override = "P3" if default == "P1" else "P1"
    _override(c, h, svc, code, override)
    return code, default, override, _export(c, h, svc)


def test_an_overridden_gap_has_the_override_on_both_xlsx_sheets(app_client) -> None:
    code, _, override, files = _overridden(app_client)
    assert _xlsx_priority(files["xlsx"], "Action Plan", code) == override
    assert _xlsx_priority(files["xlsx"], "Enterprise Profile", code) == override


def test_an_overridden_gap_is_counted_at_the_override_in_every_pdf_and_docx(app_client) -> None:
    _, default, override, files = _overridden(app_client)
    expected = _split(*(int(override == p) for p in ("P1", "P2", "P3")))
    computed = _split(*(int(default == p) for p in ("P1", "P2", "P3")))
    for kind in KINDS - {"xlsx"}:
        text = _flat_text(kind, files[kind])
        assert expected in text, f"{kind} does not count the gap at its override {override}"
        assert computed not in text, f"{kind} still counts the gap at its computed {default}"


def test_without_an_override_every_artifact_keeps_the_computed_priority(app_client) -> None:
    c, h = app_client
    svc, code = _one_gap(c, h)
    default = _default_priority(c, h, svc, code)

    files = _export(c, h, svc)

    assert _xlsx_priority(files["xlsx"], "Action Plan", code) == default
    assert _xlsx_priority(files["xlsx"], "Enterprise Profile", code) == default
    expected = _split(*(int(default == p) for p in ("P1", "P2", "P3")))
    for kind in KINDS - {"xlsx"}:
        assert expected in _flat_text(kind, files[kind]), kind


def test_an_override_stored_on_a_subcategory_that_is_not_a_gap_gives_it_no_priority(
    app_client,
) -> None:
    # The upsert route accepts any catalogue code, so an override can sit on a
    # row that is not (or is no longer) a gap. It must not turn that row into a
    # prioritised gap in the deliverable.
    c, h = app_client
    svc, code = _one_gap(c, h)
    default = _default_priority(c, h, svc, code)
    other = SUBCATEGORIES[1].code
    _override(c, h, svc, other, "P1")

    files = _export(c, h, svc)

    assert _xlsx_priority(files["xlsx"], "Enterprise Profile", other) is None
    expected = _split(*(int(default == p) for p in ("P1", "P2", "P3")))
    for kind in KINDS - {"xlsx"}:
        assert expected in _flat_text(kind, files[kind]), kind


# --- #707 round 1: the copy around the counts must stay true for an override --


def _non_core_code() -> str:
    """A subcategory that is not a Core IG metric. Its computed priority on
    three tiers is P2 by the spec (Core AND HIGH AND multi-system is P1; Core
    OR HIGH is P2), which the test asserts through the API rather than trusting."""
    from app.csf.catalog import is_core

    return next(s.code for s in SUBCATEGORIES if not is_core(s.code))


def test_a_p1_set_by_override_is_not_described_as_a_computed_p1(app_client) -> None:
    # The exec next steps described every counted P1 as "Core-metric,
    # high-impact, multi-system". A consultant can raise a non-Core gap to P1,
    # and the sentence would then describe that gap as something it is not.
    c, h = app_client
    svc, code = _one_gap(c, h, _non_core_code())
    assert _default_priority(c, h, svc, code) == "P2"
    _override(c, h, svc, code, "P1")

    files = _export(c, h, svc)

    for kind in ("exec_pdf", "exec_docx"):
        text = _flat_text(kind, files[kind])
        assert (
            "Remediate the 1 Priority 1 gap(s) first. "
            "Consultant override set Priority 1 on 1 of them." in text
        ), kind
        assert "Core-metric, high-impact, multi-system" not in text, kind


def test_a_computed_p1_keeps_its_description(app_client) -> None:
    from app.csf.catalog import is_core

    c, h = app_client
    svc, code = _one_gap(c, h, next(s.code for s in SUBCATEGORIES if is_core(s.code)))
    assert _default_priority(c, h, svc, code) == "P1"

    files = _export(c, h, svc)

    for kind in ("exec_pdf", "exec_docx"):
        assert (
            "Remediate the 1 Priority 1 gap(s) first — these are Core-metric, "
            "high-impact, multi-system weaknesses." in _flat_text(kind, files[kind])
        ), kind


def test_the_methodology_says_a_priority_can_be_overridden_and_how_many(app_client) -> None:
    # METHODOLOGY said priorities "are computed deterministically in code",
    # false once an override reaches these documents.
    c, h = app_client
    svc, code = _one_gap(c, h, _non_core_code())
    _override(c, h, svc, code, "P1")

    files = _export(c, h, svc)

    for kind in ("full_pdf", "full_docx"):
        text = _flat_text(kind, files[kind])
        assert "a consultant may override it" in text, kind
        assert "Priorities set by consultant override in this playbook: 1." in text, kind
        assert "gaps, and priorities are computed deterministically" not in text, kind


def test_the_methodology_states_zero_overrides_too(app_client) -> None:
    c, h = app_client
    svc, _ = _one_gap(c, h)

    files = _export(c, h, svc)

    for kind in ("full_pdf", "full_docx"):
        assert "Priorities set by consultant override in this playbook: 0." in _flat_text(
            kind, files[kind]
        ), kind


def test_the_gap_action_put_on_a_non_gap_row_reports_no_effective_priority(app_client) -> None:
    # #707's own behaviour change: a row that is not a gap has no priority,
    # whatever override is stored against it -- in the PUT response as in the
    # deliverable. It used to echo the override.
    c, h = app_client
    svc, _ = _one_gap(c, h)
    other = SUBCATEGORIES[1].code
    res = c.put(
        f"/csf/services/{svc}/gap-actions/{other}", headers=h, json={"priority_override": "P1"}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["priority_override"] == "P1"
    assert body["effective_priority"] is None
