"""A LIVE csf_score run leaves a row with no notes as stored (#1000, for #806).

Gene's ruling 2 (comment 6101751588) and the advisor's (6102665946): a row whose
interview answer has no notes -- none, blank, or no answer at all -- is not
assessed by the AI on a live run. Its suggestion is dropped as `no_notes`, it is
not an "omitted" row (no re-run can give it an answer), and it is disclosed:

* the consultant Run-AI panel, from `no_notes_count` (None on an offline run);
* the consultant Enterprise table, per row, from CURRENT notes (`no_notes`);
* the client Playbook files, one sentence from the most recent LIVE run's
  PERSISTED accounting (`ai_runs.result`), never from current notes.

Live runs only: fixture mode still needs canned scores on blank rows.

Every test drives the real routes. A live run is made deterministic the way
`test_csf_run_ai.py` does it: the provider's `name` is patched to a live
provider's and the response is canned, so nothing is sent anywhere.
"""

from __future__ import annotations

import io
import os
import re
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from app.ai.llm import FixtureProvider, LLMClient
from tests._ai_runs import csf_run_ai, csf_scores_by_batch

pytestmark = pytest.mark.unit

KINDS = {"xlsx", "exec_pdf", "exec_docx", "full_pdf", "full_docx"}


@pytest.fixture()
def app_client(tmp_path) -> Iterator[tuple[TestClient, FixtureProvider, dict]]:
    url = f"sqlite:///{tmp_path / 'shield-csf1000.db'}"
    os.environ["DATABASE_URL"] = url
    api_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(api_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(api_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")
    engine = create_engine(url, future=True)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    from app.db.session import get_db
    from app.main import create_app
    from app.models.client import Client as _Client
    from app.models.client_domain import ClientDomain as _ClientDomain
    from app.routes.artifacts import _storage_dep
    from app.routes.csf import _llm_dep
    from app.storage.local import LocalFilesystemStorage

    def override_get_db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    provider = FixtureProvider()
    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_llm_dep] = lambda: LLMClient(provider)
    app.dependency_overrides[_storage_dep] = lambda: LocalFilesystemStorage(tmp_path / "storage")
    seed = TestSession()
    tenant = _Client(legal_name="Test Tenant")
    seed.add(tenant)
    seed.flush()
    seed.add(_ClientDomain(client_id=tenant.id, domain="example.com"))
    seed.commit()
    with TestClient(app, headers={"X-Client-Id": str(tenant.id)}) as c:
        r = c.post(
            "/auth/register",
            json={
                "email": "admin@example.com",
                "password": "correct horse battery staple!",
                "display_name": "A",
            },
        )
        assert r.status_code == 201, r.text
        yield c, provider, {"Authorization": f"Bearer {r.json()['tokens']['access_token']}"}


# --- helpers: they build the WORLD (a profile, notes, a typed score), never the
# --- outcome under test.


def _service(c: TestClient, h: dict) -> tuple[str, dict[str, str], list[str]]:
    """A CSF service seeded on the HIGH tier only. Returns (service id, answer
    id by code, the profile's codes as the profile route lists them)."""
    svc = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()["id"]
    a = c.post(f"/csf/services/{svc}/assessments", headers=h)
    assert a.status_code == 201, a.text
    answers = {ans["subcategory_code"]: ans["id"] for ans in a.json()["answers"]}
    seeded = c.post(f"/csf/services/{svc}/profiles/seed", headers=h, json={"tiers": ["high"]})
    assert seeded.status_code in (200, 201), seeded.text
    codes = sorted(r["subcategory_code"] for r in _profile(c, h, svc))
    assert len(codes) > 3, codes
    return svc, answers, codes


def _profile(c: TestClient, h: dict, svc: str) -> list[dict]:
    r = c.get(f"/csf/services/{svc}/profile/high", headers=h)
    assert r.status_code == 200, r.text
    return r.json()["rows"]


def _row_id(c: TestClient, h: dict, svc: str, code: str) -> str:
    return next(r["id"] for r in _profile(c, h, svc) if r["subcategory_code"] == code)


def _notes(c: TestClient, h: dict, answer_id: str, notes: str | None) -> None:
    r = c.patch(f"/csf/answers/{answer_id}", headers=h, json={"notes": notes})
    assert r.status_code == 200, r.text
    assert r.json()["notes"] == notes


def _type_score(c: TestClient, h: dict, svc: str, code: str, **fields: object) -> None:
    r = c.patch(f"/csf/dimension-scores/{_row_id(c, h, svc, code)}", headers=h, json=fields)
    assert r.status_code == 200, r.text


def _go_live(provider: FixtureProvider, monkeypatch) -> None:
    """The live-name idiom: the provider answers as `anthropic` would be
    named, so the run is LIVE end to end (the route's mode, the run's stored
    `mode`), while the response stays canned."""
    monkeypatch.setattr(type(provider), "name", "anthropic", raising=False)


def _run(c: TestClient, provider: FixtureProvider, h: dict, svc: str, codes, scores) -> dict:
    provider.register("csf_score", csf_scores_by_batch(scores, tiers=["high"], codes=codes))
    serves = "offline" if provider.name == "fixture" else "live"
    return csf_run_ai(c, svc, h, serves=serves)


def _row(body: dict, code: str) -> dict:
    return next(x for x in body["rows"] if x["subcategory_code"] == code and x["tier"] == "high")


def _no_notes_drops(body: dict) -> list[dict]:
    return [d for d in body["dropped"] if d["reason"] == "no_notes"]


def _assert_invariant(body: dict) -> None:
    accounted = body["suggestions_applied"] + sum(d["values"] for d in body["dropped"])
    assert body["suggestions_received"] == accounted, body


# --- the live run --------------------------------------------------------------


def test_a_live_run_writes_a_row_that_has_notes(app_client, monkeypatch) -> None:
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    noted = codes[0]
    _notes(c, h, answers[noted], "Okta SSO covers every workforce account.")
    _go_live(provider, monkeypatch)

    body = _run(
        c,
        provider,
        h,
        svc,
        codes,
        [{"tier": "high", "subcategory_code": noted, "governance": 2, "policy": 1}],
    )

    assert _row(body, noted)["governance"] == 2
    assert _row(body, noted)["policy"] == 1
    assert body["suggestions_applied"] == 2
    assert _no_notes_drops(body) == [], body["dropped"]
    # Every OTHER profile row has no notes, so none of them was assessed.
    assert body["no_notes_count"] == len(codes) - 1
    _assert_invariant(body)


def test_a_live_run_leaves_a_hand_typed_score_on_a_row_with_no_notes(
    app_client, monkeypatch
) -> None:
    c, provider, h = app_client
    svc, _answers, codes = _service(c, h)
    bare = codes[1]
    _type_score(c, h, svc, bare, governance=2)
    _go_live(provider, monkeypatch)

    body = _run(
        c,
        provider,
        h,
        svc,
        codes,
        [{"tier": "high", "subcategory_code": bare, "governance": 0, "what_we_found": "None."}],
    )

    assert _row(body, bare)["governance"] == 2, "a live run zeroed a row with no notes"
    assert _row(body, bare)["what_we_found"] is None
    (drop,) = _no_notes_drops(body)
    assert drop["key"] == f"high|{bare}", drop
    assert drop["values"] == 2, drop
    assert body["suggestions_applied"] == 0
    assert body["no_notes_count"] == len(codes)
    assert all(d["reason"] != "protected" for d in body["dropped"]), body["dropped"]
    _assert_invariant(body)


def test_blank_notes_count_as_no_notes(app_client, monkeypatch) -> None:
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    blank = codes[2]
    _notes(c, h, answers[blank], "  \n\t ")
    _type_score(c, h, svc, blank, governance=2)
    _go_live(provider, monkeypatch)

    body = _run(
        c, provider, h, svc, codes, [{"tier": "high", "subcategory_code": blank, "governance": 0}]
    )

    assert _row(body, blank)["governance"] == 2, "whitespace notes were treated as notes"
    (drop,) = _no_notes_drops(body)
    assert drop["key"] == f"high|{blank}", drop
    assert body["no_notes_count"] == len(codes)
    _assert_invariant(body)


def test_an_omitted_row_with_no_notes_is_counted_not_alerted(app_client, monkeypatch) -> None:
    """The model leaves out every row. The one row WITH notes is an omitted
    row (a re-run can answer it); the rows without notes are not -- no re-run
    can give them an answer -- so they are counted, not alerted."""
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    noted = codes[0]
    _notes(c, h, answers[noted], "Asset inventory is reconciled monthly.")
    _go_live(provider, monkeypatch)

    body = _run(c, provider, h, svc, codes, [])

    assert body["omitted_count"] == 1, body["omitted_rows"]
    assert body["omitted_rows"] == [{"tier": "high", "subcategory_code": noted}]
    assert body["no_notes_count"] == len(codes) - 1


def test_an_offline_run_has_no_no_notes_count(app_client) -> None:
    c, provider, h = app_client
    svc, _answers, codes = _service(c, h)
    bare = codes[1]

    body = _run(
        c, provider, h, svc, codes, [{"tier": "high", "subcategory_code": bare, "governance": 2}]
    )

    # Offline output still fills a blank row (the demo and e2e rely on it).
    assert _row(body, bare)["governance"] == 2
    assert _no_notes_drops(body) == [], body["dropped"]
    assert body["no_notes_count"] is None
    assert body["omitted_count"] == len(codes) - 1


def test_the_audit_row_carries_the_no_notes_count(app_client, monkeypatch) -> None:
    from app.models.audit_entry import AuditEntry

    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    _notes(c, h, answers[codes[0]], "Policies reviewed annually.")
    _go_live(provider, monkeypatch)
    body = _run(
        c,
        provider,
        h,
        svc,
        codes,
        [
            {"tier": "high", "subcategory_code": codes[0], "governance": 1},
            {"tier": "high", "subcategory_code": codes[1], "governance": 1},
        ],
    )
    assert body["no_notes_count"] == len(codes) - 1

    engine = create_engine(os.environ["DATABASE_URL"], future=True)
    with sessionmaker(bind=engine, future=True)() as s:
        entry = s.execute(select(AuditEntry).where(AuditEntry.action == "csf.run_ai")).scalar_one()
    assert entry.details["no_notes_count"] == len(codes) - 1
    assert entry.details["dropped_by_reason"] == {"no_notes": 1}
    assert entry.details["suggestions_applied"] == 1


# --- the consultant Enterprise table ----------------------------------------------


def test_the_enterprise_profile_flags_rows_with_no_notes(app_client) -> None:
    c, _provider, h = app_client
    svc, answers, codes = _service(c, h)
    noted, bare, blank = codes[0], codes[1], codes[2]
    _notes(c, h, answers[noted], "MFA enforced for administrators.")
    _notes(c, h, answers[blank], "   ")

    r = c.get(f"/csf/services/{svc}/enterprise-profile", headers=h)
    assert r.status_code == 200, r.text
    by_code = {s["subcategory_code"]: s for s in r.json()["subcategories"]}
    assert by_code[noted]["no_notes"] is False
    assert by_code[bare]["no_notes"] is True
    assert by_code[blank]["no_notes"] is True


# --- the client Playbook files ------------------------------------------------------


def _export(c: TestClient, h: dict, svc: str) -> dict[str, bytes]:
    ex = c.post(f"/csf/services/{svc}/playbook/export", headers=h)
    assert ex.status_code == 200, ex.text
    arts = {a["kind"]: a["artifact_id"] for a in ex.json()["artifacts"]}
    assert set(arts) == KINDS
    out = {}
    for kind, art_id in arts.items():
        dl = c.get(f"/artifacts/{art_id}/download", headers=h)
        assert dl.status_code == 200, f"{kind}: {dl.status_code}"
        out[kind] = dl.content
    return out


def _flat_text(kind: str, raw: bytes) -> str:
    """Every word of the artifact, whitespace collapsed (PDF text wraps)."""
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    if kind == "xlsx":
        wb = load_workbook(io.BytesIO(raw))
        text = " ".join(
            str(v)
            for ws in wb.worksheets
            for row in ws.iter_rows(values_only=True)
            for v in row
            if v
        )
    elif kind.endswith("pdf"):
        text = " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
    else:
        doc = Document(io.BytesIO(raw))
        cells = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        text = " ".join([p.text for p in doc.paragraphs] + cells)
    return re.sub(r"\s+", " ", text)


def _texts(files: dict[str, bytes]) -> dict[str, str]:
    return {kind: _flat_text(kind, raw) for kind, raw in files.items()}


def test_the_files_carry_no_sentence_when_no_live_run_exists(app_client) -> None:
    c, provider, h = app_client
    svc, _answers, codes = _service(c, h)
    # An OFFLINE run exists; it is not a live run.
    _run(c, provider, h, svc, codes, [{"tier": "high", "subcategory_code": codes[0], "policy": 1}])

    texts = _texts(_export(c, h, svc))

    for kind, text in texts.items():
        # Positive state first: the file rendered its cover.
        assert "Working profile version: 1" in text, kind
        assert "Not assessed by AI (no notes)" not in text, kind


def test_the_files_state_the_last_live_runs_count_in_every_artifact(
    app_client, monkeypatch
) -> None:
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    _notes(c, h, answers[codes[0]], "Backups tested quarterly.")
    _notes(c, h, answers[codes[1]], "Incident plan exercised in May.")
    # {total} is the IN-SCOPE subcategories the run covered: one taken out of
    # scope is not in it.
    _type_score(c, h, svc, codes[3], in_scope=False)
    total = len(codes) - 1
    n = total - 2
    _go_live(provider, monkeypatch)
    _run(c, provider, h, svc, codes, [{"tier": "high", "subcategory_code": codes[0], "policy": 1}])

    expected = (
        f"Not assessed by AI (no notes) in the last live AI run: {n} of {total} subcategories."
    )
    for kind, text in _texts(_export(c, h, svc)).items():
        assert expected in text, f"{kind}: {text[:600]}"


def test_the_files_read_the_persisted_count_not_current_notes(app_client, monkeypatch) -> None:
    """Notes typed AFTER the live run do not change what that run did, so they
    do not change the sentence. And a later OFFLINE run is not the last LIVE
    run, so it does not change it either."""
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    _go_live(provider, monkeypatch)
    _run(c, provider, h, svc, codes, [])
    _notes(c, h, answers[codes[0]], "Typed after the run.")
    monkeypatch.setattr(type(provider), "name", "fixture", raising=False)
    _run(c, provider, h, svc, codes, [{"tier": "high", "subcategory_code": codes[1], "policy": 1}])

    expected = (
        f"Not assessed by AI (no notes) in the last live AI run: {len(codes)} of "
        f"{len(codes)} subcategories."
    )
    for kind, text in _texts(_export(c, h, svc)).items():
        assert expected in text, f"{kind}: {text[:600]}"


def test_the_files_state_zero_when_every_row_had_notes(app_client, monkeypatch) -> None:
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    for code in codes:
        _notes(c, h, answers[code], f"Interview notes for {code}.")
    _go_live(provider, monkeypatch)
    body = _run(c, provider, h, svc, codes, [])
    assert body["no_notes_count"] == 0

    expected = (
        f"Not assessed by AI (no notes) in the last live AI run: 0 of {len(codes)} subcategories."
    )
    for kind, text in _texts(_export(c, h, svc)).items():
        assert expected in text, f"{kind}: {text[:600]}"


def test_the_files_read_the_latest_live_run(app_client, monkeypatch) -> None:
    """Two live runs: the files state the SECOND, which ran after a code got
    notes, not the first."""
    c, provider, h = app_client
    svc, answers, codes = _service(c, h)
    _go_live(provider, monkeypatch)
    first = _run(c, provider, h, svc, codes, [])
    assert first["no_notes_count"] == len(codes)
    _notes(c, h, answers[codes[0]], "Vendor risk reviews are on file.")
    second = _run(c, provider, h, svc, codes, [])
    assert second["no_notes_count"] == len(codes) - 1

    expected = (
        f"Not assessed by AI (no notes) in the last live AI run: {len(codes) - 1} of "
        f"{len(codes)} subcategories."
    )
    for kind, text in _texts(_export(c, h, svc)).items():
        assert expected in text, f"{kind}: {text[:600]}"
