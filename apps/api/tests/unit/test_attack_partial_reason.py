"""The client sees WHY a technique is Partial (#554, slice R1).

Gene's H3 comment on #554 (2026-09-24): "Partial" is unusable to a client
without a reason. The seven Partial reason codes existed (`coverage.py`) with
model-facing definitions; no client surface showed them. The client wording,
approved by Gene's advisor as written on 2026-10-02 (condition 6), is the R1
proposal on #554 (issue comment 5953613758) plus the two extra rows from the
final-wording table: "Reason not recorded" and "Sub-techniques differ".

The expected strings below are COPIED from that approved text, never imported
from the module under test.

Surfaces: the XLSX Coverage sheet ("Why partial", after "Pending review", and
a legend row), the PDF and DOCX "Partial coverage, by reason" count table, and
the client dashboard (`test_attack_partial_reason_dashboard.py` and the web
test). Both rule sets carry it (option (a)): the rule-1 goldens were amended.
"""

from __future__ import annotations

import io
import uuid

import pytest

from app.attack.analytics import compute
from app.attack.catalog import TECHNIQUES
from app.attack.exporters import build_context, render_docx, render_pdf, render_xlsx
from app.attack.pending import pending_codes
from app.attack.rules import parents_computed
from app.models.attack_assessment import (
    AttackAssessment,
    AttackAssessmentStatus,
    AttackCoverage,
)

pytestmark = pytest.mark.unit

#: The approved wording, code -> (label, sentence), copied byte for byte.
APPROVED = {
    "missing_control_category": (
        "A control type is missing",
        "Something already defends against this technique, but a whole category of "
        "control, such as prevention or response, is not in place.",
    ),
    "reach_limited": (
        "Not covered everywhere",
        "Defended on most of your environment, but not on some systems, such as another "
        "operating system, a cloud or SaaS service, or unmanaged or off-network devices.",
    ),
    "detection_weak": (
        "Detection is unreliable",
        "There is a signal for this activity, but it is noisy or approximate, or depends "
        "on custom detection rules that still need to be written and tuned.",
    ),
    "prevention_limited": (
        "Detected, not blocked",
        "This activity can be detected but not prevented, sometimes because legitimate "
        "work needs the same capability.",
    ),
    "evasive_variant_uncovered": (
        "Advanced variants not covered",
        "Common forms of this technique are covered; advanced forms, such as "
        "kernel-level, firmware, encrypted or novel variants, are not.",
    ),
    "recovery_absent": (
        "No recovery evidenced",
        "The attack would be detected, but no way to recover from it, such as backups "
        "or restore procedures, is evidenced.",
    ),
    "periodic_not_continuous": (
        "Checked periodically, not continuously",
        "This is found only when a scheduled scan runs, not as it happens.",
    ),
}
NOT_RECORDED = (
    "Reason not recorded",
    "This technique was assessed as partly covered, and no reason was recorded.",
)
SUB_TECHNIQUES_DIFFER = (
    "Sub-techniques differ",
    "Its sub-techniques are covered to different degrees; each sub-technique states "
    "its own reason.",
)
WHY_LEGEND = ("Why partial", "The reason a Partial technique is only partly covered.")


def _cell(pair: tuple[str, str]) -> str:
    return f"{pair[0]}: {pair[1]}"


# --- the wording table ------------------------------------------------------------


def test_every_partial_code_has_client_wording_and_nothing_else_does() -> None:
    """Derived from the vocabulary, not listed: a new Partial code without
    client wording goes red here, before a client reads its raw code."""
    from app.attack.coverage import REASON_CODES, CoverageStatus
    from app.attack.partial_reasons import CLIENT_WORDING

    partial_codes = {r.code for r in REASON_CODES if r.status == CoverageStatus.PARTIAL}
    assert set(CLIENT_WORDING) == partial_codes


def test_the_wording_is_the_approved_text() -> None:
    from app.attack.partial_reasons import (
        CLIENT_WORDING,
        REASON_NOT_RECORDED,
    )
    from app.attack.partial_reasons import (
        SUB_TECHNIQUES_DIFFER as DIFFER,
    )

    assert {k: (v.label, v.sentence) for k, v in CLIENT_WORDING.items()} == APPROVED
    assert (REASON_NOT_RECORDED.label, REASON_NOT_RECORDED.sentence) == NOT_RECORDED
    assert (DIFFER.label, DIFFER.sentence) == SUB_TECHNIQUES_DIFFER


def test_an_unknown_code_on_a_partial_row_fails_loudly() -> None:
    """Writes validate the code against the status, so this is a writer bug;
    rendering it as "Reason not recorded" would be false, so it is refused."""
    from app.attack.partial_reasons import partial_reason

    with pytest.raises(ValueError, match="platform_absent"):
        partial_reason("partial", "platform_absent", computed_parent=False)


def test_only_a_partial_row_has_a_reason() -> None:
    from app.attack.partial_reasons import partial_reason

    for status in ("covered", "gap", "not_applicable", None):
        assert partial_reason(status, None, computed_parent=False) is None


# --- one world, every renderer ------------------------------------------------------


def _standalone_codes() -> list[str]:
    """Derived from the catalog's parent links, not from `app.attack.parents`."""
    has_children = {t.parent_id for t in TECHNIQUES if t.parent_id is not None}
    return [t.id for t in TECHNIQUES if t.parent_id is None and t.id not in has_children]


def _family() -> tuple[str, list[str]]:
    kids: dict[str, list[str]] = {}
    for t in TECHNIQUES:
        if t.parent_id is not None:
            kids.setdefault(t.parent_id, []).append(t.id)
    parent = min((p for p in kids if len(kids[p]) >= 2), key=lambda p: (len(kids[p]), p))
    return parent, sorted(kids[parent])


def _world(parent_rules: int | None):
    """Seven standalone Partials, one per code; one reasonless Partial; one
    Covered row; one pending Partial; and a family whose parent is Partial
    because one child is covered and one is a gap."""
    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=1,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=parent_rules,
    )
    standalone = iter(_standalone_codes())
    rows: list[AttackCoverage] = []

    def add(code: str, status: str, reason: str | None = None, **kw) -> AttackCoverage:
        row = AttackCoverage(
            id=uuid.uuid4(),
            assessment_id=a.id,
            technique_code=code,
            status=status,
            reason_code=reason,
            detection_tools=kw.get("tools", ["Tool A"]),
            prevention_tools=[],
            response_tools=[],
            unconfirmed_citations=kw.get("unconfirmed", []),
        )
        rows.append(row)
        return row

    by_code = {}
    for code in APPROVED:
        by_code[code] = add(next(standalone), "partial", code)
    reasonless = add(next(standalone), "partial", None)
    covered = add(next(standalone), "covered")
    pending = add(
        next(standalone),
        "partial",
        "detection_weak",
        unconfirmed=[{"tool": "Tool A", "field": "detection_tools"}],
    )
    parent, (first, second, *rest) = _family()
    add(first, "covered")
    add(second, "gap", tools=[])
    for code in rest:
        add(code, "gap", tools=[])
    parent_row = add(parent, "partial", None, tools=[])

    rollup = compute(
        {c.technique_code: c.status for c in rows},
        pending_codes(rows, parents_computed=parents_computed(a)),
    )
    ctx = build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=rows,
        rollup=rollup,
    )
    return ctx, by_code, reasonless, covered, pending, parent_row


def _coverage_sheet(raw: bytes) -> tuple[list, dict[str, list]]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))["Coverage"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    return rows[0], {r[0]: r for r in rows[1:]}


def _summary_labels(raw: bytes) -> dict:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))["Heatmap Summary"]
    return {r[0].value: r[1].value for r in ws.iter_rows() if r and r[0].value}


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return " ".join(" ".join(parts).split())


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return " ".join(
        " ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages).split()
    )


@pytest.mark.parametrize("parent_rules", [2, 1], ids=["rule-2", "rule-1"])
def test_the_xlsx_says_why_each_partial_is_partial(parent_rules) -> None:
    ctx, by_code, reasonless, covered, pending, parent_row = _world(parent_rules)
    header, rows = _coverage_sheet(render_xlsx(ctx))

    # The column sits immediately after "Pending review".
    assert header.index("Why partial") == header.index("Pending review") + 1, header
    col = header.index("Why partial")

    for code, row in by_code.items():
        assert rows[row.technique_code][col] == _cell(APPROVED[code]), code
    assert rows[reasonless.technique_code][col] == _cell(NOT_RECORDED)
    assert rows[pending.technique_code][col] == _cell(APPROVED["detection_weak"])
    assert rows[covered.technique_code][col] in (None, ""), "a Covered row has no reason"
    # A computed parent's Partial is its children's, under #620's rules; an
    # assessment approved before #620 scored its parents directly.
    expected_parent = SUB_TECHNIQUES_DIFFER if parent_rules == 2 else NOT_RECORDED
    assert rows[parent_row.technique_code][col] == _cell(expected_parent)

    assert _summary_labels(render_xlsx(ctx)).get(WHY_LEGEND[0]) == WHY_LEGEND[1]


def _expected_table(parent_rules: int) -> list[tuple[tuple[str, str], int]]:
    """(wording, count) per row, in the table's order, from the world above:
    each code once (the pending Partial is withheld, so it does not count),
    then the parent, then the reasonless row."""
    rows = [(APPROVED[code], 1) for code in APPROVED]
    if parent_rules == 2:
        rows.append((SUB_TECHNIQUES_DIFFER, 1))
        rows.append((NOT_RECORDED, 1))
    else:
        rows.append((NOT_RECORDED, 2))
    return rows


@pytest.mark.parametrize("parent_rules", [2, 1], ids=["rule-2", "rule-1"])
def test_the_docx_counts_the_partials_by_reason(parent_rules) -> None:
    ctx, *_ = _world(parent_rules)
    text = _docx_text(render_docx(ctx))
    assert "Partial coverage, by reason" in text
    assert "Reason | What it means | Techniques" in text
    for (label, sentence), n in _expected_table(parent_rules):
        assert f"{label} | {sentence} | {n}" in text, (label, n)
    # The counts add up to the rollup's own Partial figure, never to a second
    # population: the pending Partial is withheld from both.
    assert sum(n for _w, n in _expected_table(parent_rules)) == ctx.rollup.partial


@pytest.mark.parametrize("parent_rules", [2, 1], ids=["rule-2", "rule-1"])
def test_the_pdf_counts_the_partials_by_reason(parent_rules) -> None:
    ctx, *_ = _world(parent_rules)
    text = _pdf_text(render_pdf(ctx))
    assert "Partial coverage, by reason" in text
    # The counts are pinned through the DOCX, which reads table cells
    # exactly; both renderers take them from `partial_reason_counts`.
    for (label, _sentence), _n in _expected_table(parent_rules):
        # test-integrity: the needle is the approved label, a literal in
        # APPROVED above; the table is the only place a label is printed.
        assert f"{label} " in text, label
    # The sentence wraps inside a PDF cell, so its words are checked rather
    # than its line breaks.
    for (_label, sentence), _n in _expected_table(parent_rules):
        # test-integrity: the needle is the approved sentence, a literal in
        # APPROVED above, whitespace-normalised the same way as the haystack.
        assert " ".join(sentence.split()) in text, sentence


def test_a_document_with_no_partial_has_no_reason_table() -> None:
    a = AttackAssessment(
        id=uuid.UUID(int=1), service_id=uuid.UUID(int=2), version=1, parent_rules=2
    )
    rows = [
        AttackCoverage(
            id=uuid.uuid4(),
            assessment_id=a.id,
            technique_code=_standalone_codes()[0],
            status="covered",
            detection_tools=["Tool A"],
            prevention_tools=[],
            response_tools=[],
            unconfirmed_citations=[],
        )
    ]
    ctx = build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=rows,
        rollup=compute({r.technique_code: r.status for r in rows}),
    )
    assert "Partial coverage, by reason" not in _docx_text(render_docx(ctx))
    assert "Partial coverage, by reason" not in _pdf_text(render_pdf(ctx))


def test_a_reason_table_that_does_not_add_up_is_refused() -> None:
    """The table must add up to the Partial figure printed beside it. A rollup
    built without the withheld set counts the pending Partial, which the table
    does not: two populations, so finalize refuses rather than print it."""
    from app.attack.exporters import partial_reason_counts

    ctx, *_ = _world(2)
    wrong = compute({c.technique_code: c.status for c in ctx.coverage})  # no withheld set
    assert wrong.partial == ctx.rollup.partial + 1, "the world must hold one pending Partial"
    bad = build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=ctx.assessment,
        coverage=ctx.coverage,
        rollup=wrong,
    )
    with pytest.raises(ValueError, match="Partial-by-reason table counts"):
        partial_reason_counts(bad)
