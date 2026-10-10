"""`other_axes` reaches the client: the dashboard API and all three files.

Approved on #736 6094620397 (Risk E item 4), for #474:

- the column heading is "Other axes";
- a cell is the comma-joined axis display names, e.g. "Prevention, Response";
- `[]` (the claim "no other axis") is an EMPTY cell;
- NULL ("not recorded") prints "Not recorded".

Each of the three states on each surface, read from what the client receives:
the client dashboard as a client user reads it, and the PUBLISHED files as
downloaded. Expected strings are written out from the ruling, never built from
the renderer's own constants. The admin screen and the client dashboard's
rendering are the web half (`otherAxes.test.ts` and the two component tests).
"""

from __future__ import annotations

import io
import json

import pytest

from tests.unit.test_risk_publish import _client_dashboard, _post, _world
from tests.unit.test_risk_register import _pdf_text, app_client  # noqa: F401  (fixture)

pytestmark = pytest.mark.unit

LISTED = "Alpha risk"
EMPTY = "Bravo risk"
UNRECORDED = "Charlie risk"


def _entry(title: str, **over: object) -> str:
    e: dict[str, object] = {
        "title": title,
        "description": "d",
        "axis": "detection",
        "source": "coverage_finding",
        "likelihood": "high",
        "impact": "catastrophic",
        "recommended_action": "remediate",
        "rationale": "r",
    }
    e.update(over)
    return json.dumps(e)


def _published(app_client):  # noqa: F811
    c, _, _, cid, ah, ch, _ = _world(
        app_client,
        _entry(LISTED, other_axes=["prevention", "response"]),
        _entry(EMPTY, other_axes=[]),
        _entry(UNRECORDED),  # no `other_axes` at all: stored NULL
    )
    pub = _post(c, ah, cid, "publish")
    assert pub.status_code == 200, pub.text
    files = {}
    for kind in ("xlsx", "pdf", "docx"):
        d = c.get(
            f"/artifacts/{pub.json()[f'{kind}_artifact_id']}/download",
            headers={**ah, "X-Client-Id": cid},
        )
        assert d.status_code == 200, d.text
        files[kind] = d.content
    return c, ch, cid, files


def test_the_client_dashboard_sends_each_state(app_client) -> None:  # noqa: F811
    c, ch, cid, _ = _published(app_client)
    r = _client_dashboard(c, ch, cid)
    assert r.status_code == 200, r.text
    by_title = {e["title"]: e for e in r.json()["entries"]}
    assert by_title[LISTED]["other_axes"] == ["prevention", "response"]
    assert by_title[EMPTY]["other_axes"] == []
    assert by_title[UNRECORDED]["other_axes"] is None


def test_the_xlsx_has_an_other_axes_column_in_each_state(app_client) -> None:  # noqa: F811
    from openpyxl import load_workbook

    _, _, _, files = _published(app_client)
    rows = list(
        load_workbook(io.BytesIO(files["xlsx"]))["Risk Register"].iter_rows(values_only=True)
    )
    header = rows[0]
    assert "Other axes" in header
    by_title = {r["Weakness"]: r for r in (dict(zip(header, row, strict=True)) for row in rows[1:])}
    assert by_title[LISTED]["Other axes"] == "Prevention, Response"
    # An empty cell: openpyxl reads an empty string back as no value, and
    # "Not recorded" in this cell would be the defect.
    assert by_title[EMPTY]["Other axes"] is None
    assert by_title[UNRECORDED]["Other axes"] == "Not recorded"


def _docx_register(raw: bytes) -> dict[str, dict[str, str]]:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    (table,) = [t for t in doc.tables if t.rows[0].cells[1].text == "Weakness"]
    header = [cell.text for cell in table.rows[0].cells]
    rows = [dict(zip(header, (cell.text for cell in r.cells), strict=True)) for r in table.rows[1:]]
    return {r["Weakness"]: r for r in rows}


def test_the_docx_has_an_other_axes_column_in_each_state(app_client) -> None:  # noqa: F811
    _, _, _, files = _published(app_client)
    rows = _docx_register(files["docx"])
    assert rows[LISTED]["Other axes"] == "Prevention, Response"
    assert rows[EMPTY]["Other axes"] == ""
    assert rows[UNRECORDED]["Other axes"] == "Not recorded"


def test_the_pdf_has_an_other_axes_column_in_each_state(app_client) -> None:  # noqa: F811
    _, _, _, files = _published(app_client)
    text = " ".join(_pdf_text(files["pdf"]).split())
    assert "Other axes" in text
    # Each register row reads left to right: title, axis, other axes, L x I.
    assert f"{LISTED} Detection Prevention, Response High x Catastrophic" in text
    assert f"{EMPTY} Detection High x Catastrophic" in text
    assert f"{UNRECORDED} Detection Not recorded High x Catastrophic" in text
