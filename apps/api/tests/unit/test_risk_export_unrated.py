"""#844: the Risk Register deliverable states how many entries are unrated.

The defect: the PDF and Word summaries printed `Total entries` and
`Critical + High: {n}` with no untiered count, and the matrix skips any entry
without both halves of a rating, so a client could read "Critical + High: 0"
and an empty matrix over entries nobody rated. G3 option (c) on #806: export
discloses unrated entries.

Through the export endpoint and the downloaded bytes, because the client reads
the file, not the renderer. Unrated entries are produced the way production
produces them: the model returns JSON null (the approved prompt's
INSUFFICIENT INFORMATION rule), or a consultant clears a rating.
"""

from __future__ import annotations

import io

import pytest

from tests.unit.test_risk_register import (  # noqa: F401  (app_client is a fixture)
    _admin,
    _entries_payload,
    _entry,
    _generate,
    _pdf_text,
    _seed_attack_and_zt,
    app_client,
)

pytestmark = pytest.mark.unit


def _unrated(title: str, **extra: str) -> str:
    fields = ", ".join(f'"{k}": "{v}"' for k, v in extra.items())
    return (
        '{"title": "'
        + title
        + '", "likelihood": null, "impact": null'
        + (", " + fields if fields else "")
        + "}"
    )


def _export(app_client, *entries: str, edit: tuple[int, dict] | None = None) -> dict:  # noqa: F811
    """Generate, optionally PATCH one entry, export, and download all three."""
    c, provider = app_client
    bearer, cid = _admin(c)
    _seed_attack_and_zt(c, bearer, cid)
    body = _generate(c, provider, bearer, cid, _entries_payload(*entries))
    bh = {"Authorization": f"Bearer {bearer}"}
    if edit is not None:
        idx, change = edit
        r = c.patch(
            f"/risk/clients/{cid}/register/entries/{body['entries'][idx]['id']}",
            headers=bh,
            json=change,
        )
        assert r.status_code == 200, r.text
    r = c.post(f"/risk/clients/{cid}/register/export", headers=bh)
    assert r.status_code == 200, r.text
    ex = r.json()
    dh = {**bh, "X-Client-Id": cid}
    files = {}
    for kind in ("xlsx", "pdf", "docx"):
        d = c.get(f"/artifacts/{ex[f'{kind}_artifact_id']}/download", headers=dh)
        assert d.status_code == 200, d.text
        files[kind] = d.content
    return files


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _xlsx_register_rows(raw: bytes) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))["Risk Register"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, r, strict=True)) for r in rows[1:]]


def _flat(text: str) -> str:
    """PDF extraction wraps lines; compare on single-spaced text."""
    return " ".join(text.split())


UNRATED_ONE_OF_TWO = (
    "Not rated: 1 of 2 entries has no likelihood or impact, so it has no tier "
    "and is not in the matrix or in Critical + High."
)


def test_the_pdf_and_word_summaries_state_the_unrated_count(app_client) -> None:  # noqa: F811
    files = _export(app_client, _entry("Rated"), _unrated("Unrated"))
    assert UNRATED_ONE_OF_TWO in _flat(_pdf_text(files["pdf"]))
    assert UNRATED_ONE_OF_TWO in _docx_text(files["docx"])


def test_the_plural_reads_in_number(app_client) -> None:  # noqa: F811
    files = _export(app_client, _unrated("A"), _unrated("B"), _entry("C"))
    expected = (
        "Not rated: 2 of 3 entries have no likelihood or impact, so they have no "
        "tier and are not in the matrix or in Critical + High."
    )
    assert expected in _flat(_pdf_text(files["pdf"]))
    assert expected in _docx_text(files["docx"])


def test_a_fully_rated_register_says_so_rather_than_staying_silent(
    app_client,  # noqa: F811
) -> None:
    """Three values, not two: silence would read the same as "nobody counted"."""
    files = _export(app_client, _entry("Rated"))
    pdf = _flat(_pdf_text(files["pdf"]))
    assert "Every entry is rated." in pdf
    assert "Not rated:" not in pdf
    assert "Every entry is rated." in _docx_text(files["docx"])


def test_an_unrated_row_prints_not_rated_rather_than_blank(app_client) -> None:  # noqa: F811
    files = _export(app_client, _entry("Rated"), _unrated("Unrated"))
    rows = {r["Weakness"]: r for r in _xlsx_register_rows(files["xlsx"])}
    assert (rows["Unrated"]["Likelihood"], rows["Unrated"]["Impact"], rows["Unrated"]["Tier"]) == (
        "Not rated",
        "Not rated",
        "Not rated",
    )
    # And a rated row is untouched, so the marker cannot be printed everywhere.
    assert (rows["Rated"]["Likelihood"], rows["Rated"]["Tier"]) == ("High", "Critical")
    docx = _docx_text(files["docx"])
    assert "Not rated x Not rated" in docx


def test_the_pdf_matrix_names_what_it_leaves_out(app_client) -> None:  # noqa: F811
    files = _export(app_client, _entry("Rated"), _unrated("Unrated"))
    assert "1 unrated entry is not in this matrix." in _flat(_pdf_text(files["pdf"]))


def test_a_consultant_cleared_rating_is_counted_as_unrated(app_client) -> None:  # noqa: F811
    files = _export(app_client, _entry("A"), _entry("B"), edit=(0, {"impact": None}))
    assert UNRATED_ONE_OF_TWO in _flat(_pdf_text(files["pdf"]))


def test_a_consultant_set_rating_is_not_credited_to_the_model(app_client) -> None:  # noqa: F811
    files = _export(
        app_client,
        _unrated("Unrated"),
        _entry("Model rated"),
        edit=(0, {"likelihood": "low", "impact": "minor"}),
    )
    rows = {r["Weakness"]: r for r in _xlsx_register_rows(files["xlsx"])}
    assert rows["Unrated"]["Origin"] == "ai_generated; rating set by consultant"
    assert rows["Model rated"]["Origin"] == "ai_generated"
    expected = "Ratings set by a consultant: 1 of 2 entries."
    assert expected in _flat(_pdf_text(files["pdf"]))
    assert expected in _docx_text(files["docx"])


def test_entries_missing_from_the_axis_and_action_lines_are_counted(
    app_client,  # noqa: F811
) -> None:
    """#313's twin in the deliverable: the axis and action lines filter
    independently of the tier, and the client dashboard already says so."""
    files = _export(
        app_client,
        _entry("No axis", axis="mitigation"),
        _entry("No action", recommended_action="monitor"),
        _entry("Complete"),
    )
    for text in (_flat(_pdf_text(files["pdf"])), _docx_text(files["docx"])):
        assert "1 of 3 entries has no axis and is not in the line above." in text
        assert "1 of 3 entries has no recommended action and is not in the line above." in text


def test_complete_axis_and_action_lines_carry_no_note(app_client) -> None:  # noqa: F811
    files = _export(app_client, _entry("Complete"))
    pdf = _flat(_pdf_text(files["pdf"]))
    assert "has no axis" not in pdf
    assert "have no axis" not in pdf
    assert "no recommended action" not in pdf
