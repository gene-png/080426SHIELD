"""#1427: the "≥" is ON the total-cost value when the spend is a floor.

#1094 put it on the savings value; the total cost carried it only in a label
beside the number ("Annual cost (may not be complete)"), so a copied cell or a
quoted figure read as exact. Ruled on #736 (6229792065): the total cost takes
"≥" on its value, in the XLSX cell (a number format, so the cell stays numeric),
the PDF and the DOCX, and the dashboard's "Annual license spend" value.

THREE values, not two (the CHECK): the dashboard has always reported
`spend_completeness` as complete / partial / unknown, and only "partial" is a
floor -- "unknown" (the upload was never reconciled) is deliberately NOT worded
as a floor, because that would assert a defect nobody observed. The files take
the marker from the SAME function the dashboard calls, `spend_completeness`, and
this table is written from the rule, not read back from the code.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.models import CapabilityList
from app.tech_debt.exporters import DeliverableContext, render_docx, render_pdf, render_xlsx
from app.tech_debt.reconcile import spend_completeness
from tests.unit.test_tech_debt_bundle_components import _docx_text, _pdf_text

pytestmark = pytest.mark.unit

FLOOR_FORMAT = '"≥ "$#,##0'
EXACT_FORMAT = "$#,##0"
TOTAL = 123456.0

#: (id, source_rows_total, included, excluded, spend_cost_known, exclusion_state, expected)
#: Each row is a world the extraction writer can produce.
STATES = [
    ("never-reconciled", None, 3, 0, True, "not_recorded", "unknown"),
    ("never-reconciled-and-uncosted", None, 3, 0, False, "not_recorded", "unknown"),
    ("everything-accounted-for", 3, 3, 0, True, "exact", "complete"),
    ("rows-excluded", 5, 3, 2, True, "exact", "partial"),
    ("an-included-tool-has-no-cost", 3, 3, 0, False, "exact", "partial"),
    ("exclusion-count-unknown", 3, 3, 0, True, "unknown", "partial"),
    ("more-items-than-rows", 2, 3, 0, True, "exact", "partial"),
]
IDS = [s[0] for s in STATES]


def _ctx(state) -> DeliverableContext:
    _id, total_rows, included, excluded, cost_known, exclusion, _expected = state
    return DeliverableContext(
        client_legal_name="Example Co",
        service_title="Technical Debt Review",
        cap_list=CapabilityList(version=1),
        items=[],
        total_cost=TOTAL,
        estimated_savings=0.0,
        savings_cost_known=True,
        source_rows_total=total_rows,
        included_count=included,
        excluded_count=excluded,
        spend_cost_known=cost_known,
        exclusion_state=exclusion,
    )


@pytest.mark.parametrize("state", STATES, ids=IDS)
def test_the_completeness_rule_has_three_values(state) -> None:
    _id, total_rows, included, excluded, cost_known, exclusion, expected = state
    got = spend_completeness(
        source_rows_total=total_rows,
        included_count=included,
        excluded_count=excluded,
        spend_cost_known=cost_known,
        exclusion_state=exclusion,
    )
    assert got == expected


def _total_cell(raw: bytes):
    ws = load_workbook(io.BytesIO(raw)).worksheets[0]
    for row in ws.iter_rows():
        if row[0].value in (
            "Total annual cost",
            "Included annual cost",
            "Annual cost (may not be complete)",
        ):
            return row[4]
    raise AssertionError("no total-cost row in the workbook")


@pytest.mark.parametrize("state", STATES, ids=IDS)
def test_the_xlsx_total_is_numeric_and_carries_the_marker_only_on_a_floor(state) -> None:
    cell = _total_cell(render_xlsx(_ctx(state)))
    assert cell.value == TOTAL, cell.value
    assert isinstance(cell.value, (int, float)), "the cell must stay a number so sums work"
    want = FLOOR_FORMAT if state[-1] == "partial" else EXACT_FORMAT
    assert cell.number_format == want, cell.number_format


@pytest.mark.parametrize("state", STATES, ids=IDS)
def test_the_pdf_and_docx_total_carries_the_marker_only_on_a_floor(state) -> None:
    floor = state[-1] == "partial"
    for text in (_pdf_text(render_pdf(_ctx(state))), _docx_text(render_docx(_ctx(state)))):
        assert "$123,456" in text, text
        assert ("≥ $123,456" in text) is floor, text


def test_the_savings_marker_is_not_moved_by_the_total_cost_one() -> None:
    # Savings are exact in every state here, so the only "≥" in a complete
    # document is none at all.
    complete = next(s for s in STATES if s[-1] == "complete")
    assert "≥" not in _pdf_text(render_pdf(_ctx(complete)))
    assert "≥" not in _docx_text(render_docx(_ctx(complete)))
