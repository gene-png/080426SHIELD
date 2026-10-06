"""A computed parent's Partial reads "Set by its sub-techniques", whatever its children (#798 review).

`parents.computed_parent_status` returns Partial whenever the ASSESSED children
are not all Covered and not all Gap. That includes every child Partial, and one
Partial child beside N/A siblings. The first wording described the children
as mixed, which is false of both. The worlds below are the two the first test
world did not build (it built covered + gap only).

The advisor's ruling (#736, 2026-10-02 18:34Z, under Gene's delegation for
condition 6): a computed parent's Partial reads "Set by its sub-techniques",
whatever its children are -- true of every shape `computed_parent_status`
returns Partial for. The strings below are that ruling's text, copied.
"""

from __future__ import annotations

import uuid

import pytest

from app.attack.analytics import compute
from app.attack.catalog import TECHNIQUES
from app.attack.exporters import build_context, render_xlsx
from app.attack.pending import pending_codes
from app.attack.rules import parents_computed
from app.models.attack_assessment import (
    AttackAssessment,
    AttackAssessmentStatus,
    AttackCoverage,
)
from tests.unit.test_attack_catalog_version_guard import (  # noqa: F401  (fixture)
    _auth,
    _register,
    _service_and_assessment,
    env,
)

pytestmark = pytest.mark.unit

SET_BY = (
    "Set by its sub-techniques",
    "This technique's coverage is computed from its sub-techniques; see each "
    "sub-technique for its own coverage and reason.",
)

#: (children's statuses and reasons) per case: every child Partial; one
#: Partial child beside N/A siblings.
CASES = {
    "all-partial": lambda n: [("partial", "reach_limited")] * n,
    "partial-beside-na": lambda n: [("partial", "detection_weak")]
    + [("not_applicable", "platform_absent")] * (n - 1),
}


def _family() -> tuple[str, list[str]]:
    """Derived from the catalog's parent links, not from `app.attack.parents`."""
    kids: dict[str, list[str]] = {}
    for t in TECHNIQUES:
        if t.parent_id is not None:
            kids.setdefault(t.parent_id, []).append(t.id)
    parent = min((p for p in kids if len(kids[p]) >= 2), key=lambda p: (len(kids[p]), p))
    return parent, sorted(kids[parent])


def _xlsx_parent_cell(case: str) -> str:
    from io import BytesIO

    from openpyxl import load_workbook

    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=1,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=2,
        # #554 R3: models an approved/released pre-R3 assessment, backfilled by 0059.
        status_rules=1,
    )
    parent, children = _family()
    rows = []
    for code, (status, reason) in zip(children, CASES[case](len(children)), strict=True):
        rows.append(
            AttackCoverage(
                id=uuid.uuid4(),
                assessment_id=a.id,
                technique_code=code,
                status=status,
                reason_code=reason,
                detection_tools=["Tool A"] if status == "partial" else [],
                prevention_tools=[],
                response_tools=[],
                unconfirmed_citations=[],
            )
        )
    rows.append(
        AttackCoverage(
            id=uuid.uuid4(),
            assessment_id=a.id,
            technique_code=parent,
            status="partial",  # what `computed_parent_status` returns for both cases
            reason_code=None,
            detection_tools=[],
            prevention_tools=[],
            response_tools=[],
            unconfirmed_citations=[],
        )
    )
    rollup = compute(
        {c.technique_code: c.status for c in rows},
        pending_codes(rows, parents_computed=parents_computed(a)),
    )
    ctx = build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=rows,
        rollup=rollup,
    )
    ws = load_workbook(BytesIO(render_xlsx(ctx)))["Coverage"]
    header = [c.value for c in ws[1]]
    col = header.index("Why partial")
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] == parent:
            return r[col] or ""
    raise AssertionError(f"no Coverage row for {parent}")


@pytest.mark.parametrize("case", list(CASES))
def test_the_xlsx_does_not_say_the_children_differ_when_they_do_not(case) -> None:
    cell = _xlsx_parent_cell(case)
    assert cell == f"{SET_BY[0]}: {SET_BY[1]}", (case, cell)


@pytest.mark.parametrize("case", list(CASES))
def test_the_dashboard_does_not_say_the_children_differ_when_they_do_not(  # noqa: F811
    env, case  # noqa: F811
) -> None:
    c, _Sess = env
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    svc, a = _service_and_assessment(c, bearer)
    by_code = {row["technique_code"]: row for row in a["coverage"]}
    parent, children = _family()
    for code, (status, reason) in zip(children, CASES[case](len(children)), strict=True):
        r = c.patch(
            f"/attack/coverage/{by_code[code]['id']}",
            headers=_auth(bearer),
            json={
                "status": status,
                "reason_code": reason,
                "detection_tools": ["Tool A"] if status == "partial" else [],
            },
        )
        assert r.status_code == 200, r.text
    assert (
        c.post(f"/attack/assessments/{a['id']}/approve", headers=_auth(bearer)).status_code == 200
    )
    fin = c.post(f"/attack/services/{svc}/deliverables/finalize", headers=_auth(bearer))
    assert fin.status_code in (200, 201), fin.text
    rel = c.post(f"/attack/deliverables/{fin.json()['id']}/release", headers=_auth(bearer))
    assert rel.status_code == 200, rel.text
    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    dash = c.get(
        f"/clients/{client_id}/attack/{svc}/dashboard",
        headers=_auth(client["tokens"]["access_token"]),
    ).json()
    techs = {t["code"]: t for t in dash["techniques"]}
    assert techs[parent]["status"] == "partial", (case, techs[parent])
    reason = techs[parent].get("partial_reason")
    assert reason == {"label": SET_BY[0], "sentence": SET_BY[1]}, (case, reason)
