"""The shared text extractor, `app/documents/text.py` (N2, for #1071).

One reader set, one set of limits and one "not read" vocabulary, shared by
EV-1's evidence-link job (P4) and EV-2's client summary (design note rev 4,
section 15.2).

Every fixture file is GENERATED here from known text and then extracted, and
the expected text is the literal the generator was given -- never a value read
back from the module under test (#72). The corpus has a row per CLASS (section
15.2's eight), including the ones a hand-written corpus never contains:
truncated and corrupt files, encrypted PDFs, zero-byte files, non-UTF-8 text, a
file whose extension lies, and a file large enough that it must chunk.

The property every class leans on: "I could not read it" and "it was empty" are
different results. A file that cannot be read is `not_read` with a typed reason,
never `read` with empty text (CLAUDE.md, the silent-success rule).
"""

from __future__ import annotations

import io
import uuid
import zipfile

import pytest

from app.config import Settings
from app.documents.text import (
    NOT_READ_REASONS,
    ChunkPlan,
    DocText,
    ExtractLimits,
    chunk_text,
    document_text,
    extract_limits,
)
from app.models.artifact import Artifact
from app.storage.local import LocalFilesystemStorage

pytestmark = pytest.mark.unit

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TXT = "text/plain"
OCTET = "application/octet-stream"

#: Generous limits, so only the class under test decides the result.
LIMITS = ExtractLimits(
    max_bytes=5_000_000,
    max_unzipped_bytes=20_000_000,
    max_chars=1_000_000,
    chunk_chars=4_000,
    max_chunks=10,
)


# --------------------------------------------------------------------------
# Generators: known text in, file bytes out.
# --------------------------------------------------------------------------


def _pdf(pages: list[str | None]) -> bytes:
    """One PDF page per entry. `None` draws a rectangle and no text: a page
    with no text layer, which is what a scan looks like to a text extractor."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages:
        if text is None:
            c.rect(72, 72, 200, 200, fill=1)
        else:
            c.drawString(72, 720, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def _encrypted_pdf(pages: list[str], *, user_password: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_pdf(pages))))
    writer.encrypt(user_password=user_password, owner_password="owner-pw", algorithm="AES-256")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _docx(
    *,
    paragraphs: list[str],
    table: list[list[str]] | None = None,
    header: str | None = None,
    footer: str | None = None,
) -> bytes:
    from docx import Document

    doc = Document()
    if header is not None:
        doc.sections[0].header.paragraphs[0].text = header
    if footer is not None:
        doc.sections[0].footer.paragraphs[0].text = footer
    for p in paragraphs:
        doc.add_paragraph(p)
    if table is not None:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _png() -> bytes:
    # The eight-byte PNG signature and an IHDR chunk header: enough to be typed.
    return b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + b"\x00" * 17


class _Store:
    """A real LocalFilesystemStorage, plus the Artifact row pointing at it."""

    def __init__(self, tmp_path) -> None:
        self.storage = LocalFilesystemStorage(tmp_path / "storage")

    def put(self, data: bytes, *, title: str, mime: str) -> Artifact:
        key = f"test/{uuid.uuid4()}"
        self.storage.put(key, data, content_type=mime)
        return Artifact(
            id=uuid.uuid4(),
            client_id=uuid.uuid4(),
            title=title,
            file_storage_key=key,
            mime_type=mime,
            size_bytes=len(data),
            sha256="0" * 64,
            uploaded_by=uuid.uuid4(),
        )

    def read(self, data: bytes, *, title: str, mime: str, limits: ExtractLimits = LIMITS):
        return document_text(self.storage, self.put(data, title=title, mime=mime), limits=limits)


@pytest.fixture()
def store(tmp_path) -> _Store:
    return _Store(tmp_path)


def _assert_not_read(result: DocText, reason: str) -> None:
    """The not-read shape, in full: the typed reason, and no text that could be
    mistaken for an empty document."""
    assert result.status == "not_read", result
    assert result.reason == reason, result
    assert result.text == ""
    assert result.chars_total == 0
    assert result.truncated is False


def _assert_read(result: DocText) -> None:
    assert result.status == "read", result
    assert result.reason is None, result


# --------------------------------------------------------------------------
# Class 1: one happy file per type.
# --------------------------------------------------------------------------


def test_pdf_happy_reads_every_page_in_order(store):
    result = store.read(
        _pdf(["Alpha incident response plan", "Bravo access review record"]),
        title="plan.pdf",
        mime=PDF,
    )
    _assert_read(result)
    assert "Alpha incident response plan" in result.text
    assert "Bravo access review record" in result.text
    assert result.text.index("Alpha incident") < result.text.index("Bravo access")
    assert result.pages_without_text == 0
    assert result.truncated is False
    assert result.chars_total == len(result.text)


def test_docx_happy(store):
    result = store.read(
        _docx(paragraphs=["Charlie backup policy", "Delta restore test"]),
        title="policy.docx",
        mime=DOCX,
    )
    _assert_read(result)
    assert "Charlie backup policy\nDelta restore test" in result.text


def test_xlsx_happy(store):
    result = store.read(
        _xlsx({"Inventory": [["Asset", "Owner"], ["Echo firewall", "Foxtrot team"]]}),
        title="assets.xlsx",
        mime=XLSX,
    )
    _assert_read(result)
    assert "Asset\tOwner" in result.text
    assert "Echo firewall\tFoxtrot team" in result.text


def test_txt_happy_is_exact(store):
    result = store.read(
        b"Golf MFA is enforced.\nHotel logs are kept.\n",
        title="notes.txt",
        mime=TXT,
    )
    _assert_read(result)
    assert result.text == "Golf MFA is enforced.\nHotel logs are kept.\n"
    assert result.chars_total == 43


def test_type_is_decided_by_magic_when_nothing_is_declared(store):
    """octet-stream and no extension declare nothing; the bytes decide."""
    result = store.read(_pdf(["India undeclared pdf"]), title="upload", mime=OCTET)
    _assert_read(result)
    assert "India undeclared pdf" in result.text


# --------------------------------------------------------------------------
# Class 2: a scanned PDF with no text layer; encrypted PDFs.
# --------------------------------------------------------------------------


def test_scanned_pdf_is_no_text_layer_not_an_empty_read(store):
    result = store.read(_pdf([None, None]), title="scan.pdf", mime=PDF)
    _assert_not_read(result, "no_text_layer")


def test_partly_scanned_pdf_reads_and_discloses_the_pages_without_text(store):
    """One text page beside two scanned ones is `read` -- and says two pages
    gave no text, so a reader never takes the text for the whole document."""
    result = store.read(_pdf(["Juliet signed page", None, None]), title="mixed.pdf", mime=PDF)
    _assert_read(result)
    assert "Juliet signed page" in result.text
    assert result.pages_without_text == 2


def test_password_protected_pdf_is_encrypted(store):
    data = _encrypted_pdf(["Kilo secret text"], user_password="user-pw")
    _assert_not_read(store.read(data, title="locked.pdf", mime=PDF), "encrypted")


def test_owner_password_only_pdf_opens_and_reads(store):
    """An empty user password is how most "protected" PDFs ship: they open in
    any viewer. Refusing them as encrypted would refuse readable evidence."""
    data = _encrypted_pdf(["Lima permissions only"], user_password="")
    result = store.read(data, title="perm.pdf", mime=PDF)
    _assert_read(result)
    assert "Lima permissions only" in result.text


# --------------------------------------------------------------------------
# Class 3: corrupt files of each type; zero-byte files; limits.
# --------------------------------------------------------------------------


@pytest.mark.parametrize("fraction", [0.3, 0.6, 0.9])
def test_truncated_pdf_is_unreadable(store, fraction):
    data = _pdf(["Mike page one", "November page two"])
    cut = data[: int(len(data) * fraction)]
    _assert_not_read(store.read(cut, title="cut.pdf", mime=PDF), "unreadable")


def test_pdf_header_then_garbage_is_unreadable(store):
    data = b"%PDF-1.4\n" + b"\x00garbage\xff" * 64
    _assert_not_read(store.read(data, title="bad.pdf", mime=PDF), "unreadable")


def test_truncated_docx_is_unreadable(store):
    data = _docx(paragraphs=["Oscar text"])
    _assert_not_read(store.read(data[: len(data) // 2], title="cut.docx", mime=DOCX), "unreadable")


def test_xlsx_with_a_corrupt_workbook_part_is_unreadable(store):
    """A well-formed zip whose workbook XML is not XML: the zip opens, the
    reader cannot."""
    src = zipfile.ZipFile(io.BytesIO(_xlsx({"S": [["Papa"]]})))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/workbook.xml":
                body = b"<workbook><<not xml"
            dst.writestr(info, body)
    _assert_not_read(store.read(out.getvalue(), title="bad.xlsx", mime=XLSX), "unreadable")


def test_zero_byte_txt_is_read_and_empty(store):
    """The other half of the distinction: a file that was read and holds
    nothing is `read` with zero characters, not a reason."""
    result = store.read(b"", title="empty.txt", mime=TXT)
    _assert_read(result)
    assert result.text == ""
    assert result.chars_total == 0


@pytest.mark.parametrize(
    ("title", "mime"),
    [("empty.pdf", PDF), ("empty.docx", DOCX), ("empty.xlsx", XLSX)],
)
def test_zero_byte_binary_types_are_unreadable(store, title, mime):
    """A zero-byte .pdf is not an empty PDF; it is not a PDF at all."""
    _assert_not_read(store.read(b"", title=title, mime=mime), "unreadable")


def test_text_over_max_chars_is_truncated_and_counts_the_whole(store):
    limits = ExtractLimits(
        max_bytes=LIMITS.max_bytes,
        max_unzipped_bytes=LIMITS.max_unzipped_bytes,
        max_chars=10,
        chunk_chars=LIMITS.chunk_chars,
        max_chunks=LIMITS.max_chunks,
    )
    result = store.read(b"Quebec Romeo Sierra", title="long.txt", mime=TXT, limits=limits)
    _assert_read(result)
    assert result.truncated is True
    assert result.text == "Quebec Rom"
    assert result.chars_total == 19


def test_file_over_max_bytes_is_too_large(store):
    limits = ExtractLimits(
        max_bytes=100,
        max_unzipped_bytes=LIMITS.max_unzipped_bytes,
        max_chars=LIMITS.max_chars,
        chunk_chars=LIMITS.chunk_chars,
        max_chunks=LIMITS.max_chunks,
    )
    _assert_not_read(
        store.read(b"x" * 101, title="big.txt", mime=TXT, limits=limits),
        "too_large",
    )


def test_file_at_exactly_max_bytes_is_read(store):
    limits = ExtractLimits(
        max_bytes=100,
        max_unzipped_bytes=LIMITS.max_unzipped_bytes,
        max_chars=LIMITS.max_chars,
        chunk_chars=LIMITS.chunk_chars,
        max_chunks=LIMITS.max_chunks,
    )
    _assert_read(store.read(b"y" * 100, title="edge.txt", mime=TXT, limits=limits))


def test_zip_that_inflates_past_the_limit_is_too_large(store):
    """A small DOCX whose parts inflate past `max_unzipped_bytes`: refused
    before any XML parser sees it."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Tango " * 50_000)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    limits = ExtractLimits(
        max_bytes=LIMITS.max_bytes,
        max_unzipped_bytes=100_000,
        max_chars=LIMITS.max_chars,
        chunk_chars=LIMITS.chunk_chars,
        max_chunks=LIMITS.max_chunks,
    )
    assert len(data) < 100_000  # the compressed file alone would pass
    _assert_not_read(store.read(data, title="bomb.docx", mime=DOCX, limits=limits), "too_large")


# --------------------------------------------------------------------------
# Class 4: a declared type that lies; non-UTF-8 text; unsupported types.
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("make", "title", "mime"),
    [
        (lambda: _pdf(["Uniform"]), "renamed.txt", TXT),
        (lambda: b"Victor plain words", "fake.pdf", PDF),
        (lambda: _docx(paragraphs=["Whiskey"]), "wrong.xlsx", XLSX),
        (lambda: _xlsx({"S": [["Xray"]]}), "wrong.docx", DOCX),
        # The extension and the mime disagree with each other.
        (lambda: _pdf(["Yankee"]), "both.pdf", TXT),
    ],
)
def test_declared_type_that_lies_is_unreadable(store, make, title, mime):
    _assert_not_read(store.read(make(), title=title, mime=mime), "unreadable")


def test_non_utf8_text_is_unreadable_not_guessed(store):
    """cp1252 bytes are not decoded by guessing: a wrong guess puts mojibake
    into a quote that reads as the client's own words."""
    data = "Zulu caf\u00e9 \u2019quoted\u2019".encode("cp1252")
    with pytest.raises(UnicodeDecodeError):
        data.decode("utf-8")
    _assert_not_read(store.read(data, title="win.txt", mime=TXT), "unreadable")


def test_utf8_bom_is_stripped(store):
    result = store.read("\ufeffAlpha bom".encode(), title="bom.txt", mime=TXT)
    _assert_read(result)
    assert result.text == "Alpha bom"


def test_utf16_with_bom_is_read(store):
    result = store.read("Bravo wide text".encode("utf-16"), title="wide.txt", mime=TXT)
    _assert_read(result)
    assert result.text == "Bravo wide text"


def test_binary_declared_as_text_is_unreadable(store):
    _assert_not_read(store.read(b"abc\x00\x00def", title="bin.txt", mime=TXT), "unreadable")


def test_image_is_unsupported_type(store):
    _assert_not_read(store.read(_png(), title="shot.png", mime="image/png"), "unsupported_type")


def test_legacy_office_file_is_unsupported_type(store):
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
    _assert_not_read(
        store.read(ole, title="old.doc", mime="application/msword"),
        "unsupported_type",
    )


def test_zip_that_is_neither_docx_nor_xlsx_is_unsupported_type(store):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/presentation.xml", "<p/>")
    _assert_not_read(store.read(buf.getvalue(), title="deck", mime=OCTET), "unsupported_type")


# --------------------------------------------------------------------------
# Class 5: XLSX with several sheets and formulas; DOCX with tables and headers.
# --------------------------------------------------------------------------


def test_xlsx_reads_every_sheet_and_keeps_an_uncached_formula_as_its_text(store):
    """openpyxl writes formulas with no cached value, so the value a viewer
    would show is unknown here. The formula is kept as written rather than
    dropped: a blank cell would read as "nothing there"."""
    data = _xlsx(
        {
            "Controls": [["Control", "Count"], ["Charlie MFA", 3], ["Delta EDR", 4]],
            "Totals": [["Total", "=SUM(Controls!B2:B3)"]],
        }
    )
    result = store.read(data, title="multi.xlsx", mime=XLSX)
    _assert_read(result)
    assert "[Sheet: Controls]" in result.text
    assert "[Sheet: Totals]" in result.text
    assert "Charlie MFA\t3" in result.text
    assert "Total\t=SUM(Controls!B2:B3)" in result.text
    assert result.text.index("[Sheet: Controls]") < result.text.index("[Sheet: Totals]")


def test_docx_reads_headers_tables_and_footers_in_order(store):
    data = _docx(
        header="Echo CONFIDENTIAL header",
        paragraphs=["Foxtrot intro paragraph"],
        table=[["Golf control", "Hotel owner"], ["India value", "Juliet value"]],
        footer="Kilo footer text",
    )
    result = store.read(data, title="tables.docx", mime=DOCX)
    _assert_read(result)
    assert "Golf control\tHotel owner" in result.text
    assert "India value\tJuliet value" in result.text
    text = result.text
    assert (
        text.index("Echo CONFIDENTIAL")
        < text.index("Foxtrot intro")
        < text.index("Golf control")
        < text.index("Kilo footer")
    )


def test_docx_with_no_text_is_read_and_empty(store):
    result = store.read(_docx(paragraphs=[]), title="blank.docx", mime=DOCX)
    _assert_read(result)
    assert result.text.strip() == ""


# --------------------------------------------------------------------------
# Class 6: text carrying control characters.
# --------------------------------------------------------------------------


def test_control_characters_are_removed_without_joining_words(store):
    raw = "Lima\x07bell Mike\x1bescape\r\nNovember\x85next Oscar\u2028line Papa\x0cfeed\tTab\x7f!"
    result = store.read(raw.encode(), title="ctl.txt", mime=TXT)
    _assert_read(result)
    assert result.text == ("Lima bell Mike escape\nNovember\nnext Oscar\nline Papa\nfeed\tTab !")


# --------------------------------------------------------------------------
# Class 7: a storage miss, and a storage failure that is not a miss.
# --------------------------------------------------------------------------


def test_storage_miss_is_storage_missing(store):
    artifact = store.put(b"Sierra", title="gone.txt", mime=TXT)
    artifact.file_storage_key = "test/never-written"
    _assert_not_read(document_text(store.storage, artifact, limits=LIMITS), "storage_missing")


def test_a_storage_failure_that_is_not_a_miss_raises(store):
    """Only a MISS is a reason. A store that cannot be reached is not a fact
    about the file, and must not be recorded as one."""

    class _Down:
        def get(self, key: str) -> bytes:
            raise ConnectionError("store unreachable")

    artifact = store.put(b"Tango", title="t.txt", mime=TXT)
    with pytest.raises(ConnectionError, match="store unreachable"):
        document_text(_Down(), artifact, limits=LIMITS)


# --------------------------------------------------------------------------
# Class 8: chunk_text.
# --------------------------------------------------------------------------


def test_chunk_exact_boundary_is_one_chunk():
    plan = chunk_text("a" * 100, chunk_chars=100, max_chunks=5)
    assert plan == ChunkPlan(chunks=("a" * 100,), sent=1, skipped=0)


def test_chunk_one_over_the_boundary_is_two_chunks():
    plan = chunk_text("a" * 101, chunk_chars=100, max_chunks=5)
    assert plan == ChunkPlan(chunks=("a" * 100, "a"), sent=2, skipped=0)


def test_chunk_max_chunks_exceeded_counts_the_skipped():
    plan = chunk_text("b" * 1000, chunk_chars=100, max_chunks=3)
    assert plan.sent == 3
    assert plan.skipped == 7
    assert plan.chunks == ("b" * 100,) * 3


def test_chunk_empty_text_sends_nothing_and_skips_nothing():
    assert chunk_text("", chunk_chars=100, max_chunks=3) == ChunkPlan(chunks=(), sent=0, skipped=0)


def test_chunk_breaks_at_whitespace_and_loses_nothing():
    text = "Uniform victor whiskey xray yankee zulu " * 40
    plan = chunk_text(text, chunk_chars=50, max_chunks=1000)
    assert plan.skipped == 0
    assert "".join(plan.chunks) == text
    assert all(len(c) <= 50 for c in plan.chunks)
    # Every chunk but the last ends at whitespace, so no word is split.
    assert all(c[-1].isspace() for c in plan.chunks[:-1])


def test_chunk_hard_cuts_a_run_with_no_whitespace():
    plan = chunk_text("c" * 250, chunk_chars=100, max_chunks=10)
    assert [len(c) for c in plan.chunks] == [100, 100, 50]


@pytest.mark.parametrize(
    ("chunk_chars", "max_chunks"),
    [(0, 1), (-1, 1), (10, 0), (10, -2)],
)
def test_chunk_refuses_limits_that_cannot_hold_text(chunk_chars, max_chunks):
    with pytest.raises(ValueError):
        chunk_text("text", chunk_chars=chunk_chars, max_chunks=max_chunks)


def test_chunk_refuses_a_non_string():
    with pytest.raises(TypeError):
        chunk_text(b"bytes", chunk_chars=10, max_chunks=1)  # type: ignore[arg-type]


def test_a_large_file_chunks_end_to_end(store):
    """A file well past one chunk: read whole, then chunked; the recorded
    sent and skipped counts add up to every chunk the file needed."""
    line = "Victor evidence line about access reviews.\n"  # 43 characters
    data = (line * 1000).encode()
    limits = ExtractLimits(
        max_bytes=LIMITS.max_bytes,
        max_unzipped_bytes=LIMITS.max_unzipped_bytes,
        max_chars=LIMITS.max_chars,
        chunk_chars=4_300,
        max_chunks=4,
    )
    result = store.read(data, title="big.txt", mime=TXT, limits=limits)
    _assert_read(result)
    assert result.chars_total == 43_000
    plan = chunk_text(result.text, chunk_chars=limits.chunk_chars, max_chunks=limits.max_chunks)
    assert plan.sent == 4
    assert plan.skipped == 6
    assert plan.chunks[0] == line * 100


# --------------------------------------------------------------------------
# The vocabulary, the invariant, and the limits' source.
# --------------------------------------------------------------------------


def test_the_not_read_vocabulary_is_section_15s():
    assert set(NOT_READ_REASONS) == {
        "unsupported_type",
        "no_text_layer",
        "encrypted",
        "unreadable",
        "too_large",
        "storage_missing",
    }


@pytest.mark.parametrize(
    "kwargs",
    [
        # read, with a reason
        {
            "status": "read",
            "reason": "unreadable",
            "text": "x",
            "chars_total": 1,
            "truncated": False,
        },
        # not_read, with no reason
        {"status": "not_read", "reason": None, "text": "", "chars_total": 0, "truncated": False},
        # not_read, carrying text
        {
            "status": "not_read",
            "reason": "encrypted",
            "text": "leak",
            "chars_total": 4,
            "truncated": False,
        },
        # a reason outside the vocabulary
        {"status": "not_read", "reason": "lost", "text": "", "chars_total": 0, "truncated": False},
    ],
)
def test_doctext_refuses_a_shape_that_blurs_read_and_not_read(kwargs):
    with pytest.raises(ValueError):
        DocText(**kwargs)


def test_limits_are_read_from_settings():
    settings = Settings(
        shield_document_max_bytes=111,
        shield_document_max_unzipped_bytes=222,
        shield_document_max_chars=333,
        shield_document_chunk_chars=44,
        shield_document_max_chunks=5,
    )
    assert extract_limits(settings) == ExtractLimits(
        max_bytes=111,
        max_unzipped_bytes=222,
        max_chars=333,
        chunk_chars=44,
        max_chunks=5,
    )


@pytest.mark.parametrize(
    "field",
    [
        "shield_document_max_bytes",
        "shield_document_max_unzipped_bytes",
        "shield_document_max_chars",
        "shield_document_chunk_chars",
        "shield_document_max_chunks",
    ],
)
def test_a_limit_below_one_refuses_to_load(field):
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        Settings(**{field: 0})
