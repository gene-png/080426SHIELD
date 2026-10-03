"""#554 R3 in the three deliverables: what is in place, what cannot be prevented,
and the awaiting-review disclosure.

The expected text is COPIED from the copy approved on #554 (21:55Z) and from
the build plan's C1-C7, never imported from the code under test. An assessment
approved before R3 renders none of it (the advisor's (d)).
"""

from __future__ import annotations

import io
import uuid

import pytest

from app.attack.analytics import compute
from app.attack.computed import effective_coverage
from app.attack.exporters import build_context, render_docx, render_pdf, render_xlsx
from app.attack.pending import pending_codes
from app.attack.rules import parents_computed
from app.models.attack_assessment import (
    AttackAssessment,
    AttackAssessmentStatus,
    AttackCoverage,
)

pytestmark = pytest.mark.unit

NP_HEADING = "Techniques that cannot be prevented"
NP_SENTENCE = (
    "MITRE ATT&CK lists no preventive control for these techniques, so they are assessed "
    "on detection and response. A technique here is Covered when it is both detected and "
    "responded to."
)
NP_LEGEND = (
    "Cannot be prevented",
    "MITRE ATT&CK lists no preventive mitigation for this technique. It is Covered when it "
    "is detected and responded to.",
)
IN_PLACE_LEGEND = (
    "Detect, Prevent, Respond",
    "in place means at least one confirmed tool provides it; awaiting review means its "
    "tools are still to be confirmed, and the status is scored as if they were not in "
    "place.",
)
ONE_AWAITING = (
    "1 technique lists tools awaiting review; it is scored as if those tools were not in " "place."
)
SET_BY_IN_PLACE = (
    "Set by what is in place: This technique's status is computed from which of Detect, "
    "Prevent and Respond are in place."
)
REACH = (
    "Not covered everywhere: Defended on most of your environment, but not on some "
    "systems, such as another operating system, a cloud or SaaS service, or unmanaged or "
    "off-network devices."
)

#: Facts about ATT&CK 19.2 (test_attack_not_preventable.py): T1082 and T1057
#: have no mitigation; T1003.001 and T1059.001 have mitigations and no
#: sub-techniques.
UNPREVENTABLE = "T1082"
UNPREVENTABLE_PARTIAL = "T1057"
PREVENTABLE = "T1003.001"
AWAITING = "T1059.001"


def _world(status_rules: int | None):
    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=1,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=2,
        status_rules=status_rules,
    )
    rows: list[AttackCoverage] = []

    def add(code, status, d, p, r, *, reason=None, unconfirmed=()):
        rows.append(
            AttackCoverage(
                id=uuid.uuid4(),
                assessment_id=a.id,
                technique_code=code,
                status=status,
                reason_code=reason,
                detection_tools=d,
                prevention_tools=p,
                response_tools=r,
                unconfirmed_citations=[
                    {"tool": t, "cited": t, "reason": "inferred", "cleared_at": None}
                    for t in unconfirmed
                ],
            )
        )

    add(UNPREVENTABLE, "covered", ["Tool D"], [], ["Tool R"])
    add(UNPREVENTABLE_PARTIAL, "partial", ["Tool D"], [], [])
    add(PREVENTABLE, "partial", ["Tool D"], [], ["Tool R"], reason="reach_limited")
    # Detection rests on a tool awaiting review: scored as not in place.
    add(AWAITING, "covered", ["Tool X"], ["Tool P"], ["Tool R"], unconfirmed=["Tool X"])

    eff = effective_coverage(a, rows)
    rollup = compute(
        {c.technique_code: c.status for c in eff},
        pending_codes(eff, parents_computed=parents_computed(a)),
    )
    return build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=eff,
        rollup=rollup,
    )


def _sheet(raw: bytes, name: str) -> list[list]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))[name]
    return [[c.value for c in r] for r in ws.iter_rows()]


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return " ".join(" ".join(parts).split())


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return " ".join(
        " ".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages).split()
    )


@pytest.mark.parametrize("status_rules", [2, None], ids=["approved-under-r3", "draft"])
def test_the_coverage_sheet_says_what_is_in_place(status_rules) -> None:
    sheet = _sheet(render_xlsx(_world(status_rules)), "Coverage")
    header, rows = sheet[0], {r[0]: r for r in sheet[1:]}
    at = header.index("Why partial")
    assert header[at + 1 : at + 4] == ["Detect", "Prevent", "Respond"], header
    status = header.index("Status")

    def dpr(code):
        return rows[code][at + 1 : at + 4]

    assert rows[UNPREVENTABLE][status] == "Covered"
    assert dpr(UNPREVENTABLE) == ["in place", "cannot be prevented", "in place"]
    assert rows[UNPREVENTABLE_PARTIAL][status] == "Partial"
    assert dpr(UNPREVENTABLE_PARTIAL) == ["in place", "cannot be prevented", "not in place"]
    assert rows[AWAITING][status] == "Partial"
    assert dpr(AWAITING) == ["awaiting review", "in place", "in place"]
    # Q7: a stored reason stays beside the line; without one, C7.
    assert rows[PREVENTABLE][at] == REACH
    assert rows[AWAITING][at] == SET_BY_IN_PLACE
    # A row nobody scored has no line.
    assert dpr("T1001") == [None, None, None]


def test_the_summary_carries_the_legends_and_the_disclosure() -> None:
    summary = _sheet(render_xlsx(_world(2)), "Heatmap Summary")
    pairs = [tuple(r[:2]) for r in summary]
    assert IN_PLACE_LEGEND in pairs
    assert NP_LEGEND in pairs
    assert any(r[0] == ONE_AWAITING for r in summary), summary


@pytest.mark.parametrize("render", [render_docx, render_pdf], ids=["docx", "pdf"])
def test_the_documents_list_what_cannot_be_prevented_and_the_disclosure(render) -> None:
    text = (_docx_text if render is render_docx else _pdf_text)(render(_world(2)))
    assert NP_HEADING in text
    assert NP_SENTENCE in text
    assert ONE_AWAITING in text
    if render is render_docx:
        assert "Status | Techniques Covered | 1 Partial | 1" in text, text
    else:
        assert "Status Techniques Covered 1 Partial 1" in text, text


def test_an_assessment_approved_before_r3_renders_none_of_it() -> None:
    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=1,
        status=AttackAssessmentStatus.RELEASED,
        parent_rules=2,
        status_rules=1,
    )
    rows = [
        AttackCoverage(
            id=uuid.uuid4(),
            assessment_id=a.id,
            technique_code=UNPREVENTABLE,
            status="partial",
            reason_code="reach_limited",
            detection_tools=["Tool D"],
            prevention_tools=[],
            response_tools=["Tool R"],
            unconfirmed_citations=[],
        )
    ]
    ctx = build_context(
        client_legal_name="Test Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=rows,
        rollup=compute({UNPREVENTABLE: "partial"}, frozenset()),
    )
    header = _sheet(render_xlsx(ctx), "Coverage")[0]
    assert "Why partial" in header  # APPEAR before ABSENT
    assert not {"Detect", "Prevent", "Respond"} & set(header)
    text = _docx_text(render_docx(ctx))
    assert "Partial coverage, by reason" in text
    assert NP_HEADING not in text
    assert "awaiting review" not in text


def test_stored_rows_for_a_computed_assessment_are_refused() -> None:
    """The rollup and the sheet must read the same rows; a caller that forgot
    `effective_coverage` would print stored statuses beside a computed figure."""
    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=1,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=2,
        status_rules=2,
    )
    row = AttackCoverage(
        id=uuid.uuid4(), assessment_id=a.id, technique_code=PREVENTABLE, status="covered"
    )
    with pytest.raises(ValueError, match="effective_coverage"):
        build_context(
            client_legal_name="Test Client",
            service_title="ATT&CK Coverage",
            assessment=a,
            coverage=[row],
            rollup=compute({PREVENTABLE: "covered"}, frozenset()),
        )
