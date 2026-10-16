"""TD-10 (for #1048): every total that leaves out an uncosted tool says how many.

In Gene's 2026-10-11 test the Annual cost total silently left out four tools
with no cost. The admin overlap view already counted them (`no_cost_count`); the
client dashboard said "Floor - some tools lacked a cost" with no count, and the
three deliverable files said only "Annual cost (may not be complete)".

The rule for {n} is ONE helper, `no_cost_items`: source items (not a split
bundle's parts) whose `annual_cost_usd` is null. Every disposition counts, a
recorded $0 is a cost, and excluded upload rows are not items at all.

The copy is approved verbatim on #736 (6115773247; strings in 6115177462):

  files: "Note: {n} tools have no cost; the annual cost excludes them." /
         "Note: 1 tool has no cost; the annual cost excludes it."
         "Tools with no cost: {names, comma-separated}." /
         "Tool with no cost: {name}."
"""

from __future__ import annotations

import io
import os
import uuid

import pytest

from app.models.capability import CapabilityDisposition, CapabilityList, CapabilityListStatus
from app.tech_debt.exporters import build_context, render_docx, render_pdf, render_xlsx
from app.tech_debt.overlap import analyze_overlap, no_cost_items
from tests.unit.test_exporters import _item
from tests.unit.test_tech_debt_dashboard import (  # noqa: F401  (fixture)
    _register,
    _seed_release,
    app_client,
)

pytestmark = pytest.mark.unit


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    pages = PdfReader(io.BytesIO(raw)).pages
    return " ".join(" ".join((p.extract_text() or "") for p in pages).split())


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return " ".join(" ".join([p.text for p in doc.paragraphs] + cells).split())


def _xlsx_text(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    return " ".join(
        str(v)
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for v in row
        if v is not None
    )


RENDERERS = pytest.mark.parametrize(
    "render, read",
    [(render_pdf, _pdf_text), (render_docx, _docx_text), (render_xlsx, _xlsx_text)],
    ids=["pdf", "docx", "xlsx"],
)


def _list() -> CapabilityList:
    return CapabilityList(
        id=uuid.uuid4(),
        service_id=uuid.uuid4(),
        version=1,
        status=CapabilityListStatus.APPROVED,
        source_rows_total=None,
        attribution_complete=None,
    )


def _ctx(items):
    cl = _list()
    cl.source_rows_total = len([i for i in items if i.parent_item_id is None])
    cl.attribution_complete = True
    return build_context(
        client_legal_name="Atlas", service_title="Atlas - Tech Debt", cap_list=cl, items=items
    )


# ---------------------------------------------------------------------------
# The rule
# ---------------------------------------------------------------------------


def test_the_helper_counts_null_source_items_of_every_disposition() -> None:
    bundle = _item(name="Suite Bundle", annual_cost_usd=500_000)
    items = [
        _item(name="Cloud Host A", annual_cost_usd=None, disposition=None),
        _item(name="Hypervisor C", annual_cost_usd=None, disposition=CapabilityDisposition.CUT),
        _item(name="Cloud Host B", annual_cost_usd=None, disposition=CapabilityDisposition.KEEP),
        # A recorded zero is a cost, not a missing one.
        _item(name="Free Tool", annual_cost_usd=0),
        _item(name="Wiz", annual_cost_usd=350_000),
        bundle,
        # A split bundle's part has no cost of its own; its bundle carries it.
        _item(name="Bundle Part", annual_cost_usd=None, parent_item_id=bundle.id),
    ]
    assert [i.name for i in no_cost_items(items)] == [
        "Cloud Host A",
        "Hypervisor C",
        "Cloud Host B",
    ]


def test_the_admin_overlap_count_is_the_same_helper() -> None:
    items = [
        _item(name="A", annual_cost_usd=None),
        _item(name="B", annual_cost_usd=0),
        _item(name="C", annual_cost_usd=None, parent_item_id=uuid.uuid4()),
    ]
    assert analyze_overlap(items).no_cost_count == 1


# ---------------------------------------------------------------------------
# The three files
# ---------------------------------------------------------------------------


@RENDERERS
def test_the_files_state_the_plural_count_and_name_the_tools(render, read) -> None:
    items = [
        _item(name="Wiz", annual_cost_usd=350_000),
        _item(name="Cloud Host A", annual_cost_usd=None),
        _item(name="Hypervisor C", annual_cost_usd=None),
        _item(name="Free Tool", annual_cost_usd=0),
    ]
    text = read(render(_ctx(items)))
    assert "Note: 2 tools have no cost; the annual cost excludes them." in text
    assert "Tools with no cost: Cloud Host A, Hypervisor C." in text
    # The $0 row is a cost; naming it would tell the client to send a cost
    # they already sent.
    assert "Free Tool." not in text
    assert "Free Tool," not in text


@RENDERERS
def test_the_files_state_the_singular_count_and_name_the_tool(render, read) -> None:
    items = [
        _item(name="Wiz", annual_cost_usd=350_000),
        _item(name="Suite D", annual_cost_usd=None),
    ]
    text = read(render(_ctx(items)))
    assert "Note: 1 tool has no cost; the annual cost excludes it." in text
    assert "Tool with no cost: Suite D." in text
    assert "tools have no cost" not in text


@RENDERERS
def test_the_files_say_nothing_when_every_tool_has_a_cost(render, read) -> None:
    bundle = _item(name="Suite Bundle", annual_cost_usd=500_000)
    items = [
        _item(name="Wiz", annual_cost_usd=350_000),
        _item(name="Free Tool", annual_cost_usd=0),
        bundle,
        _item(name="Bundle Part", annual_cost_usd=None, parent_item_id=bundle.id),
    ]
    text = read(render(_ctx(items)))
    # POSITIVE CONTROL first: the document rendered its summary at all.
    assert "Total annual cost" in text
    assert "no cost" not in text


def test_the_pdf_prints_a_tool_name_with_markup_characters_as_written() -> None:
    # The names are client text: unescaped, "<b>" would be read as markup and
    # "&" would break the paragraph (#775).
    items = [
        _item(name="Costed Tool", annual_cost_usd=1_000),
        _item(name="R&D <b>Lab</b>", annual_cost_usd=None),
    ]
    text = _pdf_text(render_pdf(_ctx(items)))
    assert "Tool with no cost: R&D <b>Lab</b>." in text


def test_the_xlsx_names_line_holds_the_name_verbatim_and_cannot_start_a_formula() -> None:
    """What makes the names cell safe is the approved wording, not the formula
    guard: the cell always begins "Tool with no cost:", so it can never begin
    with a formula trigger, whatever the tool is called. Pinned here, with the
    name stored verbatim as text. (`safe_text_row`'s guard never fires on this
    cell, so this test does not claim to exercise it.)"""
    from openpyxl import load_workbook

    items = [
        _item(name="Costed Tool", annual_cost_usd=1_000),
        _item(name="=SUM(1,2)", annual_cost_usd=None),
    ]
    ws = load_workbook(io.BytesIO(render_xlsx(_ctx(items)))).worksheets[0]
    notes = [
        c
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("Tool with no cost:")
    ]
    assert [c.value for c in notes] == ["Tool with no cost: =SUM(1,2)."]
    assert notes[0].data_type == "s"
    assert not notes[0].value.startswith(("=", "+", "-", "@", "\t", "\r"))


@RENDERERS
def test_the_note_follows_the_cost_figure(render, read) -> None:
    items = [
        _item(name="Costed Tool", annual_cost_usd=1_000),
        _item(name="Cloud Host A", annual_cost_usd=None),
    ]
    text = read(render(_ctx(items)))
    cost = text.index("Annual cost (may not be complete)")
    note = text.index("Note: 1 tool has no cost; the annual cost excludes it.")
    names = text.index("Tool with no cost: Cloud Host A.")
    assert cost < note < names


# ---------------------------------------------------------------------------
# The client dashboard
# ---------------------------------------------------------------------------


def _null_costs(*names: str) -> None:
    import sqlalchemy as sa

    eng = sa.create_engine(os.environ["DATABASE_URL"], future=True)
    with eng.begin() as conn:
        for name in names:
            n = conn.execute(
                sa.text("UPDATE capability_items SET annual_cost_usd = NULL WHERE name = :n"),
                {"n": name},
            ).rowcount
            assert n == 1, f"expected one {name!r} row, updated {n}"


def _dashboard(c, client, svc_id: str) -> dict:
    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    r = c.get(
        f"/clients/{client_id}/tech-debt/{svc_id}/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert r.status_code == 200, r.text
    return r.json()


def _released(app_client):  # noqa: F811
    c, provider = app_client
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    svc_id = _seed_release(c, provider, admin["tokens"]["access_token"], release=True)
    return c, client, svc_id


def test_the_dashboard_counts_the_tools_with_no_cost(app_client) -> None:  # noqa: F811
    c, client, svc_id = _released(app_client)
    _null_costs("Splunk ES", "CyberArk")
    b = _dashboard(c, client, svc_id)
    assert b["no_cost_count"] == 2
    assert b["spend_completeness"] == "partial"


def test_the_dashboard_count_is_zero_when_every_cost_is_recorded(
    app_client,  # noqa: F811
) -> None:
    # The seeded Defender row is a RECORDED $0, which is a cost.
    c, client, svc_id = _released(app_client)
    b = _dashboard(c, client, svc_id)
    assert b["no_cost_count"] == 0
    assert b["spend_completeness"] == "complete"
