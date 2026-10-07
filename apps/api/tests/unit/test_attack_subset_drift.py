"""#851: ATT&CK rows that credit a tool outside the client's CURRENT security
subset are disclosed, refuse approve, and can be removed by hand.

Plan: #851 comments 5983988541 (track5) and 6022779778 (re-verified at main
82d6114). Ruling: the advisor, #736 comment 5984022081 (option A; release
discloses; the remove-tool control; S1-S5), and the coordinator's D2 (a
narrow `remove_tool` that stamps nothing else).

The world is built through the routes wherever a route exists: a Tech Debt
list whose "Legacy AV" carries the model's provisional "not security" call,
which keeps it in the ATT&CK subset until a consultant confirms the call
(`/security-classification/confirm`); that sign-off is what takes it out.
"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from typing import Any

import pytest
from sqlalchemy import select

from app.models.audit_entry import AuditEntry
from app.models.capability import (
    CapabilityDisposition,
    CapabilityItem,
    CapabilityList,
    CapabilityListStatus,
)
from app.models.service import Service, ServiceKind, ServiceStatus
from tests._attack_rows import standalone_rows
from tests.unit.test_ai_runs_attack import app_parts  # noqa: F401  (fixture)

pytestmark = pytest.mark.unit

EDR = "EDR Tool"
LEGACY = "Legacy AV"
PORTAL = "Acme Portal"  # the client's own name: shown to the AI as "[CLIENT] Portal"


@dataclass
class World:
    c: Any
    sessions: Any
    h: dict
    cid: str
    svc_id: str
    assessment_id: str
    list_id: str
    items: dict[str, str]
    rows: dict[str, str]

    def get(self) -> dict:
        r = self.c.get(f"/attack/services/{self.svc_id}/assessments/latest", headers=self.h)
        assert r.status_code == 200, r.text
        return r.json()

    def patch(self, code: str, body: dict):
        return self.c.patch(f"/attack/coverage/{self.rows[code]}", headers=self.h, json=body)

    def approve(self):
        return self.c.post(f"/attack/assessments/{self.assessment_id}/approve", headers=self.h)

    def confirm_not_security(self, name: str) -> None:
        r = self.c.post(
            f"/tech-debt/capability-items/{self.items[name]}/security-classification/confirm",
            headers=self.h,
        )
        assert r.status_code == 200, r.text

    def approve_list(self) -> None:
        r = self.c.post(f"/tech-debt/capability-lists/{self.list_id}/approve", headers=self.h)
        assert r.status_code == 200, r.text


def _world(app_parts) -> World:  # noqa: F811
    c, _app, sessions = app_parts
    bearer = c.post(
        "/auth/register",
        json={
            "email": "admin@kentro.example",
            "password": "correct horse battery staple!",
            "display_name": "A",
        },
    ).json()["tokens"]["access_token"]
    cid = c.post(
        "/admin/clients",
        headers={"Authorization": f"Bearer {bearer}"},
        json={"legal_name": "Acme"},
    ).json()["id"]
    me = c.get("/auth/me", headers={"Authorization": f"Bearer {bearer}"}).json()
    with sessions() as db:
        td = Service(
            kind=ServiceKind.TECH_DEBT,
            status=ServiceStatus.IN_PROGRESS,
            title="Acme Tech Debt",
            client_id=uuid.UUID(cid),
            opened_by=uuid.UUID(me["id"]),
        )
        db.add(td)
        db.flush()
        cl = CapabilityList(service_id=td.id, version=1, status=CapabilityListStatus.DRAFT)
        db.add(cl)
        db.flush()
        items = {}
        for name, security in ((EDR, True), (LEGACY, False), (PORTAL, True)):
            it = CapabilityItem(
                capability_list_id=cl.id,
                name=name,
                # The model's call. A False is provisional until confirmed, so
                # the row stays in the ATT&CK subset (security_scope).
                security_related=security,
                disposition=CapabilityDisposition.KEEP,
            )
            db.add(it)
            db.flush()
            items[name] = str(it.id)
        db.commit()
        list_id = str(cl.id)
    h = {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}
    svc_id = c.post(
        "/attack/services", headers=h, json={"kind": "attack_coverage", "title": "Acme ATT&CK"}
    ).json()["id"]
    a = c.post(f"/attack/services/{svc_id}/assessments", headers=h).json()
    rows = {r["technique_code"]: r["id"] for r in standalone_rows(a["coverage"], 3)}
    return World(c, sessions, h, cid, svc_id, a["id"], list_id, items, rows)


def _codes(w: World) -> list[str]:
    return sorted(w.rows)


def _outside(w: World) -> list[tuple]:
    return [
        (o["technique_code"], o["field"], o["tool"], o["locked"])
        for o in w.get()["citations_outside_subset"]
    ]


# --- disclosed on the workspace, refused at approve -----------------------------------


def test_a_tool_confirmed_not_security_is_disclosed_and_blocks_approve(
    app_parts,  # noqa: F811
) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    assert w.patch(code, {"detection_tools": [EDR, LEGACY]}).status_code == 200
    w.confirm_not_security(LEGACY)
    assert _outside(w) == [(code, "detection_tools", LEGACY, False)]
    r = w.approve()
    assert r.status_code == 409, r.text
    detail = r.json()["error"]
    assert detail["reason"] == "attack_not_release_ready"
    assert detail["cites_outside_subset"] == [
        {"technique_code": code, "field": "detection_tools", "tool": LEGACY, "locked": False}
    ]
    # S4 (the "not in" variant) and S3b, both approved by the advisor (#736).
    assert (
        "1 technique row credits a tool that is not in the client's security tool list"
        in detail["message"]
    )
    assert detail["message"].endswith(
        "Remove the tool in the technique's panel, or unlock the row and use Run AI."
    )


def test_removing_the_tool_by_hand_clears_it_and_approve_proceeds(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [EDR, LEGACY]})
    w.confirm_not_security(LEGACY)
    assert _outside(w) != []  # the positive state first
    r = w.patch(code, {"remove_tool": {"field": "detection_tools", "name": LEGACY}})
    assert r.status_code == 200, r.text
    assert r.json()["detection_tools"] == [EDR]
    assert _outside(w) == []
    assert w.approve().status_code == 200


def test_a_locked_row_is_listed_as_locked(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [LEGACY]})
    w.patch(code, {"locked": True})
    w.confirm_not_security(LEGACY)
    assert _outside(w) == [(code, "detection_tools", LEGACY, True)]


def test_every_field_and_row_is_read(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    a, b = _codes(w)[:2]
    w.patch(a, {"prevention_tools": [LEGACY]})
    w.patch(b, {"response_tools": [LEGACY], "detection_tools": [EDR]})
    w.confirm_not_security(LEGACY)
    assert sorted(_outside(w)) == sorted(
        [(a, "prevention_tools", LEGACY, False), (b, "response_tools", LEGACY, False)]
    )


def test_the_approve_refusal_counts_rows_not_tools(app_parts) -> None:  # noqa: F811
    """S4 counts technique ROWS: two outside tools on one row are one row, and
    the same tool on two rows is two (the approved singular and plural)."""
    w = _world(app_parts)
    a, b = _codes(w)[:2]
    w.patch(a, {"detection_tools": [LEGACY], "prevention_tools": ["Shadow Scanner"]})
    w.confirm_not_security(LEGACY)
    one = w.approve().json()["error"]["message"]
    assert (
        "This assessment cannot be approved: 1 technique row credits a tool that is not in "
        "the client's security tool list." in one
    ), one
    w.patch(b, {"response_tools": [LEGACY]})
    two = w.approve().json()["error"]["message"]
    assert (
        "This assessment cannot be approved: 2 technique rows credit a tool that is not in "
        "the client's security tool list." in two
    ), two


# --- which names are "outside": D1(a), behind one predicate ---------------------------


def test_a_free_text_tool_never_on_the_list_is_listed(app_parts) -> None:  # noqa: F811
    """D1(a), approved by the advisor (#736): a name the subset does not
    know is outside it, however it got onto the row."""
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": ["Shadow Scanner"]})
    assert _outside(w) == [(code, "detection_tools", "Shadow Scanner", False)]


@pytest.mark.parametrize("spelling", ["EDR TOOL", "edr  tool", "[CLIENT] Portal"])
def test_the_runs_resolver_decides_what_is_the_same_tool(app_parts, spelling) -> None:  # noqa: F811
    """Case and whitespace, and the shown form of a client-named tool, are the
    same tool: the run's own name tiers decide, never a second comparison."""
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [spelling]})
    assert _outside(w) == []
    w.patch(code, {"detection_tools": [spelling, "Shadow Scanner"]})
    assert _outside(w) == [(code, "detection_tools", "Shadow Scanner", False)]


def test_an_approved_list_answers_from_its_snapshot_until_approved_again(
    app_parts,  # noqa: F811
) -> None:  # noqa: F811
    """The snapshot is the membership: confirming after approval does not move
    it (pinned), and approving the list again does."""
    w = _world(app_parts)
    code = _codes(w)[0]
    w.approve_list()
    w.patch(code, {"detection_tools": [LEGACY]})
    w.confirm_not_security(LEGACY)
    assert _outside(w) == []
    w.approve_list()
    assert _outside(w) == [(code, "detection_tools", LEGACY, False)]


# --- after approval: disclosed, never refused -----------------------------------------


def test_an_approved_assessment_still_discloses_a_later_removal(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [LEGACY]})
    assert w.approve().status_code == 200
    w.confirm_not_security(LEGACY)
    body = w.get()
    assert body["status"] == "approved"
    assert [(o["technique_code"], o["tool"]) for o in body["citations_outside_subset"]] == [
        (code, LEGACY)
    ]


def test_the_what_if_base_carries_the_same_count(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [LEGACY, EDR]})
    assert w.approve().status_code == 200
    w.confirm_not_security(LEGACY)
    r = w.c.get(f"/attack/services/{w.svc_id}/scenarios", headers=w.h)
    assert r.status_code == 200, r.text
    assert r.json()["base"]["citations_outside_subset"] == 1


# --- the narrow remove (D2): it removes one name and stamps nothing else ---------------


def test_removing_one_tool_leaves_the_others_review_state_as_it_was(
    app_parts,  # noqa: F811
) -> None:  # noqa: F811
    """D2. A PATCH of a whole list stamps every citation in it confirmed (#102);
    the remove control must not. The inferred entry below is what the AI run
    writes for an uncleared citation (`routes/attack.py`, the run's row flags)."""
    from app.models.attack_assessment import AttackCoverage

    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [EDR, LEGACY], "prevention_tools": [LEGACY]})
    with w.sessions() as db:
        row = db.get(AttackCoverage, uuid.UUID(w.rows[code]))
        row.unconfirmed_citations = [
            {
                "tool": EDR,
                "cited": "EDR",
                "reason": "substring",
                "field": "detection_tools",
                "cleared_at": None,
            },
            {
                "tool": LEGACY,
                "cited": "Legacy",
                "reason": "substring",
                "field": "detection_tools",
                "cleared_at": None,
            },
            # The same name in ANOTHER list: a removal is one list, so it stays.
            {
                "tool": LEGACY,
                "cited": "Legacy",
                "reason": "substring",
                "field": "prevention_tools",
                "cleared_at": None,
            },
        ]
        db.commit()
    r = w.patch(code, {"remove_tool": {"field": "detection_tools", "name": LEGACY}})
    assert r.status_code == 200, r.text
    with w.sessions() as db:
        row = db.get(AttackCoverage, uuid.UUID(w.rows[code]))
        assert row.detection_tools == [EDR]
        assert row.prevention_tools == [LEGACY]
        assert [(e["tool"], e["field"], e["cleared_at"]) for e in row.unconfirmed_citations] == [
            (EDR, "detection_tools", None),
            (LEGACY, "prevention_tools", None),
        ]
        audited = (
            db.execute(
                select(AuditEntry).where(AuditEntry.action == "attack.coverage.tool_removed")
            )
            .scalars()
            .all()
        )
        # The technique and the list, never the tool's name (the client's data).
        assert [a.details for a in audited] == [
            {"technique_code": code, "field": "detection_tools"}
        ]


@pytest.mark.parametrize(
    ("remove", "reason"),
    [
        ({"field": "notes", "name": LEGACY}, "remove_tool_unknown_field"),
        ({"field": "detection_tools", "name": "Not There"}, "remove_tool_not_listed"),
    ],
)
def test_a_remove_that_cannot_apply_is_a_typed_422(app_parts, remove, reason) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": [EDR]})
    r = w.patch(code, {"remove_tool": remove})
    assert r.status_code == 422, r.text
    assert r.json()["error"]["reason"] == reason
    assert w.get()["coverage"]  # the row is untouched
    with w.sessions() as db:
        from app.models.attack_assessment import AttackCoverage

        assert db.get(AttackCoverage, uuid.UUID(w.rows[code])).detection_tools == [EDR]


# --- no list to judge against: the check does not apply (advisor, #736 6039558116) ---


def _no_list_world(app_parts) -> World:  # noqa: F811
    """The same world with its Tech Debt list removed: an ATT&CK-only client."""
    w = _world(app_parts)
    with w.sessions() as db:
        for it in db.execute(select(CapabilityItem)).scalars().all():
            db.delete(it)
        db.delete(db.get(CapabilityList, uuid.UUID(w.list_id)))
        db.commit()
    return w


def test_with_no_tech_debt_list_nothing_is_flagged_and_approve_proceeds(
    app_parts,  # noqa: F811
) -> None:
    w = _no_list_world(app_parts)
    code = _codes(w)[0]
    assert w.patch(code, {"detection_tools": ["Tool A"]}).status_code == 200
    assert w.get()["coverage"]  # the assessment is there
    assert _outside(w) == []
    assert w.approve().status_code == 200


# --- "not checked" is a THIRD state, not a pass (advisor, #736 6039558116) -------------
# The sentence is copied from the drafted wording sent for approval, never from
# the module that renders it.
NOT_CHECKED = (
    "The tools cited here were not checked against a security tool list, because "
    "the client has none."
)


def test_the_assessment_says_whether_the_tools_were_checked(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    assert w.get()["subset_checked"] is True  # a live list: checked (positive first)
    r = w.c.post(f"/tech-debt/capability-lists/{w.list_id}/discard", headers=w.h)
    assert r.status_code == 200, r.text
    assert w.get()["subset_checked"] is False


def test_an_attack_only_client_reads_not_checked(app_parts) -> None:  # noqa: F811
    assert _no_list_world(app_parts).get()["subset_checked"] is False


def _with_storage(app_parts, tmp_path) -> None:  # noqa: F811
    """Finalize writes the artifacts: a local store, as the deliverable tests use."""
    from app.routes.artifacts import _storage_dep
    from app.storage.local import LocalFilesystemStorage

    storage = LocalFilesystemStorage(tmp_path / "storage")
    app_parts[1].dependency_overrides[_storage_dep] = lambda: storage


def _finalize(w: World) -> dict:
    assert w.approve().status_code == 200
    fin = w.c.post(f"/attack/services/{w.svc_id}/deliverables/finalize", headers=w.h)
    assert fin.status_code in (200, 201), fin.text
    return fin.json()


def _format_text(w: World, fin: dict, fmt: str) -> str:
    import io

    if fmt == "summary":
        return fin["summary"]
    r = w.c.get(f"/artifacts/{fin[f'{fmt}_artifact_id']}/download", headers=w.h)
    assert r.status_code == 200, r.text
    raw = io.BytesIO(r.content)
    if fmt == "pdf":
        from pypdf import PdfReader

        text = " ".join((p.extract_text() or "") for p in PdfReader(raw).pages)
    elif fmt == "docx":
        from docx import Document

        text = " ".join(p.text for p in Document(raw).paragraphs)
    else:
        from openpyxl import load_workbook

        ws = load_workbook(raw)["Heatmap Summary"]
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    return " ".join(text.split())


FORMATS = ["pdf", "docx", "xlsx", "summary"]


@pytest.mark.parametrize("fmt", FORMATS)
def test_each_format_says_not_checked_when_the_client_has_no_list(
    app_parts, tmp_path, fmt: str  # noqa: F811
) -> None:
    _with_storage(app_parts, tmp_path)
    w = _no_list_world(app_parts)
    w.patch(_codes(w)[0], {"detection_tools": ["Tool A"]})
    text = _format_text(w, _finalize(w), fmt)
    assert NOT_CHECKED in text, text[:3000]


@pytest.mark.parametrize("fmt", FORMATS)
def test_each_format_is_silent_when_the_tools_were_checked(
    app_parts, tmp_path, fmt: str  # noqa: F811
) -> None:
    _with_storage(app_parts, tmp_path)
    w = _world(app_parts)
    w.patch(_codes(w)[0], {"detection_tools": [EDR]})
    text = _format_text(w, _finalize(w), fmt)
    # Positive first: the surface rendered.
    assert ("Coverage:" if fmt == "summary" else "ATT&CK") in text, text[:3000]
    assert "not checked against a security tool list" not in text, text[:3000]


def test_with_only_a_discarded_list_nothing_is_flagged(app_parts) -> None:  # noqa: F811
    w = _world(app_parts)
    code = _codes(w)[0]
    w.patch(code, {"detection_tools": ["Shadow Scanner"]})
    assert _outside(w) != []  # a live list: the check applies (positive first)
    r = w.c.post(f"/tech-debt/capability-lists/{w.list_id}/discard", headers=w.h)
    assert r.status_code == 200, r.text
    assert _outside(w) == []
    assert w.approve().status_code == 200
