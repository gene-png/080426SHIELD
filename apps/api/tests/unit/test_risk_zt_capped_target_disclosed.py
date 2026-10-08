"""The Risk Register says when the DoD target cap lowered a target (#915, S3).

Gene's ruling was "cap and say so", and the advisor carried it to every surface
the cap reaches (#736 comment 6048561596). S3 was approved verbatim, singular
and plural, at 6049667540; the expected sentences below are copied from the
proposal in 6049167247, never built from the code.

The world: an ATT&CK gap (the register's gate) and a released DoD Zero Trust
assessment. Plural: the client chose stage 3 at intake, so every capability with
no DoD Advanced activity (15, per `count_dod_levels.py`; pinned in
`test_risk_zt_dod_caps.py`) has its target lowered to 2. Singular: the client
chose stage 2, and only 1.1 User Inventory carries its own target of 3.

Checked through the endpoints a consultant and a client reach, and the three
files a client receives.
"""

from __future__ import annotations

import io

import pytest

from tests._ai_mode import docx_text, pdf_text
from tests._attack_rows import first_standalone
from tests._risk_inputs import release
from tests.unit.test_risk_dashboard import (  # noqa: F401  (fixture)
    _register,
    app_client,
)
from tests.unit.test_zt_dashboard import _attach_intake_target

pytestmark = pytest.mark.unit

PLURAL = (
    "In the DoD Zero Trust assessment, 15 capabilities have no DoD Advanced "
    "activities, so their target is Target (2): each is a finding only below "
    "Target. The Zero Trust deliverable names them."
)
SINGULAR = (
    "In the DoD Zero Trust assessment, 1 capability has no DoD Advanced "
    "activities, so its target is Target (2): it is a finding only below "
    "Target. The Zero Trust deliverable names it."
)
NO_ADVANCED = "DOD.USR.01"  # 1.1 User Inventory


def _world(c, *, intake_stage: int, own_target: int | None) -> tuple[str, dict, dict]:
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    cid = client["user"]["client_id"]
    c.headers["X-Client-Id"] = cid
    h = {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}

    asvc = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": "A"})
    a = c.post(f"/attack/services/{asvc.json()['id']}/assessments", headers=h).json()
    cov = first_standalone(a["coverage"])
    r = c.patch(f"/attack/coverage/{cov['id']}", headers=h, json={"status": "gap"})
    assert r.status_code == 200, r.text
    assert c.post(f"/attack/assessments/{a['id']}/approve", headers=h).status_code == 200

    zsvc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_dod", "title": "ZT"})
    _attach_intake_target(zsvc.json()["id"], zt_stage=intake_stage)
    za = c.post(f"/zt/services/{zsvc.json()['id']}/assessments", headers=h).json()
    for ans in za["answers"]:
        if ans["capability_code"] == NO_ADVANCED:
            body: dict = {"maturity_stage": 1}
            if own_target is not None:
                body["target_stage"] = own_target
            r = c.patch(f"/zt/answers/{ans['id']}", headers=h, json=body)
            assert r.status_code == 200, r.text
    assert c.post(f"/zt/assessments/{za['id']}/approve", headers=h).status_code == 200
    release(c, bearer, cid, "attack", asvc.json()["id"])
    release(c, bearer, cid, "zt", zsvc.json()["id"])
    return cid, h, client


def _flat(text: str) -> str:
    return " ".join(text.split())


def _files(c, h: dict, body: dict) -> tuple[str, str, list[str]]:
    from openpyxl import load_workbook

    def _bytes(key: str) -> bytes:
        dl = c.get(f"/artifacts/{body[key]}/download", headers=h)
        assert dl.status_code == 200, dl.text
        return dl.content

    wb = load_workbook(io.BytesIO(_bytes("xlsx_artifact_id")))
    summary = [str(v) for row in wb["Summary"].iter_rows(values_only=True) for v in row if v]
    return (
        _flat(pdf_text(_bytes("pdf_artifact_id"))),
        _flat(docx_text(_bytes("docx_artifact_id"))),
        summary,
    )


def test_the_register_persists_the_codes_and_states_them_plural(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, _client = _world(c, intake_stage=3, own_target=None)
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    assert g.json()["zt_capped_target_note"] == PLURAL
    # Read back: from the persisted provenance, not from the generate run.
    latest = c.get(f"/risk/clients/{cid}/register/latest", headers=h)
    assert latest.status_code == 200, latest.text
    body = latest.json()
    assert body["zt_capped_target_note"] == PLURAL
    (codes,) = body["capped_target_codes"].values()
    assert NO_ADVANCED in codes
    assert len(codes) == 15


def test_the_register_states_one_capability_singular(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, _client = _world(c, intake_stage=2, own_target=3)
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    assert g.json()["zt_capped_target_note"] == SINGULAR
    assert list(g.json()["capped_target_codes"].values()) == [[NO_ADVANCED]]


def test_every_published_file_and_the_client_dashboard_state_it(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, client = _world(c, intake_stage=3, own_target=None)
    assert c.post(f"/risk/clients/{cid}/register/generate", headers=h).status_code == 201
    pub = c.post(f"/risk/clients/{cid}/register/publish", headers=h)
    assert pub.status_code in (200, 201), pub.text
    pdf, docx, summary = _files(c, h, pub.json())
    assert PLURAL in pdf
    assert PLURAL in docx
    assert PLURAL in summary
    dash = c.get(
        f"/clients/{cid}/risk/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert dash.status_code == 200, dash.text
    assert dash.json()["zt_capped_target_note"] == PLURAL


def test_nothing_is_stated_when_no_target_was_lowered(app_client) -> None:  # noqa: F811
    c = app_client
    cid, h, client = _world(c, intake_stage=2, own_target=None)
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    # The positive state first: the codes were recorded, and there are none.
    assert list(g.json()["capped_target_codes"].values()) == [[]]
    assert g.json()["zt_capped_target_note"] is None
    pub = c.post(f"/risk/clients/{cid}/register/publish", headers=h)
    assert pub.status_code in (200, 201), pub.text
    pdf, docx, _summary = _files(c, h, pub.json())
    assert "Total entries:" in pdf  # a real summary, so the absence means something
    assert "In the DoD Zero Trust assessment" not in pdf
    assert "In the DoD Zero Trust assessment" not in docx
    dash = c.get(
        f"/clients/{cid}/risk/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert dash.status_code == 200, dash.text
    assert dash.json()["zt_capped_target_note"] is None
