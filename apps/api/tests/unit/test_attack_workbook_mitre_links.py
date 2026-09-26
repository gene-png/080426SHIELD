"""Every technique in the delivered ATT&CK workbook links to its MITRE page (#647).

WHERE THE EXPECTED URLS COME FROM. MITRE's own STIX objects carry each
technique's page as the `mitre-attack` external reference, so this module reads
the committed subset under `app/attack/stix/` with its OWN code and compares.
A URL rule that disagrees with MITRE for any one of the catalogue's techniques
cannot agree with this test by construction. The two literal pins are the
issue's own example, a parent and a sub-technique, each opened on
attack.mitre.org on 2026-09-26 and found to be that technique's page.
"""

from __future__ import annotations

import gzip
import io
import json
import uuid
from pathlib import Path

import pytest

from app.attack import catalog
from app.attack.analytics import compute as compute_heatmap
from app.attack.catalog import TECHNIQUES
from app.attack.coverage import CoverageStatus
from app.attack.exporters import build_context, render_xlsx
from app.models.attack_assessment import (
    AttackAssessment,
    AttackAssessmentStatus,
    AttackCoverage,
)

pytestmark = pytest.mark.unit

STIX_DIR = Path(__file__).resolve().parents[2] / "app" / "attack" / "stix"


def _mitre_urls() -> dict[str, str]:
    """{technique id: the page URL MITRE publishes for it}, live objects only."""
    bundle = json.loads(gzip.decompress((STIX_DIR / catalog.SOURCE["subset"]).read_bytes()))
    urls: dict[str, str] = {}
    for o in bundle["objects"]:
        if o["type"] != "attack-pattern" or o.get("revoked") or o.get("x_mitre_deprecated"):
            continue
        (ref,) = [r for r in o["external_references"] if r.get("source_name") == "mitre-attack"]
        urls[ref["external_id"]] = ref["url"]
    return urls


def _ctx(status: str | None, *, extra_codes: tuple[str, ...] = ()):
    a = AttackAssessment(
        id=uuid.uuid4(),
        service_id=uuid.uuid4(),
        version=1,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=2,
    )
    codes = [t.id for t in TECHNIQUES] + list(extra_codes)
    rows = [
        AttackCoverage(id=uuid.uuid4(), assessment_id=a.id, technique_code=c, status=status)
        for c in codes
    ]
    return build_context(
        client_legal_name="Example Client",
        service_title="MITRE ATT&CK Coverage",
        assessment=a,
        coverage=rows,
        rollup=compute_heatmap({r.technique_code: r.status for r in rows}),
    )


def _links(raw: bytes, sheet: str) -> dict[str, str | None]:
    """{code in column A: its hyperlink target or None}, header row skipped."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))[sheet]
    return {
        cell.value: (cell.hyperlink.target if cell.hyperlink else None)
        for (cell,) in ws.iter_rows(min_row=2, min_col=1, max_col=1)
    }


def test_every_catalogue_technique_url_is_the_one_mitre_publishes() -> None:
    mitre = _mitre_urls()
    checked = [t.id for t in TECHNIQUES if t.id in mitre]
    # The selector's count is part of the result: a catalogue id MITRE does not
    # list would otherwise be skipped silently.
    assert len(checked) == len(TECHNIQUES) > 0
    wrong = {
        c: (catalog.technique_url(c), mitre[c])
        for c in checked
        if catalog.technique_url(c) != mitre[c]
    }
    assert wrong == {}


def test_a_parent_and_a_sub_technique_link_to_their_own_pages() -> None:
    assert catalog.technique_url("T1595") == "https://attack.mitre.org/techniques/T1595"
    assert catalog.technique_url("T1595.001") == "https://attack.mitre.org/techniques/T1595/001"


def test_every_row_of_the_coverage_sheet_links_to_mitre() -> None:
    mitre = _mitre_urls()
    links = _links(render_xlsx(_ctx(CoverageStatus.COVERED.value)), "Coverage")
    assert len(links) == len(TECHNIQUES)
    assert {c: t for c, t in links.items() if t != mitre[c]} == {}


def test_every_row_of_the_gaps_sheet_links_to_mitre() -> None:
    mitre = _mitre_urls()
    links = _links(render_xlsx(_ctx(CoverageStatus.GAP.value)), "Gaps")
    assert len(links) == len(TECHNIQUES)
    assert {c: t for c, t in links.items() if t != mitre[c]} == {}


def test_every_row_of_the_unscored_sheet_links_to_mitre() -> None:
    mitre = _mitre_urls()
    links = _links(render_xlsx(_ctx(None)), "Unscored")
    assert len(links) == len(TECHNIQUES)
    assert {c: t for c, t in links.items() if t != mitre[c]} == {}


def test_a_gap_code_outside_the_catalogue_is_listed_without_a_link() -> None:
    # A stored code the catalogue does not carry has no page this workbook can
    # vouch for, so it is printed and not linked -- never linked to a guess.
    links = _links(render_xlsx(_ctx(CoverageStatus.GAP.value, extra_codes=("T9999",))), "Gaps")
    assert "T9999" in links
    assert links["T9999"] is None
    assert links["T1595"] == "https://attack.mitre.org/techniques/T1595"


def test_the_placeholder_rows_carry_no_link() -> None:
    assert _links(render_xlsx(_ctx(CoverageStatus.COVERED.value)), "Gaps") == {"—": None}
    assert _links(render_xlsx(_ctx(CoverageStatus.GAP.value)), "Unscored") == {"—": None}
