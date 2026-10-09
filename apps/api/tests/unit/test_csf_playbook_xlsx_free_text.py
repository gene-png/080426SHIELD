"""The CSF Playbook workbook stores free text as text (#972).

The Action Plan sheet carries a gap's POA&M annotation -- owner, resources,
success criteria, POA&M reference -- which a consultant types (`schemas/csf.py`
types each as free `str`). With a bare `append`, an owner of `=HYPERLINK(...)`
became a live formula in the client's workbook. Every append in
`csf/playbook_export.py::render_xlsx` now goes through the shared
`app/xlsx_export.py::safe_text_row`.

Driven through the routes a consultant uses: the annotation is saved through
`PUT /csf/services/{id}/gap-actions/{code}`, the workbook is produced by
`/playbook/export`, downloaded, and read back with openpyxl.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from tests.unit.test_csf_playbook_override_every_artifact import (  # noqa: F401  (fixture)
    _export,
    _one_gap,
    app_client,
)

pytestmark = pytest.mark.unit

ANNOTATION = {
    "owner": '=HYPERLINK("http://example.invalid","click")',
    "resources": "+resources",
    "success_criteria": "@criteria",
    "poam_ref": "-POAM-1",
}


def test_action_plan_free_text_is_text_exactly_as_typed(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, code = _one_gap(c, h)
    saved = c.put(f"/csf/services/{svc}/gap-actions/{code}", headers=h, json=ANNOTATION)
    assert saved.status_code == 200, saved.text
    ws = load_workbook(io.BytesIO(_export(c, h, svc)["xlsx"]))["Action Plan"]
    found = {c.value: c for row in ws.iter_rows() for c in row if c.value in ANNOTATION.values()}
    # Positive first: every value is in the sheet, exactly as typed.
    assert set(found) == set(ANNOTATION.values()), set(ANNOTATION.values()) - set(found)
    for value, cell in found.items():
        assert cell.data_type == "s", f"{value!r} is {cell.data_type!r}, not text"
        assert cell.quotePrefix is True, f"{value!r} is not marked as text"
