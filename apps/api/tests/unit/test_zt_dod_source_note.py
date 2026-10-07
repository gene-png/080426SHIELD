"""Every DoD deliverable carries the approved source note (#839).

The sentence is the advisor's re-ruling at #736, Oct 7 14:48Z (comment
6040458893), verbatim: the 2022 Roadmap's levels could not be matched to
activities reliably, so the note says so instead of listing differences. The
expected text is copied from that comment, never from the code under test.

The check goes through the files a client receives, and a CISA deliverable
must NOT carry it: the note is about the DoD source only.
"""

from __future__ import annotations

import io

import pytest

from tests._ai_mode import docx_text, pdf_text
from tests.unit.test_zt_dashboard import (  # noqa: F401  (fixture)
    _register,
    _seed_release,
    app_client,
)

pytestmark = pytest.mark.unit

NOTE = (
    "Levels are as published in the 2025 edition (25-T-1465). The 2022 Roadmap "
    "shows levels only as colours in its timeline, which could not be matched to "
    "activities reliably; only 1.2.1 to 1.2.5, which it tabulates, were checked, "
    "and they agree."
)


def _files(c, kind: str) -> tuple[str, str, list[str]]:
    bearer = _register(c, "admin@example.com")["tokens"]["access_token"]
    _register(c, "client@example.com")
    svc_id = _seed_release(c, bearer, release=True, target_stage=None, kind=kind)
    h = {"Authorization": f"Bearer {bearer}"}
    deliv = c.get(f"/zt/services/{svc_id}/deliverables/latest", headers=h)
    assert deliv.status_code == 200, deliv.text
    body = deliv.json()

    def _bytes(key: str) -> bytes:
        dl = c.get(f"/artifacts/{body[key]}/download", headers=h)
        assert dl.status_code == 200, dl.text
        return dl.content

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(_bytes("xlsx_artifact_id")))
    cells = [str(v) for ws in wb for row in ws.iter_rows(values_only=True) for v in row if v]
    return pdf_text(_bytes("pdf_artifact_id")), docx_text(_bytes("docx_artifact_id")), cells


def _flat(text: str) -> str:
    return " ".join(text.split())


def test_every_dod_file_carries_the_source_note(app_client) -> None:  # noqa: F811
    pdf, docx, cells = _files(app_client, "zero_trust_dod")
    assert NOTE in _flat(pdf)
    assert NOTE in _flat(docx)
    assert NOTE in cells


def test_a_cisa_file_does_not(app_client) -> None:  # noqa: F811
    pdf, docx, cells = _files(app_client, "zero_trust_cisa")
    # The positive state first: these are real deliverables with content.
    assert "CISA ZTMM 2.0" in _flat(pdf)
    assert "25-T-1465" not in _flat(pdf)
    assert "25-T-1465" not in _flat(docx)
    assert not any("25-T-1465" in v for v in cells)
