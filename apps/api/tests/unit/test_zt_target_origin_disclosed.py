"""#1343: the ZT deliverable discloses every applied per-capability target
whose origin is not recorded, with the approved sentence (#736 6220094031
item 5) and the approved placement (6225462107): in the DOCX and PDF a line
under the gap table, `{code}: <sentence>`, one per code; in the XLSX Gap Plan,
the sentence appended to that row's Notes cell, never replacing what is there,
and a line under the table for a code with no row.

Every code is disclosed, not only the rows shown (review of 1f659fb7, F1): a
target that keeps a capability OFF the gap list decided what the client reads
(the missing row and the gap count), and the client dashboard lists it too, so
the two surfaces disclose the same set.

The world, against an engagement target of 3 (CISA, Stage 4 elsewhere):
- `a`: Stage 1 under its own target of 4, origin not recorded, notes "kept";
- `b`: the same, but the consultant set it, notes "own words";
- `c`: the same as `a` with no notes;
- `d`: Stage 2 under its own target of 2, origin not recorded. It is no gap
  only because of that target (the engagement target would make it one).
Expected text is written here from the ruling, never imported.
"""

from __future__ import annotations

import io
import uuid

import pytest

from app.models.zt_assessment import (
    ZtAnswer,
    ZtAssessment,
    ZtAssessmentStatus,
    ZtFramework,
)
from app.zt.catalog import capabilities
from app.zt.exporters import build_context, render_docx, render_pdf, render_xlsx
from app.zt.maturity import ZtFrameworkCode
from app.zt.scoring import analyze_gaps, compute

pytestmark = pytest.mark.unit

SENTENCE = "Who set this target is not recorded. An earlier AI run may have set it."


def _ctx():
    fw = ZtFrameworkCode.CISA_ZTMM_2_0
    caps = list(capabilities(fw))
    a, b, c, d = (cap.code for cap in caps[:4])
    assessment = ZtAssessment(
        id=uuid.uuid4(),
        service_id=uuid.uuid4(),
        framework=ZtFramework.CISA_ZTMM_2_0,
        version=1,
        status=ZtAssessmentStatus.APPROVED,
    )
    shapes = {
        a: {"maturity_stage": 1, "target_stage": 4, "target_source": None, "notes": "kept"},
        b: {
            "maturity_stage": 1,
            "target_stage": 4,
            "target_source": "consultant",
            "notes": "own words",
        },
        c: {"maturity_stage": 1, "target_stage": 4, "target_source": None, "notes": None},
        d: {"maturity_stage": 2, "target_stage": 2, "target_source": None, "notes": None},
    }
    answers = []
    for cap in caps:
        ans = ZtAnswer(**shapes.get(cap.code, {"maturity_stage": 4}))
        ans.id = uuid.uuid4()
        ans.assessment_id = assessment.id
        ans.capability_code = cap.code
        answers.append(ans)
    stages = {x.capability_code: x.maturity_stage for x in answers}
    targets = {x.capability_code: x.target_stage for x in answers}
    notes = {x.capability_code: x.notes for x in answers}
    gap = analyze_gaps(fw, stages, notes=notes, targets=targets, target_stage=3)
    assert sorted(g.code for g in gap.gaps) == sorted([a, b, c]), gap.gaps
    ctx = build_context(
        client_legal_name="Atlas Defense Solutions",
        service_title="Zero Trust Assessment",
        framework=fw,
        assessment=assessment,
        answers=answers,
        score=compute(fw, stages),
        gap=gap,
    )
    return ctx, a, b, c, d


def test_the_docx_lists_every_unrecorded_target_under_the_gap_table() -> None:
    from docx import Document

    ctx, a, b, c, d = _ctx()
    paragraphs = [p.text for p in Document(io.BytesIO(render_docx(ctx))).paragraphs]
    for code in (a, c, d):
        assert f"{code}: {SENTENCE}" in paragraphs, (code, paragraphs)
    assert f"{b}: {SENTENCE}" not in paragraphs


def test_the_pdf_lists_every_unrecorded_target() -> None:
    from pypdf import PdfReader

    ctx, a, b, c, d = _ctx()
    raw = render_pdf(ctx)
    text = " ".join(" ".join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages).split())
    for code in (a, c, d):
        assert f"{code}: {SENTENCE}" in text, (code, text)
    assert f"{b}: {SENTENCE}" not in text


def test_the_xlsx_appends_to_notes_and_lists_a_code_with_no_row() -> None:
    from openpyxl import load_workbook

    ctx, a, b, c, d = _ctx()
    ws = load_workbook(io.BytesIO(render_xlsx(ctx)))["Gap Plan"]
    assert ws.cell(row=2, column=8).value == "Notes"
    notes = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=8).value for r in (3, 4, 5)}
    assert set(notes) == {a, b, c}, notes
    # Appended, never replacing the row's own notes.
    assert notes[a] == f"kept {SENTENCE}", notes
    # No notes of its own: exactly the sentence, with no stray separator.
    assert notes[c] == SENTENCE, notes
    # A recorded origin: the row's notes byte for byte.
    assert notes[b] == "own words", notes
    # `d` has no row, so it gets the line under the table.
    below = [ws.cell(row=r, column=1).value for r in range(6, ws.max_row + 1)]
    assert f"{d}: {SENTENCE}" in below, below
    assert f"{a}: {SENTENCE}" not in below, "a shown row is disclosed in its Notes"
