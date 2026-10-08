"""The DoD target caps reach every surface a person reads (#839).

`test_zt_dod_target_caps.py` pins the rule; this pins the WIRING, through the
endpoints a consultant and a client reach and the files a client receives: a
guard tested only at the function is not tested where it is selected.

The world: a DoD engagement whose client chose stage 3 at intake (so it is the
engagement target), every capability at stage 2 with no per-capability target,
approved, finalized and released. Under the cap, the 15 capabilities with no
Advanced activity are at their target and are no gap. The sentence checked is
the approved copy C1, verbatim, for 1.1.
"""

from __future__ import annotations

import io

import pytest

from tests._ai_mode import docx_text, pdf_text
from tests.unit.test_zt_dashboard import (  # noqa: F401  (fixture)
    _register,
    _seed_release,
    app_client,
)

pytestmark = pytest.mark.unit

C1 = "1.1 User Inventory has no DoD Advanced activities, so its target is Target (2)."
NO_ADVANCED = 15  # the extraction's count, pinned in test_zt_dod_target_caps.py
DOD_CAPABILITIES = 45


def _world(c) -> tuple[str, dict, dict, str]:
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    svc_id = _seed_release(
        c, bearer, release=True, target_stage=None, kind="zero_trust_dod", intake_stage=3
    )
    return svc_id, admin, client, bearer


def test_the_workspace_gap_endpoint_discloses_and_counts_the_caps(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _admin, _client, bearer = _world(c)
    r = c.get(
        f"/zt/services/{svc_id}/gap-analysis?target_stage=3&top_n=50",
        headers={"Authorization": f"Bearer {bearer}"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["target_stage"] == 3
    assert body["total_gap_count"] == DOD_CAPABILITIES - NO_ADVANCED
    assert C1 in body["target_cap_notes"]
    assert len(body["target_cap_notes"]) == NO_ADVANCED


def test_the_client_dashboard_discloses_the_caps(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _admin, client, _bearer = _world(c)
    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    r = c.get(
        f"/clients/{client_id}/zt/{svc_id}/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_gap_count"] == DOD_CAPABILITIES - NO_ADVANCED
    assert C1 in body["target_cap_notes"]


def test_every_released_file_states_the_caps(app_client) -> None:  # noqa: F811
    from openpyxl import load_workbook

    c = app_client
    svc_id, _admin, _client, bearer = _world(c)
    h = {"Authorization": f"Bearer {bearer}"}
    deliv = c.get(f"/zt/services/{svc_id}/deliverables/latest", headers=h)
    assert deliv.status_code == 200, deliv.text
    body = deliv.json()

    def _bytes(key: str) -> bytes:
        dl = c.get(f"/artifacts/{body[key]}/download", headers=h)
        assert dl.status_code == 200, dl.text
        return dl.content

    assert C1 in pdf_text(_bytes("pdf_artifact_id"))
    assert C1 in docx_text(_bytes("docx_artifact_id"))
    wb = load_workbook(io.BytesIO(_bytes("xlsx_artifact_id")))
    cells = [v for row in wb["Gap Plan"].iter_rows(values_only=True) for v in row if v]
    assert C1 in cells
