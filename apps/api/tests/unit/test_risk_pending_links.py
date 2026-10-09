"""#415: an ATT&CK technique pending review is not citable in the Risk Register.

Gene's ruling (#736 5986057990 item 14; built under 6072976838 item 1): ATT&CK
techniques pending review are EXCLUDED from Risk's `valid_techniques`, and the
exclusion is disclosed with its own count, separate from "unscored".

"Pending review" is `attack/pending.py::pending_codes`, the set the ATT&CK
heatmap withholds (#102). These tests take the ATT&CK API's own per-row
`pending_review` flag as the oracle for which techniques are pending, so the
Risk Register is held to what the ATT&CK screen shows, not to a second reading
of the rule.

The world is an ATT&CK assessment approved before R3 (`status_rules` 1, as
migration 0059 backfilled every such assessment): under R3 an awaiting-review
tool is scored as not in place, so a computed leaf is a gap rather than a
pending claim, and `pending_codes` is empty. The pending rows carry the shape
`resolve_citations` writes for a substring rescue.

Through generate, the admin response and the downloaded files; expected text
is the approved copy (#736 6069843323, C1 and C2; (a) and (b) in 6072976838).
"""

from __future__ import annotations

import io
import json
import uuid

import pytest

from tests._attack_rows import standalone_rows
from tests.unit.test_risk_register import (  # noqa: F401  (app_client is a fixture)
    _admin,
    _generate,
    _pdf_text,
    _session,
    app_client,
)

pytestmark = pytest.mark.unit

_PENDING_ENTRY = {
    "tool": "Tool A",
    "cited": "tool",
    "reason": "substring",
    "field": "detection_tools",
    "cleared_at": None,
}
_NOT_RECORDED = (
    "Whether any linked technique was pending review was not recorded for this register."
)


def _c1(n: int) -> str:
    if n == 1:
        return (
            "1 scored ATT&CK technique is pending review and is not linked: its status "
            "rests on evidence not yet confirmed."
        )
    return (
        f"{n} scored ATT&CK techniques are pending review and are not linked: their "
        "status rests on evidence not yet confirmed."
    )


def _h(bearer: str, cid: str) -> dict:
    return {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}


def _seed_zt(c, bearer: str, cid: str) -> None:
    """The Zero Trust half the gate needs (`has_attack and (has_csf or has_zt)`)."""
    h = _h(bearer, cid)
    zsvc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": "ZT"})
    za = c.post(f"/zt/services/{zsvc.json()['id']}/assessments", headers=h)
    zans = za.json()["answers"][0]
    r = c.patch(f"/zt/answers/{zans['id']}", headers=h, json={"maturity_stage": 1})
    assert r.status_code == 200, r.text
    zr = c.post(f"/zt/assessments/{za.json()['id']}/approve", headers=h)
    assert zr.status_code == 200, zr.text


def _set_row(row, status: str, tool: str, citations: list) -> None:
    row.status = status
    row.detection_tools = [tool]
    row.prevention_tools = [tool]
    row.response_tools = [tool]
    row.unconfirmed_citations = citations


class _World:
    """An approved ATT&CK assessment under stored statuses, one gap finding,
    one confirmed covered technique, and whichever pending rows a test asks
    for. `pending` is the ATT&CK API's own `pending_review` set, read back."""

    def __init__(self, client_fixture, *, pending_leaf: bool, pending_parent: bool) -> None:
        from app.attack.parents import PARENT_CHILDREN
        from app.models.attack_assessment import AttackAssessment, AttackCoverage

        c, self.provider = client_fixture
        self.c = c
        self.bearer, self.cid = _admin(c)
        h = _h(self.bearer, self.cid)
        asvc = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": "A"})
        assert asvc.status_code in (200, 201), asvc.text
        self.attack_service = asvc.json()["id"]
        a = c.post(f"/attack/services/{self.attack_service}/assessments", headers=h).json()
        gap, confirmed, leaf = standalone_rows(a["coverage"], 3)
        self.gap, self.confirmed, self.leaf = (
            gap["technique_code"],
            confirmed["technique_code"],
            leaf["technique_code"],
        )
        r = c.patch(f"/attack/coverage/{gap['id']}", headers=h, json={"status": "gap"})
        assert r.status_code == 200, r.text
        ap = c.post(f"/attack/assessments/{a['id']}/approve", headers=h)
        assert ap.status_code == 200, ap.text
        _seed_zt(c, self.bearer, self.cid)

        self.parent = next(iter(PARENT_CHILDREN))
        children = PARENT_CHILDREN[self.parent]
        by_code = {r["technique_code"]: r["id"] for r in a["coverage"]}
        with _session() as s:
            assessment = s.get(AttackAssessment, uuid.UUID(a["id"]))
            assessment.status_rules = 1  # approved before R3 (migration 0059)
            assessment.parent_rules = 2  # approved after #620: parents computed

            def row(code: str) -> AttackCoverage:
                return s.get(AttackCoverage, uuid.UUID(by_code[code]))

            _set_row(row(self.confirmed), "covered", "Tool A", [])
            if pending_leaf:
                _set_row(row(self.leaf), "covered", "Tool A", [dict(_PENDING_ENTRY)])
            if pending_parent:
                # The parent's OWN evidence is confirmed; it is pending only
                # because a child is, which only `pending_codes` knows (D-094).
                _set_row(row(self.parent), "covered", "Tool P", [])
                for i, child in enumerate(children):
                    cites = [dict(_PENDING_ENTRY)] if i == 0 else []
                    _set_row(row(child), "covered", "Tool A", cites)
            s.commit()

        latest = c.get(f"/attack/services/{self.attack_service}/assessments/latest", headers=h)
        assert latest.status_code == 200, latest.text
        self.pending = {
            r["technique_code"] for r in latest.json()["coverage"] if r["pending_review"]
        }

    def generate(self, *linked: str) -> dict:
        entry = {
            "title": "Credential theft exposure",
            "description": "d",
            "axis": "detection",
            "source": "coverage_finding",
            "source_id": self.gap,
            "linked_techniques": list(linked),
            "linked_controls": [],
            "likelihood": "high",
            "impact": "major",
            "recommended_action": "remediate",
            "rationale": "r",
        }
        return _generate(
            self.c, self.provider, self.bearer, self.cid, json.dumps({"entries": [entry]})
        )

    def latest_register(self) -> dict:
        r = self.c.get(
            f"/risk/clients/{self.cid}/register/latest",
            headers={"Authorization": f"Bearer {self.bearer}"},
        )
        assert r.status_code == 200, r.text
        return r.json()

    def provenance(self) -> dict:
        from sqlalchemy import select

        from app.models.risk_register import RiskRegister

        with _session() as s:
            return dict(s.execute(select(RiskRegister)).scalar_one().provenance)

    def exports(self) -> dict:
        from docx import Document
        from openpyxl import load_workbook

        h = _h(self.bearer, self.cid)
        ex = self.c.post(f"/risk/clients/{self.cid}/register/export", headers=h)
        assert ex.status_code == 200, ex.text
        raw = {
            k: self.c.get(f"/artifacts/{ex.json()[f'{k}_artifact_id']}/download", headers=h).content
            for k in ("pdf", "docx", "xlsx")
        }
        doc = Document(io.BytesIO(raw["docx"]))
        wb = load_workbook(io.BytesIO(raw["xlsx"]))
        return {
            "pdf": " ".join(_pdf_text(raw["pdf"]).split()),
            "docx": " ".join(" ".join(p.text for p in doc.paragraphs).split()),
            "xlsx_summary": [r[0] for r in wb["Summary"].iter_rows(values_only=True) if r[0]],
            "xlsx_scope": list(wb["Scored coverage"].iter_rows(values_only=True)),
        }


def _linked(body: dict) -> set[str]:
    (entry,) = body["entries"]
    return set(entry["linked_techniques"])


def test_a_pending_technique_is_not_linked_and_a_confirmed_one_is(app_client) -> None:  # noqa: F811
    w = _World(app_client, pending_leaf=True, pending_parent=False)
    assert w.pending == {w.leaf}, w.pending  # the world as the ATT&CK screen shows it
    body = w.generate(w.confirmed, w.leaf)
    assert _linked(body) == {w.confirmed}


def test_a_computed_parent_pending_through_its_child_is_not_linked(
    app_client,  # noqa: F811
) -> None:
    """A parent whose own evidence is confirmed but whose child's is not is
    pending under D-094 (`pending_codes`). A per-row reading of the parent's
    own citations would let it through."""
    w = _World(app_client, pending_leaf=False, pending_parent=True)
    assert w.parent in w.pending, w.pending
    body = w.generate(w.confirmed, w.parent)
    assert _linked(body) == {w.confirmed}


def test_the_count_is_recorded_and_reaches_the_admin_response(app_client) -> None:  # noqa: F811
    w = _World(app_client, pending_leaf=True, pending_parent=True)
    n = len(w.pending)
    assert n >= 3, w.pending  # the leaf, the parent and its pending child
    w.generate(w.confirmed)
    assert w.provenance()["link_scope"]["attack"]["pending_review"] == n
    rows = {r["service"]: r for r in w.latest_register()["excluded_unscored_links"]}
    assert rows["attack"]["pending_review"] == n
    # Not an ATT&CK review queue, so not a count: CSF and ZT have none.
    assert rows["zt"]["pending_review"] is None


def test_every_file_states_the_pending_count(app_client) -> None:  # noqa: F811
    w = _World(app_client, pending_leaf=True, pending_parent=True)
    n = len(w.pending)
    w.generate(w.confirmed)
    files = w.exports()
    assert _c1(n) in files["pdf"]
    assert _c1(n) in files["docx"]
    assert _c1(n) in files["xlsx_summary"]
    header, *rows = files["xlsx_scope"]
    assert header == ("Assessment", "Rows scored", "Rows total", "Not citable", "Pending review")
    by_label = {r[0]: r for r in rows if r and r[0]}
    assert by_label["ATT&CK coverage"][4] == n
    assert by_label["Zero Trust"][4] == "n/a"
    assert any(r and r[0] == _c1(n) for r in rows), rows
    assert _NOT_RECORDED not in files["pdf"]


def test_one_pending_technique_reads_in_the_singular(app_client) -> None:  # noqa: F811
    w = _World(app_client, pending_leaf=True, pending_parent=False)
    assert len(w.pending) == 1, w.pending
    w.generate(w.confirmed)
    files = w.exports()
    assert _c1(1) in files["pdf"]
    assert _c1(1) in files["docx"]


def test_none_pending_prints_nothing_and_the_workbook_shows_zero(app_client) -> None:  # noqa: F811
    """Ruling (a): n = 0 prints nothing in the summary; the XLSX shows 0."""
    w = _World(app_client, pending_leaf=False, pending_parent=False)
    assert w.pending == set(), w.pending
    w.generate(w.confirmed)
    files = w.exports()
    assert "Scored coverage available to link" in files["pdf"]  # the positive state first
    assert "pending review" not in files["pdf"]
    assert "pending review" not in files["docx"]
    by_label = {r[0]: r for r in files["xlsx_scope"][1:] if r and r[0]}
    assert by_label["ATT&CK coverage"][4] == 0


def test_a_register_from_before_this_says_the_count_was_not_recorded(
    app_client,  # noqa: F811
) -> None:
    """Ruling (b). Built from a real register with the count taken out of its
    record, the state every register generated before this is in."""
    from sqlalchemy import select

    from app.models.risk_register import RiskRegister

    w = _World(app_client, pending_leaf=True, pending_parent=False)
    w.generate(w.confirmed)
    with _session() as s:
        reg = s.execute(select(RiskRegister)).scalar_one()
        prov = dict(reg.provenance)
        scope = {k: dict(v) for k, v in prov["link_scope"].items()}
        assert scope["attack"].pop("pending_review") == 1  # the real record first
        prov["link_scope"] = scope
        reg.provenance = prov
        s.commit()
    rows = {r["service"]: r for r in w.latest_register()["excluded_unscored_links"]}
    assert rows["attack"]["pending_review"] is None
    files = w.exports()
    assert _NOT_RECORDED in files["pdf"]
    assert _NOT_RECORDED in files["docx"]
    assert "pending review and is not linked" not in files["pdf"]
    by_label = {r[0]: r for r in files["xlsx_scope"][1:] if r and r[0]}
    assert by_label["ATT&CK coverage"][4] is None
    assert any(r and r[0] == _NOT_RECORDED for r in files["xlsx_scope"]), files["xlsx_scope"]


@pytest.mark.parametrize(
    "bad",
    [-1, 999999, "2", True, 1.0, None],
    ids=["negative", "above-scored", "string", "bool", "float", "null"],
)
def test_an_unreadable_pending_count_makes_the_scope_unreadable(bad, capsys) -> None:
    """The reader fails closed on the new key the way it does on `scored`."""
    from app.routes.risk import _link_scope_fields

    stored = {"link_scope": {"attack": {"scored": 5, "total": 10, "pending_review": bad}}}
    got = _link_scope_fields(stored)
    assert got["excluded_unscored_links_recorded"] is False
    assert got["excluded_unscored_links"] == []
    assert "risk_register_link_scope_unreadable" in capsys.readouterr().out
