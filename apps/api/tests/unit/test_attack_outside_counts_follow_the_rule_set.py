"""#621's additions appear only where #620's rule set says the new rules apply.

Option (a), decided by Gene's advisor on 2026-09-26: the "Not verified N,
Outside control surface M" counts, the columns, cards and chips that carry them,
and #621's rewording of the coverage definition render only for an assessment
that `attack/rules.py::parents_computed` says is under D-094's rules -- approved
after #620 (parent_rules = 2) or a draft (NULL). An assessment approved before
#620 (parent_rules = 1) renders exactly what was delivered.

BOTH HALVES, from one world. The rule-1 half compares against text RECORDED from
`main` at 22c47a4, before #621 (`tests/golden/outside_counts_rule1/`, see its
RECORDED.md), never regenerated. The rule-2 and draft halves assert the counts
appear, so the rule-1 half cannot pass for a change that removed them everywhere.

The client dashboard JSON and the finalize workbook and summary under rule 1 are
pinned by #620's `test_attack_parent_rules_golden.py`, unchanged.
"""

from __future__ import annotations

import io
import json
import uuid
from pathlib import Path

import pytest
from sqlalchemy import update

from app.attack.analytics import compute
from app.attack.catalog import TECHNIQUES
from app.attack.exporters import build_context, render_docx, render_pdf, render_xlsx
from app.attack.pending import pending_codes
from app.models.attack_assessment import (
    AttackAssessment,
    AttackAssessmentStatus,
    AttackCoverage,
)
from tests._attack_rows import standalone_rows
from tests.unit.test_attack_catalog_version_guard import (  # noqa: F401  (fixture)
    _auth,
    _register,
    _service_and_assessment,
    env,
)

pytestmark = pytest.mark.unit

GOLDEN = Path(__file__).resolve().parents[1] / "golden" / "outside_counts_rule1"
SENTENCE_ZERO = "Not verified 0, Outside control surface 0"


# --- the renderers -----------------------------------------------------------


def world(parent_rules: int | None):
    """A world a rule-1 assessment can hold: the four legacy statuses, a pending
    claim, unscored techniques. Built from the catalogue in its own order, so
    the same world is built on `main` for the recording."""
    a = AttackAssessment(
        id=uuid.UUID(int=1),
        service_id=uuid.UUID(int=2),
        version=3,
        status=AttackAssessmentStatus.APPROVED,
        parent_rules=parent_rules,
    )
    cycle = ["covered", "partial", "gap", "not_applicable", None]
    coverage = []
    for i, t in enumerate(TECHNIQUES):
        row = AttackCoverage(
            id=uuid.UUID(int=1000 + i),
            assessment_id=a.id,
            technique_code=t.id,
            status=cycle[i % len(cycle)],
            rationale=f"rationale {i}",
            detection_tools=["Tool A"] if i % 3 == 0 or i == 5 else [],
            prevention_tools=[],
            response_tools=[],
            unconfirmed_citations=[],
        )
        if i == 5:  # a covered claim whose only tool is inferred: pending (#102)
            row.unconfirmed_citations = [{"tool": "Tool A", "field": "detection_tools"}]
        coverage.append(row)
    from app.attack.rules import parents_computed

    rollup = compute(
        {c.technique_code: c.status for c in coverage},
        pending_codes(coverage, parents_computed=parents_computed(a)),
    )
    return build_context(
        client_legal_name="Golden Client",
        service_title="ATT&CK Coverage",
        assessment=a,
        coverage=coverage,
        rollup=rollup,
    )


def docx_text(raw: bytes) -> list[str]:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    return [p.text for p in doc.paragraphs] + [
        " | ".join(cell.text for cell in row.cells) for t in doc.tables for row in t.rows
    ]


def pdf_text(raw: bytes) -> list[str]:
    from pypdf import PdfReader

    return [page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages]


def xlsx_cells(raw: bytes) -> dict[str, list[list]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    return {
        ws.title: [[cell.value for cell in row] for row in ws.iter_rows()] for ws in wb.worksheets
    }


def record(out: Path) -> None:
    """Run ONCE, on `main` before #621, to write the golden files."""
    ctx = world(1)
    out.mkdir(parents=True, exist_ok=True)
    for name, value in (
        ("docx.json", docx_text(render_docx(ctx))),
        ("pdf.json", pdf_text(render_pdf(ctx))),
        ("xlsx.json", xlsx_cells(render_xlsx(ctx))),
    ):
        (out / name).write_text(
            json.dumps(value, indent=1, ensure_ascii=False, default=str) + "\n", encoding="utf-8"
        )


def _golden(name: str):
    return json.loads((GOLDEN / name).read_text(encoding="utf-8"))


def _jsonable(value):
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))


def test_under_rule_1_the_docx_is_what_main_delivered() -> None:
    assert docx_text(render_docx(world(1))) == _golden("docx.json")


def _words(pages: list[str]) -> list[str]:
    return [" ".join(page.split()) for page in pages]


def test_under_rule_1_the_pdf_is_what_main_delivered() -> None:
    # Compared word for word, not byte for byte: pypdf's layout whitespace
    # differs between the container the golden was recorded in and the CI
    # runner (" Tactic" vs "Tactic" on CI, first run of 1b81c60), with the
    # words identical. The words are what the client reads.
    assert _words(pdf_text(render_pdf(world(1)))) == _words(_golden("pdf.json"))


def test_under_rule_1_the_xlsx_is_what_main_delivered() -> None:
    assert _jsonable(xlsx_cells(render_xlsx(world(1)))) == _golden("xlsx.json")


@pytest.mark.parametrize("parent_rules", [2, None], ids=["rule-2", "draft"])
def test_under_the_new_rules_every_renderer_states_the_counts(parent_rules) -> None:
    ctx = world(parent_rules)
    assert SENTENCE_ZERO in " ".join(" ".join(docx_text(render_docx(ctx))).split())
    assert SENTENCE_ZERO in " ".join(" ".join(pdf_text(render_pdf(ctx))).split())
    summary = {r[0]: r[1] for r in xlsx_cells(render_xlsx(ctx))["Heatmap Summary"] if r}
    assert (summary.get("Not verified"), summary.get("Outside control surface")) == (0, 0)


# --- the stored summary, the admin heatmap, the client dashboard, the card ---


def _unverified_world(c, Sess, bearer: str) -> tuple[str, dict]:
    """A draft with one Not verified row, written straight to the database (no
    writer may produce the status yet, D-092 Decision 5)."""
    svc, a = _service_and_assessment(c, bearer)
    (row,) = standalone_rows(a["coverage"], 1)
    with Sess() as s:
        s.execute(
            update(AttackCoverage)
            .where(AttackCoverage.id == uuid.UUID(row["id"]))
            .values(status="unable_to_determine", unconfirmed_citations=[])
        )
        s.commit()
    return svc, a


def _set_rule(Sess, assessment_id: str, value: int | None) -> None:
    with Sess() as s:
        s.execute(
            update(AttackAssessment)
            .where(AttackAssessment.id == uuid.UUID(assessment_id))
            .values(parent_rules=value)
        )
        s.commit()


def test_the_admin_heatmap_carries_the_counts_for_a_draft_and_not_under_rule_1(
    env,  # noqa: F811
) -> None:
    c, Sess = env
    bearer = _register(c, "admin@example.com")["tokens"]["access_token"]
    svc, a = _unverified_world(c, Sess, bearer)

    draft = c.get(f"/attack/services/{svc}/heatmap", headers=_auth(bearer)).json()
    assert (draft["unable_to_determine"], draft["outside_control_surface"]) == (1, 0)
    assert all(t["unable_to_determine"] is not None for t in draft["by_tactic"])

    _set_rule(Sess, a["id"], 1)
    old = c.get(f"/attack/services/{svc}/heatmap", headers=_auth(bearer)).json()
    assert (old["unable_to_determine"], old["outside_control_surface"]) == (None, None)
    assert {t["unable_to_determine"] for t in old["by_tactic"]} == {None}
    assert {t["outside_control_surface"] for t in old["by_tactic"]} == {None}


def _release(c, bearer: str, svc: str, a: dict) -> str:
    assert (
        c.post(f"/attack/assessments/{a['id']}/approve", headers=_auth(bearer)).status_code == 200
    )
    fin = c.post(f"/attack/services/{svc}/deliverables/finalize", headers=_auth(bearer))
    assert fin.status_code in (200, 201), fin.text
    rel = c.post(f"/attack/deliverables/{fin.json()['id']}/release", headers=_auth(bearer))
    assert rel.status_code == 200, rel.text
    return fin.json()["summary"]


@pytest.mark.parametrize("rule", [2, 1], ids=["rule-2", "rule-1"])
def test_the_summary_dashboard_and_value_card_follow_the_rule(env, rule) -> None:  # noqa: F811
    c, Sess = env
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    svc, a = _unverified_world(c, Sess, bearer)
    if rule == 1:
        # Approve writes 2; a rule-1 assessment is one approved before #620,
        # which the migration marked 1. Set it between approve and finalize.
        assert (
            c.post(f"/attack/assessments/{a['id']}/approve", headers=_auth(bearer)).status_code
            == 200
        )
        _set_rule(Sess, a["id"], 1)
        fin = c.post(f"/attack/services/{svc}/deliverables/finalize", headers=_auth(bearer))
        assert fin.status_code in (200, 201), fin.text
        rel = c.post(f"/attack/deliverables/{fin.json()['id']}/release", headers=_auth(bearer))
        assert rel.status_code == 200, rel.text
        summary = fin.json()["summary"]
    else:
        summary = _release(c, bearer, svc, a)

    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    ch = _auth(client["tokens"]["access_token"])
    dash = c.get(f"/clients/{client_id}/attack/{svc}/dashboard", headers=ch)
    assert dash.status_code == 200, dash.text
    rollup = dash.json()["rollup"]
    card = c.get(f"/clients/{client_id}/value-summary", headers=ch)
    assert card.status_code == 200, card.text

    if rule == 2:
        assert "Not verified 1, Outside control surface 0" in summary
        assert (rollup["unable_to_determine"], rollup["outside_control_surface"]) == (1, 0)
        assert all("unable_to_determine" in t for t in rollup["by_tactic"])
        assert card.json()["attack_not_verified_count"] == 1
    else:
        assert "Not verified" not in summary
        assert "unable_to_determine" not in rollup
        assert "outside_control_surface" not in rollup
        assert not any("unable_to_determine" in t for t in rollup["by_tactic"])
        assert card.json()["attack_not_verified_count"] is None


# --- a rule-2 response missing the counts still fails loudly -------------------


def _response(*, parents_computed: bool | None, counts: int | None, omit: bool = False):
    from datetime import UTC, datetime

    from app.schemas.clients import (
        AttackDashboardResponse,
        AttackDashboardRollup,
        AttackTacticCoverage,
    )

    extra = {} if omit else {"outside_control_surface": counts, "unable_to_determine": counts}
    tactic = AttackTacticCoverage(
        tactic_id="TA0001",
        tactic_name="Initial Access",
        covered=1,
        partial=0,
        gap=0,
        not_applicable=0,
        unscored=0,
        coverage_pct=100.0,
        **extra,
    )
    return AttackDashboardResponse(
        service_id=uuid.UUID(int=1),
        service_title="ATT&CK Coverage",
        released_at=datetime(2026, 9, 1, tzinfo=UTC),
        deliverable_version=1,
        parents_computed=parents_computed,
        rollup=AttackDashboardRollup(
            total_evaluated=1,
            covered=1,
            partial=0,
            gap=0,
            not_applicable=0,
            coverage_pct=100.0,
            by_tactic=[tactic],
            **extra,
        ),
        techniques=[],
    )


def test_a_response_built_without_the_count_fields_is_refused() -> None:
    # REQUIRED fields, no default: forgetting to wire them fails at build.
    from pydantic import ValidationError

    with pytest.raises(ValidationError, match="unable_to_determine"):
        _response(parents_computed=True, counts=None, omit=True)


def test_a_rule_2_response_with_the_counts_missing_is_refused() -> None:
    # Omitting under rule 1 must not become a way for rule 2 to drop them.
    from pydantic import ValidationError

    with pytest.raises(ValidationError, match="missing on a rule-2 response"):
        _response(parents_computed=True, counts=None)


def test_a_rule_1_response_carrying_the_counts_is_refused() -> None:
    from pydantic import ValidationError

    with pytest.raises(ValidationError, match="present on a rule-1 response"):
        _response(parents_computed=None, counts=0)


def test_both_well_formed_responses_build() -> None:
    new = _response(parents_computed=True, counts=0).model_dump(mode="json")
    old = _response(parents_computed=None, counts=None).model_dump(mode="json")
    assert new["rollup"]["unable_to_determine"] == 0
    assert "unable_to_determine" not in old["rollup"]
