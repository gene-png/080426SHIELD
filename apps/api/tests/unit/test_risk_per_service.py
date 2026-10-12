"""#876, Gene's ruling (#736 6018510340, advisor 6019425290): the Risk Register
takes findings from EVERY engaged service, not one per kind.

The case the ruling names is Zero Trust: a client engaged for both CISA ZTMM
2.0 and DoD ZT Reference Architecture gets both frameworks' findings. Until
#876 synthesis read one assessment per kind, so one framework was dropped in
silence -- a client-number gap, not a label.

Two engaged services of the same kind AND framework would produce the same
finding codes, and `source_id` is the key every per-finding record joins on, so
generate refuses that case (Q2 option (a)) with the advisor's sentence.

Everything goes through the real routes; inputs are released through the real
finalize -> release endpoints (`tests/_risk_inputs.py`).

RR-1 (#1122, reversing #876 Q1): a register draws on ONE Zero Trust framework,
chosen at generate. A client with both frameworks now generates once per
framework. Tests that pinned BOTH frameworks in one register are REPLACED with
RR-1's claim, and say so (CLAUDE.md core principle 3). Tests that only needed a
two-framework client generate twice, once per framework, and assert on each
real register. No test builds a merged register: no writer produces one.
"""

from __future__ import annotations

import io
import json

import pytest

from app.ai.llm import LLMResponse
from tests._risk_inputs import release, seed_attack_and_zt
from tests.unit.test_risk_register import (  # noqa: F401  (app_client is a fixture)
    _admin,
    _pdf_text,
    _session,
    app_client,
)

pytestmark = pytest.mark.unit

_DUPLICATE = (
    "Two engaged services of the same kind would produce the same findings: {titles}. "
    "The register cannot be generated while both are engaged."
)


def _h(bearer: str, cid: str) -> dict:
    return {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}


def _seed_dod(c, bearer: str, cid: str, *, approve: bool = True) -> tuple[str, str]:
    """A DoD ZT service with one capability at stage 1 (below every target)."""
    h = _h(bearer, cid)
    svc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_dod", "title": "ZT DoD"})
    assert svc.status_code in (200, 201), svc.text
    a = c.post(f"/zt/services/{svc.json()['id']}/assessments", headers=h)
    assert a.status_code in (200, 201), a.text
    ans = a.json()["answers"][0]
    r = c.patch(f"/zt/answers/{ans['id']}", headers=h, json={"maturity_stage": 1})
    assert r.status_code == 200, r.text
    if approve:
        ap = c.post(f"/zt/assessments/{a.json()['id']}/approve", headers=h)
        assert ap.status_code == 200, ap.text
    return svc.json()["id"], ans["capability_code"]


def _per_finding(seen: list[dict]):
    """One entry per finding, as the real prompt asks; records each payload."""

    def fn(payload: dict) -> LLMResponse:
        seen.append(payload)
        entries = [
            {
                "title": f"Risk for {sid}",
                "description": "d",
                "axis": "detection",
                "source": f["source"],
                "source_id": sid,
                "linked_techniques": [],
                "linked_controls": [],
                "likelihood": "high",
                "impact": "major",
                "recommended_action": "remediate",
                "rationale": "r",
            }
            for sid, f in payload.get("findings", {}).items()
        ]
        return LLMResponse(json.dumps({"entries": entries}))

    return fn


def _generate(c, provider, bearer: str, cid: str, seen: list[dict]):
    provider.register("risk_synthesize", _per_finding(seen))
    return c.post(f"/risk/clients/{cid}/register/generate", headers=_h(bearer, cid))


CISA = "cisa_ztmm_2_0"
DOD = "dod_ztra"


def _generate_zt(c, provider, bearer: str, cid: str, seen: list[dict], framework: str):
    """One generate naming its Zero Trust framework (RR-1), no CSF."""
    provider.register("risk_synthesize", _per_finding(seen))
    return c.post(
        f"/risk/clients/{cid}/register/generate",
        headers=_h(bearer, cid),
        json={"include": {"csf": False, "zero_trust": framework}},
    )


def _both_frameworks(c, bearer: str, cid: str):
    s = seed_attack_and_zt(c, bearer, cid)
    dod_svc, dod_code = _seed_dod(c, bearer, cid)
    return s, dod_svc, dod_code


def _export_texts(c, bearer: str, cid: str) -> dict[str, str]:
    from docx import Document
    from openpyxl import load_workbook

    h = _h(bearer, cid)
    ex = c.post(f"/risk/clients/{cid}/register/export", headers=h)
    assert ex.status_code == 200, ex.text
    body = ex.json()
    raw = {
        k: c.get(f"/artifacts/{body[f'{k}_artifact_id']}/download", headers=h).content
        for k in ("pdf", "docx", "xlsx")
    }
    wb = load_workbook(io.BytesIO(raw["xlsx"]))
    xlsx = "\n".join(
        str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v
    )
    doc = Document(io.BytesIO(raw["docx"]))
    docx = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    )
    return {"pdf": " ".join(_pdf_text(raw["pdf"]).split()), "docx": docx, "xlsx": xlsx}


def test_with_both_frameworks_each_register_holds_only_its_frameworks_findings(
    app_client,  # noqa: F811
) -> None:
    """REPLACES #876's `test_both_zero_trust_frameworks_feed_the_register`,
    whose premise (both frameworks' findings in ONE register) RR-1 reverses.
    CISA and DoD both engaged: each register, generated with one framework,
    holds that framework's finding and none of the other's, and every export
    carries it."""
    c, provider = app_client
    bearer, cid = _admin(c)
    s, _dod_svc, dod_code = _both_frameworks(c, bearer, cid)
    assert s.capability.startswith("CISA.") and dod_code.startswith("DOD."), (s, dod_code)
    for framework, mine, theirs in ((CISA, s.capability, dod_code), (DOD, dod_code, s.capability)):
        seen: list[dict] = []
        r = _generate_zt(c, provider, bearer, cid, seen, framework)
        assert r.status_code == 201, r.text
        sources = {e["source_id"] for e in r.json()["entries"]}
        assert s.technique in sources, sources  # the positive state first
        assert mine in sources, (framework, sources)
        assert theirs not in sources, (framework, sources)
        texts = _export_texts(c, bearer, cid)
        for kind, text in texts.items():
            assert mine in text, (framework, kind)


def test_one_zero_trust_framework_feeds_only_itself(app_client) -> None:  # noqa: F811
    """CISA only: the register has CISA findings and no DoD finding at all."""
    c, provider = app_client
    bearer, cid = _admin(c)
    s = seed_attack_and_zt(c, bearer, cid)
    seen: list[dict] = []
    r = _generate(c, provider, bearer, cid, seen)
    assert r.status_code == 201, r.text
    sources = {e["source_id"] for e in r.json()["entries"]}
    assert s.capability in sources, sources
    assert not [x for x in sources if x.startswith("DOD.")], sources


def test_each_framework_is_held_to_its_own_target(app_client) -> None:  # noqa: F811
    """The CISA engagement targets stage 4; DoD has no target, so it takes the
    default 3, the top of its three-stage ladder. A capability at stage 3 is
    therefore a finding under CISA and none under DoD. Resolving the target
    once per KIND would hold DoD to CISA's 4 and invent a DoD finding."""
    from tests.unit.test_risk_register import _set_zt_stage, _set_zt_target

    c, provider = app_client
    bearer, cid = _admin(c)
    s, _dod_svc, dod_code = _both_frameworks(c, bearer, cid)
    _set_zt_target(cid, 4, kind="zero_trust_cisa")
    _set_zt_stage(cid, s.capability, 3)
    _set_zt_stage(cid, dod_code, 3)
    # RR-1: one register per framework, each held to its own target.
    seen: list[dict] = []
    gen = _generate_zt(c, provider, bearer, cid, seen, CISA)
    assert gen.status_code == 201, gen.text
    sources = {e["source_id"] for e in gen.json()["entries"]}
    assert s.capability in sources, sources  # the positive state first
    gen = _generate_zt(c, provider, bearer, cid, seen, DOD)
    assert gen.status_code == 201, gen.text
    sources = {e["source_id"] for e in gen.json()["entries"]}
    assert s.technique in sources, sources  # the positive state first
    assert dod_code not in sources, sources


def test_with_both_frameworks_no_request_carries_the_other_framework(
    app_client,  # noqa: F811
) -> None:
    """REPLACES #876's `test_no_batch_mixes_two_zero_trust_services`. Its
    premise (both frameworks' findings in one generate, kept in separate
    batches) RR-1 reverses: a generate reads ONE framework, so no request of
    that generate carries the other framework's findings at all."""
    c, provider = app_client
    bearer, cid = _admin(c)
    _both_frameworks(c, bearer, cid)
    for framework, mine, theirs in ((CISA, "CISA.", "DOD."), (DOD, "DOD.", "CISA.")):
        seen: list[dict] = []
        assert _generate_zt(c, provider, bearer, cid, seen, framework).status_code == 201
        sids = [sid for p in seen for sid in p["findings"]]
        assert [x for x in sids if x.startswith(mine)], (framework, sids)  # positive first
        assert not [x for x in sids if x.startswith(theirs)], (framework, sids)


def test_a_finding_carries_its_own_services_state(app_client) -> None:  # noqa: F811
    """CISA released, DoD approved: only the DoD finding is labelled, because
    the label comes from the finding's own service, never from its kind."""
    c, provider = app_client
    bearer, cid = _admin(c)
    s, _dod_svc, dod_code = _both_frameworks(c, bearer, cid)
    release(c, bearer, cid, "zt", s.zt_service)
    # RR-1: one register per framework; each finding still carries its own
    # service's state.
    seen: list[dict] = []
    r = _generate_zt(c, provider, bearer, cid, seen, DOD)
    assert r.status_code == 201, r.text
    states = {e["source_id"]: e["source_state"] for e in r.json()["entries"]}
    assert states[dod_code] == "approved", states
    r = _generate_zt(c, provider, bearer, cid, seen, CISA)
    assert r.status_code == 201, r.text
    states = {e["source_id"]: e["source_state"] for e in r.json()["entries"]}
    assert s.capability in states, states  # the positive state first
    assert states[s.capability] is None, states


@pytest.mark.parametrize(
    "framework, name",
    [(CISA, "Zero Trust (CISA ZTMM 2.0)"), (DOD, "Zero Trust (DoD ZT Reference Architecture)")],
)
def test_the_provenance_and_scored_coverage_name_the_chosen_framework(
    app_client, framework: str, name: str  # noqa: F811
) -> None:
    """REPLACES #876's `test_the_provenance_and_scored_coverage_name_each_framework`,
    whose register held both frameworks' inputs and scored-coverage rows; RR-1
    reverses that. Per register: ONE Zero Trust input, of the chosen framework;
    its scored-coverage key names the framework, because both frameworks have
    a record (design note §3.3 step 5); and the XLSX "Scored coverage" sheet
    labels that row with the framework (§6.2). The label alone also appears in
    copy 8, so the needle is the sheet's ROW, not the bare label."""
    from openpyxl import load_workbook
    from sqlalchemy import select

    from app.models.risk_register import RiskRegister

    c, provider = app_client
    bearer, cid = _admin(c)
    _both_frameworks(c, bearer, cid)
    seen: list[dict] = []
    assert _generate_zt(c, provider, bearer, cid, seen, framework).status_code == 201
    with _session() as db:
        prov = db.execute(select(RiskRegister)).scalar_one().provenance
    zt_inputs = sorted(i["framework"] for i in prov["inputs"] if i["kind"] == "zt")
    assert zt_inputs == [framework], prov["inputs"]
    assert sorted(prov["link_scope"]) == ["attack", f"zt:{framework}"], prov["link_scope"]
    h = _h(bearer, cid)
    ex = c.post(f"/risk/clients/{cid}/register/export", headers=h)
    assert ex.status_code == 200, ex.text
    raw = c.get(f"/artifacts/{ex.json()['xlsx_artifact_id']}/download", headers=h).content
    rows = list(load_workbook(io.BytesIO(raw))["Scored coverage"].iter_rows(values_only=True))
    labels = [row[0] for row in rows]
    assert name in labels, rows


def test_a_single_zero_trust_keeps_its_unqualified_label(app_client) -> None:  # noqa: F811
    """The framework is named only when a kind has more than one row."""
    from sqlalchemy import select

    from app.models.risk_register import RiskRegister

    c, provider = app_client
    bearer, cid = _admin(c)
    seed_attack_and_zt(c, bearer, cid)
    seen: list[dict] = []
    assert _generate(c, provider, bearer, cid, seen).status_code == 201
    with _session() as db:
        prov = db.execute(select(RiskRegister)).scalar_one().provenance
    assert sorted(prov["link_scope"]) == ["attack", "zt"], prov["link_scope"]
    g = c.get(f"/risk/clients/{cid}/gate", headers=_h(bearer, cid)).json()
    assert [r["qualifier"] for r in g["inputs"] if r["kind"] == "zt"] == [None]


def test_the_inputs_panel_names_each_zero_trust_framework(app_client) -> None:  # noqa: F811
    c, _ = app_client
    bearer, cid = _admin(c)
    _both_frameworks(c, bearer, cid)
    g = c.get(f"/risk/clients/{cid}/gate", headers=_h(bearer, cid)).json()
    assert [r["qualifier"] for r in g["inputs"] if r["kind"] == "zt"] == [
        "CISA ZTMM 2.0",
        "DoD ZT Reference Architecture",
    ]


def test_two_services_of_one_kind_and_framework_refuse_generate(app_client) -> None:  # noqa: F811
    """Q2 (a): two CISA services would produce the same finding codes, and
    `source_id` is the key every per-finding record joins on. Refused, typed,
    with the advisor's sentence; the gate carries the same sentence."""
    c, provider = app_client
    bearer, cid = _admin(c)
    seed_attack_and_zt(c, bearer, cid)
    h = _h(bearer, cid)
    svc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": "ZT 2"})
    assert c.post(f"/zt/services/{svc.json()['id']}/assessments", headers=h).status_code in (
        200,
        201,
    )
    expected = _DUPLICATE.format(titles="ZT and ZT 2")
    g = c.get(f"/risk/clients/{cid}/gate", headers=h).json()
    assert g["duplicate_inputs"] == expected, g
    seen: list[dict] = []
    r = _generate(c, provider, bearer, cid, seen)
    assert r.status_code == 409, r.text
    assert r.json()["error"]["reason"] == "risk_register_duplicate_inputs", r.json()
    assert r.json()["error"]["message"] == expected
    assert seen == [], "a refused generate must not reach the model"


def test_one_service_per_kind_is_not_a_duplicate(app_client) -> None:  # noqa: F811
    c, _ = app_client
    bearer, cid = _admin(c)
    _both_frameworks(c, bearer, cid)
    g = c.get(f"/risk/clients/{cid}/gate", headers=_h(bearer, cid)).json()
    assert g["unlocked"] is True and g["duplicate_inputs"] is None, g


# ---------------------------------------------------------------------------
# #891 review round 1
# ---------------------------------------------------------------------------

#: RR-1 (#1122), design note §8 copy 19 (pending the advisor): the publish
#: refusal for a register whose `selection` is absent or fails §4.
_SELECTION_NOT_RECORDED = (
    "This register has no usable record of how its services were chosen or "
    "which were left out, so it cannot be published. Regenerate it, then "
    "publish the new version."
)


def _publish(c, bearer: str, cid: str):
    return c.post(f"/risk/clients/{cid}/register/publish", headers=_h(bearer, cid))


def _released_both(c, bearer: str, cid: str):
    s, dod_svc, dod_code = _both_frameworks(c, bearer, cid)
    release(c, bearer, cid, "attack", s.attack_service)
    release(c, bearer, cid, "zt", s.zt_service)
    release(c, bearer, cid, "zt", dod_svc)
    return s, dod_svc, dod_code


def test_a_per_service_register_with_both_frameworks_publishes(app_client) -> None:  # noqa: F811
    """The positive half of B1: CISA and DoD both released. Under RR-1 each
    register draws on one framework and records the other as deselected, so
    each publishes: the deselected framework does not hold it back."""
    c, provider = app_client
    bearer, cid = _admin(c)
    _released_both(c, bearer, cid)
    for framework in (CISA, DOD):
        seen: list[dict] = []
        assert _generate_zt(c, provider, bearer, cid, seen, framework).status_code == 201
        r = _publish(c, bearer, cid)
        assert r.status_code == 200, (framework, r.text)


def test_a_register_from_before_per_service_reading_is_refused(app_client) -> None:  # noqa: F811
    """B1. A register generated BEFORE #876 for a CISA-and-DoD client recorded
    both services in `current_inputs` but drew findings from ONE, and its
    `inputs` rows carry no `service_id`. Every other check passes for it, so
    publishing it would deliver the gap #876 closes. Refused at PUBLISH,
    typed, and the consultant regenerates.

    The legacy `inputs` rows are written into the stored provenance: that is
    what the pre-#876 writer stored (one `inputs` row per kind, no service),
    and no current writer can produce it.

    RR-1 (#1122): the register KEEPS the `selection` its generate wrote, so
    the refusal comes from design note §4's check 2 (`included` must equal
    the CSF and ZT service ids in `inputs`, and these rows name none), which
    this test is now what pins. It is a register-level `not_recorded`, before
    any row is read, with §8's copy 19. The pre-#876 per-input check that
    used to answer here (`kind in SYNTHESIS_KINDS and str(service_id) not in
    drawn_from`) is now reached by no register, so no test reaches it: it is
    kept in the code as a ratchet (design note §0g)."""
    from sqlalchemy import select

    from app.models.risk_register import RiskRegister

    c, provider = app_client
    bearer, cid = _admin(c)
    _released_both(c, bearer, cid)
    seen: list[dict] = []
    assert _generate_zt(c, provider, bearer, cid, seen, CISA).status_code == 201
    with _session() as db:
        reg = db.execute(select(RiskRegister)).scalar_one()
        prov = dict(reg.provenance)
        zt_rows = [i for i in prov["inputs"] if i["kind"] == "zt"]
        assert len(zt_rows) == 1, prov["inputs"]  # the per-service shape first
        assert "selection" in prov, sorted(prov)  # kept: check 2 must refuse
        legacy = []
        for kind in ("attack", "zt"):
            row = dict(next(i for i in prov["inputs"] if i["kind"] == kind))
            row.pop("service_id")
            row.pop("framework")
            legacy.append(row)
        prov["inputs"] = legacy
        reg.provenance = prov
        db.commit()
    r = _publish(c, bearer, cid)
    assert r.status_code == 409, r.text
    err = r.json()["error"]
    assert err["reason"] == "risk_register_inputs_not_final", err
    assert [b["input"] for b in err["blockers"]] == ["register"], err
    assert {b["reason"] for b in err["blockers"]} == {"not_recorded"}, err
    assert err["message"] == _SELECTION_NOT_RECORDED


def test_a_single_service_client_sends_what_it_always_sent(app_client) -> None:  # noqa: F811
    """A4. One service per kind: one request, as before #876 (kinds still
    share a batch), and every finding carries exactly the four keys it always
    carried -- no service or source key reaches the model."""
    c, provider = app_client
    bearer, cid = _admin(c)
    s = seed_attack_and_zt(c, bearer, cid)
    seen: list[dict] = []
    assert _generate(c, provider, bearer, cid, seen).status_code == 201
    assert len(seen) == 1, len(seen)
    findings = seen[0]["findings"]
    assert set(findings) >= {s.technique, s.capability}, findings
    for f in findings.values():
        assert set(f) == {"source", "kind", "label", "evidence"}, f


def test_a_second_service_not_started_is_listed_on_the_draft(app_client) -> None:  # noqa: F811
    """A2, the advisor's condition (#736 6024072042): a second engaged ZT
    service with no assessment does not stop a draft, and the Inputs panel
    lists it, named, as not started.

    RR-1 (#1122) REPLACES A2's last clause, "publish refuses it": that
    service is recorded as INELIGIBLE at generate, and RR-1 Q4 (#736
    6131606581) discloses it (copy 9c) and does not block publish. The panel
    row says so through `in_register` (design note §6.5)."""
    c, provider = app_client
    bearer, cid = _admin(c)
    seed_attack_and_zt(c, bearer, cid)
    h = _h(bearer, cid)
    dod = c.post("/zt/services", headers=h, json={"kind": "zero_trust_dod", "title": "ZT DoD"})
    assert dod.status_code in (200, 201), dod.text
    seen: list[dict] = []
    assert _generate(c, provider, bearer, cid, seen).status_code == 201
    g = c.get(f"/risk/clients/{cid}/gate", headers=h).json()
    zt_rows = [r for r in g["inputs"] if r["kind"] == "zt"]
    assert zt_rows == [
        {
            "kind": "zt",
            "engaged": True,
            "status": "approved",
            "version": 1,
            "qualifier": "CISA ZTMM 2.0",
            "in_register": "input",
        },
        {
            "kind": "zt",
            "engaged": True,
            "status": None,
            "version": None,
            "qualifier": "DoD ZT Reference Architecture",
            "in_register": "ineligible",
        },
    ]
    blockers = _publish(c, bearer, cid).json()["error"]["blockers"]
    # Positive first: CISA is approved, not released, so publish still refuses.
    assert {"input": "zt", "reason": "not_released", "status": "approved"} in blockers, blockers
    assert {"input": "zt", "reason": "not_started", "status": None} not in blockers, blockers


def _add_cisa(c, bearer: str, cid: str, title: str) -> None:
    h = _h(bearer, cid)
    svc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": title})
    assert svc.status_code in (200, 201), svc.text
    a = c.post(f"/zt/services/{svc.json()['id']}/assessments", headers=h)
    assert a.status_code in (200, 201), a.text


def test_three_services_of_one_framework_are_counted(app_client) -> None:  # noqa: F811
    """A1 (wording pending the advisor): "Two ... both" is false of three."""
    c, _ = app_client
    bearer, cid = _admin(c)
    seed_attack_and_zt(c, bearer, cid)
    _add_cisa(c, bearer, cid, "ZT 2")
    _add_cisa(c, bearer, cid, "ZT 3")
    g = c.get(f"/risk/clients/{cid}/gate", headers=_h(bearer, cid)).json()
    assert g["duplicate_inputs"] == (
        "3 engaged services of the same kind would produce the same findings: "
        "ZT, ZT 2 and ZT 3. The register cannot be generated while more than one is engaged."
    )


def test_two_duplicate_groups_get_a_sentence_each(app_client) -> None:  # noqa: F811
    """A1: two CISA services and two ATT&CK services are two pairs, not one
    group of four, so each pair gets the approved sentence."""
    c, _ = app_client
    bearer, cid = _admin(c)
    seed_attack_and_zt(c, bearer, cid)
    h = _h(bearer, cid)
    _add_cisa(c, bearer, cid, "ZT 2")
    a2 = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": "A 2"})
    assert c.post(f"/attack/services/{a2.json()['id']}/assessments", headers=h).status_code in (
        200,
        201,
    )
    g = c.get(f"/risk/clients/{cid}/gate", headers=h).json()
    assert g["duplicate_inputs"] == (
        _DUPLICATE.format(titles="A and A 2") + " " + _DUPLICATE.format(titles="ZT and ZT 2")
    )
