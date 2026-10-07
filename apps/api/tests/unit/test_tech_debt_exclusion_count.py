"""#193 on every Tech Debt surface, from #177's one reader.

When the extraction could not attribute every item to one uploaded row, the
true number of excluded rows is unknown: `received - included` is only a FLOOR
(the rows the items actually claimed are fewer than the items), and it floors
to ZERO when items are as many as rows -- #193's silent case, which printed
"Total annual cost" over a real exclusion.

`reconcile.exclusion_count_state` decides it once; the deliverable, the client
dashboard and the admin list all call it. The worlds here are built by the
REAL extraction writer, so every state is one a list can actually reach.
"""

from __future__ import annotations

import io
import json
import uuid

import pytest
from sqlalchemy import update

from app.ai.llm import LLMResponse
from app.models.capability import CapabilityList
from tests._ai_runs import tech_debt_extract
from tests.unit.test_tech_debt_dashboard import (  # noqa: F401  (fixture)
    _register,
    app_client,
)

pytestmark = pytest.mark.unit

THREE_ROWS = b"Tool,Cost\nWiz,1\nSplunk,1\nLacework,1\n"


def _item(name: str, row: int | None) -> dict:
    return {
        "name": name,
        "vendor": name,
        "category": name,
        "function": "x",
        "annual_cost_usd": 1000,
        "license_count": 1,
        "notes": None,
        "confidence_pct": 90,
        "source_row_index": row,
    }


def _released(
    c, provider, items: list[dict], *, null_flag: bool = False, include_all: bool = False
) -> dict:
    """Extract, keep everything, approve, finalize, release. Returns the ids
    and every surface's reading of the list."""
    provider.register(
        "extract.capabilities", lambda _p: LLMResponse(content=json.dumps({"items": items}))
    )
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    cid = client["user"]["client_id"]
    h = {"Authorization": f"Bearer {bearer}"}
    svc = c.post("/tech-debt/services", headers=h, json={"title": "TD"}).json()["id"]
    art = c.post(
        "/artifacts", headers=h, files={"file": ("inv.csv", io.BytesIO(THREE_ROWS), "text/csv")}
    ).json()["id"]
    ext = tech_debt_extract(c, svc, h, art)
    if null_flag:
        # The PRE-0058 state, made the way it was made: the same writer, whose
        # `excluded_rows` logic is unchanged since 0036, with the flag the
        # column did not yet exist to hold.
        with _sessions()() as s:
            s.execute(
                update(CapabilityList)
                .where(CapabilityList.id == uuid.UUID(ext["id"]))
                .values(attribution_complete=None)
            )
            s.commit()
    if include_all:
        # The consultant pulls every named excluded row back in, through the
        # include-row endpoint, one at a time.
        for row in ext["excluded_rows"]:
            r = c.post(
                f"/tech-debt/capability-lists/{ext['id']}/excluded-rows/{row['index']}/include",
                headers=h,
                json={"name": f"row {row['index']}", "annual_cost_usd": 1000},
            )
            assert r.status_code == 201, r.text
        ext = c.get(f"/tech-debt/services/{svc}/capability-lists/latest", headers=h).json()
        assert ext["excluded_rows"] == []
    for it in ext["items"]:
        c.patch(f"/tech-debt/capability-items/{it['id']}", headers=h, json={"disposition": "keep"})
    if not include_all:
        # #850: approve refuses an excluded row nobody confirmed. A confirmed
        # row stays in `excluded_rows`, so every count asserted here is the same.
        for row in ext["excluded_rows"]:
            r = c.post(
                f"/tech-debt/capability-lists/{ext['id']}/excluded-rows/{row['index']}/confirm",
                headers=h,
            )
            assert r.status_code == 200, r.text
    assert c.post(f"/tech-debt/capability-lists/{ext['id']}/approve", headers=h).status_code == 200
    fin = c.post(f"/tech-debt/services/{svc}/deliverables/finalize", headers=h)
    assert fin.status_code == 201, fin.text
    body = fin.json()
    rel = c.post(f"/tech-debt/deliverables/{body['id']}/release", headers=h)
    assert rel.status_code == 200, rel.text

    def _bytes(key: str) -> bytes:
        return c.get(f"/artifacts/{body[key]}/download", headers=h).content

    admin_list = c.get(f"/tech-debt/services/{svc}/capability-lists/latest", headers=h).json()
    c.headers["X-Client-Id"] = cid
    dash = c.get(
        f"/clients/{cid}/tech-debt/{svc}/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    ).json()
    return {
        "pdf": pdf_text(_bytes("pdf_artifact_id")),
        "docx": docx_text(_bytes("docx_artifact_id")),
        "xlsx": xlsx_text(_bytes("xlsx_artifact_id")),
        "admin": admin_list,
        "dash": dash,
    }


UNKNOWN_TAIL = "the AI could not match every extracted capability to one uploaded row"


def _sessions():
    import os

    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    return sessionmaker(bind=create_engine(os.environ["DATABASE_URL"], future=True))


def pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return " ".join(" ".join((p.extract_text() or "") for p in reader.pages).split())


def xlsx_text(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    cells = (str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row)
    return " ".join(" ".join(c for c in cells if c != "None").split())


def docx_text(raw: bytes) -> str:
    from docx import Document

    return " ".join(" ".join(p.text for p in Document(io.BytesIO(raw)).paragraphs).split())


def test_the_silent_case_says_the_count_is_unknown_everywhere(app_client) -> None:  # noqa: F811
    """#193: three rows, three items, two claiming row 0. K floors to 0."""
    c, provider = app_client
    got = _released(c, provider, [_item("Wiz", 0), _item("Wiz CNAPP", 0), _item("Splunk", 1)])
    line = f"3 rows received · 3 included · excluded count unknown: {UNKNOWN_TAIL}"
    for surface in ("pdf", "docx", "xlsx"):
        assert line in got[surface], surface
        assert "Total annual cost" not in got[surface], surface
        assert "Annual cost (may not be complete)" in got[surface], surface
    assert got["admin"]["exclusion_count_state"] == "unknown"
    assert got["dash"]["excluded_count_exact"] is False
    assert got["dash"]["spend_completeness"] == "partial"


def test_k_is_a_floor_when_the_count_is_unknown(app_client) -> None:  # noqa: F811
    """Three rows, two items, one naming no row: at least one row was excluded,
    and maybe two. K is printed as a floor, never as the count."""
    c, provider = app_client
    got = _released(c, provider, [_item("Wiz", 0), _item("Splunk", None)])
    line = f"3 rows received · 2 included · excluded count unknown (at least 1): {UNKNOWN_TAIL}"
    for surface in ("pdf", "docx", "xlsx"):
        assert line in got[surface], surface
        assert "1 excluded" not in got[surface].replace("(at least 1)", ""), surface
    assert got["admin"]["exclusion_count_state"] == "unknown"
    assert got["dash"]["excluded_count_exact"] is False
    assert got["dash"]["excluded_count"] == 1


def test_an_exact_count_reads_exact_everywhere(app_client) -> None:  # noqa: F811
    """The positive control: attribution complete, one row excluded and named."""
    c, provider = app_client
    got = _released(c, provider, [_item("Wiz", 0), _item("Splunk", 1)])
    for surface in ("pdf", "docx", "xlsx"):
        assert "3 rows received · 2 included · 1 excluded" in got[surface], surface
        assert UNKNOWN_TAIL not in got[surface], surface
    assert got["admin"]["exclusion_count_state"] == "exact"
    assert got["dash"]["excluded_count_exact"] is True


def test_a_clean_exact_list_is_a_total_everywhere(app_client) -> None:  # noqa: F811
    c, provider = app_client
    got = _released(c, provider, [_item("Wiz", 0), _item("Splunk", 1), _item("Lacework", 2)])
    for surface in ("pdf", "docx", "xlsx"):
        assert "Total annual cost" in got[surface], surface
        assert "rows received" not in got[surface], surface
    assert got["admin"]["exclusion_count_state"] == "exact"
    assert got["dash"]["excluded_count_exact"] is True
    assert got["dash"]["spend_completeness"] == "complete"


@pytest.mark.parametrize(
    ("items", "expected"),
    [
        # The old writer named the excluded row because attribution was
        # complete: `reconcile_rows` fills `excluded_rows` only under
        # `if attribution_complete`. So a non-empty list proves it.
        ([("Wiz", 0), ("Splunk", 1)], "exact"),
        # The old writer stored [] because attribution failed. An empty list
        # proves nothing.
        ([("Wiz", 0), ("Wiz CNAPP", 0), ("Splunk", 1)], "unknown"),
        # A clean run with nothing excluded ALSO stored [] before 0058, and is
        # indistinguishable from the failure above: unknown, never complete.
        ([("Wiz", 0), ("Splunk", 1), ("Lacework", 2)], "unknown"),
    ],
)
def test_a_pre_0058_list_reads_what_its_stored_record_proves(
    app_client, items, expected  # noqa: F811
) -> None:
    c, provider = app_client
    got = _released(c, provider, [_item(n, r) for n, r in items], null_flag=True)
    assert got["admin"]["attribution_complete"] is None
    assert got["admin"]["exclusion_count_state"] == expected
    assert got["dash"]["excluded_count_exact"] is (expected == "exact")


@pytest.mark.parametrize("null_flag", [False, True], ids=["post-0058", "pre-0058"])
def test_an_unbalanced_list_is_never_exact(app_client, null_flag) -> None:  # noqa: F811
    """More items than source rows, built by the real writer. Two items share
    a row, so attribution cannot be complete and the excluded count cannot be
    known. The web's spendSub has no "does not reconcile" branch because of
    this test: an unbalanced list always arrives with `excluded_count_exact`
    false."""
    c, provider = app_client
    items = [_item("Wiz", 0), _item("Wiz CNAPP", 0), _item("Splunk", 1), _item("Lacework", 2)]
    got = _released(c, provider, items, null_flag=null_flag)
    assert got["admin"]["source_rows_total"] == 3
    assert len(got["admin"]["items"]) == 4
    assert got["admin"]["exclusion_count_state"] == "unknown"
    assert got["dash"]["excluded_count_exact"] is False
    assert got["dash"]["spend_completeness"] == "partial"


def test_including_every_named_row_keeps_a_pre_0058_list_exact(app_client) -> None:  # noqa: F811
    """A pre-0058 list proves its count exact only by its NAMED rows. Including
    every one of them empties that list, and NULL with [] reads unknown, so the
    include-row route stamps the proof (`attribution_complete=True`) before it
    consumes it. Built by the real writer with the flag nulled."""
    c, provider = app_client
    got = _released(c, provider, [_item("Wiz", 0)], null_flag=True, include_all=True)
    assert got["admin"]["attribution_complete"] is True
    assert got["admin"]["exclusion_count_state"] == "exact"
    assert got["dash"]["excluded_count_exact"] is True
    for surface in ("pdf", "docx", "xlsx"):
        assert "Total annual cost" in got[surface], surface
        assert UNKNOWN_TAIL not in got[surface], surface
