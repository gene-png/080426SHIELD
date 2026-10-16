"""#1242 review, F1: the priority's precision through the finalize route,
across a release (advisor, #736 6220638526 item 3, option (A): every new
render floors to one decimal; a released file is stored bytes and is never
re-rendered, so it keeps its numbers without a check).

Two sequences, each through the real routes:

(a) finalize v1, release it, re-finalize v2. A client then holds v1 and, once
    it is released, v2 for the same assessment version: the two should state
    the same priority for the same row.
(b) the same, with the assessment set back to APPROVED after the release by
    direct SQL, so the render reads an APPROVED parent beside a released
    deliverable. No route produces exactly this state: a real pre-0041
    release also has `parent_version` NULL. What the render reads (the
    assessment's status) is the part that matters, and it is ordinary.

One Protect subcategory at tier 1 against the client's target 4, every other
at 4: exactly one gap, 3 tiers x 1.15 = 3.45, which floors to 3.4.
"""

from __future__ import annotations

import io
import os
import uuid

import pytest
from sqlalchemy import create_engine, update
from sqlalchemy.orm import sessionmaker

from app.csf.catalog import SUBCATEGORIES
from tests.unit.test_export_targets import (
    _attach_intake_target,
    _register,
    app_client,  # noqa: F401
)

pytestmark = pytest.mark.unit

PR_CODE = next(s.code for s in SUBCATEGORIES if s.code.startswith("PR."))


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _priority_cells(text: str) -> list[str]:
    """The gap table's Priority cells: the line after "Partial → Adaptive"."""
    lines = text.split("\n")
    return [lines[i + 1] for i, line in enumerate(lines) if line == "Partial → Adaptive"]


def _world(c) -> tuple[dict, str, str]:
    admin = _register(c, "admin@example.com")
    h = {"Authorization": f"Bearer {admin['tokens']['access_token']}"}
    svc = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()["id"]
    _attach_intake_target(svc, csf_tier=4)
    a = c.post(f"/csf/services/{svc}/assessments", headers=h).json()
    for ans in a["answers"]:
        tier = 1 if ans["subcategory_code"] == PR_CODE else 4
        r = c.patch(f"/csf/answers/{ans['id']}", headers=h, json={"maturity_tier": tier})
        assert r.status_code == 200, r.text
    r = c.post(f"/csf/assessments/{a['id']}/approve", headers=h)
    assert r.status_code == 200, r.text
    return h, svc, a["id"]


def _finalize(c, h: dict, svc: str) -> tuple[str, str]:
    fin = c.post(f"/csf/services/{svc}/deliverables/finalize", headers=h)
    assert fin.status_code == 201, fin.text
    body = fin.json()
    dl = c.get(f"/artifacts/{body['docx_artifact_id']}/download", headers=h)
    assert dl.status_code == 200, dl.text
    return body["id"], _docx_text(dl.content)


def _release(c, h: dict, deliverable_id: str) -> None:
    r = c.post(f"/csf/deliverables/{deliverable_id}/release", headers=h)
    assert r.status_code == 200, r.text


def _set_approved(assessment_id: str) -> None:
    from app.models.csf_assessment import CsfAssessment, CsfAssessmentStatus

    eng = create_engine(os.environ["DATABASE_URL"], future=True)
    with sessionmaker(bind=eng, future=True)() as s:
        s.execute(
            update(CsfAssessment)
            .where(CsfAssessment.id == uuid.UUID(assessment_id))
            .values(status=CsfAssessmentStatus.APPROVED)
        )
        s.commit()


def test_a_re_finalize_after_a_release_states_the_released_priority(
    app_client,  # noqa: F811
) -> None:
    c = app_client
    h, svc, _aid = _world(c)
    v1, text1 = _finalize(c, h, svc)
    assert _priority_cells(text1) == ["3.4"], _priority_cells(text1)  # positive first
    _release(c, h, v1)
    _v2, text2 = _finalize(c, h, svc)
    assert _priority_cells(text2) == _priority_cells(text1), (
        _priority_cells(text1),
        _priority_cells(text2),
    )


def test_a_re_finalize_beside_a_release_whose_parent_reads_approved(
    app_client,  # noqa: F811
) -> None:
    c = app_client
    h, svc, aid = _world(c)
    v1, text1 = _finalize(c, h, svc)
    assert _priority_cells(text1) == ["3.4"], _priority_cells(text1)
    _release(c, h, v1)
    _set_approved(aid)  # the pre-0041 shape: released, parent still APPROVED
    _v2, text2 = _finalize(c, h, svc)
    assert _priority_cells(text2) == _priority_cells(text1), (
        _priority_cells(text1),
        _priority_cells(text2),
    )
