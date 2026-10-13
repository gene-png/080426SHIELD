"""RR-1 (for #1122), design note §7 and §9.1: the demo seed's Risk Register
files say what the register was generated from (copy 8) and what was left out
(copy 9a), through `selection_context`, like the export route's.

Runs the whole `main()` against SQLite, as `test_seed_demo_risk_register.py`
does, and reads the three files the seed stored. Expected strings are §8's,
verbatim. The seed selects CSF and CISA Zero Trust and leaves DoD out, so the
register mentions two Zero Trust services and §6.2 names each framework.
"""

from __future__ import annotations

import io

import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.models.artifact import Artifact
from app.models.risk_register import RiskRegister
from app.storage.local import LocalFilesystemStorage

pytestmark = pytest.mark.unit

COPY_8 = (
    "This register was generated from ATT&CK coverage, NIST CSF and Zero Trust (CISA ZTMM 2.0)."
)
COPY_9A = "Zero Trust (DoD ZT Reference Architecture) was not included in this register."
COPY_11 = (
    "This register has no usable record of how its services were chosen or which were left out."
)


def _texts(storage: LocalFilesystemStorage, db: Session, reg: RiskRegister) -> dict[str, str]:
    from docx import Document
    from openpyxl import load_workbook

    from tests.unit.test_risk_register import _pdf_text

    def raw(artifact_id) -> bytes:
        art = db.get(Artifact, artifact_id)
        assert art is not None, artifact_id
        return storage.get(art.file_storage_key)

    wb = load_workbook(io.BytesIO(raw(reg.xlsx_artifact_id)))
    xlsx = "\n".join(
        str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v
    )
    doc = Document(io.BytesIO(raw(reg.docx_artifact_id)))
    docx = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    )
    pdf = _pdf_text(raw(reg.pdf_artifact_id))
    return {"pdf": pdf, "docx": docx, "xlsx": xlsx}


def test_the_seeded_files_print_copy_8_and_9a(tmp_path, monkeypatch) -> None:
    from scripts import seed_demo

    url = f"sqlite:///{tmp_path / 'seed.db'}"
    storage = LocalFilesystemStorage(tmp_path / "storage")
    monkeypatch.setenv("DATABASE_URL", url)
    monkeypatch.setattr(seed_demo, "get_storage", lambda: storage)
    seed_demo.main()

    with Session(create_engine(url, future=True)) as db:
        reg = db.execute(select(RiskRegister)).scalar_one()
        texts = _texts(storage, db, reg)
    for kind, text in texts.items():
        flat = " ".join(text.split())
        assert COPY_8 in flat, kind  # the positive state first
        assert COPY_9A in flat, kind
        assert COPY_11 not in flat, kind
