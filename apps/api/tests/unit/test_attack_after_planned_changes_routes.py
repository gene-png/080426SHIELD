"""#801 through the routes: coverage after planned changes, on every surface.

The world: a client whose approved Tech Debt plan cuts one tool and keeps
another, and an R3 draft with three standalone techniques, all covered today:
  * one rests on the CUT tool alone           -> after: gap   (scores lower)
  * one rests on the KEPT tool                -> after: covered
  * one detects with a tool on NO plan        -> after: partial (scores lower,
    and its retirement status is unknown)
Today 100.0%; after (1 + 0.5) / 3 = 50.0%.

Every expected string is COPIED from the proposal the advisor approved on #801
(05:05Z), never imported from the code under test.
"""

from __future__ import annotations

import uuid

import pytest
from sqlalchemy import select, update

from app.models.attack_assessment import AttackAssessment, AttackCoverage
from app.models.audit_entry import AuditEntry
from app.models.capability import CapabilityDisposition, CapabilityItem, CapabilityListStatus
from tests._attack_rows import standalone_rows
from tests.unit.test_attack_catalog_version_guard import (  # noqa: F401  (fixture)
    _auth,
    _register,
    _service_and_assessment,
    env,
)
from tests.unit.test_attack_planned_retirement import (
    _docx_text,
    _download,
    _pdf_text,
    _tech_debt_list,
)

pytestmark = pytest.mark.unit

CUT_TOOL, KEPT_TOOL, UNLISTED = "Cut Tool", "Kept Tool", "Unlisted Tool"

D1 = (
    "After planned changes: 50.0%, if the tools marked for planned retirement are cut "
    "and nothing else changes."
)
D2 = "This uses the current consolidation plan."
A1 = "2 techniques would score lower."
A3 = (
    "1 technique cites tools whose retirement status is unknown; it is scored after "
    "planned changes as if those tools were not in place."
)
P1 = (
    "Coverage after planned changes: 50.0%, if the tools marked for planned retirement "
    "are cut and nothing else changes."
)
X1 = "Coverage % after planned changes"
X2 = (
    "Coverage % after planned changes means",
    "The same calculation with the tools marked for planned retirement removed: if these "
    "tools are cut and nothing else changes.",
)


def _three(tools: list[str]) -> dict:
    return {"detection_tools": tools, "prevention_tools": tools, "response_tools": tools}


#: The three techniques of the module docstring.
_BODIES = [
    _three([CUT_TOOL]),
    _three([KEPT_TOOL]),
    {
        "detection_tools": [UNLISTED],
        "prevention_tools": [KEPT_TOOL],
        "response_tools": [KEPT_TOOL],
    },
]


def _world(env, *, with_plan: bool = True, bodies: list[dict] | None = None):  # noqa: F811
    c, Sess = env
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    client_id = client["user"]["client_id"]
    ids = {}
    if with_plan:
        ids = _tech_debt_list(
            Sess,
            client_id,
            admin["user"]["id"],
            status=CapabilityListStatus.APPROVED,
            items=[
                (CUT_TOOL, CapabilityDisposition.CUT),
                (KEPT_TOOL, CapabilityDisposition.KEEP),
            ],
        )
    svc, a = _service_and_assessment(c, bearer)
    bodies = _BODIES if bodies is None else bodies
    rows = standalone_rows(a["coverage"], len(bodies))
    for row, body in zip(rows, bodies, strict=True):
        r = c.patch(
            f"/attack/coverage/{row['id']}",
            headers=_auth(bearer),
            json={"status": "covered", **body},
        )
        assert r.status_code == 200, r.text
        assert r.json()["computed_status"] == "covered"  # today, retirement ignored
    return c, Sess, bearer, client, client_id, svc, a, ids, rows


def _finalize(c, bearer, svc, a) -> dict:
    r = c.post(f"/attack/assessments/{a['id']}/approve", headers=_auth(bearer))
    assert r.status_code == 200, r.text
    fin = c.post(f"/attack/services/{svc}/deliverables/finalize", headers=_auth(bearer))
    assert fin.status_code in (200, 201), fin.text
    return fin.json()


def test_the_admin_card_states_both_figures(env) -> None:  # noqa: F811
    c, _Sess, bearer, *_rest, svc, _a, _ids, _rows = _world(env)
    heat = c.get(f"/attack/services/{svc}/heatmap", headers=_auth(bearer)).json()
    assert heat["coverage_pct"] == 100.0  # today's figure, unchanged
    assert heat["after_planned_changes"] == [D1, A1, A3]  # H1: no current-plan note


def test_the_finalize_summary_and_its_audit_carry_the_figure(env) -> None:  # noqa: F811
    c, Sess, bearer, *_rest, svc, a, _ids, _rows = _world(env)
    fin = _finalize(c, bearer, svc, a)
    assert fin["summary"].startswith("Coverage: 100.0%. "), fin["summary"]
    assert fin["summary"].endswith(f" After planned changes: 50.0%. {A1} {A3}"), fin["summary"]
    with Sess() as s:
        (event,) = (
            s.execute(select(AuditEntry).where(AuditEntry.action == "attack.deliverable.finalized"))
            .scalars()
            .all()
        )
    assert event.details["coverage_pct_after_planned_changes"] == 50.0


def test_the_documents_carry_the_figure(env) -> None:  # noqa: F811
    import io

    from openpyxl import load_workbook

    c, _Sess, bearer, *_rest, svc, a, _ids, _rows = _world(env)
    fin = _finalize(c, bearer, svc, a)
    for text in (
        _pdf_text(_download(c, bearer, fin["pdf_artifact_id"])),
        _docx_text(_download(c, bearer, fin["docx_artifact_id"])),
    ):
        # test-integrity: the needle is three full approved sentences, written out above.
        assert f"{P1} {A1} {A3}" in text, text
    ws = load_workbook(io.BytesIO(_download(c, bearer, fin["xlsx_artifact_id"])))["Heatmap Summary"]
    rows = [tuple(cell.value for cell in r[:2]) for r in ws.iter_rows()]
    at = rows.index(("Coverage %", 100.0))
    assert rows[at + 1] == (X1, 50.0)
    assert [r[0] for r in rows[at + 2 : at + 4]] == [A1, A3]
    assert X2 in rows


def test_the_client_dashboard_reads_the_current_plan_and_the_documents_do_not(
    env,  # noqa: F811
) -> None:
    """D1 + D2 on the dashboard. The documents keep the plan as of finalize; the
    dashboard reads it live (#686's join), and says so."""
    c, Sess, bearer, client, client_id, svc, a, ids, _rows = _world(env)
    fin = _finalize(c, bearer, svc, a)
    rel = c.post(f"/attack/deliverables/{fin['id']}/release", headers=_auth(bearer))
    assert rel.status_code == 200, rel.text
    c.headers["X-Client-Id"] = client_id
    url = f"/clients/{client_id}/attack/{svc}/dashboard"
    client_auth = _auth(client["tokens"]["access_token"])
    body = c.get(url, headers=client_auth).json()
    assert body["rollup"]["coverage_pct"] == 100.0
    assert body["after_planned_changes"] == [D1, D2, A1, A3]

    # The plan changes after release: the kept tool is cut too.
    with Sess() as s:
        s.execute(
            update(CapabilityItem)
            .where(CapabilityItem.id == ids[KEPT_TOOL])
            .values(disposition=CapabilityDisposition.CUT)
        )
        s.commit()
    body = c.get(url, headers=client_auth).json()
    assert body["after_planned_changes"][0] == (
        "After planned changes: 0.0%, if the tools marked for planned retirement are cut "
        "and nothing else changes."
    )
    pdf = _pdf_text(_download(c, bearer, fin["pdf_artifact_id"]))
    assert P1 in pdf  # the delivered document keeps 50.0%


def test_a3_counts_an_unknown_retirement_and_not_a_pending_citation(env) -> None:  # noqa: F811
    """Two techniques, both covered today and both lower after planned changes:
      * Detect lists the CUT tool and the KEPT tool, the kept tool's citation
        still pending (as a run that inferred it leaves it). Every retirement
        verdict is known, so this is not an unknown retirement;
      * Detect lists a tool on NO plan, which is.
    After: (0 + 0.5) / 2 = 25.0%, and A3 counts one technique, not two."""
    c, Sess, bearer, *_rest, svc, _a, _ids, rows = _world(
        env,
        bodies=[
            {
                "detection_tools": [CUT_TOOL, KEPT_TOOL],
                "prevention_tools": [CUT_TOOL],
                "response_tools": [CUT_TOOL],
            },
            _BODIES[2],
        ],
    )
    with Sess() as s:
        s.execute(
            update(AttackCoverage)
            .where(AttackCoverage.id == uuid.UUID(rows[0]["id"]))
            .values(
                unconfirmed_citations=[
                    {
                        "tool": KEPT_TOOL,
                        "cited": KEPT_TOOL,
                        "reason": "inferred",
                        "cleared_at": None,
                    }
                ]
            )
        )
        s.commit()
    heat = c.get(f"/attack/services/{svc}/heatmap", headers=_auth(bearer)).json()
    assert heat["coverage_pct"] == 100.0
    assert heat["after_planned_changes"] == [
        "After planned changes: 25.0%, if the tools marked for planned retirement are cut "
        "and nothing else changes.",
        "2 techniques would score lower.",
        A3,
    ]


def test_no_plan_no_second_figure(env) -> None:  # noqa: F811
    c, _Sess, bearer, *_rest, svc, a, _ids, _rows = _world(env, with_plan=False)
    heat = c.get(f"/attack/services/{svc}/heatmap", headers=_auth(bearer)).json()
    assert heat["coverage_pct"] == 100.0
    assert heat["after_planned_changes"] is None
    fin = _finalize(c, bearer, svc, a)
    assert "After planned changes" not in fin["summary"]


def test_an_assessment_approved_before_r3_has_no_second_figure(env) -> None:  # noqa: F811
    c, Sess, bearer, *_rest, svc, a, _ids, _rows = _world(env)
    with Sess() as s:
        s.execute(
            update(AttackAssessment)
            .where(AttackAssessment.id == uuid.UUID(a["id"]))
            .values(status_rules=1)
        )
        s.commit()
    heat = c.get(f"/attack/services/{svc}/heatmap", headers=_auth(bearer)).json()
    assert heat["coverage_pct"] == 100.0  # APPEAR before ABSENT
    assert heat["after_planned_changes"] is None
