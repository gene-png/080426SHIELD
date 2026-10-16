"""CSF Options 1 and 3 in the Playbook files, through the export route.

- The cap sentences stay ONLY on a frozen assessment (Q2); every other export
  says evidence does not change a level (C8, as proposed in #736 6204461502).
- A subcategory with no scored tier is "Not scored", never "L1", and is in no
  average: one subcategory scored at level 3 beside 105 nobody scored
  averages Level 3. Before, the 105 seed rows read level 1 and the average
  rounded to 1.
- The C3 coverage line is stated on the files.

Every expected string is typed here, never imported from the renderer.
"""

from __future__ import annotations

import io

import pytest

from app.csf.catalog import SUBCATEGORIES
from tests.unit.test_csf_playbook_override_every_artifact import (  # noqa: F401
    _export,
    _flat_text,
    app_client,
)
from tests.unit.test_csf_self_reported_and_uncapped import (
    LEVEL_3,
    _approve,
    _ok_patch,
    _set_rollup_rules,
    _world,
)

pytestmark = pytest.mark.unit

CODE = SUBCATEGORIES[0].code
OTHER = SUBCATEGORIES[1].code
CAP_SENTENCE = (
    "An evidence cap applies: where producible evidence is absent, Implementation is "
    "held to at most 1 and the maturity level to at most 2."
)
NO_CAP_SENTENCE = "Evidence does not change a score or a level."


def _sheets(raw: bytes) -> dict[str, list[tuple]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}


def _table_row(sheet: list[tuple], code: str) -> dict:
    (header_at,) = [i for i, r in enumerate(sheet) if r and r[0] == "Subcategory"]
    header = sheet[header_at]
    (row,) = [r for r in sheet[header_at + 1 :] if r[0] == code]
    return dict(zip(header, row, strict=True))


def _cover_text(sheets: dict[str, list[tuple]]) -> str:
    return " ".join(str(c) for r in sheets["About"] for c in r if c is not None)


def test_a_draft_export_has_no_cap_and_states_not_scored(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, _aid = _world(c, h, ["high"])
    _ok_patch(c, h, svc, "high", CODE, LEVEL_3)
    files = _export(c, h, svc)
    sheets = _sheets(files["xlsx"])
    scored = _table_row(sheets["Enterprise Profile"], CODE)
    assert scored["Enterprise"] == "L3", scored  # positive first
    other = _table_row(sheets["Enterprise Profile"], OTHER)
    # An empty cell reads back as None.
    assert (other["Enterprise"], other["Rule"]) == ("Not scored", None), other
    total = len(sheets["Enterprise Profile"]) - 2  # banner and header rows
    cover = _cover_text(sheets)
    assert (
        f"1 of {total} subcategories scored; 0 with notes, not yet scored; "
        f"0 self-reported only, not validated; {total - 1} not answered."
    ) in cover, cover
    assert (
        "Of the 1 scored subcategories: 0 with evidence on file for every scored tier; "
        "0 confirmed in discussion (every scored tier on file or confirmed, at least one "
        "confirmed); 0 mixed (at least one scored tier with neither, and at least one with "
        "evidence on file or confirmed); 1 with neither on any scored tier; 0 at Managed or "
        "Optimized are notional."
    ) in cover, cover
    tier = _table_row(sheets["High"], CODE)
    assert tier["Evidence"] == "Neither", tier
    full = _flat_text("full_pdf", files["full_pdf"])
    assert NO_CAP_SENTENCE in full
    assert CAP_SENTENCE not in full
    assert "averages Level 3 of 5." in full, full[:2000]
    for kind in ("full_docx", "exec_docx", "exec_pdf"):
        assert CAP_SENTENCE not in _flat_text(kind, files[kind]), kind


@pytest.mark.parametrize("rules", [1, 2])
def test_a_frozen_assessment_export_keeps_the_cap_copy(
    app_client,  # noqa: F811
    rules: int,
) -> None:
    c, h = app_client
    svc, aid = _world(c, h, ["high"])
    _ok_patch(c, h, svc, "high", CODE, LEVEL_3)
    _approve(c, h, aid)
    _set_rollup_rules(aid, rules)
    files = _export(c, h, svc)
    full = _flat_text("full_pdf", files["full_pdf"])
    assert CAP_SENTENCE in full
    assert NO_CAP_SENTENCE not in full
    assert CAP_SENTENCE in _flat_text("full_docx", files["full_docx"])
    sheets = _sheets(files["xlsx"])
    assert "Evidence capped" in sheets["High"][1], "tier sheet header"
    cover = _cover_text(sheets)
    assert " subcategories scored; " in cover, cover  # positive first: the C3 line
    assert "Of the " not in cover, "no E3 line where evidence caps the level"
