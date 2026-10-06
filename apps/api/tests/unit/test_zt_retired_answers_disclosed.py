"""Answers kept on rows CISA ZTMM 2.0 does not have are disclosed (#838).

Migration 0062 KEEPS the answers on the 13 rows the corrected catalog dropped
(decision 4, option B), so an assessment that predates it holds rows the
catalog no longer has. That migration is the writer of the state seeded here:
the rows are inserted after the assessment is created, exactly as 0062 leaves
them. The expected sentences are the approved copy (#838 comment 5982917066),
written out literally, and every check goes through the surface a person
reaches: the workspace's GET and the finalized files.
"""

from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook

from app.models.zt_assessment import ZtAnswer, ZtAssessment
from tests._ai_mode import docx_text, env_sessions, pdf_text
from tests.unit.test_zt_acceptance import (  # noqa: F401  (fixture)
    _register,
    app_client,
)

pytestmark = pytest.mark.unit

TWO = "2 recorded answers belong to rows that CISA ZTMM 2.0 does not have, so they are not scored."
ONE = "1 recorded answer belongs to a row that CISA ZTMM 2.0 does not have, so it is not scored."


def _assessment(c) -> tuple[dict, str, dict]:
    bearer = _register(c, "admin@example.com")["tokens"]["access_token"]
    h = {"Authorization": f"Bearer {bearer}"}
    svc_id = c.post(
        "/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": "ZT"}
    ).json()["id"]
    a = c.post(f"/zt/services/{svc_id}/assessments", headers=h).json()
    return h, svc_id, a


def _keep(assessment_id: str, rows: list[tuple[str, int | None, str | None]]) -> None:
    """Rows on retired codes, as migration 0062 leaves them."""
    with env_sessions()() as db:
        a = db.get(ZtAssessment, uuid.UUID(assessment_id))
        for code, stage, notes in rows:
            db.add(
                ZtAnswer(
                    assessment_id=a.id,
                    client_id=a.client_id,
                    capability_code=code,
                    maturity_stage=stage,
                    notes=notes,
                )
            )
        db.commit()


def _latest(c, h: dict, svc_id: str) -> dict:
    r = c.get(f"/zt/services/{svc_id}/assessments/latest", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def test_the_workspace_counts_recorded_answers_only(app_client) -> None:  # noqa: F811
    c = app_client
    h, svc_id, a = _assessment(c)
    _keep(
        a["id"],
        [
            ("CISA.VA.01", 2, None),  # a stage: counted
            ("CISA.GV.01", None, "notes only"),  # notes: counted
            ("CISA.AO.01", None, "   "),  # blank: not an answer
        ],
    )
    body = _latest(c, h, svc_id)
    assert body["id"] == a["id"]
    assert body["retired_answers"] == 2
    assert body["retired_answers_note"] == TWO


def test_one_retired_answer_reads_in_the_singular(app_client) -> None:  # noqa: F811
    c = app_client
    h, svc_id, a = _assessment(c)
    _keep(a["id"], [("CISA.ID.05", 1, None)])  # a mapped code whose target existed
    body = _latest(c, h, svc_id)
    assert body["retired_answers"] == 1
    assert body["retired_answers_note"] == ONE


def test_no_retired_answer_means_no_sentence(app_client) -> None:  # noqa: F811
    c = app_client
    h, svc_id, a = _assessment(c)
    body = _latest(c, h, svc_id)
    assert body["id"] == a["id"]
    assert len(body["answers"]) == 37
    assert body["retired_answers"] == 0
    assert body["retired_answers_note"] is None


def _finalize(c, h: dict, svc_id: str, a: dict) -> tuple[str, str, list[list[object]]]:
    for ans in a["answers"]:
        r = c.patch(f"/zt/answers/{ans['id']}", headers=h, json={"maturity_stage": 3})
        assert r.status_code == 200, r.text
    assert c.post(f"/zt/assessments/{a['id']}/approve", headers=h).status_code == 200
    fin = c.post(f"/zt/services/{svc_id}/deliverables/finalize", headers=h)
    assert fin.status_code == 201, fin.text

    def _bytes(key: str) -> bytes:
        dl = c.get(f"/artifacts/{fin.json()[key]}/download", headers=h)
        assert dl.status_code == 200, dl.text
        return dl.content

    wb = load_workbook(io.BytesIO(_bytes("xlsx_artifact_id")))
    summary = [list(r) for r in wb["Score Summary"].iter_rows(values_only=True)]
    return pdf_text(_bytes("pdf_artifact_id")), docx_text(_bytes("docx_artifact_id")), summary


def test_every_finalized_file_states_the_retired_answers(app_client) -> None:  # noqa: F811
    c = app_client
    h, svc_id, a = _assessment(c)
    _keep(a["id"], [("CISA.VA.01", 2, None), ("CISA.AO.02", 3, "kept")])
    pdf, docx, summary = _finalize(c, h, svc_id, a)
    assert TWO in pdf
    assert TWO in docx
    assert ["Not scored", TWO] in [row[:2] for row in summary]


def test_a_file_with_nothing_retired_says_nothing_about_it(app_client) -> None:  # noqa: F811
    c = app_client
    h, svc_id, a = _assessment(c)
    pdf, docx, summary = _finalize(c, h, svc_id, a)
    assert "Maturity summary" in docx  # the document rendered
    assert "does not have" not in pdf
    assert "does not have" not in docx
    assert "Not scored" not in [row[0] for row in summary]
