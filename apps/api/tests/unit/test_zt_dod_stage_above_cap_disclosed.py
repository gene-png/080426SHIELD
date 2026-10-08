"""A stage stored above its capability's maximum is disclosed everywhere (#839, S1).

The advisor's ruling (#736 comment 6048561596): such a stage is LEFT as stored,
counted at the capped level in every figure, and disclosed on the screen and in
all three files. S1 was approved verbatim at 6049667540; the expected sentence
below is the example in 6049167247, copied from that comment.

The world: a DoD engagement with every capability at stage 2 through the API,
then 1.1 User Inventory (no Advanced activity) set to 3 straight in the store,
as a row written before the write-path guard would be. Then approve, finalize
and release. Checked through the endpoints a consultant and a client reach and
the files a client receives.
"""

from __future__ import annotations

import io
import uuid

import pytest
from sqlalchemy import update

from app.models.zt_assessment import ZtAnswer
from tests._ai_mode import docx_text, env_sessions, pdf_text
from tests.unit.test_zt_dashboard import (  # noqa: F401  (fixture)
    _register,
    app_client,
)

pytestmark = pytest.mark.unit

S1 = (
    "1.1 User Inventory is recorded at stage 3, but it has no DoD Advanced "
    "activities, so every figure here counts it as Target (2)."
)


def _world(c, *, over_cap: bool) -> tuple[str, str, dict, dict]:
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    bearer = admin["tokens"]["access_token"]
    h = {"Authorization": f"Bearer {bearer}"}
    svc_id = c.post(
        "/zt/services", headers=h, json={"kind": "zero_trust_dod", "title": "ZT"}
    ).json()["id"]
    a = c.post(f"/zt/services/{svc_id}/assessments", headers=h).json()
    for ans in a["answers"]:
        r = c.patch(f"/zt/answers/{ans['id']}", headers=h, json={"maturity_stage": 2})
        assert r.status_code == 200, r.text
    if over_cap:
        with env_sessions()() as db:
            db.execute(
                update(ZtAnswer)
                .where(
                    ZtAnswer.assessment_id == uuid.UUID(a["id"]),
                    ZtAnswer.capability_code == "DOD.USR.01",
                )
                .values(maturity_stage=3)
            )
            db.commit()
    assert c.post(f"/zt/assessments/{a['id']}/approve", headers=h).status_code == 200
    deliv = c.post(f"/zt/services/{svc_id}/deliverables/finalize", headers=h)
    assert deliv.status_code in (200, 201), deliv.text
    rel = c.post(f"/zt/deliverables/{deliv.json()['id']}/release", headers=h)
    assert rel.status_code == 200, rel.text
    return svc_id, bearer, h, client


def _files(c, svc_id: str, h: dict) -> tuple[str, str, object]:
    from openpyxl import load_workbook

    body = c.get(f"/zt/services/{svc_id}/deliverables/latest", headers=h).json()

    def _bytes(key: str) -> bytes:
        dl = c.get(f"/artifacts/{body[key]}/download", headers=h)
        assert dl.status_code == 200, dl.text
        return dl.content

    wb = load_workbook(io.BytesIO(_bytes("xlsx_artifact_id")))
    return pdf_text(_bytes("pdf_artifact_id")), docx_text(_bytes("docx_artifact_id")), wb


def _flat(text: str) -> str:
    return " ".join(text.split())


def test_the_workspace_gap_endpoint_discloses_it(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _bearer, h, _client = _world(c, over_cap=True)
    r = c.get(f"/zt/services/{svc_id}/gap-analysis?target_stage=3", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["stage_above_max_notes"] == [S1]


def test_the_client_dashboard_discloses_it(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _bearer, _h, client = _world(c, over_cap=True)
    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    r = c.get(
        f"/clients/{client_id}/zt/{svc_id}/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["stage_above_max_notes"] == [S1]


def test_every_released_file_discloses_it_and_keeps_the_stored_stage(
    app_client,  # noqa: F811
) -> None:
    c = app_client
    svc_id, _bearer, h, _client = _world(c, over_cap=True)
    pdf, docx, wb = _files(c, svc_id, h)
    assert S1 in _flat(pdf)
    assert S1 in _flat(docx)
    summary = [str(v) for row in wb["Score Summary"].iter_rows(values_only=True) for v in row if v]
    assert S1 in summary
    # The Answers sheet keeps the stored stage: never rewritten (6049667540).
    answers = {row[0]: row for row in wb["Answers"].iter_rows(values_only=True)}
    assert answers["DOD.USR.01"][4] == 3


def test_nothing_is_disclosed_when_every_stage_is_reachable(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _bearer, h, _client = _world(c, over_cap=False)
    r = c.get(f"/zt/services/{svc_id}/gap-analysis?target_stage=3", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["stage_above_max_notes"] == []
    pdf, docx, _wb = _files(c, svc_id, h)
    # A real rendered file, so the absences below mean something.
    assert "DoD ZT Reference Architecture" in _flat(pdf)
    assert "DoD ZT Reference Architecture" in _flat(docx)
    assert "is recorded at stage" not in _flat(pdf)
    assert "is recorded at stage" not in _flat(docx)


# --- The figures S1 promises (review round 2, finding 1) ---------------------
#
# S1 says "every figure here counts it as Target (2)". In this world every
# capability is at 2 and only DOD.USR.01 is stored at 3, so if the 3 is counted
# as 2 every figure is the all-2 figure. Derived by hand from the world, never
# from `compute`:
#
#   overall average  = 2.00            (45 rows at 2)      -- uncapped: 91/45 = 2.02
#   overall percent  = 2/3   -> 66.7   (DoD's top stage 3) -- uncapped: 2.02/3 -> 67.3
#   User pillar      = 2.00, 66.7%     (9 rows at 2)       -- uncapped: 19/9 = 2.11 -> 70.3
#   overall label    = "Target"        (stage 2)

ALL_TWO_AVERAGE = 2.0
ALL_TWO_PCT = 66.7


def test_the_score_endpoint_counts_it_at_the_maximum(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _bearer, h, _client = _world(c, over_cap=True)
    r = c.get(f"/zt/services/{svc_id}/score", headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["answered_capabilities"] == 45  # every row is counted, first
    assert body["average_stage"] == ALL_TWO_AVERAGE
    assert body["overall_stage_label"] == "Target"
    usr = next(p for p in body["by_pillar"] if p["pillar_code"] == "USR")
    assert usr["average_stage"] == ALL_TWO_AVERAGE


def test_the_client_dashboard_counts_it_at_the_maximum(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _bearer, _h, client = _world(c, over_cap=True)
    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    r = c.get(
        f"/clients/{client_id}/zt/{svc_id}/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["stage_above_max_notes"] == [S1]  # the disclosure is there, first
    assert body["current_pct"] == ALL_TWO_PCT
    usr = next(p for p in body["pillars"] if p["code"] == "USR")
    assert usr["answered_count"] == 9
    assert usr["current_pct"] == ALL_TWO_PCT


def test_the_released_files_count_it_at_the_maximum(app_client) -> None:  # noqa: F811
    c = app_client
    svc_id, _bearer, h, _client = _world(c, over_cap=True)
    _pdf, _docx, wb = _files(c, svc_id, h)
    rows = [list(r) for r in wb["Score Summary"].iter_rows(values_only=True)]
    by_label = {r[0]: r[1] for r in rows if r and r[0]}
    assert by_label["Coverage"] == "45/45"  # every row is counted, first
    assert by_label["Average stage"] == "2.00"
    assert by_label["Overall stage"] == "Target"
    usr = next(r for r in rows if r and r[0] == "USR")
    assert usr[5] == "2.00"  # the pillar row's Average stage
