"""The CSF and ZT Gap Plans say how many rows were NOT scored (#699).

`analyze_gaps` skips unscored rows, and no exporter read `unscored_codes`, so
a mostly-unscored approved assessment could deliver "All 0 gaps listed" (ZT)
or "All 0 gaps at target T3" (CSF) with nothing saying what was never looked
at. Partial self-assessment is allowed, so this DISCLOSES rather than blocks.

Driven through every renderer a client receives (XLSX, DOCX, PDF), reading
the rendered bytes back, and asserted as literal sentences with their counts,
never a bare number (`str(n) in blob` would match any unrelated digit).
"""

from __future__ import annotations

import io
import re
import uuid

import pytest

from app.csf.catalog import SUBCATEGORIES
from app.csf.exporters import build_context as csf_context
from app.csf.exporters import render_docx as csf_docx
from app.csf.exporters import render_pdf as csf_pdf
from app.csf.exporters import render_xlsx as csf_xlsx
from app.csf.gap import analyze as csf_gaps
from app.csf.scoring import compute as csf_score
from app.models.csf_assessment import CsfAnswer, CsfAssessment, CsfAssessmentStatus
from app.models.zt_assessment import ZtAnswer, ZtAssessment, ZtAssessmentStatus, ZtFramework
from app.zt.catalog import capabilities
from app.zt.exporters import build_context as zt_context
from app.zt.exporters import render_docx as zt_docx
from app.zt.exporters import render_pdf as zt_pdf
from app.zt.exporters import render_xlsx as zt_xlsx
from app.zt.maturity import ZtFrameworkCode, level_count
from app.zt.scoring import analyze_gaps as zt_gaps
from app.zt.scoring import compute as zt_score


def _flat(text: str) -> str:
    """Whitespace collapsed: PDF extraction breaks lines inside a sentence."""
    return re.sub(r"\s+", " ", text)


def _xlsx_caption(raw: bytes) -> str:
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(raw))["Gap Plan"].cell(row=1, column=1).value or ""


def _docx_text(raw: bytes) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs)


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return "".join(page.extract_text() for page in PdfReader(io.BytesIO(raw)).pages)


def _every_format(render_xlsx, render_docx, render_pdf, ctx) -> dict[str, str]:
    return {
        "xlsx": _flat(_xlsx_caption(render_xlsx(ctx))),
        "docx": _flat(_docx_text(render_docx(ctx))),
        "pdf": _flat(_pdf_text(render_pdf(ctx))),
    }


# ---------------------------------------------------------------------------
# CSF
# ---------------------------------------------------------------------------


def _csf(unscored: int):
    a = CsfAssessment(
        id=uuid.uuid4(), service_id=uuid.uuid4(), version=1, status=CsfAssessmentStatus.APPROVED
    )
    answers = [
        CsfAnswer(
            id=uuid.uuid4(),
            assessment_id=a.id,
            subcategory_code=sc.code,
            maturity_tier=None if i < unscored else 4,
        )
        for i, sc in enumerate(SUBCATEGORIES)
    ]
    tiers = {ans.subcategory_code: ans.maturity_tier for ans in answers}
    gap = csf_gaps(tiers, target_tier=3)
    assert len(gap.unscored_codes) == unscored, "setup: the engine must see the unscored rows"
    return csf_context(
        client_legal_name="Atlas Defense Solutions",
        service_title="NIST CSF 2.0 Assessment",
        assessment=a,
        answers=answers,
        score=csf_score(tiers),
        gap=gap,
    )


@pytest.mark.unit
def test_csf_gap_plan_states_the_unscored_count_in_every_format() -> None:
    # Every scored row is past the target, so the plan lists no gaps at all:
    # exactly the deliverable #699 describes, "no gaps" over unscored rows.
    ctx = _csf(unscored=60)
    assert ctx.gap.total_gap_count == 0, "setup: nothing scored may be a gap"
    expected = "60 subcategories were not scored and are not in this plan."
    for fmt, text in _every_format(csf_xlsx, csf_docx, csf_pdf, ctx).items():
        assert expected in text, f"{fmt} does not state the unscored rows: {text[:400]!r}"


@pytest.mark.unit
def test_csf_one_unscored_subcategory_is_one_row() -> None:
    ctx = _csf(unscored=1)
    expected = "1 subcategory was not scored and is not in this plan."
    for fmt, text in _every_format(csf_xlsx, csf_docx, csf_pdf, ctx).items():
        assert expected in text, f"{fmt}: {text[:400]!r}"


@pytest.mark.unit
def test_csf_says_nothing_about_unscored_rows_when_every_row_was_scored() -> None:
    ctx = _csf(unscored=0)
    for fmt, text in _every_format(csf_xlsx, csf_docx, csf_pdf, ctx).items():
        assert "not scored" not in text, f"{fmt} claims unscored rows that do not exist"
        # The caption itself still renders: this is not a test of silence.
        assert "All 0 gaps at target T3." in text, f"{fmt}: {text[:400]!r}"


# ---------------------------------------------------------------------------
# ZT
# ---------------------------------------------------------------------------


def _zt(unscored: int, framework: ZtFrameworkCode = ZtFrameworkCode.CISA_ZTMM_2_0):
    a = ZtAssessment(
        id=uuid.uuid4(),
        service_id=uuid.uuid4(),
        framework=ZtFramework.CISA_ZTMM_2_0,
        version=1,
        status=ZtAssessmentStatus.APPROVED,
    )
    top = level_count(framework)
    answers = [
        ZtAnswer(
            id=uuid.uuid4(),
            assessment_id=a.id,
            capability_code=cap.code,
            maturity_stage=None if i < unscored else top,
        )
        for i, cap in enumerate(capabilities(framework))
    ]
    stages = {ans.capability_code: ans.maturity_stage for ans in answers}
    gap = zt_gaps(framework, stages, target_stage=3)
    assert len(gap.unscored_codes) == unscored, "setup: the engine must see the unscored rows"
    return zt_context(
        client_legal_name="Atlas Defense Solutions",
        service_title="Zero Trust Assessment",
        framework=framework,
        assessment=a,
        answers=answers,
        score=zt_score(framework, stages),
        gap=gap,
    )


@pytest.mark.unit
def test_zt_gap_plan_states_the_unscored_count_in_every_format() -> None:
    ctx = _zt(unscored=20)
    assert ctx.gap.total_gap_count == 0, "setup: nothing scored may be a gap"
    expected = "All 0 gaps listed. 20 capabilities were not scored and are not in this plan."
    for fmt, text in _every_format(zt_xlsx, zt_docx, zt_pdf, ctx).items():
        assert expected in text, f"{fmt} does not state the unscored rows: {text[:400]!r}"


@pytest.mark.unit
def test_zt_one_unscored_capability_is_one_row() -> None:
    ctx = _zt(unscored=1)
    expected = "1 capability was not scored and is not in this plan."
    for fmt, text in _every_format(zt_xlsx, zt_docx, zt_pdf, ctx).items():
        assert expected in text, f"{fmt}: {text[:400]!r}"


@pytest.mark.unit
def test_zt_says_nothing_about_unscored_rows_when_every_row_was_scored() -> None:
    ctx = _zt(unscored=0)
    for fmt, text in _every_format(zt_xlsx, zt_docx, zt_pdf, ctx).items():
        assert "not scored" not in text, f"{fmt} claims unscored rows that do not exist"
        assert "All 0 gaps listed." in text, f"{fmt}: {text[:400]!r}"
