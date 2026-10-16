"""The CSF self-assessment deliverable reads tiers by name, and its priority
has one decimal, rounded down, on an unreleased deliverable only (advisor,
#736 6216024281 item 3c, answering 6211388486).

- "target T4" becomes "target Adaptive", and "T1 → T4" becomes
  "Partial → Adaptive", in the PDF, the DOCX and the XLSX caption. Names
  apply to a released re-export too.
- The Priority cell: 3 tiers below target in Protect (weight 1.15) is 3.45.
  Unreleased, it shows 3.4 (rounded DOWN; rounding to nearest gives 3.5). In
  Detect (weight 1.2) it is 3.6, not the 3.5 a float 3 x 1.2 = 3.5999...
  floors to. A RELEASED re-export keeps the number it was delivered with,
  3.45 and 3.60.

The expected values are typed here from the weights the gap list states
(PR 1.15, DE 1.2), never read from `FUNCTION_WEIGHTS` or the renderer.
"""

from __future__ import annotations

import io
import uuid
import zipfile

import pytest

from app.csf.catalog import SUBCATEGORIES
from app.csf.exporters import build_context, render_docx, render_pdf, render_xlsx
from app.csf.gap import analyze as analyze_gaps
from app.csf.scoring import compute as compute_score
from app.models.csf_assessment import CsfAnswer, CsfAssessment, CsfAssessmentStatus

pytestmark = pytest.mark.unit

PR_CODE = next(s.code for s in SUBCATEGORIES if s.code.startswith("PR."))
DE_CODE = next(s.code for s in SUBCATEGORIES if s.code.startswith("DE."))


def _ctx(status: CsfAssessmentStatus):
    """Every subcategory at tier 4 except one Protect and one Detect row at
    tier 1, against target 4: exactly two gaps of 3."""
    a = CsfAssessment(id=uuid.uuid4(), service_id=uuid.uuid4(), version=1, status=status)
    answers = [
        CsfAnswer(
            id=uuid.uuid4(),
            assessment_id=a.id,
            subcategory_code=sc.code,
            maturity_tier=1 if sc.code in (PR_CODE, DE_CODE) else 4,
        )
        for sc in SUBCATEGORIES
    ]
    tier_map = {ans.subcategory_code: ans.maturity_tier for ans in answers}
    gap = analyze_gaps(tier_map, target_tier=4)
    assert gap.total_gap_count == 2, gap.total_gap_count
    return build_context(
        client_legal_name="Test Tenant",
        service_title="NIST CSF 2.0 Assessment",
        assessment=a,
        answers=answers,
        score=compute_score(tier_map),
        gap=gap,
    )


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return " ".join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def _gap_plan(raw: bytes) -> tuple[str, dict[str, dict]]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))["Gap Plan"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1]
    return rows[0][0], {r[0]: dict(zip(header, r, strict=True)) for r in rows[2:]}


@pytest.mark.parametrize("render", [render_pdf, render_docx], ids=["pdf", "docx"])
def test_an_unreleased_deliverable_names_tiers_and_floors_the_priority(render) -> None:
    raw = render(_ctx(CsfAssessmentStatus.APPROVED))
    text = _pdf_text(raw) if render is render_pdf else _docx_text(raw)
    assert "Top remediation gaps (target Adaptive)" in text, text  # positive first
    assert "All 2 gaps at target Adaptive." in text, text
    assert "Partial → Adaptive" in text, text
    assert "3.4" in text and "3.6" in text, text
    assert "3.45" not in text and "3.5" not in text, text
    assert "T4" not in text and "T1 →" not in text, text


@pytest.mark.parametrize("render", [render_pdf, render_docx], ids=["pdf", "docx"])
def test_a_released_re_export_keeps_its_numbers_and_takes_the_names(render) -> None:
    raw = render(_ctx(CsfAssessmentStatus.RELEASED))
    text = _pdf_text(raw) if render is render_pdf else _docx_text(raw)
    assert "Partial → Adaptive" in text, text  # names apply to a re-export
    assert "3.45" in text and "3.60" in text, text  # the delivered numbers
    assert "T4" not in text, text


def test_the_xlsx_gap_plan_names_the_target_and_floors_unreleased_only() -> None:
    caption, rows = _gap_plan(render_xlsx(_ctx(CsfAssessmentStatus.APPROVED)))
    assert caption.startswith("All 2 gaps at target Adaptive."), caption
    assert (rows[PR_CODE]["Priority"], rows[DE_CODE]["Priority"]) == (3.4, 3.6), rows
    _caption, released = _gap_plan(render_xlsx(_ctx(CsfAssessmentStatus.RELEASED)))
    assert (released[PR_CODE]["Priority"], released[DE_CODE]["Priority"]) == (3.45, 3.6)


def test_the_docx_reader_reads_the_document() -> None:
    """Guard on the helper: `_docx_text` reads a real DOCX, so an empty
    extraction cannot pass the absences above."""
    raw = render_docx(_ctx(CsfAssessmentStatus.APPROVED))
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        assert "word/document.xml" in z.namelist()
    assert "NIST CSF 2.0 Assessment" in _docx_text(raw)
