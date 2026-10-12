"""#1031 / #1077, the advisor's rulings (#736 6116519635), through the routes:

- ruling 3: an assessment approved before #1031 keeps the every-tier roll-up
  (`rollup_rules = 1`, migration 0070's backfill); approve writes 2;
- a POA&M entry never vanishes silently: an entry on a code that stops being
  a gap is listed, and the Action Plan sheet says "No longer a gap: kept for
  the record";
- the Playbook export copy for the scored-tiers rule, approved verbatim (as
  proposed in 6116347717), and the old copy for a frozen assessment.

Every expected string is written out here, never imported from the module
that renders it.
"""

from __future__ import annotations

import io
import uuid

import pytest
from sqlalchemy import create_engine, text

from app.csf.catalog import SUBCATEGORIES
from tests.unit.test_csf_playbook_override_every_artifact import (  # noqa: F401
    _export,
    _flat_text,
    app_client,
)

pytestmark = pytest.mark.unit

CODE = SUBCATEGORIES[0].code
LEVEL_3 = {
    "governance": 2,
    "policy": 2,
    "implementation": 2,
    "monitoring": 1,
    "improvement": 0,
    "has_evidence": True,
}
LEVEL_5 = {
    "governance": 2,
    "policy": 2,
    "implementation": 2,
    "monitoring": 2,
    "improvement": 2,
    "has_evidence": True,
}

KEPT = "No longer a gap: kept for the record"
METHOD_NEW = (
    "Levels are code-computed: total = sum of the five dimensions (0-10); L1 0-2, L2 3-5, "
    "L3 6-7, L4 8-9, L5 10. Enterprise = weighted-floor roll-up across the in-scope tiers "
    'that have been scored. A tier with no score recorded is left out and marked "Not '
    'scored".'
)
METHOD_OLD = (
    "Levels are code-computed: total = sum of the five dimensions (0-10); L1 0-2, L2 3-5, "
    "L3 6-7, L4 8-9, L5 10. Enterprise = weighted-floor roll-up across the tiers in use."
)
SUMMARY_NEW = "Enterprise maturity, rolled up across the scored impact tiers in use, averages Level"
SUMMARY_OLD = "Enterprise maturity, rolled up across the impact tiers in use, averages Level"
TWO_UNSCORED = "2 in-scope tier rows were not scored and are not part of the Enterprise level."
ONE_UNSCORED = "1 in-scope tier row was not scored and is not part of the Enterprise level."


def _world(c, h: dict, *, tiers: list[str], high: dict) -> tuple[str, str]:
    """A CSF service with `tiers` seeded and CODE's HIGH row PATCHed with
    `high`; returns (service id, assessment id)."""
    svc = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": "CSF"}).json()["id"]
    a = c.post(f"/csf/services/{svc}/assessments", headers=h)
    assert a.status_code in (200, 201), a.text
    r = c.post(f"/csf/services/{svc}/profiles/seed", headers=h, json={"tiers": tiers})
    assert r.status_code == 200, r.text
    _patch(c, h, svc, "high", high)
    return svc, a.json()["id"]


def _patch(c, h: dict, svc: str, tier: str, body: dict) -> None:
    rows = c.get(f"/csf/services/{svc}/profile/{tier}", headers=h).json()["rows"]
    sid = next(r["id"] for r in rows if r["subcategory_code"] == CODE)
    r = c.patch(f"/csf/dimension-scores/{sid}", headers=h, json=body)
    assert r.status_code == 200, r.text


def _sub(c, h: dict, svc: str) -> dict:
    r = c.get(f"/csf/services/{svc}/enterprise-profile", headers=h)
    assert r.status_code == 200, r.text
    return next(s for s in r.json()["subcategories"] if s["subcategory_code"] == CODE)


def _rollup_rules(aid: str) -> int | None:
    import os

    engine = create_engine(os.environ["DATABASE_URL"], future=True)
    with engine.connect() as db:
        return (
            db.execute(
                text("SELECT rollup_rules FROM csf_assessments WHERE id = :id"),
                {"id": str(uuid.UUID(aid)).replace("-", "")},
            ).scalar_one_or_none()
            or db.execute(
                text("SELECT rollup_rules FROM csf_assessments WHERE id = :id"), {"id": aid}
            ).scalar_one_or_none()
        )


def _set_rollup_rules(aid: str, value: int) -> None:
    """Setup only: the state migration 0070's backfill writes for an
    assessment approved before #1031 (that writer is pinned in
    `test_csf_rollup_rules.py`)."""
    import os

    engine = create_engine(os.environ["DATABASE_URL"], future=True)
    with engine.begin() as db:
        n = db.execute(
            text("UPDATE csf_assessments SET rollup_rules = :v WHERE id IN (:a, :b)"),
            {"v": value, "a": aid, "b": str(uuid.UUID(aid)).replace("-", "")},
        ).rowcount
    assert n == 1, n


def _xlsx(raw: bytes) -> dict[str, list[tuple]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw))
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}


def _row(sheet: list[tuple], code: str) -> dict:
    (header_at,) = [i for i, r in enumerate(sheet) if r and r[0] == "Subcategory"]
    header = sheet[header_at]
    (row,) = [r for r in sheet[header_at + 1 :] if r[0] == code]
    return dict(zip(header, row, strict=True))


# --- Ruling 3: approve records the rule; a frozen assessment keeps the old one ---------


def test_approve_records_the_scored_tiers_rule(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, aid = _world(c, h, tiers=["high", "moderate", "low"], high={**LEVEL_3, "target_level": 3})
    assert _rollup_rules(aid) is None  # positive first: a draft records nothing
    r = c.post(f"/csf/assessments/{aid}/approve", headers=h)
    assert r.status_code == 200, r.text
    assert _rollup_rules(aid) == 2
    assert _sub(c, h, svc)["enterprise_level"] == 3


def test_an_assessment_approved_before_1031_keeps_the_every_tier_roll_up(
    app_client,  # noqa: F811
) -> None:
    c, h = app_client
    svc, aid = _world(c, h, tiers=["high", "moderate", "low"], high={**LEVEL_3, "target_level": 3})
    assert c.post(f"/csf/assessments/{aid}/approve", headers=h).status_code == 200
    new = _sub(c, h, svc)
    assert (new["tier_levels"], new["enterprise_level"], new["gap"]) == ({"high": 3}, 3, False)
    _set_rollup_rules(aid, 1)
    old = _sub(c, h, svc)
    assert old["tier_levels"] == {"high": 3, "moderate": 1, "low": 1}, old
    assert (old["enterprise_level"], old["rollup_rule"]) == (1, 4), old
    assert (old["target_level"], old["gap"]) == (3, True), old


# --- POA&M entries never vanish silently --------------------------------------------------


def _gap_actions(c, h: dict, svc: str) -> dict[str, dict]:
    r = c.get(f"/csf/services/{svc}/gap-actions", headers=h)
    assert r.status_code == 200, r.text
    return {a["subcategory_code"]: a for a in r.json()["actions"]}


def test_a_poam_entry_on_a_code_that_stops_being_a_gap_is_kept(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high"], high={**LEVEL_3, "target_level": 4})
    before = _gap_actions(c, h, svc)
    assert before[CODE]["no_longer_gap"] is False, before[CODE]  # positive first: a gap
    r = c.put(f"/csf/services/{svc}/gap-actions/{CODE}", headers=h, json={"owner": "Ops lead"})
    assert r.status_code == 200, r.text
    _patch(c, h, svc, "high", LEVEL_5)  # level 5 meets target 4: no longer a gap
    assert _sub(c, h, svc)["gap"] is False
    after = _gap_actions(c, h, svc)
    assert CODE in after, after
    assert (after[CODE]["owner"], after[CODE]["no_longer_gap"]) == ("Ops lead", True)
    plan = _row(_xlsx(_export(c, h, svc)["xlsx"])["Action Plan"], CODE)
    assert (plan["Owner"], plan["Priority"]) == ("Ops lead", KEPT), plan


def test_an_empty_annotation_on_a_non_gap_is_not_listed(app_client) -> None:  # noqa: F811
    """Only a recorded entry is kept: an annotation row with nothing in it is
    not a POA&M entry, and listing it would invent one."""
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high"], high={**LEVEL_3, "target_level": 4})
    r = c.put(f"/csf/services/{svc}/gap-actions/{CODE}", headers=h, json={"owner": "x"})
    assert r.status_code == 200, r.text
    r = c.put(f"/csf/services/{svc}/gap-actions/{CODE}", headers=h, json={"owner": None})
    assert r.status_code == 200, r.text
    _patch(c, h, svc, "high", LEVEL_5)
    assert CODE not in _gap_actions(c, h, svc)


# --- The approved export copy -------------------------------------------------------------


def test_the_export_describes_the_scored_tiers_rule(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high", "moderate", "low"], high={**LEVEL_3, "target_level": 3})
    files = _export(c, h, svc)
    book = _xlsx(files["xlsx"])
    about = [r[0] for r in book["About"] if r and r[0]]
    assert METHOD_NEW in about, about
    ent = _row(book["Enterprise Profile"], CODE)
    assert (ent["High"], ent["Mod"], ent["Low"], ent["Enterprise"]) == (
        3,
        "Not scored",
        "Not scored",
        "L3",
    ), ent
    # Every other subcategory has no scored tier, keeps its rows (#1073), and
    # leaves nothing out: the count is CODE's two rows.
    assert TWO_UNSCORED in about, about
    for kind in ("exec_pdf", "exec_docx", "full_pdf", "full_docx"):
        flat = _flat_text(kind, files[kind])
        assert SUMMARY_NEW in flat, kind
        assert TWO_UNSCORED in flat, kind
        assert SUMMARY_OLD not in flat, kind
    for kind in ("full_pdf", "full_docx"):
        assert "Not scored" in _flat_text(kind, files[kind]), kind
    assert METHOD_OLD not in about


def test_the_count_line_has_a_singular(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high", "moderate"], high={**LEVEL_3, "target_level": 3})
    about = [r[0] for r in _xlsx(_export(c, h, svc)["xlsx"])["About"] if r and r[0]]
    assert ONE_UNSCORED in about, about


def test_a_frozen_assessment_exports_the_old_copy(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, aid = _world(c, h, tiers=["high", "moderate", "low"], high={**LEVEL_3, "target_level": 3})
    assert c.post(f"/csf/assessments/{aid}/approve", headers=h).status_code == 200
    _set_rollup_rules(aid, 1)
    files = _export(c, h, svc)
    book = _xlsx(files["xlsx"])
    about = [r[0] for r in book["About"] if r and r[0]]
    assert METHOD_OLD in about, about  # positive first
    ent = _row(book["Enterprise Profile"], CODE)
    assert (ent["High"], ent["Mod"], ent["Low"], ent["Enterprise"]) == (3, 1, 1, "L1"), ent
    assert METHOD_NEW not in about
    for kind in ("exec_pdf", "exec_docx", "full_pdf", "full_docx"):
        flat = _flat_text(kind, files[kind])
        assert SUMMARY_OLD in flat, kind
        assert "not part of the Enterprise level" not in flat, kind
        assert "Not scored" not in flat, kind


# --- "Not measured": a target only on unscored tiers (#736 6127667643) -------------------


def test_a_target_only_on_an_unscored_tier_reads_not_measured(app_client) -> None:  # noqa: F811
    """HIGH scored with no target, MODERATE target-only. Under ruling 2 (a)
    the target is not counted, and the Target and Gap cells say why."""
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high", "moderate"], high=LEVEL_3)
    _patch(c, h, svc, "moderate", {"target_level": 4})
    sub = _sub(c, h, svc)
    assert (sub["enterprise_level"], sub["target_level"], sub["gap"]) == (3, None, False), sub
    assert (sub["target_not_measured"], sub["unscored_tiers"]) == (True, ["moderate"]), sub
    files = _export(c, h, svc)
    ent = _row(_xlsx(files["xlsx"])["Enterprise Profile"], CODE)
    assert (ent["Target"], ent["Gap"], ent["Mod"]) == (
        "Not measured",
        "Not measured",
        "Not scored",
    ), ent
    for kind in ("full_pdf", "full_docx"):
        assert "Not measured" in _flat_text(kind, files[kind]), kind


def test_a_row_with_no_target_at_all_still_reads_no_target(app_client) -> None:  # noqa: F811
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high", "moderate"], high=LEVEL_3)
    sub = _sub(c, h, svc)
    assert (sub["target_level"], sub["target_not_measured"]) == (None, False), sub
    ent = _row(_xlsx(_export(c, h, svc)["xlsx"])["Enterprise Profile"], CODE)
    assert (ent["Target"], ent["Gap"]) == (None, "No target"), ent


# --- Review of #1077: a kept entry keeps "Not measured" (F1) -------------------------------


def test_a_kept_entry_on_an_unmeasured_target_reads_not_measured(app_client) -> None:  # noqa: F811
    """HIGH scored with no target, MODERATE target-only at 4, and a POA&M entry
    on the code: the Enterprise Profile reads "Not measured", and so do the
    gap-actions response and the Action Plan sheet."""
    c, h = app_client
    svc, _aid = _world(c, h, tiers=["high", "moderate"], high=LEVEL_3)
    _patch(c, h, svc, "moderate", {"target_level": 4})
    r = c.put(f"/csf/services/{svc}/gap-actions/{CODE}", headers=h, json={"owner": "Ops lead"})
    assert r.status_code == 200, r.text
    kept = _gap_actions(c, h, svc)[CODE]
    assert (kept["no_longer_gap"], kept["target_not_measured"]) == (True, True), kept
    files = _export(c, h, svc)
    book = _xlsx(files["xlsx"])
    assert _row(book["Enterprise Profile"], CODE)["Target"] == "Not measured"
    plan = _row(book["Action Plan"], CODE)
    assert (plan["Owner"], plan["Target"]) == ("Ops lead", "Not measured"), plan
