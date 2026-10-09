"""Every service's deliverables print free text as text: PDF markup (#775) and
XLSX formulas (#972).

PDF. reportlab's `Paragraph` parses its input as markup. Unescaped, a service
title "ATT&CK Coverage" printed "ATT&CK; Coverage", a client named
"R&D <Labs> Co" printed "R&D; Co", and a client named "Acme </b> Inc" made the
whole finalize fail with ValueError -- in ATT&CK, CSF, Zero Trust and Tech Debt
alike (#736 comment 6069694559, measured). Each `render_pdf` now prints the
title and the legal name through `app.pdf_export.pdf_text`.

XLSX. openpyxl stores a string starting with "=" as a live formula, and a
leading "+", "-", "@", tab or carriage return is the same attack when the cell
is re-read. Every free-text cell in CSF, Zero Trust, Tech Debt and Risk now goes
through `app.xlsx_export.safe_text_row`, ATT&CK's guard made shared: the cell
is TEXT (data type "s" with Excel's quote prefix) and reads back exactly as
typed.

Every check goes through the surface the client reaches: the legal name is set
through the admin route that creates a client, the title through the route
that opens the service, the documents through each service's finalize (or the
Risk export) route, and the files are downloaded and read back -- PDF with
pypdf, XLSX with openpyxl. The positive state is asserted first.
"""

from __future__ import annotations

import io
import json
import os
import re
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.ai.llm import FixtureProvider, LLMClient, LLMResponse
from app.pdf_export import pdf_escape
from tests._ai_mode import pdf_text as read_pdf
from tests._ai_runs import tech_debt_extract

pytestmark = pytest.mark.unit

#: Any HTML/XML entity: what a markup round trip leaves behind.
_ENTITY = re.compile(r"&\w+;")

NAME = "R&D <Labs> Co"
TITLE = "Ops&Sec <Review> 2026"
CLOSING_TAG_NAME = "Acme </b> Inc"
FORMULA_NAME = '=HYPERLINK("http://example.invalid","click")'
SERVICES = ("attack", "csf", "zt", "tech_debt")


@pytest.fixture()
def app_client(tmp_path) -> Iterator[tuple[TestClient, FixtureProvider]]:
    url = f"sqlite:///{tmp_path / 'shield-escape.db'}"
    os.environ["DATABASE_URL"] = url
    api_root = Path(__file__).resolve().parents[2]
    cfg = Config(str(api_root / "alembic.ini"))
    cfg.set_main_option("script_location", str(api_root / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")
    engine = create_engine(url, future=True)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    from app.db.session import get_db
    from app.main import create_app
    from app.routes.artifacts import _storage_dep
    from app.routes.risk import _llm_dep as risk_llm
    from app.routes.tech_debt import _llm_dep as td_llm
    from app.storage.local import LocalFilesystemStorage

    def override_get_db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    provider = FixtureProvider()
    storage = LocalFilesystemStorage(tmp_path / "storage")
    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[risk_llm] = lambda: LLMClient(provider)
    app.dependency_overrides[td_llm] = lambda: LLMClient(provider)
    app.dependency_overrides[_storage_dep] = lambda: storage
    with TestClient(app) as c:
        yield c, provider


def _tenant(c: TestClient, legal_name: str) -> tuple[str, str, dict]:
    """An admin, and a client created with this legal name through the admin
    route. Returns (bearer, client id, headers naming that client)."""
    r = c.post(
        "/auth/register",
        json={
            "email": "admin@kentro.example",
            "password": "correct horse battery staple!",
            "display_name": "Kentro Admin",
        },
    )
    assert r.status_code == 201, r.text
    bearer = r.json()["tokens"]["access_token"]
    created = c.post(
        "/admin/clients",
        headers={"Authorization": f"Bearer {bearer}"},
        json={"legal_name": legal_name},
    )
    assert created.status_code in (200, 201), created.text
    cid = created.json()["id"]
    return bearer, cid, {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}


def _finalize_attack(c, provider, bearer, h, title: str) -> dict:
    svc = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": title})
    assert svc.status_code in (200, 201), svc.text
    a = c.post(f"/attack/services/{svc.json()['id']}/assessments", headers=h)
    assert a.status_code in (200, 201), a.text
    ok = c.post(f"/attack/assessments/{a.json()['id']}/approve", headers=h)
    assert ok.status_code == 200, ok.text
    return c.post(f"/attack/services/{svc.json()['id']}/deliverables/finalize", headers=h)


def _finalize_csf(c, provider, bearer, h, title: str) -> dict:
    svc = c.post("/csf/services", headers=h, json={"kind": "nist_csf", "title": title})
    assert svc.status_code in (200, 201), svc.text
    a = c.post(f"/csf/services/{svc.json()['id']}/assessments", headers=h)
    assert a.status_code in (200, 201), a.text
    first = a.json()["answers"][0]
    r = c.patch(
        f"/csf/answers/{first['id']}", headers=h, json={"maturity_tier": 2, "notes": "@notes"}
    )
    assert r.status_code == 200, r.text
    ok = c.post(f"/csf/assessments/{a.json()['id']}/approve", headers=h)
    assert ok.status_code == 200, ok.text
    return c.post(f"/csf/services/{svc.json()['id']}/deliverables/finalize", headers=h)


def _finalize_zt(c, provider, bearer, h, title: str) -> dict:
    svc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": title})
    assert svc.status_code in (200, 201), svc.text
    a = c.post(f"/zt/services/{svc.json()['id']}/assessments", headers=h)
    assert a.status_code in (200, 201), a.text
    first = a.json()["answers"][0]
    r = c.patch(f"/zt/answers/{first['id']}", headers=h, json={"maturity_stage": 2, "notes": "@n"})
    assert r.status_code == 200, r.text
    ok = c.post(f"/zt/assessments/{a.json()['id']}/approve", headers=h)
    assert ok.status_code == 200, ok.text
    return c.post(f"/zt/services/{svc.json()['id']}/deliverables/finalize", headers=h)


#: What the extraction model returns for the uploaded row, in the extraction
#: prompt's own item shape. Its free text opens with formula triggers.
_TD_ITEM = {
    "name": "=SUM(A1:A9)",
    "vendor": "+Vendor",
    "category": None,
    "function": "-function",
    "annual_cost_usd": 100,
    "license_count": 5,
    "notes": "@notes",
    "security_related": True,
    "security_functions": ["detect"],
    "confidence_pct": 90,
    "source_row_index": 0,
}


def _finalize_tech_debt(c, provider, bearer, h, title: str) -> dict:
    provider.register(
        "extract.capabilities", lambda _p: LLMResponse(json.dumps({"items": [_TD_ITEM]}))
    )
    svc = c.post("/tech-debt/services", headers=h, json={"title": title})
    assert svc.status_code in (200, 201), svc.text
    up = c.post(
        "/artifacts",
        headers=h,
        files={"file": ("x.csv", io.BytesIO(b"name\nrow0\n"), "text/csv")},
    )
    assert up.status_code == 201, up.text
    body = tech_debt_extract(c, svc.json()["id"], h, up.json()["id"])
    r = c.patch(
        f"/tech-debt/capability-items/{body['items'][0]['id']}",
        headers=h,
        json={"disposition": "keep"},
    )
    assert r.status_code == 200, r.text
    ok = c.post(f"/tech-debt/capability-lists/{body['id']}/approve", headers=h)
    assert ok.status_code == 200, ok.text
    return c.post(f"/tech-debt/services/{svc.json()['id']}/deliverables/finalize", headers=h)


_FINALIZE = {
    "attack": _finalize_attack,
    "csf": _finalize_csf,
    "zt": _finalize_zt,
    "tech_debt": _finalize_tech_debt,
}


def _deliver(app_client, service: str, *, name: str, title: str) -> tuple[object, dict]:
    """Finalize `service` for a client named `name` with a service titled
    `title`. Returns (the finalize response, {"pdf": bytes, "xlsx": bytes})
    -- the files only when finalize succeeded."""
    c, provider = app_client
    bearer, _cid, h = _tenant(c, name)
    fin = _FINALIZE[service](c, provider, bearer, h, title)
    files: dict = {}
    if fin.status_code == 201:
        for kind in ("pdf", "xlsx"):
            d = c.get(f"/artifacts/{fin.json()[f'{kind}_artifact_id']}/download", headers=h)
            assert d.status_code == 200, d.text
            files[kind] = d.content
    return fin, files


# --- the shared helper ---------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "escaped"),
    [
        ("ATT&CK Coverage", "ATT&amp;CK Coverage"),
        ("R&D <Labs> Co", "R&amp;D &lt;Labs&gt; Co"),
        ("Acme </b> Inc", "Acme &lt;/b&gt; Inc"),
        ('say "hi"', 'say "hi"'),  # quote=False: quotes are not markup here
    ],
)
def test_pdf_escape_escapes_exactly_the_three_markup_characters(raw, escaped) -> None:
    assert pdf_escape(raw) == escaped


# --- PDF: per service and per site ------------------------------------------------


@pytest.mark.parametrize("service", SERVICES)
def test_the_client_legal_name_prints_exactly_in_the_pdf(app_client, service) -> None:
    fin, files = _deliver(app_client, service, name=NAME, title="Plain title")
    assert fin.status_code == 201, fin.text
    text = read_pdf(files["pdf"])
    assert NAME in text  # the positive state first
    assert _ENTITY.search(text) is None, _ENTITY.findall(text)


@pytest.mark.parametrize("service", SERVICES)
def test_the_service_title_prints_exactly_in_the_pdf(app_client, service) -> None:
    fin, files = _deliver(app_client, service, name="Plain Client", title=TITLE)
    assert fin.status_code == 201, fin.text
    text = read_pdf(files["pdf"])
    assert TITLE in text
    assert _ENTITY.search(text) is None, _ENTITY.findall(text)


@pytest.mark.parametrize("service", SERVICES)
def test_an_attandck_title_prints_without_an_entity(app_client, service) -> None:
    fin, files = _deliver(app_client, service, name="Plain Client", title="ATT&CK Coverage")
    assert fin.status_code == 201, fin.text
    text = read_pdf(files["pdf"])
    assert "ATT&CK Coverage" in text
    assert "ATT&CK;" not in text


@pytest.mark.parametrize("service", SERVICES)
def test_a_legal_name_with_a_closing_tag_finalizes(app_client, service) -> None:
    """Unescaped, `Paragraph` raised ValueError and finalize failed for the
    tenant: no deliverable at all."""
    fin, files = _deliver(app_client, service, name=CLOSING_TAG_NAME, title="Plain title")
    assert fin.status_code == 201, fin.text
    assert CLOSING_TAG_NAME in read_pdf(files["pdf"])


# --- XLSX: free text is text --------------------------------------------------------


def _cell_with(wb, value: str) -> list:
    """Every cell, on any sheet, whose value is exactly `value` -- at least
    one, so the value is first shown to read back exactly as typed."""
    found = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value == value]
    assert found, f"{value!r} is in no cell"
    return found


def _assert_text(cells: list) -> None:
    for cell in cells:
        where = f"{cell.parent.title}!{cell.coordinate}"
        assert cell.data_type == "s", f"{where} is {cell.data_type!r}, not text"
        assert cell.quotePrefix is True, f"{where} is not marked as text"


@pytest.mark.parametrize("service", ("attack", "csf", "zt"))
def test_a_formula_legal_name_is_text_in_the_workbook(app_client, service) -> None:
    fin, files = _deliver(app_client, service, name=FORMULA_NAME, title="+Title")
    assert fin.status_code == 201, fin.text
    wb = load_workbook(io.BytesIO(files["xlsx"]))
    _assert_text(_cell_with(wb, FORMULA_NAME))  # reads back exactly as typed
    _assert_text(_cell_with(wb, "+Title"))


@pytest.mark.parametrize(("service", "note"), [("csf", "@notes"), ("zt", "@n")])
def test_a_note_opening_with_at_is_text_in_the_workbook(app_client, service, note) -> None:
    fin, files = _deliver(app_client, service, name="Plain Client", title="Plain title")
    assert fin.status_code == 201, fin.text
    wb = load_workbook(io.BytesIO(files["xlsx"]))
    _assert_text(_cell_with(wb, note))


def test_tech_debt_free_text_cells_are_text_in_the_workbook(app_client) -> None:
    fin, files = _deliver(app_client, "tech_debt", name="Plain Client", title="Plain title")
    assert fin.status_code == 201, fin.text
    wb = load_workbook(io.BytesIO(files["xlsx"]))
    for value in ("=SUM(A1:A9)", "+Vendor", "-function", "@notes"):
        _assert_text(_cell_with(wb, value))


def test_risk_free_text_cells_are_text_in_the_workbook(app_client) -> None:
    from tests.unit.test_risk_register import _entries_payload, _entry, _seed_attack_and_zt

    c, provider = app_client
    bearer, cid, h = _tenant(c, FORMULA_NAME)
    _seed_attack_and_zt(c, bearer, cid)
    payload = _entries_payload(
        _entry("=Risk title", description="+description", rationale="@rationale")
    )
    provider.register_static("risk_synthesize", LLMResponse(payload))
    g = c.post(f"/risk/clients/{cid}/register/generate", headers=h)
    assert g.status_code == 201, g.text
    ex = c.post(f"/risk/clients/{cid}/register/export", headers=h)
    assert ex.status_code == 200, ex.text
    d = c.get(f"/artifacts/{ex.json()['xlsx_artifact_id']}/download", headers=h)
    assert d.status_code == 200, d.text
    wb = load_workbook(io.BytesIO(d.content))
    for value in ("=Risk title", "+description", "@rationale"):
        _assert_text(_cell_with(wb, value))


# --- Zero Trust's catalog-name sentences (#775 (c)) -------------------------------
#
# No current capability name carries "&" or "<", so these change no output
# today. The DoD activity text already does ("decides A&O modifications"), so
# the sites are pinned with a sentence that does, through the real finalize.


@pytest.mark.parametrize("hook", ["target_cap_sentences", "stage_above_max_sentences"])
def test_zt_catalog_sentences_print_as_text(app_client, monkeypatch, hook) -> None:
    import app.zt.exporters as zt_exporters

    sentence = "DoD 1.2 decides A&O modifications <draft>."
    monkeypatch.setattr(zt_exporters, hook, lambda *a, **kw: [sentence])
    fin, files = _deliver(app_client, "zt", name="Plain Client", title="Plain title")
    assert fin.status_code == 201, fin.text
    text = read_pdf(files["pdf"])
    assert sentence in text
    assert _ENTITY.search(text) is None, _ENTITY.findall(text)


# --- control characters: stripped, and every strip logged (advisor, #972) --------


def test_a_control_character_is_stripped_and_the_strip_is_logged(app_client, capsys) -> None:
    """The cell is clean, and ONE warning names the sheet, the column header,
    the row and how many characters went -- never the value. `capsys`, not
    `caplog`: structlog renders to stdout (`app/logging.py`), so stdlib log
    capture sees nothing (the lesson `test_risk_link_scope.py` records).

    `app_client` comes FIRST and is otherwise unused: the app's lifespan is
    what calls `configure_logging`, so without it this test passed only when
    an earlier test had started the app, and failed run alone. Listed before
    `capsys`, the lifespan binds the stream structlog treats as stdout, which
    it resolves at call time, so the capture sees the line."""
    from openpyxl import Workbook

    from app.xlsx_export import safe_text_row

    wb = Workbook()
    ws = wb.active
    ws.title = "Answers"
    safe_text_row(ws, ["Code", "Notes"])
    capsys.readouterr()  # the header row removed nothing
    secret = "Sensitive\x00note\x07here\x1b"
    safe_text_row(ws, ["ID.AM-01", secret])
    assert ws["B2"].value == "Sensitivenotehere"  # the positive state first
    out = capsys.readouterr().out
    events = [json.loads(line) for line in out.splitlines() if line.startswith("{")]
    removed = [e for e in events if e.get("event") == "xlsx_export.control_characters_removed"]
    assert len(removed) == 1, events
    (e,) = removed
    assert e["level"] == "warning"
    assert (e["sheet"], e["column"], e["row"], e["removed"]) == ("Answers", "Notes", 2, 3)
    for fragment in ("Sensitive", "note", "here", "\x00", "\\u0000"):
        assert fragment not in out, f"the log carries part of the value: {fragment!r}"


def test_a_sheet_without_a_header_names_the_column_letter(app_client, capsys) -> None:
    # `app_client` first: the lifespan configures logging (see above).
    from openpyxl import Workbook

    from app.xlsx_export import safe_text_row

    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    safe_text_row(ws, ["clean", "x\x01"])
    out = capsys.readouterr().out
    (e,) = [
        json.loads(line)
        for line in out.splitlines()
        if line.startswith("{") and "control_characters_removed" in line
    ]
    assert (e["sheet"], e["column"], e["row"], e["removed"]) == ("Cover", "B", 1, 1)


def test_a_clean_row_logs_nothing(app_client, capsys) -> None:
    # `app_client` first, or this passes vacuously on an unconfigured logger.
    from openpyxl import Workbook

    from app.xlsx_export import safe_text_row

    ws = Workbook().active
    safe_text_row(ws, ["Header"])
    safe_text_row(ws, ["clean value", 3, None])
    assert "control_characters_removed" not in capsys.readouterr().out
