"""Gene's condition on #620, through the client-facing routes (D-094).

A released assessment keeps rendering what was delivered. `main`'s rendering
of a fixed world was recorded ONCE (`tests/golden/parent_rules_old/`, see its
RECORDED.md) and is compared here, never regenerated. This test loads that
world's database -- recorded at `main`'s migration head -- and runs `alembic
upgrade head` on it, so migration 0054's backfill is the path under test, as it
will be in production.

The same world with the rule set flipped to D-094's is the other half: it must
render DIFFERENTLY, or the byte-identical half proves nothing.
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
from collections.abc import Iterator
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session, sessionmaker

from app.ai.llm import FixtureProvider, LLMClient, LLMResponse

API_ROOT = Path(__file__).resolve().parents[2]
GOLDEN = API_ROOT / "tests" / "golden" / "parent_rules_old"
PASSWORD = "correct horse battery staple!"


@pytest.fixture()
def world(tmp_path) -> Iterator[tuple[TestClient, sessionmaker, FixtureProvider, Path]]:
    db_path = tmp_path / "golden.db"
    con = sqlite3.connect(db_path)
    con.executescript((GOLDEN / "world.sql").read_text(encoding="utf-8"))
    con.close()
    url = f"sqlite:///{db_path}"
    os.environ["DATABASE_URL"] = url
    cfg = Config(str(API_ROOT / "alembic.ini"))
    cfg.set_main_option("script_location", str(API_ROOT / "alembic"))
    cfg.set_main_option("sqlalchemy.url", url)
    command.upgrade(cfg, "head")  # 0054's backfill runs HERE
    Sess = sessionmaker(bind=create_engine(url, future=True), autoflush=False, future=True)

    from app.db.session import get_db
    from app.main import create_app
    from app.routes.artifacts import _storage_dep
    from app.routes.risk import _llm_dep
    from app.storage.local import LocalFilesystemStorage

    def override_get_db() -> Iterator[Session]:
        db = Sess()
        try:
            yield db
        finally:
            db.close()

    storage_dir = tmp_path / "storage"
    storage = LocalFilesystemStorage(storage_dir)
    provider = FixtureProvider()
    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_storage_dep] = lambda: storage
    app.dependency_overrides[_llm_dep] = lambda: LLMClient(provider)
    with TestClient(app) as c:
        yield c, Sess, provider, storage_dir


def _ids(Sess) -> dict[str, str]:
    with Sess() as s:
        out = {"cid": s.execute(text("SELECT id FROM client")).scalar_one()}
        for title in ("Golden A", "Golden B"):
            out[title] = s.execute(
                text("SELECT id FROM services WHERE title = :t"), {"t": title}
            ).scalar_one()
    return {k: str(v) for k, v in out.items()}


def _login(c: TestClient, email: str, cid: str) -> dict:
    r = c.post(
        "/auth/login", json={"email": email, "password": PASSWORD}, headers={"X-Client-Id": cid}
    )
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}", "X-Client-Id": cid}


def _dashboard(c: TestClient, ids: dict) -> bytes:
    h = _login(c, "client@example.com", ids["cid"])
    r = c.get(f"/clients/{ids['cid']}/attack/{ids['Golden A']}/dashboard", headers=h)
    assert r.status_code == 200, r.text
    return r.content


@pytest.mark.unit
def test_the_migration_marked_both_approved_assessments_as_old_rules(world) -> None:
    _, Sess, _, _ = world
    with Sess() as s:
        got = s.execute(text("SELECT status, parent_rules FROM attack_assessments")).all()
    assert sorted(got) == [("APPROVED", 1), ("RELEASED", 1)]


@pytest.mark.unit
def test_a_released_assessment_renders_byte_identically_to_main(world) -> None:
    c, Sess, _, _ = world
    assert _dashboard(c, _ids(Sess)) == (GOLDEN / "dashboard.json").read_bytes()


@pytest.mark.unit
def test_finalizing_an_assessment_approved_before_620_renders_as_main_did(world) -> None:
    c, Sess, _, storage_dir = world
    ids = _ids(Sess)
    h = _login(c, "admin@example.com", ids["cid"])
    before = set(storage_dir.rglob("*.xlsx"))
    r = c.post(f"/attack/services/{ids['Golden B']}/deliverables/finalize", headers=h)
    assert r.status_code == 201, r.text
    (xlsx,) = set(storage_dir.rglob("*.xlsx")) - before
    wb = load_workbook(io.BytesIO(xlsx.read_bytes()))
    cells = {
        ws.title: [[cell.value for cell in row] for row in ws.iter_rows()] for ws in wb.worksheets
    }
    got = json.loads(json.dumps({"summary": r.json()["summary"], "xlsx": cells}, default=str))
    assert got == json.loads((GOLDEN / "finalize_b.json").read_text(encoding="utf-8"))


@pytest.mark.unit
def test_the_risk_register_takes_the_findings_main_took(world) -> None:
    c, Sess, provider, _ = world
    ids = _ids(Sess)
    sent: list[str] = []

    def spy(payload: dict) -> LLMResponse:
        sent.extend(
            f["source_id"] for f in payload.get("findings", []) if f.get("kind") == "attack"
        )
        return LLMResponse(json.dumps({"entries": []}))

    provider.register("risk_synthesize", spy)
    c.post(
        f"/risk/clients/{ids['cid']}/register/generate",
        headers=_login(c, "admin@example.com", ids["cid"]),
    )
    assert sorted(sent) == json.loads((GOLDEN / "risk_findings.json").read_text(encoding="utf-8"))


@pytest.mark.unit
def test_the_same_world_under_the_new_rules_renders_differently(world) -> None:
    """The other half. Without it, the byte-identical test above would also
    pass for a change that did nothing at all."""
    c, Sess, _, _ = world
    ids = _ids(Sess)
    with Sess() as s:
        s.execute(text("UPDATE attack_assessments SET parent_rules = 2 WHERE status = 'RELEASED'"))
        s.commit()
    body = json.loads(_dashboard(c, ids))
    parent = next(t for t in body["techniques"] if t["code"] == "T1001")
    assert body["parents_computed"] is True
    assert parent["computed_parent"] is True
    assert (parent["detection_tools"], parent["rationale"]) == ([], None)
    # Main showed the parent pending on its own citation; under D-094 it is
    # pending only through a child, and none of these children is.
    assert parent["pending_review"] is False
