"""A bundle part's stored cost is not a saving (for #929, advisor ruling on #736).

A split bundle's part (`parent_item_id` set) carries no cost of its own: the
bundle row holds the licence value. Since #927 a cost can no longer be PATCHed
onto a part, but one stored before that refusal is still in the row. #1044
stopped every screen showing it ("Included in {bundle}"), and the advisor ruled
(#736 6115773247) that it stops counting toward savings too, so the savings
figures stop contradicting that cell.

The rule, which keeps the earlier ruling that a Cut part with no cost is an
UNKNOWN saving: a part's stored cost is read as no cost. So a Cut part still
makes savings a floor ("≥ $…"), and its legacy figure never adds to it.

The plan: CrowdStrike Falcon (120,000) Cut, and Microsoft Defender for Endpoint,
a part of Microsoft 365 E5 holding a legacy 20,000, Cut; everything else Keep.
Before: $140,000, known. After: at least $120,000.

One test per surface, each through that surface's own endpoint or file, as
`test_tech_debt_savings_surfaces.py` pins the six for #804.

Two more halves of the same rule, from review of #1044:

  - each file's capability list names the bundle in a part's cost cell
    ("Included in {bundle}", the dashboard's approved wording), so a row the
    reader can see never carries the figure the savings left out;
  - the per-category redundancy savings carry their own `savings_cost_known`,
    so a category whose Cut part has no usable cost reads as a floor there too.
"""

from __future__ import annotations

import io

import pytest

from tests.unit.test_tech_debt_bundle_components import (
    BUNDLE,
    SPLIT_BESIDE_TOOLS,
    _docx_text,
    _extract_and_split,
    _pdf_text,
    _store_legacy_part_cost,
)
from tests.unit.test_tech_debt_dashboard import (  # noqa: F401  (fixture)
    _register,
    app_client,
)

pytestmark = pytest.mark.unit

LEGACY = 20000
PLAN = {"CrowdStrike Falcon": "cut", "Microsoft Defender for Endpoint": "cut"}
#: CrowdStrike alone: the part's legacy 20,000 is not a saving.
SAVINGS = 120_000.0


PARTS = [c["name"] for c in SPLIT_BESIDE_TOOLS]
#: The XLSX floor note beside the savings cell (`exporters.build_xlsx`).
XLSX_FLOOR_NOTE = "≥ (one or more cut rows missing a cost)"


def _release(app_client, bundle_name: str | None = None) -> dict:  # noqa: F811
    c, provider = app_client
    h, svc, lst, items, bundle_id = _extract_and_split(c, provider, SPLIT_BESIDE_TOOLS)
    client = _register(c, "client@example.com")
    if bundle_name is not None:
        r = c.patch(
            f"/tech-debt/capability-items/{bundle_id}", headers=h, json={"name": bundle_name}
        )
        assert r.status_code == 200, r.text
    _store_legacy_part_cost(bundle_id, LEGACY, expected=len(SPLIT_BESIDE_TOOLS))
    for it in items:
        r = c.patch(
            f"/tech-debt/capability-items/{it['id']}",
            headers=h,
            json={"disposition": PLAN.get(it["name"], "keep")},
        )
        assert r.status_code == 200, r.text
    plan_card = c.get(f"/tech-debt/services/{svc}/consolidation-plan", headers=h)
    assert plan_card.status_code == 200, plan_card.text
    preview = c.post(
        f"/tech-debt/capability-lists/{lst['id']}/savings-preview",
        headers=h,
        json={"dispositions": {}},
    )
    assert preview.status_code == 200, preview.text
    assert c.post(f"/tech-debt/capability-lists/{lst['id']}/approve", headers=h).status_code == 200
    fin = c.post(f"/tech-debt/services/{svc}/deliverables/finalize", headers=h)
    assert fin.status_code == 201, fin.text
    deliv = fin.json()
    rel = c.post(f"/tech-debt/deliverables/{deliv['id']}/release", headers=h)
    assert rel.status_code == 200, rel.text
    files = {}
    for kind in ("xlsx", "pdf", "docx"):
        got = c.get(f"/artifacts/{deliv[f'{kind}_artifact_id']}/download", headers=h)
        assert got.status_code == 200, got.text
        files[kind] = got.content
    cid = client["user"]["client_id"]
    ch = {"Authorization": f"Bearer {client['tokens']['access_token']}", "X-Client-Id": cid}
    dash = c.get(f"/clients/{cid}/tech-debt/{svc}/dashboard", headers=ch)
    assert dash.status_code == 200, dash.text
    home = c.get(f"/clients/{cid}/value-summary", headers=ch)
    assert home.status_code == 200, home.text
    return {
        "summary": deliv["summary"],
        "plan": plan_card.json(),
        "preview": preview.json(),
        "dash": dash.json(),
        "home": home.json(),
        "files": files,
    }


@pytest.fixture()
def released(app_client):  # noqa: F811
    return _release(app_client)


def _xlsx_rows(raw: bytes) -> dict[str, list]:
    """The capability sheet's rows by name (column A), header included."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw)).worksheets[0]
    header = [c.value for c in ws[1]]
    rows = {str(r[0].value): [c.value for c in r] for r in ws.iter_rows(min_row=2) if r[0].value}
    rows["__header__"] = header
    return rows


def _docx_all_text(raw: bytes) -> str:
    """Paragraphs AND table cells: `_docx_text` reads paragraphs only, and the
    capability list is a table."""
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return " ".join(" ".join(parts).split())


def _xlsx_cost(raw: bytes, name: str):
    rows = _xlsx_rows(raw)
    col = rows["__header__"].index("Annual Cost (USD)")
    return rows[name][col]


def test_the_deliverable_summary_leaves_out_a_parts_legacy_cost(released) -> None:
    assert "≥ $120,000 estimated annual savings" in released["summary"], released["summary"]


def test_the_pdf_and_docx_leave_out_a_parts_legacy_cost(released) -> None:
    for text in (_pdf_text(released["files"]["pdf"]), _docx_text(released["files"]["docx"])):
        assert "Estimated annual savings: ≥ $120,000" in text, text
        assert "$140,000" not in text


def test_the_plan_card_leaves_out_a_parts_legacy_cost(released) -> None:
    assert released["plan"]["estimated_annual_savings"] == SAVINGS
    assert released["plan"]["savings_cost_known"] is False


def test_the_what_if_leaves_out_a_parts_legacy_cost(released) -> None:
    assert released["preview"]["estimated_annual_savings"] == SAVINGS
    assert released["preview"]["savings_cost_known"] is False


def test_the_client_dashboard_headline_leaves_out_a_parts_legacy_cost(released) -> None:
    assert released["dash"]["identified_savings_usd"] == SAVINGS
    assert released["dash"]["savings_cost_known"] is False


def test_the_client_dashboard_redundancy_leaves_out_a_parts_legacy_cost(released) -> None:
    by_cat = {r["category"]: r["savings_usd"] for r in released["dash"]["redundancies"]}
    assert by_cat["EDR"] == SAVINGS, by_cat


def test_the_home_value_card_leaves_out_a_parts_legacy_cost(released) -> None:
    assert released["home"]["tech_debt_savings_usd"] == SAVINGS
    assert released["home"]["tech_debt_savings_cost_known"] is False


def test_the_xlsx_savings_is_a_floor_without_the_parts_legacy_cost(released) -> None:
    rows = _xlsx_rows(released["files"]["xlsx"])
    savings = rows["Estimated annual savings"]
    assert savings[4] == SAVINGS, savings
    assert savings[5] == XLSX_FLOOR_NOTE, savings


# --- a part's cost cell in the files ----------------------------------------


def test_the_xlsx_names_the_bundle_in_a_parts_cost_cell(released) -> None:
    raw = released["files"]["xlsx"]
    for part in PARTS:
        assert _xlsx_cost(raw, part) == f"Included in {BUNDLE}", part
    # A source row keeps its number.
    assert _xlsx_cost(raw, "CrowdStrike Falcon") == 120000


def test_the_pdf_and_docx_name_the_bundle_in_a_parts_cost_cell(released) -> None:
    for text in (_pdf_text(released["files"]["pdf"]), _docx_all_text(released["files"]["docx"])):
        for part in PARTS:
            assert f"{part} Microsoft" in text, text
        assert text.count(f"Included in {BUNDLE}") == len(PARTS), text
        # The legacy figure appears nowhere, and a source row keeps its cost.
        assert "$20,000" not in text, text
        assert "$120,000 Cut" in text, text


def test_a_blank_bundle_name_reads_included_in_its_bundle_in_every_file(
    app_client,
) -> None:  # noqa: F811
    got = _release(app_client, bundle_name="   ")
    raw = got["files"]["xlsx"]
    for part in PARTS:
        assert _xlsx_cost(raw, part) == "Included in its bundle", part
    for text in (_pdf_text(got["files"]["pdf"]), _docx_all_text(got["files"]["docx"])):
        assert text.count("Included in its bundle") == len(PARTS), text


# --- per-category savings carry their own floor flag -------------------------


def test_each_redundancys_savings_says_whether_it_is_a_floor(released) -> None:
    by_cat = {
        r["category"]: (r["savings_usd"], r["savings_cost_known"])
        for r in released["dash"]["redundancies"]
    }
    # EDR: CrowdStrike Cut at 120,000, and a Cut part whose cost is the bundle's.
    # IAM: Okta and Entra ID, both kept, so nothing is unknown.
    assert by_cat == {"EDR": (SAVINGS, False), "IAM": (0.0, True)}, by_cat
