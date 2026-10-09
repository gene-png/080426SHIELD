"""#737: exporting a Risk Register no longer publishes it to the client.

`finalized_at` is the single condition `clients.py::risk_dashboard` gates the
client's dashboard on, and `export` used to set it, so exporting a register to
review it internally put it in front of the client. Export now renders the
files only; `publish` is the one writer of `finalized_at`, and it refuses an
unrated entry (#844 D1), an already-published register, and the inputs export
already refuses.

Every assertion is made through the surface that decides it: the client
dashboard as a client user reads it, and the admin endpoints.
"""

from __future__ import annotations

import pytest

from tests._risk_inputs import seed_released
from tests.unit.test_risk_register import (  # noqa: F401  (app_client is a fixture)
    _admin,
    _entries_payload,
    _entry,
    _generate,
    _pdf_text,
    _session,
    app_client,
)

pytestmark = pytest.mark.unit


def _unrated(title: str) -> str:
    return '{"title": "' + title + '", "likelihood": null, "impact": null}'


def _world(app_client, *entries: str):  # noqa: F811
    """An admin, a client user of the same tenant, and a generated register."""
    c, provider = app_client
    bearer, cid = _admin(c)
    ah = {"Authorization": f"Bearer {bearer}"}
    r = c.post(f"/admin/clients/{cid}/domains", headers=ah, json={"domain": "acme.example"})
    assert r.status_code in (200, 201), r.text
    user = c.post(
        "/auth/register",
        json={
            "email": "user@acme.example",
            "password": "correct horse battery staple!",
            "display_name": "U",
        },
    )
    assert user.status_code == 201, user.text
    ch = {"Authorization": f"Bearer {user.json()['tokens']['access_token']}", "X-Client-Id": cid}
    # RELEASED, not only approved: publish needs every engaged input final
    # (Gene's rule, #737). The input refusals are in test_risk_publish_inputs.
    seed_released(c, bearer, cid)
    body = _generate(c, provider, bearer, cid, _entries_payload(*(entries or (_entry("R"),))))
    return c, provider, bearer, cid, ah, ch, body


def _client_dashboard(c, ch: dict, cid: str):
    return c.get(f"/clients/{cid}/risk/dashboard", headers=ch)


def _post(c, ah: dict, cid: str, action: str):
    return c.post(f"/risk/clients/{cid}/register/{action}", headers=ah)


def test_export_does_not_publish(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, ch, _ = _world(app_client)
    ex = _post(c, ah, cid, "export")
    assert ex.status_code == 200, ex.text
    assert ex.json()["finalized_at"] is None
    assert ex.json()["pdf_artifact_id"] is not None
    r = _client_dashboard(c, ch, cid)
    assert r.status_code == 404, r.text
    assert r.json()["error"]["reason"] == "dashboard_not_released"


def test_publish_puts_the_register_in_front_of_the_client(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, ch, _ = _world(app_client)
    pub = _post(c, ah, cid, "publish")
    assert pub.status_code == 200, pub.text
    assert pub.json()["finalized_at"] is not None
    # Publish renders its own files; nothing had to be exported first.
    assert pub.json()["pdf_artifact_id"] is not None
    r = _client_dashboard(c, ch, cid)
    assert r.status_code == 200, r.text
    assert r.json()["version"] == 1


def test_publish_refuses_an_unrated_entry_and_publishes_nothing(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, ch, _ = _world(app_client, _entry("Rated"), _unrated("Unrated"))
    pub = _post(c, ah, cid, "publish")
    assert pub.status_code == 409, pub.text
    err = pub.json()["error"]
    assert err["reason"] == "risk_register_unrated_entries"
    assert err["unrated"] == 1
    assert err["message"] == (
        "1 entry has no likelihood or impact. Rate it in the Register table before publishing."
    )
    assert _client_dashboard(c, ch, cid).status_code == 404


def test_rating_the_entry_lets_it_publish(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, ch, body = _world(app_client, _unrated("Unrated"))
    entry_id = body["entries"][0]["id"]
    r = c.patch(
        f"/risk/clients/{cid}/register/entries/{entry_id}",
        headers=ah,
        json={"likelihood": "low", "impact": "minor"},
    )
    assert r.status_code == 200, r.text
    assert _post(c, ah, cid, "publish").status_code == 200
    dash = _client_dashboard(c, ch, cid)
    assert dash.status_code == 200, dash.text
    assert dash.json()["entries_without_tier"] == 0


def test_a_published_register_cannot_be_published_again(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, _, _ = _world(app_client)
    assert _post(c, ah, cid, "publish").status_code == 200
    again = _post(c, ah, cid, "publish")
    assert again.status_code == 409, again.text
    assert again.json()["error"]["reason"] == "risk_register_already_published"


def test_publish_refuses_inputs_that_cannot_be_certified(app_client) -> None:  # noqa: F811
    """The export guards, shared rather than copied. A register whose provenance
    records excluded inputs but not used ones is the shape `seed_demo.py`
    writes; unpublished, it must not reach the client."""
    from sqlalchemy import select

    from app.models.risk_register import RiskRegister

    c, _, _, cid, ah, ch, _ = _world(app_client)
    with _session() as s:
        reg = s.execute(select(RiskRegister)).scalar_one()
        prov = dict(reg.provenance)
        prov.pop("inputs")
        reg.provenance = prov
        s.commit()
    pub = _post(c, ah, cid, "publish")
    assert pub.status_code == 409, pub.text
    assert pub.json()["error"]["reason"] == "register_inputs_not_recorded"
    assert _client_dashboard(c, ch, cid).status_code == 404


def test_ratings_stay_editable_after_export_and_lock_at_publish(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, _, body = _world(app_client)
    entry_id = body["entries"][0]["id"]
    url = f"/risk/clients/{cid}/register/entries/{entry_id}"
    assert _post(c, ah, cid, "export").status_code == 200
    after_export = c.patch(url, headers=ah, json={"impact": "minor"})
    assert after_export.status_code == 200, after_export.text
    assert _post(c, ah, cid, "publish").status_code == 200
    after_publish = c.patch(url, headers=ah, json={"impact": "major"})
    assert after_publish.status_code == 409, after_publish.text
    assert after_publish.json()["error"]["reason"] == "risk_register_published"


def test_published_files_are_rendered_at_publication(app_client) -> None:  # noqa: F811
    """A rating edited after an export must reach the client's files: publish
    renders its own, so the published PDF cannot predate the published data."""
    c, _, _, cid, ah, _, body = _world(app_client, _unrated("Was unrated"))
    ex = _post(c, ah, cid, "export").json()
    entry_id = body["entries"][0]["id"]
    c.patch(
        f"/risk/clients/{cid}/register/entries/{entry_id}",
        headers=ah,
        json={"likelihood": "low", "impact": "minor"},
    )
    pub = _post(c, ah, cid, "publish").json()
    assert pub["pdf_artifact_id"] != ex["pdf_artifact_id"]
    pdf = c.get(f"/artifacts/{pub['pdf_artifact_id']}/download", headers={**ah, "X-Client-Id": cid})
    text = " ".join(_pdf_text(pdf.content).split())
    assert "Every entry is rated." in text
    assert "Not rated:" not in text


def test_nothing_to_publish_is_not_found(app_client) -> None:  # noqa: F811
    c, _ = app_client
    bearer, cid = _admin(c)
    r = c.post(
        f"/risk/clients/{cid}/register/publish", headers={"Authorization": f"Bearer {bearer}"}
    )
    assert r.status_code == 404, r.text


def test_a_client_user_cannot_publish(app_client) -> None:  # noqa: F811
    c, _, _, cid, _, ch, _ = _world(app_client)
    r = c.post(f"/risk/clients/{cid}/register/publish", headers=ch)
    assert r.status_code == 403, r.text
    assert _client_dashboard(c, ch, cid).status_code == 404


def test_a_client_user_cannot_download_an_unpublished_export(app_client) -> None:  # noqa: F811
    """THE GUARD THE CLIENT RELIES ON, since export dropped its input guards
    (advisor, #736 5998764095). A draft export renders freely -- here from
    inputs that are approved and NOT released, so the register is a draft --
    and what keeps it from the client is that a Risk export is not a
    `Deliverable`: `/artifacts` serves it to an admin and 404s a client user.

    POSITIVE FIRST: the admin download succeeds, so the client's 404 is about
    the reader, not a missing or unrendered file.
    """
    from tests._risk_inputs import seed_attack_and_zt

    c, provider = app_client
    bearer, cid = _admin(c)
    ah = {"Authorization": f"Bearer {bearer}"}
    r = c.post(f"/admin/clients/{cid}/domains", headers=ah, json={"domain": "acme.example"})
    assert r.status_code in (200, 201), r.text
    user = c.post(
        "/auth/register",
        json={
            "email": "user@acme.example",
            "password": "correct horse battery staple!",
            "display_name": "U",
        },
    )
    assert user.status_code == 201, user.text
    ch = {"Authorization": f"Bearer {user.json()['tokens']['access_token']}", "X-Client-Id": cid}
    seed_attack_and_zt(c, bearer, cid)  # approved, NOT released
    _generate(c, provider, bearer, cid, _entries_payload(_entry("R")))
    ex = _post(c, ah, cid, "export")
    assert ex.status_code == 200, ex.text
    body = ex.json()
    assert body["finalized_at"] is None
    for kind in ("pdf", "xlsx", "docx"):
        ok = c.get(
            f"/artifacts/{body[f'{kind}_artifact_id']}/download",
            headers={**ah, "X-Client-Id": cid},
        )
        assert ok.status_code == 200, (kind, ok.text)
    for kind in ("pdf", "xlsx", "docx"):
        r = c.get(f"/artifacts/{body[f'{kind}_artifact_id']}/download", headers=ch)
        assert r.status_code == 404, (kind, r.text)


def _texts(c, ah: dict, cid: str, body: dict) -> dict[str, str]:
    import io

    from docx import Document
    from openpyxl import load_workbook

    dh = {**ah, "X-Client-Id": cid}
    raw = {
        k: c.get(f"/artifacts/{body[f'{k}_artifact_id']}/download", headers=dh).content
        for k in ("pdf", "docx", "xlsx")
    }
    wb = load_workbook(io.BytesIO(raw["xlsx"]))
    return {
        "pdf": " ".join(_pdf_text(raw["pdf"]).split()),
        "docx": "\n".join(p.text for p in Document(io.BytesIO(raw["docx"])).paragraphs),
        "xlsx": "\n".join(
            str(r[0]) for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]
        ),
    }


def test_an_export_of_an_unpublished_register_is_marked_draft(app_client) -> None:  # noqa: F811
    """#737, Gene's marker, verbatim, in all three files."""
    c, _, _, cid, ah, _, _ = _world(app_client)
    texts = _texts(c, ah, cid, _post(c, ah, cid, "export").json())
    for kind, text in texts.items():
        assert "Draft: not published" in text, kind


def test_the_published_files_carry_no_draft_marker(app_client) -> None:  # noqa: F811
    c, _, _, cid, ah, _, _ = _world(app_client)
    texts = _texts(c, ah, cid, _post(c, ah, cid, "publish").json())
    for kind, text in texts.items():
        assert "Draft: not published" not in text, kind


@pytest.mark.parametrize(
    "state, label",
    [
        ("draft", " (from a draft assessment)"),
        ("submitted", " (from a submitted assessment)"),
        ("approved", " (from an approved assessment)"),
    ],
)
def test_every_source_state_label_reads_correctly(state: str, label: str) -> None:
    """#737: every label the export's source cell can render, as literals --
    "an" before a vowel. A released input carries none."""
    from types import SimpleNamespace

    from app.risk.exporters import _source

    e = SimpleNamespace(source="coverage_finding", source_id="T1078")
    assert _source(e, {"T1078": state}, csf_playbook=False) == f"coverage_finding:T1078{label}"
    assert _source(e, {}, csf_playbook=False) == "coverage_finding:T1078"
