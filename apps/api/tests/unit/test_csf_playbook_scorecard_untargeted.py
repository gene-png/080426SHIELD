"""The CSF Playbook's PER-FUNCTION figures state their untargeted rows (#765).

#764 made the assessment-level sentences say when in-scope subcategories have
no target (#762). The per-function figures still did not: the scorecard (exec
and full, PDF and DOCX) took Target as the max over the TARGETED rows and Gaps
as a count of targeted gaps, so a function with one row targeted at L3 and met
plus two untargeted rows printed "Subcategories 3 · Target L3 · Gaps 0" -- the
whole function meeting L3. The full files' per-function line said the same.
Three states per function (`CLAUDE.md`'s three-value CHECK): has gaps / meets
its target / NO TARGET.

Item 2 of #765 rides along: with no row targeted, the overview no longer reads
"0 subcategories fall short of their target maturity — 0 Priority 1 ...".

Every case reads all five rendered files back. The XLSX has no per-function
figure, so for it the check is that its per-row Gap cells are unchanged.
"""

from __future__ import annotations

import io
import re
from types import SimpleNamespace

import pytest

from app.csf.playbook_export import (
    render_exec_docx,
    render_exec_pdf,
    render_full_docx,
    render_full_pdf,
    render_xlsx,
)


def _row(code: str, level: int, target: int | None) -> SimpleNamespace:
    gap = target is not None and level < target
    return SimpleNamespace(
        subcategory_code=code,
        name=f"Outcome {code}",
        function=code[:2],
        tier_levels={"moderate": level},
        enterprise_level=level,
        rollup_rule=1,
        target_level=target,
        gap=gap,
        priority="P2" if gap else None,
    )


def _three_functions() -> list[SimpleNamespace]:
    """One assessment, each per-function state once, so a value cannot bleed
    from one function into the next unseen."""
    return [
        # Govern: #765's own example -- one row targeted at L3 and met, two
        # untargeted.
        _row("GV.OC-01", 3, 3),
        _row("GV.OC-02", 2, None),
        _row("GV.OC-03", 1, None),
        # Identify: every row targeted, one gap. The pre-#765 wording is true.
        _row("ID.AM-01", 4, 3),
        _row("ID.AM-02", 2, 3),
        # Protect: no row targeted.
        _row("PR.AA-01", 4, None),
    ]


def _untargeted_world() -> list[SimpleNamespace]:
    return [_row("GV.OC-01", 2, None), _row("ID.AM-02", 1, None), _row("PR.AA-01", 4, None)]


def _targeted_world() -> list[SimpleNamespace]:
    return [_row("GV.OC-01", 4, 3), _row("ID.AM-02", 4, 4)]


def _flat(text: str) -> str:
    return re.sub(r"\s+", " ", text)


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    return _flat("".join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages))


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return _flat("\n".join(parts))


def _xlsx_gap_column(raw: bytes) -> dict[str, str]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(raw))["Enterprise Profile"]
    header_row = next(
        r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Subcategory"
    )
    gap_col = [c.value for c in ws[header_row]].index("Gap") + 1
    return {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=gap_col).value or ""
        for r in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row=r, column=1).value
    }


def _render(rows: list[SimpleNamespace]) -> dict[str, object]:
    kw = {"client_name": "Atlas", "version": 1, "enterprise_rows": rows, "approved": True}
    return {
        "xlsx": _xlsx_gap_column(render_xlsx(**kw, tier_profiles={}, gap_actions={})),
        "exec_pdf": _pdf_text(render_exec_pdf(**kw)),
        "exec_docx": _docx_text(render_exec_docx(**kw)),
        "full_pdf": _pdf_text(render_full_pdf(**kw)),
        "full_docx": _docx_text(render_full_docx(**kw)),
    }


PDFS = ("exec_pdf", "full_pdf")
DOCXS = ("exec_docx", "full_docx")
FULL = ("full_pdf", "full_docx")


def _scorecard_row(name: str, cells: tuple[str, ...]) -> str:
    """A scorecard row as the file's text reads it back: DOCX table cells are
    joined with " | " by `_docx_text`, PDF table cells with a space."""
    return (" | " if name in DOCXS else " ").join(cells)


def _assert_scorecard(out: dict[str, object], rows: list[tuple[str, ...]]) -> None:
    for name in (*PDFS, *DOCXS):
        text = out[name]
        assert isinstance(text, str)
        for cells in rows:
            assert (
                _scorecard_row(name, cells) in text
            ), f"{name}: scorecard row {cells!r} not found in {text[:1500]!r}"


# ---------------------------------------------------------------------------
# One assessment: a mixed function, an all-targeted one, an all-untargeted one
# ---------------------------------------------------------------------------


@pytest.mark.unit
def test_the_scorecard_states_each_functions_untargeted_rows_and_nothing_bleeds() -> None:
    out = _render(_three_functions())
    _assert_scorecard(
        out,
        [
            ("Govern", "3", "L2", "L3", "0 (2 no target)"),
            ("Identify", "2", "L3", "L3", "1"),
            ("Protect", "1", "L4", "No target", "—"),
        ],
    )


@pytest.mark.unit
def test_the_full_files_function_line_states_each_functions_untargeted_rows() -> None:
    out = _render(_three_functions())
    for name in FULL:
        text = out[name]
        assert isinstance(text, str)
        for line in (
            "3 subcategories · average Level 2 · 0 gaps among the 1 with a target · "
            "2 with no target set.",
            "2 subcategories · average Level 3 · 1 gap.",
            "1 subcategory · average Level 4 · no target set, so not assessed for gaps.",
        ):
            assert line in text, f"{name}: {line!r} not in {text[:3000]!r}"


@pytest.mark.unit
def test_one_untargeted_row_beside_two_targeted_is_counted_in_number() -> None:
    rows = [_row("GV.OC-01", 3, 3), _row("GV.OC-02", 1, 3), _row("GV.OC-03", 2, None)]
    out = _render(rows)
    _assert_scorecard(out, [("Govern", "3", "L2", "L3", "1 (1 no target)")])
    for name in FULL:
        assert (
            "3 subcategories · average Level 2 · 1 gap among the 2 with a target · "
            "1 with no target set." in out[name]
        ), name


@pytest.mark.unit
def test_the_xlsx_gap_cells_are_unchanged() -> None:
    out = _render(_three_functions())
    assert out["xlsx"] == {
        "GV.OC-01": "",
        "GV.OC-02": "No target",
        "GV.OC-03": "No target",
        "ID.AM-01": "",
        "ID.AM-02": "Yes",
        "PR.AA-01": "No target",
    }


# ---------------------------------------------------------------------------
# No row targeted anywhere (the default after seed + Run AI)
# ---------------------------------------------------------------------------


@pytest.mark.unit
def test_with_no_target_anywhere_every_function_says_so() -> None:
    out = _render(_untargeted_world())
    _assert_scorecard(
        out,
        [
            ("Govern", "1", "L2", "No target", "—"),
            ("Identify", "1", "L1", "No target", "—"),
            ("Protect", "1", "L4", "No target", "—"),
        ],
    )
    assert set(out["xlsx"].values()) == {"No target"}  # type: ignore[union-attr]


@pytest.mark.unit
def test_with_no_target_anywhere_the_overview_does_not_count_zero_shortfalls() -> None:
    # #765 item 2.
    out = _render(_untargeted_world())
    for name in (*PDFS, *DOCXS):
        text = out[name]
        assert isinstance(text, str)
        assert (
            "No gaps or priorities are assessed: 3 in-scope subcategories have no target set."
            in text
        ), f"{name}: {text[:1500]!r}"
        assert "fall short of their target maturity" not in text, name


@pytest.mark.unit
def test_with_one_untargeted_row_only_the_overview_is_in_number() -> None:
    out = _render([_row("GV.OC-01", 2, None)])
    for name in (*PDFS, *DOCXS):
        assert (
            "No gaps or priorities are assessed: 1 in-scope subcategory has no target set."
            in out[name]
        ), name


@pytest.mark.unit
def test_with_some_targets_the_overview_keeps_its_shortfall_count() -> None:
    out = _render(_three_functions())
    for name in (*PDFS, *DOCXS):
        text = out[name]
        assert isinstance(text, str)
        assert "1 subcategory falls short of its target maturity" in text, name
        assert "No gaps or priorities are assessed" not in text, name


# ---------------------------------------------------------------------------
# Control: every row targeted. The pre-#765 figures are the true ones.
# ---------------------------------------------------------------------------


@pytest.mark.unit
def test_a_fully_targeted_assessment_keeps_its_figures() -> None:
    out = _render(_targeted_world())
    _assert_scorecard(
        out,
        [("Govern", "1", "L4", "L3", "0"), ("Identify", "1", "L4", "L4", "0")],
    )
    for name in FULL:
        assert "1 subcategory · average Level 4 · 0 gaps." in out[name], name
    for name in (*PDFS, *DOCXS):
        text = out[name]
        assert isinstance(text, str)
        assert "0 subcategories fall short of their target maturity" in text, name
        assert "no target" not in text.lower(), name
    assert set(out["xlsx"].values()) == {""}  # type: ignore[union-attr]
