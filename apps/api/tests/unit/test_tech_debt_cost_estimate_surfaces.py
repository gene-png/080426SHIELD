"""TD-5 (for #1060): an AI-estimated cost is labelled everywhere it appears.

Copy approved on #736 (6205600298, ruled 6206818759 item 5), typed here:

  E5  PDF/DOCX cost cell "$12,000 (AI estimate)"; XLSX keeps the number and
      gains a LAST column "Cost source" reading "AI estimate".
  E6  files: "Note: {k} annual costs are AI estimates; the annual cost includes
      them." / "Note: 1 annual cost is an AI estimate; the annual cost includes
      it."
  E7  any estimated cost makes the dashboard's `spend_completeness` "partial".

The rule for "is an estimate" is `cost_estimate.is_ai_estimate`: a source item
whose `cost_source` is ai_estimate and whose cost is set. An overridden cost
(`consultant`) and an uploaded one (None) are not estimates.
"""

from __future__ import annotations

import io
import os

import pytest

from app.tech_debt.exporters import render_docx, render_pdf, render_xlsx
from tests.unit.test_exporters import _item
from tests.unit.test_tech_debt_dashboard import (  # noqa: F401  (fixture)
    _register,
    _seed_release,
    app_client,
)
from tests.unit.test_tech_debt_no_cost_count import RENDERERS, _ctx, _pdf_text

pytestmark = pytest.mark.unit

PLURAL = "Note: 2 annual costs are AI estimates; the annual cost includes them."
SINGULAR = "Note: 1 annual cost is an AI estimate; the annual cost includes it."


def _estimated(name: str, cost: float = 12_000) -> object:
    return _item(name=name, annual_cost_usd=cost, cost_source="ai_estimate")


@RENDERERS
def test_the_files_state_how_many_costs_are_estimates(render, read) -> None:
    items = [
        _item(name="Tool A", annual_cost_usd=1_000),
        _estimated("Tool B"),
        _estimated("Tool C"),
        _item(name="Tool D", annual_cost_usd=500, cost_source="consultant"),
    ]
    text = read(render(_ctx(items)))
    assert PLURAL in text
    assert "is an AI estimate" not in text


@RENDERERS
def test_the_files_use_the_singular_for_one_estimate(render, read) -> None:
    text = read(render(_ctx([_item(name="Tool A", annual_cost_usd=1_000), _estimated("Tool B")])))
    assert SINGULAR in text
    assert "are AI estimates" not in text


@RENDERERS
def test_an_overridden_or_uploaded_cost_is_not_an_estimate(render, read) -> None:
    items = [
        _item(name="Tool A", annual_cost_usd=1_000),
        _item(name="Tool B", annual_cost_usd=2_000, cost_source="consultant"),
    ]
    text = read(render(_ctx(items)))
    # POSITIVE CONTROL first: the summary rendered and states the total.
    assert "Total annual cost" in text
    assert "AI estimate" not in text


@pytest.mark.parametrize("render", [render_pdf, render_docx], ids=["pdf", "docx"])
def test_the_pdf_and_docx_cost_cell_says_ai_estimate(render) -> None:
    from tests.unit.test_tech_debt_no_cost_count import _docx_text

    read = _pdf_text if render is render_pdf else _docx_text
    text = read(render(_ctx([_item(name="Tool A", annual_cost_usd=1_000), _estimated("Tool B")])))
    assert "$12,000 (AI estimate)" in text
    assert "$1,000 (AI estimate)" not in text


def test_the_xlsx_keeps_the_number_and_adds_a_last_cost_source_column() -> None:
    from openpyxl import load_workbook

    raw = render_xlsx(_ctx([_item(name="Tool A", annual_cost_usd=1_000), _estimated("Tool B")]))
    ws = load_workbook(io.BytesIO(raw)).worksheets[0]
    header = [c.value for c in ws[1]]
    assert header[4] == "Annual Cost (USD)"  # an existing column, not moved
    assert header[-1] == "Cost source"
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2) if r[0].value in ("Tool A", "Tool B")}
    assert rows["Tool B"][4].value == 12000  # the number stays a number
    assert rows["Tool B"][len(header) - 1].value == "AI estimate"
    assert rows["Tool A"][len(header) - 1].value in (None, "")


@RENDERERS
def test_a_total_including_an_estimate_is_not_called_a_total(render, read) -> None:
    """The deliverable's twin of E7: the dashboard calls this list "partial", so
    the document must not call its figure a "Total" (#126's rule: two surfaces,
    one list, no contradictory claims)."""
    text = read(render(_ctx([_item(name="Tool A", annual_cost_usd=1_000), _estimated("Tool B")])))
    assert "Total annual cost" not in text
    assert "Annual cost (may not be complete)" in text


# ---------------------------------------------------------------------------
# The client dashboard
# ---------------------------------------------------------------------------


def _mark_estimated(*names: str) -> None:
    import sqlalchemy as sa

    eng = sa.create_engine(os.environ["DATABASE_URL"], future=True)
    with eng.begin() as conn:
        for name in names:
            n = conn.execute(
                sa.text(
                    "UPDATE capability_items SET cost_source = 'ai_estimate', "
                    "cost_estimate_basis = 'b' WHERE name = :n AND annual_cost_usd IS NOT NULL"
                ),
                {"n": name},
            ).rowcount
            assert n == 1, f"expected one costed {name!r} row, updated {n}"


def _dashboard(c, client, svc_id: str) -> dict:
    client_id = client["user"]["client_id"]
    c.headers["X-Client-Id"] = client_id
    r = c.get(
        f"/clients/{client_id}/tech-debt/{svc_id}/dashboard",
        headers={"Authorization": f"Bearer {client['tokens']['access_token']}"},
    )
    assert r.status_code == 200, r.text
    return r.json()


def test_an_estimate_makes_the_dashboard_partial_and_is_counted(app_client) -> None:  # noqa: F811
    c, provider = app_client
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    svc_id = _seed_release(c, provider, admin["tokens"]["access_token"], release=True)
    # POSITIVE CONTROL: the seeded list reads complete before anything is estimated.
    before = _dashboard(c, client, svc_id)
    assert before["spend_completeness"] == "complete"
    assert before["ai_estimate_count"] == 0
    _mark_estimated("Splunk ES")
    after = _dashboard(c, client, svc_id)
    assert after["ai_estimate_count"] == 1
    assert after["spend_completeness"] == "partial"
    # The estimate is INCLUDED in spend, and nothing is uncosted.
    assert after["annual_spend_usd"] == before["annual_spend_usd"]
    assert after["no_cost_count"] == 0


def test_the_dashboard_marks_each_estimated_row(app_client) -> None:  # noqa: F811
    """E3: each estimate is labelled wherever it appears, the client's
    inventory and redundancy rows included."""
    c, provider = app_client
    admin = _register(c, "admin@example.com")
    client = _register(c, "client@example.com")
    svc_id = _seed_release(c, provider, admin["tokens"]["access_token"], release=True)
    _mark_estimated("Splunk ES")
    body = _dashboard(c, client, svc_id)
    flags = {i["name"]: i["ai_estimate"] for i in body["items"]}
    assert flags["Splunk ES"] is True
    # POSITIVE CONTROL: a supplied cost is not marked.
    assert flags["CrowdStrike Falcon"] is False
