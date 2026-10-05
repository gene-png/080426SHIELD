"""#554 R3 in the Risk Register, under the advisor's ruling on #737 (#736
5998764095, option (b)): a register DRAFTED while the ATT&CK review queue holds
rows is generated, and each finding for a technique awaiting review is
labelled; publication is what refuses it.

It used to be refused at generate (the 2026-10-03 00:35Z fail-closed ruling).
That protection moved, and it is still there: an assessment with unreviewed
computed statuses cannot be RELEASED (`release_readiness.blocking_condition`,
pinned by `test_attack_computed_status_routes.py::test_release_is_refused_until_the_queue_is_reviewed`),
and publish needs every engaged input released. Each test here declares the
change it made, old vs new, on #860.
"""

from __future__ import annotations

import pytest

from app.ai.llm import LLMResponse
from tests._attack_rows import standalone_rows
from tests.unit.test_risk_register import _admin, app_client  # noqa: F401  (fixture)

pytestmark = pytest.mark.unit

ALL_THREE = {
    "detection_tools": ["Tool D"],
    "prevention_tools": ["Tool P"],
    "response_tools": ["Tool R"],
}


def _approved_attack_and_zt(c, bearer: str, cid: str, n: int = 1) -> tuple[dict, str]:
    """An approved R3 ATT&CK assessment whose one scored row the AI called Gap
    with all three in place (computed: Covered, so it awaits review), and an
    approved ZT assessment so the register is unlocked."""
    h = {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}
    svc = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": "A"})
    a = c.post(f"/attack/services/{svc.json()['id']}/assessments", headers=h).json()
    rows = standalone_rows(a["coverage"], n)
    for cov in rows:
        r = c.patch(f"/attack/coverage/{cov['id']}", headers=h, json={"status": "gap", **ALL_THREE})
        assert r.status_code == 200, r.text
        assert r.json()["in_review_queue"] is True
    cov = rows[0]
    zsvc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": "ZT"})
    za = c.post(f"/zt/services/{zsvc.json()['id']}/assessments", headers=h).json()
    zans = za["answers"][0]
    c.patch(f"/zt/answers/{zans['id']}", headers=h, json={"maturity_stage": 1})
    assert c.post(f"/attack/assessments/{a['id']}/approve", headers=h).status_code == 200
    assert c.post(f"/zt/assessments/{za['id']}/approve", headers=h).status_code == 200
    return {"assessment": a, "code": cov["technique_code"]}, zans["capability_code"]


def _covered_with_no_tools(c, bearer: str, cid: str, n: int = 1) -> tuple[dict, str]:
    """Like `_approved_attack_and_zt`, but the AI called each row Covered with no
    tools: the computed status is Gap, which DIFFERS from the AI's (so the row
    awaits review) and IS a finding (so the register has an entry to label)."""
    h = {"Authorization": f"Bearer {bearer}", "X-Client-Id": cid}
    svc = c.post("/attack/services", headers=h, json={"kind": "attack_coverage", "title": "A"})
    a = c.post(f"/attack/services/{svc.json()['id']}/assessments", headers=h).json()
    rows = standalone_rows(a["coverage"], n)
    for cov in rows:
        r = c.patch(f"/attack/coverage/{cov['id']}", headers=h, json={"status": "covered"})
        assert r.status_code == 200, r.text
        assert r.json()["in_review_queue"] is True
    zsvc = c.post("/zt/services", headers=h, json={"kind": "zero_trust_cisa", "title": "ZT"})
    za = c.post(f"/zt/services/{zsvc.json()['id']}/assessments", headers=h).json()
    zans = za["answers"][0]
    c.patch(f"/zt/answers/{zans['id']}", headers=h, json={"maturity_stage": 1})
    assert c.post(f"/attack/assessments/{a['id']}/approve", headers=h).status_code == 200
    assert c.post(f"/zt/assessments/{za['id']}/approve", headers=h).status_code == 200
    return {
        "assessment": a,
        "service": svc.json()["id"],
        "codes": [cov["technique_code"] for cov in rows],
    }, zans["capability_code"]


def _entry_for(code: str) -> str:
    return (
        '{"title": "Exposure ' + code + '", "description": "d", "axis": "detection",'
        ' "source_id": "' + code + '", "linked_techniques": [], "linked_controls": [],'
        ' "likelihood": "high", "impact": "major", "recommended_action": "remediate",'
        ' "rationale": "r"}'
    )


def test_generate_drafts_an_unreviewed_assessment_and_publish_refuses(
    app_client,  # noqa: F811
) -> None:
    """DECLARED (#860, accepted). OLD: generate 409
    `attack_computed_status_unreviewed`. NEW: generate 201; the finding for the
    unreviewed technique is labelled; publish 409, because the ATT&CK input is
    not released -- and cannot be: release refuses while the queue holds it."""
    c, provider = app_client
    bearer, cid = _admin(c)
    world, _capability = _covered_with_no_tools(c, bearer, cid)
    [code] = world["codes"]
    bh = {"Authorization": f"Bearer {bearer}"}
    provider.register_static(
        "risk_synthesize", LLMResponse('{"entries": [' + _entry_for(code) + "]}")
    )

    gen = c.post(f"/risk/clients/{cid}/register/generate", headers=bh)
    assert gen.status_code == 201, gen.text
    [entry] = gen.json()["entries"]
    assert entry["source_id"] == code
    assert entry["source_review_pending"] is True

    pub = c.post(f"/risk/clients/{cid}/register/publish", headers=bh)
    assert pub.status_code == 409, pub.text
    assert pub.json()["error"]["reason"] == "risk_register_inputs_not_final"
    assert {"input": "attack", "reason": "not_released", "status": "approved"} in pub.json()[
        "error"
    ]["blockers"]

    # And the input cannot become final while the queue holds the code.
    h = {**bh, "X-Client-Id": cid}
    fin = c.post(f"/attack/services/{world['service']}/deliverables/finalize", headers=h)
    assert fin.status_code in (200, 201), fin.text
    rel = c.post(f"/attack/deliverables/{fin.json()['id']}/release", headers=h)
    assert rel.status_code == 409, rel.text
    assert rel.json()["error"]["reason"] == "attack_computed_status_unreviewed"


def test_the_export_labels_a_finding_awaiting_review(app_client) -> None:  # noqa: F811
    """DECLARED (#860, accepted: repointed, not deleted). OLD: pinned the
    generate refusal's plural sentence. NEW: there is no refusal at generate,
    so it pins the per-finding label -- a literal, in the downloaded XLSX."""
    import io

    from openpyxl import load_workbook

    c, provider = app_client
    bearer, cid = _admin(c)
    world, _capability = _covered_with_no_tools(c, bearer, cid, n=2)
    bh = {"Authorization": f"Bearer {bearer}"}
    entries = ", ".join(_entry_for(code) for code in world["codes"])
    provider.register_static("risk_synthesize", LLMResponse('{"entries": [' + entries + "]}"))
    assert c.post(f"/risk/clients/{cid}/register/generate", headers=bh).status_code == 201
    ex = c.post(f"/risk/clients/{cid}/register/export", headers=bh).json()
    raw = c.get(
        f"/artifacts/{ex['xlsx_artifact_id']}/download", headers={**bh, "X-Client-Id": cid}
    ).content
    rows = list(load_workbook(io.BytesIO(raw))["Risk Register"].iter_rows(values_only=True))
    sources = sorted(dict(zip(rows[0], r, strict=True))["Source"] for r in rows[1:])
    assert sources == sorted(
        f"coverage_finding:{code} (from a approved assessment) (computed status awaiting review)"
        for code in world["codes"]
    )


def test_the_gate_no_longer_blocks_on_unreviewed_statuses(app_client) -> None:  # noqa: F811
    """DECLARED (#860, accepted; the coordinator's (a): drop the sentence). OLD:
    the gate carried the generate refusal's sentence, so the page offered no
    Generate. NEW: nothing blocks Generate any more, so the gate reports null;
    the per-finding label and the Inputs panel carry the fact."""
    c, _provider = app_client
    bearer, cid = _admin(c)
    _covered_with_no_tools(c, bearer, cid)
    gate = c.get(f"/risk/clients/{cid}/gate", headers={"Authorization": f"Bearer {bearer}"}).json()
    assert gate["attack_computed_status_unreviewed"] is None
