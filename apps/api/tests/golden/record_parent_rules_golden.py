"""Record the OLD rules' client-facing output ONCE, from `main` (#620, D-094).

Gene's condition, 2026-09-25: a released assessment keeps rendering what was
delivered. Gene's addition 2: the "before" is recorded ONCE from `main` at a
named SHA and committed; the test compares against the committed files and
never regenerates them. This script is committed beside them so the
recording can be re-derived by hand.

HOW TO RE-DERIVE (Git Bash, from the repo root):

    git worktree add --detach ../golden-main <MAIN_SHA>
    cp apps/api/tests/golden/record_parent_rules_golden.py ../golden-main/apps/api/
    MSYS_NO_PATHCONV=1 docker run --rm \\
      -v "<abs path>/golden-main:/repo" -v "<abs path>/out:/out" \\
      -w /repo/apps/api shield-v2-api:latest \\
      python record_parent_rules_golden.py /out

It must run on MAIN's code, where a parent is not computed and PATCH accepts
a parent's status and tools. On #620's code the world cannot be built this
way, which is the point.

WHAT IT WRITES to <out>:
  world.sql            the database after the world is built (SQLite dump,
                       at main's migration head)
  dashboard.json       GET /clients/{cid}/attack/{svc}/dashboard, RELEASED A
  finalize_b.json      POST finalize on APPROVED-unreleased B: summary text
                       and every XLSX cell value, sheet by sheet
  risk_findings.json   the ATT&CK findings risk synthesis sends, sorted

THE WORLD: one tenant, an admin and a client user, and three services.
  A (RELEASED)  T1001 hand-scored `covered` WITH its own tools, rationale and
                an uncleared citation of its own (pending under main's rule);
                its sub-techniques mixed; one standalone covered.
  B (APPROVED)  T1003 hand-scored `gap` with its own rationale; its
                sub-techniques mixed. Not finalized, so #620's code finalizes
                it -- the exporter path for "approved before, finalized after".
  ZT            one answer, approved, so the Risk Register can be generated.
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
import sys
import tempfile
import uuid
from pathlib import Path

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
work = Path(tempfile.mkdtemp())
DB = work / "world.db"
os.environ["DATABASE_URL"] = f"sqlite:///{DB}"

from alembic import command  # noqa: E402
from alembic.config import Config  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

API_ROOT = Path.cwd()
cfg = Config(str(API_ROOT / "alembic.ini"))
cfg.set_main_option("script_location", str(API_ROOT / "alembic"))
cfg.set_main_option("sqlalchemy.url", os.environ["DATABASE_URL"])
command.upgrade(cfg, "head")
Sess = sessionmaker(bind=create_engine(os.environ["DATABASE_URL"], future=True))

from app.ai.llm import FixtureProvider, LLMClient, LLMResponse  # noqa: E402
from app.db.session import get_db  # noqa: E402
from app.main import create_app  # noqa: E402
from app.models.attack_assessment import AttackCoverage  # noqa: E402
from app.models.client import Client  # noqa: E402
from app.models.client_domain import ClientDomain  # noqa: E402
from app.routes.artifacts import _storage_dep  # noqa: E402
from app.routes.risk import _llm_dep  # noqa: E402
from app.storage.local import LocalFilesystemStorage  # noqa: E402

storage = LocalFilesystemStorage(work / "storage")
provider = FixtureProvider()


def _db():
    s = Sess()
    try:
        yield s
    finally:
        s.close()


app = create_app()
app.dependency_overrides[get_db] = _db
app.dependency_overrides[_storage_dep] = lambda: storage
app.dependency_overrides[_llm_dep] = lambda: LLMClient(provider)

with Sess() as s:
    tenant = Client(legal_name="Golden Tenant")
    s.add(tenant)
    s.flush()
    s.add(ClientDomain(client_id=tenant.id, domain="example.com"))
    s.commit()
    CID = str(tenant.id)

PASSWORD = "correct horse battery staple!"


def ok(r, code=200):
    assert r.status_code == code, (r.status_code, r.text)
    return r.json()


with TestClient(app, headers={"X-Client-Id": CID}) as c:
    admin = ok(
        c.post(
            "/auth/register",
            json={"email": "admin@example.com", "password": PASSWORD, "display_name": "A"},
        ),
        201,
    )
    client = ok(
        c.post(
            "/auth/register",
            json={"email": "client@example.com", "password": PASSWORD, "display_name": "C"},
        ),
        201,
    )
    H = {"Authorization": f"Bearer {admin['tokens']['access_token']}"}
    CH = {"Authorization": f"Bearer {client['tokens']['access_token']}"}

    def assessment(title: str) -> tuple[str, str, dict]:
        svc = ok(
            c.post("/attack/services", headers=H, json={"kind": "attack_coverage", "title": title}),
            201,
        )["id"]
        a = ok(c.post(f"/attack/services/{svc}/assessments", headers=H), 201)
        return svc, a["id"], {r["technique_code"]: r["id"] for r in a["coverage"]}

    def patch(rows: dict, code: str, body: dict) -> None:
        ok(c.patch(f"/attack/coverage/{rows[code]}", headers=H, json=body))

    # --- A: released ---------------------------------------------------------
    svc_a, a_id, rows = assessment("Golden A")
    patch(rows, "T1001", {"status": "covered", "detection_tools": [], "rationale": "Parent text."})
    patch(rows, "T1001.001", {"status": "gap", "rationale": "Child one."})
    patch(rows, "T1001.002", {"status": "covered", "rationale": "Child two."})
    patch(rows, "T1001.003", {"status": "partial", "rationale": "Child three."})
    patch(rows, "T1005", {"status": "covered", "rationale": "Standalone."})
    with Sess() as s:
        # Legacy state no PATCH can write: the parent's own tools and an
        # uncleared citation of its own -- pending under main's per-row rule.
        row = s.get(AttackCoverage, uuid.UUID(rows["T1001"]))
        row.detection_tools = ["Tool A"]
        row.prevention_tools = ["Tool A"]
        row.unconfirmed_citations = [
            {
                "tool": "Tool A",
                "cited": "tool",
                "reason": "substring",
                "field": "detection_tools",
                "cleared_at": None,
            }
        ]
        s.commit()
    ok(c.post(f"/attack/assessments/{a_id}/approve", headers=H))
    deliv_a = ok(c.post(f"/attack/services/{svc_a}/deliverables/finalize", headers=H), 201)["id"]
    ok(c.post(f"/attack/deliverables/{deliv_a}/release", headers=H))

    # --- B: approved, not finalized ------------------------------------------
    svc_b, b_id, rows_b = assessment("Golden B")
    patch(rows_b, "T1003", {"status": "gap", "rationale": "Parent gap text."})
    patch(rows_b, "T1003.001", {"status": "gap"})
    patch(rows_b, "T1003.002", {"status": "covered"})
    ok(c.post(f"/attack/assessments/{b_id}/approve", headers=H))

    # --- ZT, for the Risk Register gate ---------------------------------------
    zsvc = ok(
        c.post("/zt/services", headers=H, json={"kind": "zero_trust_cisa", "title": "ZT"}), 201
    )["id"]
    za = ok(c.post(f"/zt/services/{zsvc}/assessments", headers=H), 201)
    ok(c.patch(f"/zt/answers/{za['answers'][0]['id']}", headers=H, json={"maturity_stage": 1}))
    ok(c.post(f"/zt/assessments/{za['id']}/approve", headers=H))

    # The world is complete: dump it BEFORE reading anything, so the new code
    # starts from exactly this state.
    con = sqlite3.connect(DB)
    (OUT / "world.sql").write_text("\n".join(con.iterdump()) + "\n", encoding="utf-8")
    con.close()

    # --- outputs ---------------------------------------------------------------
    dash = c.get(f"/clients/{CID}/attack/{svc_a}/dashboard", headers={**CH, "X-Client-Id": CID})
    assert dash.status_code == 200, dash.text
    (OUT / "dashboard.json").write_bytes(dash.content)

    fin = ok(c.post(f"/attack/services/{svc_b}/deliverables/finalize", headers=H), 201)
    # B's XLSX is the newest file in storage: A's was written before the dump.
    wb = load_workbook(
        io.BytesIO(
            sorted((work / "storage").rglob("*.xlsx"), key=lambda p: p.stat().st_mtime)[
                -1
            ].read_bytes()
        )
    )
    cells = {
        ws.title: [[cell.value for cell in row] for row in ws.iter_rows()] for ws in wb.worksheets
    }
    (OUT / "finalize_b.json").write_text(
        json.dumps(
            {"summary": fin["summary"], "xlsx": cells}, indent=1, sort_keys=True, default=str
        )
        + "\n",
        encoding="utf-8",
    )

    sent: list[str] = []

    def spy(payload: dict) -> LLMResponse:
        sent.extend(
            f["source_id"] for f in payload.get("findings", []) if f.get("kind") == "attack"
        )
        return LLMResponse(json.dumps({"entries": []}))

    provider.register("risk_synthesize", spy)
    c.post(f"/risk/clients/{CID}/register/generate", headers=H)
    (OUT / "risk_findings.json").write_text(
        json.dumps(sorted(sent), indent=1) + "\n", encoding="utf-8"
    )

print("recorded to", OUT)
